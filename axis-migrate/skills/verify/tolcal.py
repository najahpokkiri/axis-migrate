#!/usr/bin/env python3
"""tolcal.py — measure the nearest-wrong side of a tolerance by deterministic mutation of a legacy file, and turn
two measured numbers into a PROPOSED tolerance record (06 §3.2, DESIGN §4.6).

  mutate  --legacy L --kind raster|vector|table --out DIR [--spec spec.json] [--recipe rasterize.json]
          [--only id,id] [--key col]… [--json]
      Writes each mutant into DIR, compares it with compare.compare_pair (the comparator under calibration) and
      writes DIR/tolcal.json: per mutant the verdict, reason codes and the measured metric of every check, plus a
      per-check summary of the smallest failing value (the nearest wrong recipe among these mutants).
      Raster mutants: shift_k<k> (transform moved by k/2 cells on both axes, values copied 1:1), bilinear_roundtrip
      (bilinear to a half-cell-shifted grid of the same resolution and back), all_touched_toggle (only with a rasterize
      recipe), scale_k<k> (values × 10^k). Vector/table mutants: unstrip_key (a trailing space on one key),
      drop_rows (first rows by key order), scale_k<k> (numeric columns × 10^k).
  propose --dataset ds --check chk --correct-spread X --nearest-wrong Y [--form rtol|atol|value|max_mismatch] …
      Prints one YAML record with agreed_by: null. Code never signs: only a human fills agreed_by/agreed_at.

The correct-spread number is not produced here: it comes from an independent correct implementation of the step.
Without it there is no record and the check stays exact (DESIGN §9).
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import shutil
import sys
from pathlib import Path

os.environ["GDAL_PAM_ENABLED"] = "NO"
os.environ.setdefault("OGR_CURRENT_DATE", "2000-01-01T00:00:00Z")    # byte-stable GPKG mutants

import numpy as np  # noqa: E402

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "lib"))
sys.path.insert(0, str(HERE))
import axis_common as ac  # noqa: E402
import compare  # noqa: E402

TOLCAL_SCHEMA = "axis-migrate.tolcal.v1"


# ------------------------------------------------------------------------------------------------ raster mutants
def _raster_profile(path):
    import rasterio
    with rasterio.open(path) as ds:
        return ds.profile.copy(), ds.read(), ds.transform


def _write_raster(path: Path, profile: dict, data: np.ndarray, transform=None) -> None:
    import rasterio
    p = dict(profile)
    if transform is not None:
        p["transform"] = transform
    p["count"], p["height"], p["width"] = data.shape
    with rasterio.open(path, "w", **p) as dst:
        dst.write(data.astype(p["dtype"]))


def _valid_mask(data: np.ndarray, nodata) -> np.ndarray:
    v = np.ones(data.shape, dtype=bool)
    if data.dtype.kind == "f":
        v &= np.isfinite(data)
    if nodata is not None and not (isinstance(nodata, float) and math.isnan(nodata)):
        v &= data != nodata
    return v


def raster_mutants(legacy: str, out: Path, *, recipe: dict | None, shift_k=(1, 2), scale_k=(-3, 3)) -> list[dict]:
    from affine import Affine
    from rasterio.warp import Resampling, reproject
    prof, data, tr = _raster_profile(legacy)
    nd = prof.get("nodata")
    valid = _valid_mask(data, nd)
    muts = []
    for k in shift_k:
        p = out / f"shift_k{k}.tif"
        t2 = Affine(tr.a, tr.b, tr.c + tr.a * k / 2.0, tr.d, tr.e, tr.f + tr.e * k / 2.0)
        _write_raster(p, prof, data, t2)
        muts.append({"id": f"shift_k{k}", "recipe": f"transform moved by {k}/2 cells on both axes, values 1:1",
                     "path": p})
    # bilinear round trip on the same grid (via a half-cell-shifted grid of the same resolution)
    src = np.where(valid, data, np.nan).astype("float64")
    mid_t = Affine(tr.a, tr.b, tr.c + tr.a / 2.0, tr.d, tr.e, tr.f + tr.e / 2.0)
    mid = np.full(src.shape, np.nan)
    back = np.full(src.shape, np.nan)
    for b in range(src.shape[0]):
        reproject(src[b], mid[b], src_transform=tr, src_crs=prof["crs"], dst_transform=mid_t, dst_crs=prof["crs"],
                  src_nodata=np.nan, dst_nodata=np.nan, resampling=Resampling.bilinear)
        reproject(mid[b], back[b], src_transform=mid_t, src_crs=prof["crs"], dst_transform=tr, dst_crs=prof["crs"],
                  src_nodata=np.nan, dst_nodata=np.nan, resampling=Resampling.bilinear)
    fill = nd if nd is not None else np.nan
    rt = np.where(np.isfinite(back), back, fill)
    if np.dtype(prof["dtype"]).kind in "iu":
        rt = np.round(rt)
    p = out / "bilinear_roundtrip.tif"
    _write_raster(p, prof, rt)
    muts.append({"id": "bilinear_roundtrip", "recipe": "bilinear to a half-cell-shifted grid and back", "path": p})
    if recipe:
        from rasterio.features import rasterize
        import pyogrio
        import pyproj
        df = pyogrio.read_dataframe(recipe["vector"], layer=recipe.get("layer"))
        if recipe.get("where"):
            col, val = recipe["where"]
            df = df[df[col] == val]
        if df.crs is not None and not pyproj.CRS(df.crs).equals(pyproj.CRS.from_wkt(prof["crs"].to_wkt())):
            df = df.to_crs(pyproj.CRS.from_wkt(prof["crs"].to_wkt()))
        burn = recipe.get("burn", 1)
        shapes = [(g, (row[burn] if isinstance(burn, str) else burn)) for g, (_, row) in
                  zip(df.geometry, df.iterrows()) if g is not None]
        arr = rasterize(shapes, out_shape=(prof["height"], prof["width"]), transform=tr,
                        fill=recipe.get("fill", 0), all_touched=not bool(recipe.get("all_touched", False)),
                        dtype=prof["dtype"])
        p = out / "all_touched_toggle.tif"
        _write_raster(p, prof, arr[None, ...])
        muts.append({"id": "all_touched_toggle",
                     "recipe": f"rasterize with all_touched={not bool(recipe.get('all_touched', False))}", "path": p})
    for k in scale_k:
        p = out / f"scale_k{k}.tif"
        sc = np.where(valid, data.astype("float64") * (10.0 ** k), data)
        _write_raster(p, prof, sc)
        muts.append({"id": f"scale_k{k}", "recipe": f"values × 10^{k}", "path": p})
    return muts


# ------------------------------------------------------------------------------------------- vector/table mutants
def _numeric_cols(df, exclude):
    import pandas as pd
    return [c for c in df.columns if c not in exclude and c != getattr(df, "_geometry_column_name", None)
            and pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c])]


def vector_mutants(legacy: str, out: Path, *, key: list[str], layer: str | None, drop_n=1, scale_k=(-3, 3)):
    import pyogrio
    lay = layer or str(pyogrio.list_layers(legacy)[0][0])
    df = pyogrio.read_dataframe(legacy, layer=lay)
    order = df.sort_values(key).index
    muts = []

    def write(name, frame, recipe):
        p = out / f"{name}.gpkg"
        if p.exists():
            p.unlink()
        pyogrio.write_dataframe(frame, p, layer=lay, driver="GPKG")
        muts.append({"id": name, "recipe": recipe, "path": p})
    if key:
        m = df.copy()
        k0 = key[0]
        m[k0] = m[k0].astype(object)
        m.loc[order[0], k0] = f"{m.loc[order[0], k0]} "
        write("unstrip_key", m, f"trailing space on {k0} of the first key")
    write("drop_rows", df.drop(index=list(order[:drop_n])), f"drop the first {drop_n} rows by key order")
    for k in scale_k:
        m = df.copy()
        for c in _numeric_cols(m, set(key)):
            m[c] = m[c].astype("float64") * (10.0 ** k)
        write(f"scale_k{k}", m, f"numeric columns × 10^{k}")
    return muts


def table_mutants(legacy: str, out: Path, *, key: list[str], spec: dict | None, drop_n=1, scale_k=(-3, 3)):
    low = legacy.lower()
    if low.endswith(".gpkg"):
        return vector_mutants(legacy, out, key=key, layer=(spec or {}).get("layer"), drop_n=drop_n, scale_k=scale_k)
    cmp_ = (spec or {}).get("compare") or {}
    hdr = int(cmp_.get("header_row") or 1)
    if low.endswith(".csv"):
        with open(legacy, newline="", encoding="utf-8-sig") as fh:
            rows = [list(r) for r in csv.reader(fh)]
        header, body = rows[hdr - 1], rows[hdr:]
        ki = header.index(key[0]) if key else 0
        body_sorted = sorted(range(len(body)), key=lambda i: body[i][ki])
        muts = []

        def write(name, rws, recipe):
            p = out / f"{name}.csv"
            with open(p, "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh)
                for r in rows[: hdr - 1] + [header] + rws:
                    w.writerow(r)
            muts.append({"id": name, "recipe": recipe, "path": p})
        m = [list(r) for r in body]
        m[body_sorted[0]][ki] = m[body_sorted[0]][ki] + " "
        write("unstrip_key", m, f"trailing space on {header[ki]} of the first key")
        drop = set(body_sorted[:drop_n])
        write("drop_rows", [r for i, r in enumerate(body) if i not in drop], f"drop the first {drop_n} rows")
        num = [j for j in range(len(header)) if j != ki and all(_isnum(r[j]) for r in body if j < len(r) and r[j] != "")]
        for k in scale_k:
            m = [list(r) for r in body]
            for r in m:
                for j in num:
                    if j < len(r) and r[j] != "":
                        r[j] = repr(float(r[j]) * (10.0 ** k))
            write(f"scale_k{k}", m, f"numeric columns × 10^{k}")
        return muts
    if low.endswith((".xlsx", ".xlsm")):
        import openpyxl
        sheet = (spec or {}).get("sheet") or cmp_.get("sheet")
        data_rows = cmp_.get("data_rows")
        base = openpyxl.load_workbook(legacy, data_only=True)
        ws = base[sheet] if sheet else base.worksheets[0]
        header = [c.value for c in ws[hdr]]
        ki = header.index(key[0]) + 1 if key else 1
        last = data_rows[1] if data_rows else ws.max_row
        first = data_rows[0] if data_rows else hdr + 1
        rows_idx = [r for r in range(first, last + 1) if ws.cell(r, ki).value is not None]
        rows_idx.sort(key=lambda r: str(ws.cell(r, ki).value))
        muts = []

        def variant(name, fn, recipe):
            wb = openpyxl.load_workbook(legacy, data_only=True)
            w = wb[sheet] if sheet else wb.worksheets[0]
            fn(w)
            p = out / f"{name}.xlsx"
            wb.save(p)
            muts.append({"id": name, "recipe": recipe, "path": p})
        variant("unstrip_key", lambda w: w.cell(rows_idx[0], ki, f"{w.cell(rows_idx[0], ki).value} "),
                "trailing space on the first key")

        def drop(w):
            for r in rows_idx[:drop_n]:
                for c in range(1, w.max_column + 1):
                    w.cell(r, c).value = None
        variant("drop_rows", drop, f"blank the first {drop_n} rows by key order")
        for k in scale_k:
            def scale(w, k=k):
                for r in rows_idx:
                    for c in range(1, w.max_column + 1):
                        v = w.cell(r, c).value
                        if c != ki and isinstance(v, (int, float)) and not isinstance(v, bool):
                            w.cell(r, c).value = float(v) * (10.0 ** k)
            variant(f"scale_k{k}", scale, f"numeric cells × 10^{k}")
        return muts
    raise ValueError(f"table format of {legacy} not supported")


def _isnum(s) -> bool:
    try:
        float(s)
        return True
    except (TypeError, ValueError):
        return False


# ---------------------------------------------------------------------------------------------------- measuring
def _metrics(result: dict) -> dict:
    out = {}
    for c in result.get("checks", []):
        out[c["id"]] = {"state": c["state"], "metric": c.get("metric"), "value": c.get("value"),
                        "reason_codes": c.get("reason_codes", [])}
        if c.get("stats"):
            out[c["id"]]["stats"] = {k: v for k, v in c["stats"].items() if k != "quantiles_abs"}
    return out


def mutate(legacy: str, kind: str, out_dir: str, *, spec: dict | None = None, recipe: dict | None = None,
           only: list[str] | None = None, key: list[str] | None = None, shift_k=(1, 2), scale_k=(-3, 3),
           drop_n: int = 1) -> dict:
    out = Path(out_dir)
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    spec = dict(spec or {"dataset_id": "ds.tolcal", "kind": kind, "role": "intermediate", "compare": {}})
    spec.setdefault("kind", kind)
    key = key or list(((spec.get("compare") or {}).get("key")) or [])
    if kind == "raster":
        muts = raster_mutants(legacy, out, recipe=recipe, shift_k=shift_k, scale_k=scale_k)
    elif kind == "vector":
        muts = vector_mutants(legacy, out, key=key, layer=spec.get("layer"), drop_n=drop_n, scale_k=scale_k)
    elif kind == "table":
        muts = table_mutants(legacy, out, key=key, spec=spec, drop_n=drop_n, scale_k=scale_k)
    else:
        raise ValueError(f"kind {kind!r} not in raster|vector|table")
    if only:
        muts = [m for m in muts if m["id"] in only or any(m["id"].startswith(o) for o in only)]
    rows, summary = [], {}
    for m in muts:
        r = compare.compare_pair(str(legacy), str(m["path"]), spec)
        met = _metrics(r)
        rows.append({"id": m["id"], "recipe": m["recipe"], "path": m["path"].name, "sha256": ac.sha256_file(m["path"]),
                     "verdict": r["verdict"], "reason_codes": r["reason_codes"],
                     "diagnostic_codes": r["diagnostic_codes"], "metrics": met})
        for cid, x in met.items():
            if x["state"] == "FAIL" and isinstance(x["value"], (int, float)) and not isinstance(x["value"], bool):
                cur = summary.get(cid)
                if cur is None or x["value"] < cur["value"]:
                    summary[cid] = {"value": x["value"], "metric": x["metric"], "mutant": m["id"]}
    res = {"schema": TOLCAL_SCHEMA, "legacy": {"name": Path(legacy).name, "sha256": ac.sha256_file(legacy)
                                               if Path(legacy).is_file() else None},
           "kind": kind, "comparator": compare.comparator_block(), "mutants": rows,
           "nearest_wrong_by_check": dict(sorted(summary.items())),
           "note": "nearest_wrong among these deterministic mutants only; a plausible wrong recipe of the step "
                   "itself (code variant) may be nearer and must be measured separately"}
    ac.dump_json(out / "tolcal.json", res)
    return res


def _nice_threshold(s: float, w: float, form: str) -> float:
    if form == "max_mismatch":
        t = math.floor((s + w) / 2)
        if not (s < t < w):
            raise ValueError(f"no integer threshold strictly between {s} and {w}")
        return t
    lo = s if s > 0 else w * 1e-3
    t = 10.0 ** round((math.log10(lo) + math.log10(w)) / 2) if s > 0 else 10.0 ** math.floor(math.log10(lo))
    if not (s < t < w):
        t = math.sqrt(max(lo, 1e-300) * w)
    if not (s < t < w):
        raise ValueError(f"no threshold strictly between {s} and {w}")
    for digits in range(1, 16):                      # shortest decimal that stays strictly inside the interval
        r = float(f"{t:.{digits}g}")
        if s < r < w:
            return r
    return t


def propose(dataset: str, check: str, *, correct_spread: float, nearest_wrong: float, metric: str = "abs+rel",
            form: str | None = None, units: str = "1", spread_how: str | None = None,
            spread_evidence: str | None = None, wrong_variant: str | None = None, wrong_evidence: str | None = None,
            version: int = 1, threshold: float | None = None) -> dict:
    """{correct_spread, nearest_wrong} → a PROPOSED record (agreed_by: null). Never signed by code."""
    s, w = float(correct_spread), float(nearest_wrong)
    if not (math.isfinite(s) and math.isfinite(w)) or s < 0 or w <= s:
        raise ValueError(f"no tolerance possible: correct_spread {s} must be < nearest_wrong {w}")
    form = form or ("value" if check.startswith(("vector.symdiff", "vector.hausdorff")) else "rtol")
    t = threshold if threshold is not None else _nice_threshold(s, w, form)
    if not (s < t < w):
        raise ValueError(f"threshold {t} not strictly between {s} and {w}")
    rec = {"id": f"tol.{dataset.removeprefix('ds.')}.{check.replace(':', '.')}.v{version}", "dataset": dataset, "check": check,
           "metric": metric, "atol": 0.0, "rtol": 0.0, "max_mismatch": 0, "units": units, "basis": "measured_spread",
           "correct_spread": {"value": s, "how": spread_how or "independent correct implementation",
                              "evidence": spread_evidence},
           "nearest_wrong": {"value": w, "variant": wrong_variant or "nearest wrong recipe",
                             "evidence": wrong_evidence},
           "agreed_by": None, "agreed_at": None, "decision_ref": None, "fixed_before": None}
    if form == "rtol":
        rec["rtol"] = t
    elif form == "atol":
        rec["atol"] = t
    elif form == "max_mismatch":
        rec["max_mismatch"] = int(t)
        rec["correct_spread"]["mismatch"] = s
        rec["nearest_wrong"]["mismatch"] = w
    else:
        rec["value"] = t
    assert set(rec) <= compare.RECORD_KEYS
    return rec


def main(argv=None) -> int:
    import yaml
    ap = argparse.ArgumentParser(prog="tolcal.py", description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("mutate")
    m.add_argument("--legacy", required=True)
    m.add_argument("--kind", required=True, choices=["raster", "vector", "table"])
    m.add_argument("--out", required=True)
    m.add_argument("--spec")
    m.add_argument("--recipe")
    m.add_argument("--only")
    m.add_argument("--key", action="append", default=[])
    m.add_argument("--json", action="store_true")
    p = sub.add_parser("propose")
    p.add_argument("--dataset", required=True)
    p.add_argument("--check", required=True)
    p.add_argument("--correct-spread", type=float, required=True)
    p.add_argument("--nearest-wrong", type=float, required=True)
    p.add_argument("--metric", default="abs+rel")
    p.add_argument("--form", choices=["rtol", "atol", "value", "max_mismatch"])
    p.add_argument("--units", default="1")
    p.add_argument("--spread-how")
    p.add_argument("--spread-evidence")
    p.add_argument("--wrong-variant")
    p.add_argument("--wrong-evidence")
    p.add_argument("--version", type=int, default=1)
    a = ap.parse_args(argv)
    if a.cmd == "mutate":
        spec = json.loads(Path(a.spec).read_text()) if a.spec else None
        recipe = json.loads(Path(a.recipe).read_text()) if a.recipe else None
        res = mutate(a.legacy, a.kind, a.out, spec=spec, recipe=recipe,
                     only=a.only.split(",") if a.only else None, key=a.key or None)
        if a.json:
            print(json.dumps(res, sort_keys=True, indent=1))
        else:
            for r in res["mutants"]:
                print(f"{r['id']:22s} {r['verdict']:16s} {','.join(r['reason_codes'])}")
        return 0
    try:
        rec = propose(a.dataset, a.check, correct_spread=a.correct_spread, nearest_wrong=a.nearest_wrong,
                      metric=a.metric, form=a.form, units=a.units, spread_how=a.spread_how,
                      spread_evidence=a.spread_evidence, wrong_variant=a.wrong_variant,
                      wrong_evidence=a.wrong_evidence, version=a.version)
    except ValueError as e:
        print(f"tolcal.py: {e}", file=sys.stderr)
        return 2
    print(yaml.safe_dump(rec, sort_keys=False), end="")
    return 0


if __name__ == "__main__":
    sys.exit(main())

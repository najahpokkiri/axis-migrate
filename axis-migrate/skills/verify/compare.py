#!/usr/bin/env python3
"""compare.py — per-dataset parity of a new pipeline run against the legacy run's own files (axis-migrate verify).

Subcommands (CONTRACTS §7):
  run       --run <run_dir> [--dataset ds]…            pair from graph lineage, freeze the contract, compare, write
                                                        verify/compare/<ds>.json (+ .mismatches.parquet)
  aggregate --run <run_dir> [--isolation-run <dir>]…   dataset → step → pipeline verdicts, visual receipts, blame;
                                                        writes verify/verdict.json (no rung: rungs are ladder.py's)
  pair      --legacy L --new N --spec spec.json [--out x.json]   standalone (T0 calibration, eval scoring, dry run)

Rules this file enforces (DESIGN §4.6, 06 §4): pairing from lineage only; structure gates values; grid deviation in
pixels at origin and far corner ≤ 1e-6; CRS compared with pyproj CRS.equals (never by EPSG code); masks before values;
NaN never equals a number; nothing is ever resampled, reprojected or cropped to make a comparison possible (vector
geometry is projected into measure_crs for measurement only); default exact; a tolerance is a signed record with two
measured numbers; evaluator faults are never candidate failures. Exit 0 only when every requested dataset is PASS_*;
1 when any is not; 2 when the comparison could not be set up at all (usage, unreadable run directory).
"""
from __future__ import annotations

import argparse
import collections
import datetime as _dt
import hashlib
import itertools
import json
import math
import os
import re
import sys
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

os.environ["GDAL_PAM_ENABLED"] = "NO"          # a stats read must never write .aux.xml next to a legacy file

import numpy as np  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "lib"))
import axis_common as ac  # noqa: E402
import visual_receipt as vrc  # noqa: E402

COMPARATOR_ID = "axis-migrate.compare"
COMPARATOR_VERSION = "1.0.0"
RESULT_SCHEMA = "axis-migrate.compare-result.v1"
CONTRACT_SCHEMA = "axis-migrate.verify-contract.v1"
VERDICT_SCHEMA = "axis-migrate.verdict.v1"
GRID_TOL_PX = 1e-6
TILE = 512
SAMPLE_N = 20
MISMATCH_ROWS_MAX = 2_000_000
DIAG_MAX_CELLS = 25_000_000
EXIT_PASS, EXIT_NOT_PASS, EXIT_SETUP = 0, 1, 2

PASS = ac.PASS_STATES
DIAGNOSTIC_CODES = {"NUMERIC_SCALE_FACTOR", "NUMERIC_OFFSET", "NUMERIC_SIGN_FLIP", "DATA_OFFSET_SUSPECTED",
                    "SMOOTHING_SUSPECTED", "ACCUMULATED_UPSTREAM_DIFFERENCE", "COMPENSATING_DEFECT",
                    "CANDIDATE_IDENTICAL_TO_REFERENCE", "GEOMETRY_AREA_RATIO", "VERTEX_COUNT_DROP"}

RECORD_KEYS = {"id", "dataset", "check", "metric", "atol", "rtol", "value", "max_mismatch", "units", "basis",
               "correct_spread", "nearest_wrong", "agreed_by", "agreed_at", "decision_ref", "fixed_before"}
SPREAD_KEYS = {"value", "how", "evidence", "metric", "variant", "mismatch"}
TOL_DOC_KEYS = {"version", "defaults", "records"}
BUILTIN_DEFAULTS = {
    "raster.grid.transform": {"metric": "transform_dev_px", "value": GRID_TOL_PX, "basis": "float_storage"},
    "raster.mask": {"metric": "xor_cells", "value": 0},
    "raster.values": {"metric": "abs+rel", "atol": 0, "rtol": 0, "max_mismatch": 0},
    "vector.geometry_type": {"normalise": "multi_equivalence"},
}
KD_KEYS = {"id", "kind", "datasets", "checks", "rule", "bound", "reason", "evidence", "owner", "decided_at",
           "decision_ref", "status", "supersedes"}
KD_KINDS = {"normalisation", "exclusion", "accepted_change", "legacy_defect_preserved"}
KD_NORMALISATION_RULES = {"nodata_equivalent", "dtype_widening", "multi_equivalence", "string_strip", "casefold",
                          "rename", "content_hash_only"}
SIDE_LEGACY, SIDE_NEW = "legacy", "new"


# =================================================================================================================
# Errors
# =================================================================================================================
class SetupError(Exception):
    """The comparison could not be set up at all (CLI exit 2)."""


class EvaluatorFault(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


class LegacyUnreadable(EvaluatorFault):
    def __init__(self, message: str):
        super().__init__("CHECKER_EXCEPTION", f"legacy reference unreadable: {message}")


class OutputUnreadable(Exception):
    pass


# =================================================================================================================
# Small helpers
# =================================================================================================================
def _self_sha256() -> str:
    return ac.sha256_file(Path(__file__).resolve())


def comparator_block() -> dict:
    return {"id": COMPARATOR_ID, "version": COMPARATOR_VERSION, "sha256": _self_sha256()}


def _jsonable(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (np.generic,)):
        v = v.item()
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    if isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()
    if isinstance(v, bytes):
        return v.hex()
    if isinstance(v, (list, tuple)):
        return [_jsonable(x) for x in v]
    if isinstance(v, dict):
        return {str(k): _jsonable(x) for k, x in v.items()}
    try:
        import pandas as pd
        if v is pd.NaT or (not isinstance(v, str) and pd.isna(v)):
            return None
    except (TypeError, ValueError):
        return str(v)
    return str(v)


def _sanitise(obj: Any, nulls: list, path: str = "$") -> Any:
    """Replace every non-finite float by None and record where (NaN is never written, 06 D6-5)."""
    if isinstance(obj, (float, np.floating)):
        f = float(obj)
        if not math.isfinite(f):
            nulls.append(path)
            return None
        return f
    if isinstance(obj, np.integer):
        return int(obj)
    if isinstance(obj, np.bool_):
        return bool(obj)
    if isinstance(obj, dict):
        return {str(k): _sanitise(v, nulls, f"{path}.{k}") for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_sanitise(v, nulls, f"{path}[{i}]") for i, v in enumerate(obj)]
    return obj


def _fnum(x: Any) -> float | None:
    try:
        f = float(x)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def parse_time(x: Any, *, end_of_day: bool) -> _dt.datetime | None:
    """ISO time → aware UTC datetime. A date without a time is read as the end of that day when `end_of_day`
    (a record or KD dated the run day cannot be shown to precede the run: fail closed)."""
    if x is None or x == "":
        return None
    if isinstance(x, _dt.datetime):
        return x if x.tzinfo else x.replace(tzinfo=_dt.timezone.utc)
    if isinstance(x, _dt.date):
        t = _dt.time(23, 59, 59, 999999) if end_of_day else _dt.time(0, 0)
        return _dt.datetime.combine(x, t, tzinfo=_dt.timezone.utc)
    s = str(x).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return parse_time(_dt.date.fromisoformat(s), end_of_day=end_of_day)
    s = s.replace("Z", "+00:00")
    d = _dt.datetime.fromisoformat(s)
    return d if d.tzinfo else d.replace(tzinfo=_dt.timezone.utc)


def _is_human(who: Any) -> bool:
    if not isinstance(who, str) or not who.strip():
        return False
    w = who.strip().lower()
    return not (w.startswith(("agent", "model", "claude", "ai:", "ai ", "codex", "gpt", "bot")) or w in ("ai",))


def _check_matches(listed: str, check_id: str) -> bool:
    return check_id == listed or check_id.startswith(listed + ":") or check_id.startswith(listed + ".") \
        or check_id.startswith(listed + "@")


def _sha_any(p: Path) -> str:
    return ac.tree_hash(p).split(":", 1)[1] if p.is_dir() else ac.sha256_file(p)


def _content_sha(p: Path) -> str | None:
    if p.is_dir():
        return None
    return ac.content_sha256(p)


def _env_block() -> dict:
    out = {}
    try:
        import rasterio
        out["rasterio"] = rasterio.__version__
        out["gdal"] = rasterio.__gdal_version__
    except ImportError as e:            # recorded, never hidden
        out["rasterio"] = f"unavailable: {e}"
    try:
        import pyogrio
        out["pyogrio"] = pyogrio.__version__
        out["pyogrio_gdal"] = pyogrio.__gdal_version_string__
    except ImportError as e:
        out["pyogrio"] = f"unavailable: {e}"
    import pyproj
    import shapely
    out.update({"proj": pyproj.proj_version_str, "pyproj": pyproj.__version__, "shapely": shapely.__version__,
                "geos": shapely.geos_version_string, "numpy": np.__version__})
    try:
        import geopandas
        import pandas
        out.update({"geopandas": geopandas.__version__, "pandas": pandas.__version__})
    except ImportError as e:
        out["geopandas"] = f"unavailable: {e}"
    return out


# =================================================================================================================
# Checks
# =================================================================================================================
def new_check(cid: str, order: int, *, gate: bool = False) -> dict:
    return {"id": cid, "order": order, "gate": gate, "state": None, "reason_codes": [], "diagnostic_codes": [],
            "metric": None, "value": None, "threshold": None, "tolerance_id": None, "legacy": None, "new": None,
            "diagnostics": {}, "samples": [], "by_construction": False, "kd": [], "detail": {}}


def gated(cid: str, order: int, by: str) -> dict:
    c = new_check(cid, order)
    c["state"] = "NOT_COMPARED"
    c["reason_codes"] = ["GATED"]
    c["gated_by"] = by
    return c


def set_state(c: dict, state: str, *codes: str) -> dict:
    c["state"] = state
    for code in codes:
        if code and code not in c["reason_codes"]:
            c["reason_codes"].append(code)
    return c


def add_diag(c: dict, code: str, **info) -> None:
    if code not in c["diagnostic_codes"]:
        c["diagnostic_codes"].append(code)
    if info:
        c["diagnostics"].update(info)


def counts_for_verdict(c: dict) -> bool:
    if c.get("kd_check"):
        return False                                     # accepted_change never changes the compare state
    return not (c["state"] == "NOT_COMPARED" and c["reason_codes"] == ["GATED"])


def verdict_of_checks(checks: list[dict]) -> str | None:
    return ac.worst(c["state"] for c in checks if counts_for_verdict(c))


# =================================================================================================================
# Tolerance records (06 §3.2; CONTRACTS §8)
# =================================================================================================================
@dataclass
class Rule:
    atol: float = 0.0
    rtol: float = 0.0
    max_mismatch: int = 0
    value: float | None = None            # single-metric threshold (geometry, counts)
    id: str | None = None

    @property
    def exact(self) -> bool:
        return self.atol == 0 and self.rtol == 0 and self.max_mismatch == 0 and not self.value

    def over(self, l: np.ndarray, n: np.ndarray) -> np.ndarray:
        d = np.abs(n - l)
        if self.atol == 0 and self.rtol == 0:
            return d != 0
        return d > (self.atol + self.rtol * np.abs(l))

    def close(self, l: float, n: float) -> bool:
        if not (math.isfinite(l) and math.isfinite(n)):
            return False
        return abs(n - l) <= self.atol + self.rtol * abs(l)

    def scalar_ok(self, metric: float | None) -> bool:
        if metric is None or not math.isfinite(metric):
            return False
        return metric <= (self.value if self.value is not None else self.atol)

    def describe(self) -> dict:
        return {"atol": self.atol, "rtol": self.rtol, "max_mismatch": self.max_mismatch, "value": self.value}


EXACT = Rule(id=None)


@dataclass
class TolRes:
    rule: Rule = field(default_factory=lambda: EXACT)
    proposed: Rule | None = None
    proposed_info: dict | None = None
    notes: list = field(default_factory=list)
    error: tuple | None = None
    considered: list = field(default_factory=list)


def _record_version(rid: str) -> int:
    m = re.search(r"\.v(\d+)$", rid or "")
    return int(m.group(1)) if m else 0


def _threshold(rec: dict) -> tuple[float | None, str | None]:
    if rec.get("value") not in (None, 0):
        return _fnum(rec["value"]), None
    rt, at = _fnum(rec.get("rtol") or 0) or 0.0, _fnum(rec.get("atol") or 0) or 0.0
    if rt and at:
        return None, "both atol and rtol are non-zero: two numbers can justify only one threshold"
    if rt:
        return rt, None
    if at:
        return at, None
    mm = _fnum(rec.get("max_mismatch") or 0) or 0.0
    if mm:
        return mm, None
    return 0.0, None


def _crs_invalid_for_measure(crs_text: Any) -> str | None:
    if not crs_text:
        return "measure_crs is not set"
    import pyproj
    try:
        crs = pyproj.CRS.from_user_input(crs_text)
    except pyproj.exceptions.CRSError as e:
        return f"measure_crs unreadable: {e}"
    if crs.is_geographic:
        return f"measure_crs {crs_text} is geographic (degrees)"
    op = crs.coordinate_operation
    name = (op.method_name if op is not None else "") or ""
    if "pseudo-mercator" in name.lower() or "popular visualisation" in name.lower():
        return f"measure_crs {crs_text} is Web Mercator (distances and areas are scaled)"
    return None


class ToleranceBook:
    """Validates every record once; lookups give the rule to apply, an unsigned proposal, or an error."""

    def __init__(self, doc: Any, run_started_at: _dt.datetime, measure_crs: Any = None, proposals: Any = None):
        self.started_at = run_started_at
        self.measure_crs = measure_crs
        self.doc_errors: list[str] = []
        self.defaults_ignored: list[str] = []
        self.records: list[dict] = []
        self.status: dict[int, dict] = {}
        # tolerances.proposed.yaml: unsigned proposals, used ONLY to tell NEEDS_DECISION/TOLERANCE_UNSIGNED from FAIL;
        # they never make a check pass and never enter the contract hash
        self.proposals: list[dict] = []
        self.proposal_status: dict[int, dict] = {}
        self.proposal_notes: list[dict] = []
        self._load_proposals(proposals)
        if doc is None:
            return
        if not isinstance(doc, dict):
            self.doc_errors.append("tolerances document is not a mapping")
            return
        unknown = set(doc) - TOL_DOC_KEYS
        if unknown:
            self.doc_errors.append(f"unknown top-level keys {sorted(unknown)}")
        for k, v in (doc.get("defaults") or {}).items():
            if BUILTIN_DEFAULTS.get(k) != v and not (k in BUILTIN_DEFAULTS and _same_default(BUILTIN_DEFAULTS[k], v)):
                self.defaults_ignored.append(k)
        for i, rec in enumerate(doc.get("records") or []):
            self.records.append(rec)
            self.status[i] = self._validate(rec)

    def _load_proposals(self, pdoc: Any) -> None:
        if pdoc is None:
            return
        recs = pdoc.get("records") if isinstance(pdoc, dict) else None
        if not isinstance(recs, list):
            self.proposal_notes.append({"id": None, "why": "proposals file has no records list; ignored"})
            return
        for rec in recs:
            if not isinstance(rec, dict):
                self.proposal_notes.append({"id": None, "why": "proposal is not a mapping; ignored"})
                continue
            if rec.get("agreed_by") not in (None, "") or rec.get("agreed_at") not in (None, ""):
                self.proposal_notes.append({"id": rec.get("id"), "dataset": rec.get("dataset"),
                                            "why": "a signed record in tolerances.proposed.yaml is ignored: a "
                                                   "signature counts only in tolerances.yaml"})
                continue
            st = self._validate(rec)
            if st["status"] != "unsigned":
                self.proposal_notes.append({"id": rec.get("id"), "dataset": rec.get("dataset"),
                                            "why": f"proposal ignored: {st.get('why')}"})
                continue
            self.proposal_status[len(self.proposals)] = st
            self.proposals.append(rec)

    def _validate(self, rec: Any) -> dict:
        if not isinstance(rec, dict):
            return {"status": "invalid", "why": "record is not a mapping"}
        rid = rec.get("id")
        unknown = set(rec) - RECORD_KEYS
        for k in ("correct_spread", "nearest_wrong"):
            if isinstance(rec.get(k), dict):
                unknown |= {f"{k}.{x}" for x in set(rec[k]) - SPREAD_KEYS}
        if unknown:
            return {"id": rid, "status": "invalid", "why": f"unknown keys {sorted(unknown)}"}
        if not rid or not rec.get("dataset") or not rec.get("check"):
            return {"id": rid, "status": "invalid", "why": "id, dataset and check are required"}
        units = str(rec.get("units") or "").lower()
        if "deg" in units:
            return {"id": rid, "status": "measure_crs_invalid", "why": f"units {rec.get('units')!r} are degrees"}
        if rec["check"].startswith(("vector.symdiff", "vector.hausdorff", "vector.geom")):
            bad = _crs_invalid_for_measure(self.measure_crs)
            if bad:
                return {"id": rid, "status": "measure_crs_invalid", "why": bad}
        thr, why = _threshold(rec)
        cs = rec.get("correct_spread") if isinstance(rec.get("correct_spread"), dict) else None
        nw = rec.get("nearest_wrong") if isinstance(rec.get("nearest_wrong"), dict) else None
        justified, jwhy = True, None
        if why:
            justified, jwhy = False, why
        elif cs is None or nw is None or _fnum(cs.get("value")) is None or _fnum(nw.get("value")) is None:
            justified, jwhy = False, "correct_spread and nearest_wrong must both be measured numbers"
        elif not (_fnum(cs["value"]) < thr < _fnum(nw["value"])):
            justified, jwhy = False, (f"not correct_spread ({cs['value']}) < threshold ({thr}) < nearest_wrong "
                                      f"({nw['value']})")
        elif (_fnum(rec.get("max_mismatch") or 0) or 0) > 0 and thr != _fnum(rec.get("max_mismatch")):
            mm = _fnum(rec["max_mismatch"])
            if not (_fnum(cs.get("mismatch")) is not None and _fnum(nw.get("mismatch")) is not None
                    and _fnum(cs["mismatch"]) < mm < _fnum(nw["mismatch"])):
                justified, jwhy = False, "max_mismatch > 0 needs its own two numbers (correct_spread.mismatch, nearest_wrong.mismatch)"
        signed = _is_human(rec.get("agreed_by")) and rec.get("agreed_at") not in (None, "")
        info = {"id": rid, "threshold": thr, "correct_spread": _fnum((cs or {}).get("value")),
                "nearest_wrong": _fnum((nw or {}).get("value"))}
        if not justified:
            return dict(info, status="unjustified", why=jwhy)
        if not signed:
            return dict(info, status="unsigned", why="agreed_by is not a human or agreed_at is missing")
        try:
            at = parse_time(rec.get("agreed_at"), end_of_day=True)
        except ValueError as e:
            return dict(info, status="unjustified", why=f"agreed_at unreadable: {e}")
        if at is None or at >= self.started_at:
            return dict(info, status="unjustified",
                        why=f"agreed_at {rec.get('agreed_at')} is not before the run start {self.started_at.isoformat()} "
                            "(a date without a time is read as the end of that day)")
        return dict(info, status="signed", why=None)

    @staticmethod
    def _rule(rec: dict) -> Rule:
        mm = int(_fnum(rec.get("max_mismatch") or 0) or 0)
        val = _fnum(rec.get("value")) if rec.get("value") not in (None,) else None
        return Rule(atol=_fnum(rec.get("atol") or 0) or 0.0, rtol=_fnum(rec.get("rtol") or 0) or 0.0,
                    max_mismatch=mm, value=val, id=rec.get("id"))

    def table(self) -> list[dict]:
        out = []
        for i, rec in enumerate(self.records):
            st = self.status[i]
            out.append({"record": rec, "status": st["status"], "why": st.get("why"), "threshold": st.get("threshold")})
        return out

    def _resolve(self, idxs: list[int], pidxs: Iterable[int] = ()) -> TolRes:
        res = TolRes()
        if self.doc_errors:
            res.error = ("TOLERANCE_UNJUSTIFIED", "; ".join(self.doc_errors))
            return res
        idxs = sorted(idxs, key=lambda i: -_record_version(self.records[i].get("id", "")))
        signed = unsigned = None
        for i in idxs:
            st = self.status[i]
            res.considered.append({"id": st.get("id"), "status": st["status"], "why": st.get("why")})
            if st["status"] == "invalid":
                res.error = ("TOLERANCE_UNJUSTIFIED", f"record {st.get('id')}: {st['why']}")
                return res
            if st["status"] == "measure_crs_invalid":
                res.error = ("GEOMETRY_MEASURE_CRS_INVALID", f"record {st.get('id')}: {st['why']}")
                return res
            if st["status"] == "signed" and signed is None:
                signed = i
            elif st["status"] == "unsigned" and unsigned is None:
                unsigned = i
            if st["status"] in ("unjustified", "unsigned"):
                if "TOLERANCE_UNJUSTIFIED" not in res.notes:
                    res.notes.append("TOLERANCE_UNJUSTIFIED")
        if signed is not None:
            res.rule = self._rule(self.records[signed])
            res.notes = []
            return res
        if unsigned is not None:
            st = self.status[unsigned]
            res.proposed = self._rule(self.records[unsigned])
            res.proposed_info = {"record_id": st["id"], "threshold": st["threshold"], "source": "tolerances.yaml",
                                 "correct_spread": st["correct_spread"], "nearest_wrong": st["nearest_wrong"]}
            return res
        pidxs = sorted(pidxs, key=lambda i: -_record_version(self.proposals[i].get("id", "")))
        for i in pidxs:
            st = self.proposal_status[i]
            res.considered.append({"id": st.get("id"), "status": "proposed", "why": "tolerances.proposed.yaml"})
        if pidxs:
            st = self.proposal_status[pidxs[0]]
            res.proposed = self._rule(self.proposals[pidxs[0]])
            res.proposed_info = {"record_id": st["id"], "threshold": st["threshold"],
                                 "source": "tolerances.proposed.yaml", "correct_spread": st["correct_spread"],
                                 "nearest_wrong": st["nearest_wrong"]}
        return res

    def lookup(self, dataset: str, check: str) -> TolRes:
        return self._resolve([i for i, r in enumerate(self.records)
                              if isinstance(r, dict) and r.get("dataset") == dataset and r.get("check") == check],
                             [i for i, r in enumerate(self.proposals)
                              if r.get("dataset") == dataset and r.get("check") == check])

    def lookup_id(self, rid: str | None) -> TolRes:
        if not rid:
            return TolRes()
        idxs = [i for i, r in enumerate(self.records) if isinstance(r, dict) and r.get("id") == rid]
        pidxs = [i for i, r in enumerate(self.proposals) if r.get("id") == rid]
        if not idxs and pidxs:
            return self._resolve([], pidxs)
        if not idxs:
            res = TolRes()
            res.notes.append("TOLERANCE_UNJUSTIFIED")
            res.considered.append({"id": rid, "status": "missing", "why": "no such record"})
            return res
        return self._resolve(idxs)


def _same_default(a: dict, b: Any) -> bool:
    if not isinstance(b, dict):
        return False
    for k, v in b.items():
        if k not in a:
            return False
        if isinstance(v, (int, float)) and isinstance(a[k], (int, float)):
            if float(v) != float(a[k]):
                return False
        elif v != a[k]:
            return False
    return True


def apply_tol(c: dict, res: TolRes, passes: Callable[[Rule], bool], any_diff: bool, fail_code: str,
              metric_desc: dict | None = None) -> None:
    """Default exact; a signed record widens; an unsigned record that would pass gives NEEDS_DECISION."""
    for n in res.notes:
        if n not in c["reason_codes"]:
            c["reason_codes"].append(n)
    if res.considered:
        c["detail"]["tolerance_records"] = res.considered
    if res.error:
        set_state(c, "EVALUATOR_ERROR", res.error[0])
        c["detail"]["tolerance_error"] = res.error[1]
        return
    rule = res.rule
    c["tolerance_id"] = rule.id or f"default.{c['id']}.exact"
    if metric_desc:
        c["detail"]["rule"] = rule.describe()
    if passes(rule):
        set_state(c, "PASS_WITHIN_TOLERANCE" if (any_diff and not rule.exact) else "PASS_EXACT")
        return
    if res.proposed is not None and passes(res.proposed):
        set_state(c, "NEEDS_DECISION", "TOLERANCE_UNSIGNED")
        c["proposed"] = dict(res.proposed_info or {}, rule=res.proposed.describe())
        return
    set_state(c, "FAIL", fail_code)


# =================================================================================================================
# Known differences (DESIGN §4.6; 06 §3.3)
# =================================================================================================================
_KD_HEAD = re.compile(r"^##\s+(KD-\d+)\b.*$", re.M)
_YAML_FENCE = re.compile(r"```ya?ml\s*\n(.*?)```", re.S)


def parse_known_differences(text: str | None) -> tuple[list[dict], list[dict]]:
    """Each `## KD-nnn` section ends with one fenced yaml block; prose is ignored."""
    import yaml
    if not text:
        return [], []
    heads = list(_KD_HEAD.finditer(text))
    kds, errs = [], []
    for i, h in enumerate(heads):
        end = heads[i + 1].start() if i + 1 < len(heads) else len(text)
        body = text[h.end():end]
        m = _YAML_FENCE.search(body)
        if not m:
            errs.append({"id": h.group(1), "why": "no fenced yaml block"})
            continue
        try:
            d = yaml.safe_load(m.group(1))
        except yaml.YAMLError as e:
            errs.append({"id": h.group(1), "why": f"yaml: {e}"})
            continue
        if not isinstance(d, dict):
            errs.append({"id": h.group(1), "why": "yaml block is not a mapping"})
            continue
        d.setdefault("id", h.group(1))
        d["_head"] = h.group(1)
        d["_line"] = text.count("\n", 0, h.start() + h.end() - h.start() + m.start()) + 2
        kds.append(d)
    # append-only correction: a later entry `supersedes: KD-nnn` withdraws that entry (even a malformed one)
    sup = set()
    for d in kds:
        s = d.get("supersedes")
        sup.update(map(str, s if isinstance(s, (list, tuple)) else ([s] if s else [])))
    kds = [d for d in kds if str(d.get("id")) not in sup]
    errs = [e for e in errs if str(e.get("id")) not in sup]
    return kds, errs


def _kd_wellformed(kd: dict, role: str | None) -> str | None:
    unknown = set(kd) - KD_KEYS - {"_line", "_head"}
    if unknown:
        return f"unknown keys {sorted(unknown)}"
    if kd.get("kind") not in KD_KINDS:
        return f"kind {kd.get('kind')!r} not in {sorted(KD_KINDS)}"
    rule = kd.get("rule")
    kind = kd["kind"]
    if kind == "normalisation":
        if not isinstance(rule, dict) or not rule or set(rule) - KD_NORMALISATION_RULES:
            return f"normalisation rule must use {sorted(KD_NORMALISATION_RULES)}"
    elif kind == "exclusion":
        if not isinstance(rule, dict) or rule.get("reason") not in ("NO_REFERENCE", "FAMILY_UNSUPPORTED"):
            return "exclusion rule must be {reason: NO_REFERENCE|FAMILY_UNSUPPORTED}"
        if role == "final":
            return "an exclusion never applies to a role: final dataset"
    elif kind == "accepted_change":
        if not isinstance(rule, dict) or "expected" not in rule or _parse_expected(rule["expected"]) is None:
            return "accepted_change rule must be {expected: 'new = legacy <op> <number>'}"
    return None


class KDBook:
    def __init__(self, kds: list[dict], started_at: _dt.datetime, dataset_id: str, role: str | None):
        self.active: list[dict] = []
        self.candidates: list[dict] = []
        self.ignored: list[dict] = []
        self.applied: dict[str, dict] = {}
        for kd in kds or []:
            if not isinstance(kd, dict):
                continue
            if dataset_id not in (kd.get("datasets") or []):
                continue
            kid = kd.get("id")
            bad = _kd_wellformed(kd, role)
            if bad:
                self.ignored.append({"id": kid, "why": bad})
                continue
            status = kd.get("status")
            human = isinstance(kd.get("owner"), str) and kd["owner"].startswith("human:") and \
                _is_human(kd["owner"][6:])
            if status == "withdrawn":
                self.ignored.append({"id": kid, "why": "withdrawn"})
                continue
            if status == "accepted" and human:
                try:
                    when = parse_time(kd.get("decided_at"), end_of_day=True)
                except ValueError as e:
                    self.ignored.append({"id": kid, "why": f"decided_at unreadable: {e}"})
                    continue
                if when is None or when >= started_at:
                    self.ignored.append({"id": kid, "why": f"decided after the run started "
                                                            f"({kd.get('decided_at')}); enters from the next run"})
                    continue
                self.active.append(kd)
                continue
            why = "status proposed" if status == "proposed" else f"status {status!r}"
            if not human:
                why = f"owner {kd.get('owner')!r} is not human:<name> (model owners are ignored)"
            self.ignored.append({"id": kid, "why": why})
            if status in ("proposed", "accepted"):
                self.candidates.append(kd)

    @staticmethod
    def _for(kds, check_id, rule_key=None, kind="normalisation"):
        out = []
        for kd in kds:
            if kd.get("kind") != kind:
                continue
            if rule_key and not (isinstance(kd.get("rule"), dict) and rule_key in kd["rule"]):
                continue
            if any(_check_matches(str(x), check_id) for x in (kd.get("checks") or [])):
                out.append(kd)
        return out

    def norm(self, check_id: str, rule_key: str) -> tuple[list[dict], list[dict]]:
        return self._for(self.active, check_id, rule_key), self._for(self.candidates, check_id, rule_key)

    def exclusion(self) -> dict | None:
        ex = [kd for kd in self.active if kd["kind"] == "exclusion"]
        return ex[0] if ex else None

    def accepted_changes(self) -> list[tuple[dict, bool]]:
        return [(kd, True) for kd in self.active if kd["kind"] == "accepted_change"] + \
               [(kd, False) for kd in self.candidates if kd["kind"] == "accepted_change"]

    def preserved(self) -> list[dict]:
        return [kd for kd in self.active if kd["kind"] == "legacy_defect_preserved"]

    def mark(self, kd: dict, check_id: str, **extra) -> None:
        e = self.applied.setdefault(kd["id"], {"id": kd["id"], "kind": kd["kind"], "checks": [],
                                                "owner": kd.get("owner"), "decision_ref": kd.get("decision_ref")})
        if check_id not in e["checks"]:
            e["checks"].append(check_id)
        e.update(extra)


def apply_norm(c: dict, raw_ok: bool, kdb: KDBook, rule_key: str, fixed_by: Callable[[dict], bool],
               fail_code: str) -> None:
    if raw_ok:
        set_state(c, "PASS_EXACT")
        return
    active, cands = kdb.norm(c["id"], rule_key)
    ok = [kd for kd in active if fixed_by(kd)]
    if ok:
        set_state(c, "PASS_WITHIN_TOLERANCE")
        c["kd"] = [kd["id"] for kd in ok]
        for kd in ok:
            kdb.mark(kd, c["id"])
        return
    cok = [kd for kd in cands if fixed_by(kd)]
    if cok:
        set_state(c, "NEEDS_DECISION", "KD_UNSIGNED")
        c["proposed_kd"] = [kd["id"] for kd in cok]
        return
    set_state(c, "FAIL", fail_code)


_EXPECTED = re.compile(r"^\s*new\s*=\s*legacy\s*(?:([*/+-])\s*([-+0-9.eE]+))?\s*$")


def _parse_expected(expr: Any) -> tuple[str, float] | None:
    if not isinstance(expr, str):
        return None
    m = _EXPECTED.match(expr)
    if not m:
        return None
    op, num = m.group(1), m.group(2)
    if op is None:
        return ("*", 1.0)
    try:
        v = float(num)
    except ValueError:
        return None
    if op == "/" and v == 0:
        return None
    return (op, v)


def _apply_expected(op: tuple[str, float], legacy: np.ndarray) -> np.ndarray:
    o, v = op
    if o == "*":
        return legacy * v
    if o == "/":
        return legacy / v
    if o == "+":
        return legacy + v
    return legacy - v


def _bound_metric(metric: str, expected: np.ndarray, new: np.ndarray) -> float | None:
    if expected.size == 0:
        return 0.0
    d = np.abs(new - expected)
    if metric == "max_abs":
        return float(d.max())
    if metric == "max_rel":
        nz = expected != 0
        if np.any(d[~nz] > 0):
            return math.inf
        return float((d[nz] / np.abs(expected[nz])).max()) if nz.any() else 0.0
    if metric == "count_over":
        return float(np.count_nonzero(d))
    return None


def kd_change_checks(kdb: KDBook, order: int, values_for: Callable[[str], tuple[np.ndarray, np.ndarray] | None]
                     ) -> list[dict]:
    """accepted_change: the KD's rule/bound evaluated as an extra check; the compare state never changes."""
    out = []
    for kd, active in kdb.accepted_changes():
        c = new_check(f"kd:{kd['id']}", order)
        c["kd_check"] = True
        c["kd"] = [kd["id"]]
        op = _parse_expected(kd["rule"]["expected"])
        bound = kd.get("bound") or {}
        metric = bound.get("metric", "max_abs")
        bmax = _fnum(bound.get("max"))
        covers = [str(x) for x in (kd.get("checks") or [])]
        held, value = None, None
        for chk in covers:
            vals = values_for(chk)
            if vals is None:
                continue
            l, n = vals
            mv = _bound_metric(metric, _apply_expected(op, l.astype("float64")), n.astype("float64"))
            if mv is None:
                break
            value = mv if value is None else max(value, mv)
        if value is not None and bmax is not None:
            held = bool(value <= bmax)
        c["metric"], c["value"], c["threshold"] = metric, value, bmax
        c["held"] = held
        c["detail"] = {"expected": kd["rule"]["expected"], "covers_checks": covers, "signed": active}
        set_state(c, "PASS_WITHIN_TOLERANCE" if held else "FAIL")
        if not active:
            c["reason_codes"].append("KD_UNSIGNED")
        if active:
            kdb.mark(kd, c["id"], held=held, metric=metric, value=value, max=bmax, covers_checks=covers)
        out.append(c)
    return out


# =================================================================================================================
# Raster family (06 §4b raster; never warps)
# =================================================================================================================
def _open_raster(path: str, side: str):
    import rasterio
    try:
        return rasterio.open(path)
    except Exception as e:                                # typed re-raise: legacy → evaluator, new → candidate
        if side == SIDE_LEGACY:
            raise LegacyUnreadable(f"{path}: {e}") from e
        raise OutputUnreadable(f"{path}: {e}") from e


def _xy(tr, col: int, row: int) -> tuple[float, float]:
    """Map coordinates of a cell centre."""
    x, y = col + 0.5, row + 0.5
    return tr.a * x + tr.b * y + tr.c, tr.d * x + tr.e * y + tr.f


def _windows(h: int, w: int):
    from rasterio.windows import Window
    for r0 in range(0, h, TILE):
        for c0 in range(0, w, TILE):
            yield Window(c0, r0, min(TILE, w - c0), min(TILE, h - r0))


def _read_valid(ds, band: int, win):
    a = ds.read(band, window=win)
    m = ds.read_masks(band, window=win) > 0
    if a.dtype.kind == "f":
        m &= np.isfinite(a)
    nd = ds.nodatavals[band - 1]
    if nd is not None and not (isinstance(nd, float) and math.isnan(nd)):
        m &= a != np.asarray(nd).astype(a.dtype) if a.dtype.kind in "iu" and float(nd).is_integer() else a != nd
    return a, m


def _transform_dev(lt, nt, W: int, H: int) -> tuple[float, list[float]]:
    a = abs(lt.a) or math.hypot(lt.a, lt.d) or 1.0
    e = abs(lt.e) or math.hypot(lt.b, lt.e) or 1.0
    terms = [abs(nt.c - lt.c) / a, abs(nt.f - lt.f) / e, W * abs(nt.a - lt.a) / a, H * abs(nt.e - lt.e) / e,
             H * abs(nt.b - lt.b) / a, W * abs(nt.d - lt.d) / e]
    shift = [round((nt.c - lt.c) / lt.a, 6) if lt.a else None, round((nt.f - lt.f) / lt.e, 6) if lt.e else None]
    return max(terms), shift


def _crs_equal(a, b) -> bool:
    import pyproj
    # never to_epsg(): two unknown CRSs both give None and would compare equal
    return pyproj.CRS.from_wkt(a.to_wkt()).equals(pyproj.CRS.from_wkt(b.to_wkt()))


def _nodata_equal(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    fa, fb = float(a), float(b)
    if math.isnan(fa) or math.isnan(fb):
        return math.isnan(fa) and math.isnan(fb)
    return fa == fb


class _BandAccum:
    """Streaming stats of one band (lifted from the archive comparator with the 06 §5 fixes: float64
    accumulation, Chan/Welford merge for the variance, valid and non-finite counts)."""

    def __init__(self) -> None:
        self.count = 0
        self.mean = 0.0
        self.m2 = 0.0
        self.min = math.inf
        self.max = -math.inf
        self.non_zero = 0
        self.invalid = 0
        self._parts: list[float] = []

    def update(self, arr: np.ndarray, valid: np.ndarray) -> None:
        self.invalid += int(valid.size - np.count_nonzero(valid))
        v = arr[valid].astype("float64", copy=False)
        if v.size == 0:
            return
        n_b = int(v.size)
        mean_b = float(v.mean(dtype="float64"))
        m2_b = float(((v - mean_b) ** 2).sum(dtype="float64"))
        n_a = self.count
        delta = mean_b - self.mean
        tot = n_a + n_b
        self.mean += delta * n_b / tot
        self.m2 += m2_b + delta * delta * n_a * n_b / tot
        self.count = tot
        self.min = min(self.min, float(v.min()))
        self.max = max(self.max, float(v.max()))
        self.non_zero += int(np.count_nonzero(v))
        self._parts.append(math.fsum(v.tolist()))

    def exact_sum(self) -> float:
        return math.fsum(self._parts)

    def finalise(self) -> dict:
        if self.count == 0:
            return {"count": 0, "invalid": self.invalid, "sum": 0.0, "min": None, "max": None, "mean": None,
                    "std": None, "non_zero": 0}
        return {"count": self.count, "invalid": self.invalid, "sum": self.exact_sum(), "min": self.min,
                "max": self.max, "mean": self.mean, "std": math.sqrt(max(0.0, self.m2 / self.count)),
                "non_zero": self.non_zero}


def _exact_total(ds, band: int) -> tuple[float, int, bool]:
    """Order-independent exact float sum of valid cells on the raster's own grid (works when grids differ)."""
    chunks, n, integral = [], 0, True
    for win in _windows(ds.height, ds.width):
        a, m = _read_valid(ds, band, win)
        v = a[m].astype("float64")
        n += int(v.size)
        if integral and v.size and not np.all(np.floor(v) == v):
            integral = False
        chunks.append(v.tolist())
    return math.fsum(itertools.chain.from_iterable(chunks)), n, integral


HIST_EDGES = np.concatenate([[0.0], np.logspace(-15, 15, 301)])


class _ValueAccum:
    def __init__(self, rules: list[Rule], ds, band: int, check_id: str, categorical: bool):
        self.rules = rules
        self.ds, self.band, self.check_id, self.categorical = ds, band, check_id, categorical
        self.n = 0
        self.ndiff = 0
        self.over = [0] * len(rules)
        self.max_abs = 0.0
        self.max_rel = 0.0
        self.rel_undefined = False
        self.sum_d: list[float] = []
        self.hist = np.zeros(len(HIST_EDGES), dtype="int64")
        self.samples: list[dict] = []
        self.rows: list[dict] = []
        self.transitions: collections.Counter = collections.Counter()

    def update(self, win, la, lm, na, nm, mismatch_rows: bool) -> None:
        both = lm & nm
        if not both.any():
            return
        rr, cc = np.nonzero(both)
        l = la[both].astype("float64")
        n = na[both].astype("float64")
        d = n - l
        self.n += int(l.size)
        nz = d != 0
        self.ndiff += int(nz.sum())
        if not nz.any():
            return
        ad = np.abs(d)
        self.max_abs = max(self.max_abs, float(ad.max()))
        lnz = nz & (l != 0)
        if lnz.any():
            self.max_rel = max(self.max_rel, float((ad[lnz] / np.abs(l[lnz])).max()))
        if (nz & (l == 0)).any() and not lnz.any() and self.max_rel == 0.0:
            self.rel_undefined = True
        self.sum_d.append(math.fsum(d[nz].tolist()))
        idx = np.searchsorted(HIST_EDGES, ad[nz], side="right") - 1
        np.add.at(self.hist, np.clip(idx, 0, len(HIST_EDGES) - 1), 1)
        over0 = None
        for i, r in enumerate(self.rules):
            o = r.over(l, n)
            self.over[i] += int(o.sum())
            if i == 0:
                over0 = o
        if self.categorical:
            for a_, b_ in zip(l[nz].tolist(), n[nz].tolist()):
                self.transitions[(a_, b_)] += 1
        sel = np.nonzero(over0)[0]
        if sel.size:
            tr = self.ds.transform
            for k in sel[: max(0, SAMPLE_N - len(self.samples))]:
                r, c = int(rr[k] + win.row_off), int(cc[k] + win.col_off)
                x, y = _xy(tr, c, r)
                self.samples.append({"row": r, "col": c, "x": x, "y": y, "legacy": float(l[k]), "new": float(n[k])})
            if mismatch_rows:
                room = MISMATCH_ROWS_MAX - len(self.rows)
                for k in sel[: max(0, room)]:
                    r, c = int(rr[k] + win.row_off), int(cc[k] + win.col_off)
                    x, y = _xy(tr, c, r)
                    self.rows.append({"check": self.check_id, "band": self.band, "row": r, "col": c, "x": x, "y": y,
                                      "legacy": float(l[k]), "new": float(n[k]), "kind": "value"})

    def quantiles(self) -> dict:
        tot = int(self.hist.sum())
        out = {}
        if tot == 0:
            return {"p50": 0.0, "p90": 0.0, "p99": 0.0, "p99.9": 0.0, "max": 0.0}
        cum = np.cumsum(self.hist)
        # quantiles of |d| over ALL compared cells (zero differences included)
        zeros = self.n - tot
        for name, q in (("p50", 0.5), ("p90", 0.9), ("p99", 0.99), ("p99.9", 0.999)):
            target = q * self.n
            if target <= zeros:
                out[name] = 0.0
                continue
            k = int(np.searchsorted(cum, target - zeros))
            k = min(k, len(HIST_EDGES) - 1)
            out[name] = float(HIST_EDGES[min(k + 1, len(HIST_EDGES) - 1)])   # upper bin edge (conservative)
        out["max"] = self.max_abs
        return out


class _MaskAccum:
    def __init__(self, ds, check_id: str, band: int):
        self.ds, self.check_id, self.band = ds, check_id, band
        self.xor = self.legacy_only = self.new_only = self.valid_l = self.valid_n = 0
        self.samples: list[dict] = []
        self.rows: list[dict] = []

    def update(self, win, lm, nm, mismatch_rows: bool):
        self.valid_l += int(lm.sum())
        self.valid_n += int(nm.sum())
        x = lm ^ nm
        if not x.any():
            return
        lo = lm & ~nm
        self.xor += int(x.sum())
        self.legacy_only += int(lo.sum())
        self.new_only += int((nm & ~lm).sum())
        rr, cc = np.nonzero(x)
        tr = self.ds.transform
        take = rr.size if mismatch_rows else max(0, SAMPLE_N - len(self.samples))
        for k in range(min(rr.size, max(take, SAMPLE_N - len(self.samples)))):
            r, c = int(rr[k] + win.row_off), int(cc[k] + win.col_off)
            xx, yy = _xy(tr, c, r)
            kind = "mask_legacy_only" if lo[rr[k], cc[k]] else "mask_new_only"
            if len(self.samples) < SAMPLE_N:
                self.samples.append({"row": r, "col": c, "x": xx, "y": yy, "kind": kind})
            if mismatch_rows and len(self.rows) < MISMATCH_ROWS_MAX:
                self.rows.append({"check": self.check_id, "band": self.band, "row": r, "col": c, "x": xx, "y": yy,
                                  "legacy": None, "new": None, "kind": kind})


def _raster_values(ctx: "PairCtx", lds, nds, band: int, suffix: str, gate_ok: bool, order: int) -> list[dict]:
    """Masks then values on cells valid on both sides; one streaming pass over the common grid."""
    spec = ctx.spec
    cmp_ = spec.get("compare") or {}
    mid, vid = f"raster.mask{suffix}", f"raster.values{suffix}"
    if not gate_ok:
        return [gated(mid, order, ctx.gate_name), gated(vid, order + 1, ctx.gate_name)]
    ldt, ndt = lds.dtypes[band - 1], nds.dtypes[band - 1]
    integer = np.dtype(ldt).kind in "iub" and np.dtype(ndt).kind in "iub"
    categorical = bool(cmp_.get("categorical"))
    vres = ctx.book.lookup(ctx.ds_id, "raster.values")
    rules = [vres.rule] + ([vres.proposed] if vres.proposed is not None else [])
    if integer or categorical:
        rules = [EXACT]
    mres = ctx.book.lookup(ctx.ds_id, "raster.mask")
    macc = _MaskAccum(lds, mid, band)
    vacc = _ValueAccum(rules, lds, band, vid, categorical)
    lacc, nacc = _BandAccum(), _BandAccum()
    want_rows = ctx.mismatches_path is not None
    count_frac = False
    for win in _windows(lds.height, lds.width):
        la, lm = _read_valid(lds, band, win)
        na, nm = _read_valid(nds, band, win)
        lacc.update(la, lm)
        nacc.update(na, nm)
        macc.update(win, lm, nm, want_rows)
        vacc.update(win, la, lm, na, nm, want_rows)
        if cmp_.get("count") and nm.any():
            nv = na[nm].astype("float64")
            if not np.all(np.floor(nv) == nv):
                count_frac = True
    ctx.mismatch_rows.extend(macc.rows)
    ctx.mismatch_rows.extend(vacc.rows)
    ctx.band_stats[band] = {"legacy": lacc.finalise(), "new": nacc.finalise()}
    out = []
    # ---- M: mask
    m = new_check(mid, order)
    m.update(metric="xor_cells", value=macc.xor, legacy=macc.valid_l, new=macc.valid_n, samples=macc.samples)
    m["detail"] = {"legacy_only": macc.legacy_only, "new_only": macc.new_only, "valid_legacy": macc.valid_l,
                   "valid_new": macc.valid_n}
    mrule_val = lambda r: macc.xor <= (int(r.value) if r.value is not None else r.max_mismatch)  # noqa: E731
    apply_tol(m, mres, mrule_val, macc.xor > 0, "NODATA_MASK_MISMATCH")
    m["threshold"] = (mres.rule.value if mres.rule.value is not None else mres.rule.max_mismatch)
    if macc.valid_n == 0 and macc.valid_l > 0 and not cmp_.get("may_be_empty"):
        set_state(m, "FAIL", "OUTPUT_EMPTY")
    out.append(m)
    # ---- V: values
    v = new_check(vid, order + 1)
    stats = {"compared": vacc.n, "count_diff": vacc.ndiff, "count_over": vacc.over[0], "max_abs": vacc.max_abs,
             "max_rel": None if vacc.rel_undefined else vacc.max_rel,
             "mean_bias": (math.fsum(vacc.sum_d) / vacc.n) if vacc.n else 0.0, "quantiles_abs": vacc.quantiles()}
    if vacc.rel_undefined:
        v.setdefault("nulls", {})["stats.max_rel"] = "UNDEFINED_ZERO_REFERENCE"
    v["stats"] = stats
    v["samples"] = vacc.samples
    if categorical:
        v["metric"], v["value"] = "cells_changed", vacc.ndiff
        v["detail"]["transitions"] = [{"legacy": a, "new": b, "count": k}
                                      for (a, b), k in sorted(vacc.transitions.items(), key=lambda t: (-t[1], t[0]))[:20]]
        set_state(v, "PASS_EXACT" if vacc.ndiff == 0 else "FAIL",
                  None if vacc.ndiff == 0 else "RASTER_CATEGORY_MISMATCH")
        v["tolerance_id"] = "default.raster.values.exact"
    elif integer:
        v["metric"], v["value"] = "count_over", vacc.ndiff
        set_state(v, "PASS_EXACT" if vacc.ndiff == 0 else "FAIL",
                  None if vacc.ndiff == 0 else "RASTER_VALUE_MISMATCH")
        v["tolerance_id"] = "default.raster.values.exact"
        if vres.rule.id or vres.proposed is not None:
            v["detail"]["note"] = "integer dtypes compare exactly; tolerance records do not apply"
    else:
        v["metric"], v["value"] = "count_over", vacc.over[0]
        v["threshold"] = vres.rule.describe()

        def passes(r: Rule) -> bool:
            i = rules.index(r) if r in rules else None
            return i is not None and vacc.over[i] <= r.max_mismatch
        apply_tol(v, vres, passes, vacc.ndiff > 0, "RASTER_VALUE_MISMATCH")
    out.append(v)
    if cmp_.get("count"):
        c = new_check(f"raster.count_integer{suffix}", order + 2)
        _, _, leg_int = (None, None, ctx.legacy_integral(lds, band))
        c["metric"], c["value"] = "new_has_fraction", bool(count_frac)
        if leg_int and count_frac:
            set_state(c, "FAIL", "RASTER_NON_INTEGER_COUNTS")
        else:
            set_state(c, "PASS_EXACT")
        out.append(c)
    ctx.value_arrays_hint[band] = (vacc.over[0], macc.xor)
    return out


def _zones_frame(zspec: dict, target_crs) -> tuple[list, list]:
    """Zone polygons as (geometry, id) in the raster's CRS (zones are measurement aids; candidates never move)."""
    import pyogrio
    import pyproj
    df = pyogrio.read_dataframe(zspec["path"], layer=zspec.get("layer"))
    key = zspec["key"]
    if key not in df.columns:
        raise EvaluatorFault("CHECKER_EXCEPTION", f"zone key {key!r} not in {zspec['path']}")
    if df.crs is None:
        raise EvaluatorFault("CRS_MISSING", f"zone layer {zspec['path']} has no CRS")
    if target_crs is None:
        raise EvaluatorFault("CRS_MISSING", "raster has no CRS for zone rasterisation")
    tcrs = pyproj.CRS.from_wkt(target_crs.to_wkt())
    if not pyproj.CRS(df.crs).equals(tcrs):
        df = df.to_crs(tcrs)
    keys = sorted({str(_jsonable(k)) for k in df[key].tolist()})
    kid = {k: i + 1 for i, k in enumerate(keys)}
    shapes = [(g, kid[str(_jsonable(k))]) for g, k in zip(df.geometry, df[key]) if g is not None and not g.is_empty]
    return shapes, keys


def _zone_sums(ds, band: int, shapes: list, nz: int, all_touched: bool) -> tuple[list[float], dict]:
    from rasterio.features import rasterize
    from rasterio.windows import transform as wtransform
    parts: list[list[float]] = [[] for _ in range(nz + 1)]
    out_cells, out_sum = 0, []
    for win in _windows(ds.height, ds.width):
        a, m = _read_valid(ds, band, win)
        zid = rasterize(shapes, out_shape=(int(win.height), int(win.width)), transform=wtransform(win, ds.transform),
                        fill=0, all_touched=all_touched, dtype="int32")
        v = a[m].astype("float64")
        z = zid[m]
        if v.size == 0:
            continue
        sums = np.bincount(z, weights=v, minlength=nz + 1)
        for i in np.nonzero(np.bincount(z, minlength=nz + 1))[0]:
            parts[i].append(float(sums[i]))
        outside = (z == 0) & (v != 0)
        out_cells += int(outside.sum())
        out_sum.append(math.fsum(v[outside].tolist()))
    return [math.fsum(p) for p in parts], {"cells_nonzero_outside_zones": out_cells, "sum_outside": math.fsum(out_sum)}


def _conservation(ctx: "PairCtx", lds, nds, band: int, order: int) -> list[dict]:
    """Totals overall and per zone, each on its own grid: runs even when the grid gate failed (06 §4c)."""
    cmp_ = ctx.spec.get("compare") or {}
    if cmp_.get("measure") != "extensive":
        return []
    out = []
    T_l, n_l, _ = _exact_total(lds, band)
    T_n, n_n, _ = _exact_total(nds, band)
    c = new_check("conservation.total", order)
    c.update(metric="rel_drift", legacy=T_l, new=T_n)
    rel = (abs(T_n - T_l) / abs(T_l)) if T_l else (0.0 if T_n == T_l else None)
    c["value"] = rel
    if rel is None:
        c.setdefault("nulls", {})["value"] = "UNDEFINED_ZERO_REFERENCE"
    res = ctx.book.lookup(ctx.ds_id, "conservation.total")
    apply_tol(c, res, lambda r: r.close(T_l, T_n), T_l != T_n, "TOTAL_DRIFT")
    c["threshold"] = res.rule.describe()
    out.append(c)
    ctx.totals = (T_l, T_n)
    zones = cmp_.get("zones")
    if zones:
        z = new_check("conservation.zone_totals", order + 1)
        z["zones"] = zones
        rule = cmp_.get("zone_rule")
        zspec = ctx.spec.get("zones")
        if rule not in ("centre", "all_touched"):
            set_state(z, "EVALUATOR_ERROR", "CHECKER_EXCEPTION")
            z["detail"]["error"] = "compare.zone_rule must be declared (centre | all_touched) and equal the legacy rule"
            out.append(z)
            return out
        if not zspec or not zspec.get("path"):
            set_state(z, "EVALUATOR_ERROR", "CHECKER_EXCEPTION")
            z["detail"]["error"] = f"zones {zones} could not be resolved to a file"
            out.append(z)
            return out
        zspec = dict(zspec, key=zspec.get("key") or zones.split(":", 1)[-1])
        shapes_l, keys = _zones_frame(zspec, lds.crs)
        shapes_n, keys_n = _zones_frame(zspec, nds.crs) if nds.crs is not None else (None, keys)
        if shapes_n is None:
            set_state(z, "FAIL", "CRS_MISSING")
            out.append(z)
            return out
        sl, ol = _zone_sums(lds, band, shapes_l, len(keys), rule == "all_touched")
        sn, on = _zone_sums(nds, band, shapes_n, len(keys), rule == "all_touched")
        zres = ctx.book.lookup(ctx.ds_id, "conservation.zone_totals")
        table, maxrel = [], 0.0
        for i, k in enumerate(keys, start=1):
            r = (abs(sn[i] - sl[i]) / abs(sl[i])) if sl[i] else (0.0 if sn[i] == sl[i] else math.inf)
            maxrel = max(maxrel, r)
            table.append({"zone": k, "legacy": sl[i], "new": sn[i], "rel": r})
        z.update(metric="max_rel_zone_drift", value=maxrel, table=table)
        z["detail"] = {"rule": rule, "outside_legacy": ol, "outside_new": on}
        apply_tol(z, zres, lambda r: all(r.close(sl[i], sn[i]) for i in range(1, len(keys) + 1)),
                  any(sl[i] != sn[i] for i in range(1, len(keys) + 1)), "ZONE_TOTAL_DRIFT")
        z["threshold"] = zres.rule.describe()
        out.append(z)
    return out


def _full_arrays(ds, band: int):
    a = ds.read(band)
    m = ds.read_masks(band) > 0
    if a.dtype.kind == "f":
        m &= np.isfinite(a)
    nd = ds.nodatavals[band - 1]
    if nd is not None and not (isinstance(nd, float) and math.isnan(nd)):
        m &= a != nd
    return a.astype("float64"), m


def _numeric_diagnostics(c: dict, l: np.ndarray, n: np.ndarray, over: np.ndarray | None) -> None:
    """NUMERIC_SCALE_FACTOR (median ratio ≈ 10^k, k≠0), NUMERIC_SIGN_FLIP, NUMERIC_OFFSET (constant shift)."""
    nz = (l != 0) & (n != 0) & np.isfinite(l) & np.isfinite(n)
    if nz.any():
        med = float(np.median(n[nz] / l[nz]))
        if med != 0 and math.isfinite(med):
            if abs(med + 1) <= 0.01:
                add_diag(c, "NUMERIC_SIGN_FLIP", sign_flip={"median_ratio": med})
            elif med > 0:
                k = round(math.log10(med))
                if k != 0 and abs(med / (10.0 ** k) - 1) <= 0.01:
                    add_diag(c, "NUMERIC_SCALE_FACTOR", scale_factor={"k": int(k), "median_ratio": med})
    if over is not None and over.any():
        d = (n - l)[over]
        d0 = float(np.median(d))
        if d0 != 0 and np.all(np.abs(d - d0) <= 1e-9 * max(1.0, abs(d0))):
            add_diag(c, "NUMERIC_OFFSET", offset={"value": d0, "cells": int(d.size)})


def _raster_diagnostics(ctx: "PairCtx", lds, nds, band: int, checks: list[dict]) -> None:
    """Diagnostics on FAIL only; never change a state (06 §4b.7)."""
    sfx = "" if band == 1 else f":b{band}"
    vchk = next((c for c in checks if c["id"] == f"raster.values{sfx}"), None)
    mchk = next((c for c in checks if c["id"] == f"raster.mask{sfx}"), None)
    target = vchk if (vchk and vchk["state"] == "FAIL") else (mchk if (mchk and mchk["state"] == "FAIL") else None)
    if target is None:
        return
    if lds.width * lds.height > DIAG_MAX_CELLS:
        target["diagnostics"]["skipped"] = f"grid larger than {DIAG_MAX_CELLS} cells"
        return
    L, Lm = _full_arrays(lds, band)
    Nn, Nm = _full_arrays(nds, band)
    res = ctx.book.lookup(ctx.ds_id, "raster.values")
    rule = res.rule

    def mism(lsl, nsl):
        lm_, nm_ = Lm[lsl], Nm[nsl]
        both = lm_ & nm_
        o = rule.over(L[lsl][both], Nn[nsl][both])
        return int((lm_ ^ nm_).sum()) + int(o.sum())
    full = (slice(None), slice(None))
    m0 = mism(full, full)
    best = None
    H, W = L.shape
    if m0 > 0:
        for dy in range(-2, 3):
            for dx in range(-2, 3):
                if dy == 0 and dx == 0:
                    continue
                lr = slice(max(0, -dy), H - max(0, dy))
                lc = slice(max(0, -dx), W - max(0, dx))
                nr = slice(max(0, dy), H - max(0, -dy))
                nc = slice(max(0, dx), W - max(0, -dx))
                mm = mism((lr, lc), (nr, nc))
                if best is None or mm < best[0]:
                    best = (mm, [dy, dx])
        if best and best[0] <= 0.1 * m0:
            add_diag(target, "DATA_OFFSET_SUSPECTED",
                     data_offset={"shift": best[1], "mismatches": m0, "after_shift": best[0]})
    both = Lm & Nm
    lv, nv = L[both], Nn[both]
    _numeric_diagnostics(target, lv, nv, rule.over(lv, nv))
    T_l = math.fsum(L[Lm].tolist())
    T_n = math.fsum(Nn[Nm].tolist())
    if T_l and abs(T_n - T_l) <= 0.01 * abs(T_l) and H >= 3 and W >= 3:
        core = both[1:-1, 1:-1] & both[:-2, 1:-1] & both[2:, 1:-1] & both[1:-1, :-2] & both[1:-1, 2:]
        if core.any():
            def lap(A):
                return np.abs(4 * A[1:-1, 1:-1] - A[:-2, 1:-1] - A[2:, 1:-1] - A[1:-1, :-2] - A[1:-1, 2:])[core]
            ll, ln = float(lap(L).mean()), float(lap(Nn).mean())
            if ll > 0:
                ratio = ln / ll
                target["diagnostics"]["laplacian_ratio"] = ratio
                if ratio < 0.8:
                    add_diag(target, "SMOOTHING_SUSPECTED", smoothing={"laplacian_ratio": ratio,
                                                                       "total_rel_drift": abs(T_n - T_l) / abs(T_l)})


def compare_raster(ctx: "PairCtx") -> list[dict]:
    lds = _open_raster(ctx.legacy, SIDE_LEGACY)
    try:
        nds = _open_raster(ctx.new, SIDE_NEW)
    except OutputUnreadable:
        lds.close()
        raise
    with lds, nds:
        checks: list[dict] = []
        # G1 band count
        g1 = new_check("raster.bands", 1, gate=True)
        g1.update(metric="band_count", legacy=lds.count, new=nds.count)
        set_state(g1, *(("PASS_EXACT",) if lds.count == nds.count else ("FAIL", "RASTER_BAND_COUNT_MISMATCH")))
        checks.append(g1)
        # G2 CRS
        g2 = new_check("raster.grid.crs", 2, gate=True)
        g2["legacy"] = lds.crs.to_string() if lds.crs else None
        g2["new"] = nds.crs.to_string() if nds.crs else None
        if lds.crs is None or nds.crs is None:
            set_state(g2, "FAIL", "CRS_MISSING")
        elif _crs_equal(lds.crs, nds.crs):
            set_state(g2, "PASS_EXACT")
        else:
            set_state(g2, "FAIL", "CRS_MISMATCH")
        checks.append(g2)
        # G3 shape
        g3 = new_check("raster.grid.shape", 3, gate=True)
        g3.update(metric="shape", legacy=[lds.height, lds.width], new=[nds.height, nds.width])
        if (lds.height, lds.width) == (nds.height, nds.width):
            set_state(g3, "PASS_EXACT")
        else:
            set_state(g3, "FAIL", "RASTER_SHAPE_MISMATCH")
            lt, nt = lds.transform, nds.transform
            g3["diagnostics"]["offset_cells"] = [round((nt.c - lt.c) / lt.a, 6) if lt.a else None,
                                                 round((nt.f - lt.f) / lt.e, 6) if lt.e else None]
        checks.append(g3)
        # G4 transform: origin and far corner, in pixels
        g4 = new_check("raster.grid.transform", 4, gate=True)
        dev, shift = _transform_dev(lds.transform, nds.transform, lds.width, lds.height)
        g4.update(metric="transform_dev_px", value=dev, threshold=GRID_TOL_PX,
                  tolerance_id="default.raster.grid.transform", legacy=list(lds.transform)[:6],
                  new=list(nds.transform)[:6])
        if dev <= GRID_TOL_PX:
            set_state(g4, "PASS_EXACT")
        else:
            set_state(g4, "FAIL", "RASTER_TRANSFORM_MISMATCH")
            g4["diagnostics"]["shift_px"] = shift
        checks.append(g4)
        failed_gates = [g["id"] for g in (g1, g2, g3, g4) if g["state"] != "PASS_EXACT"]
        gate_ok = not failed_gates
        ctx.gate_name = failed_gates[0] if failed_gates else None
        # S1 dtype / S2 nodata (hard, never gate)
        nb = min(lds.count, nds.count)
        for b in range(1, nb + 1):
            sfx = "" if b == 1 else f":b{b}"
            s1 = new_check(f"raster.dtype{sfx}", 5)
            ldt, ndt = lds.dtypes[b - 1], nds.dtypes[b - 1]
            s1.update(metric="dtype", legacy=ldt, new=ndt)
            apply_norm(s1, ldt == ndt, ctx.kdb, "dtype_widening",
                       lambda kd: kd["rule"]["dtype_widening"].get("from") == ldt
                       and kd["rule"]["dtype_widening"].get("to") == ndt, "DTYPE_MISMATCH")
            checks.append(s1)
            s2 = new_check(f"raster.nodata_value{sfx}", 6)
            lnd, nnd = lds.nodatavals[b - 1], nds.nodatavals[b - 1]
            s2.update(metric="nodata", legacy=_jsonable(lnd), new=_jsonable(nnd))
            if lnd is not None and math.isnan(lnd):
                s2["legacy"] = "nan"
            if nnd is not None and math.isnan(nnd):
                s2["new"] = "nan"

            def nd_fix(kd, lnd=lnd, nnd=nnd):
                eq = kd["rule"]["nodata_equivalent"]
                return isinstance(eq, list) and any(_nodata_equal(lnd, x) for x in eq) and \
                    any(_nodata_equal(nnd, x) for x in eq)
            apply_norm(s2, _nodata_equal(lnd, nnd), ctx.kdb, "nodata_equivalent", nd_fix, "NODATA_VALUE_MISMATCH")
            checks.append(s2)
            checks.extend(_raster_values(ctx, lds, nds, b, sfx, gate_ok, 7))
        if not gate_ok and nb == 0:
            checks.extend([gated("raster.mask", 7, ctx.gate_name), gated("raster.values", 8, ctx.gate_name)])
        checks.extend(_conservation(ctx, lds, nds, 1, 20))
        if gate_ok:
            for b in range(1, nb + 1):
                _raster_diagnostics(ctx, lds, nds, b, checks)
            ctx.raster_pair_arrays = lambda: (_full_arrays(lds, 1), _full_arrays(nds, 1))

            def values_for(chk: str):
                if not chk.startswith("raster.values") or lds.width * lds.height > DIAG_MAX_CELLS:
                    return None
                (L, Lm), (Nn, Nm) = _full_arrays(lds, 1), _full_arrays(nds, 1)
                both = Lm & Nm
                return L[both], Nn[both]
            checks.extend(kd_change_checks(ctx.kdb, 30, values_for))
        else:
            checks.extend(kd_change_checks(ctx.kdb, 30, lambda chk: None))
        checks.extend(_invariant_checks(ctx, 40))
        return checks


# =================================================================================================================
# Vector and table families (06 §4b vector/table; declared key, never row order)
# =================================================================================================================
def _family_of(values: list) -> str:
    fams = set()
    for v in values:
        if v is None:
            continue
        if isinstance(v, float) and math.isnan(v):
            continue
        if isinstance(v, (bool, np.bool_)):
            fams.add("bool")
        elif isinstance(v, (int, float, np.integer, np.floating)):
            fams.add("numeric")
        elif isinstance(v, (_dt.datetime, _dt.date, np.datetime64)):
            fams.add("datetime")
        elif isinstance(v, str):
            fams.add("string")
        elif isinstance(v, bytes):
            fams.add("bytes")
        else:
            try:
                import pandas as pd
                if isinstance(v, pd.Timestamp):
                    fams.add("datetime")
                    continue
            except ImportError as e:
                raise EvaluatorFault("CHECKER_EXCEPTION", str(e)) from e
            fams.add(type(v).__name__)
    if not fams:
        return "empty"
    if len(fams) == 1:
        return fams.pop()
    return "mixed:" + ",".join(sorted(fams))


def _is_null(v) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, np.floating) and np.isnan(v):
        return True
    try:
        import pandas as pd
        if v is pd.NaT or v is pd.NA:
            return True
    except ImportError as e:
        raise EvaluatorFault("CHECKER_EXCEPTION", str(e)) from e
    return False


def _canon_key(v, strip=False, casefold=False) -> str:
    if _is_null(v):
        return ""
    if isinstance(v, np.generic):
        v = v.item()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (int, float)):
        return repr(v)
    s = str(v)
    if strip:
        s = s.strip()
    if casefold:
        s = s.casefold()
    return s


@dataclass
class Frame:
    columns: list
    data: dict                     # column -> list of python values
    nrows: int
    geoms: list | None = None
    crs: Any = None
    layer: str | None = None
    layout: dict | None = None
    total_row: dict | None = None


def _frame_from_df(df, geom: bool) -> Frame:
    cols = [c for c in df.columns if not (geom and c == df.geometry.name)] if geom else list(df.columns)
    data = {}
    for c in cols:
        vals = df[c].tolist()
        data[c] = [None if _is_null(v) else (v.item() if isinstance(v, np.generic) else v) for v in vals]
    fr = Frame(columns=cols, data=data, nrows=len(df))
    if geom:
        fr.geoms = list(df.geometry.values)
        fr.crs = df.crs
    return fr


def _list_layers(path: str, side: str) -> list[str]:
    import pyogrio
    try:
        return [str(x[0]) for x in pyogrio.list_layers(path)]
    except Exception as e:                                # typed re-raise
        if side == SIDE_LEGACY:
            raise LegacyUnreadable(f"{path}: {e}") from e
        raise OutputUnreadable(f"{path}: {e}") from e


def _read_vector(path: str, layer: str, side: str) -> Frame:
    import pyogrio
    try:
        df = pyogrio.read_dataframe(path, layer=layer)
    except Exception as e:                                # typed re-raise
        if side == SIDE_LEGACY:
            raise LegacyUnreadable(f"{path}|{layer}: {e}") from e
        raise OutputUnreadable(f"{path}|{layer}: {e}") from e
    has_geom = hasattr(df, "geometry") and getattr(df, "_geometry_column_name", None) in df.columns
    fr = _frame_from_df(df, has_geom)
    fr.layer = layer
    return fr


def _resolve_layers(ctx: "PairCtx") -> tuple[str, str | None, dict]:
    spec = ctx.spec
    ll = spec.get("layer") or spec.get("legacy_layer")
    llayers = _list_layers(ctx.legacy, SIDE_LEGACY)
    nlayers = _list_layers(ctx.new, SIDE_NEW)
    if ll is None:
        if len(llayers) != 1:
            raise EvaluatorFault("LAYER_AMBIGUOUS", f"{len(llayers)} layers in the legacy file and none declared: "
                                                    f"{llayers}")
        ll = llayers[0]
    elif ll not in llayers:
        raise EvaluatorFault("LAYER_AMBIGUOUS", f"declared legacy layer {ll!r} not in {llayers}")
    nl = spec.get("new_layer")
    if nl is None:
        if spec.get("layer") or spec.get("legacy_layer"):
            nl = ll
        elif len(nlayers) == 1:
            nl = nlayers[0]
        else:
            nl = ll
    info = {"legacy_layers": llayers, "new_layers": nlayers, "legacy_layer": ll, "new_layer": nl}
    return ll, (nl if nl in nlayers else None), info


def _geom_family(g, multi_equiv: bool) -> str:
    if g is None or g.is_empty:
        return "none"
    t = g.geom_type
    if multi_equiv:
        t = {"MultiPolygon": "Polygon", "MultiLineString": "LineString", "MultiPoint": "Point"}.get(t, t)
    return t


def _canon_geom(g):
    import shapely
    if g is None:
        return None
    if g.geom_type.startswith("Multi") and len(g.geoms) == 1:
        g = g.geoms[0]
    return shapely.normalize(g)


def _read_table(ctx: "PairCtx", path: str, side: str) -> Frame:
    cmp_ = ctx.spec.get("compare") or {}
    p = str(path).lower()
    header_row = int(cmp_.get("header_row") or 1)
    data_rows = cmp_.get("data_rows")
    total_row = cmp_.get("total_row")
    if p.endswith((".xlsx", ".xlsm")):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:                            # typed re-raise
            if side == SIDE_LEGACY:
                raise LegacyUnreadable(f"{path}: {e}") from e
            raise OutputUnreadable(f"{path}: {e}") from e
        try:
            sheet = ctx.spec.get("sheet") or cmp_.get("sheet")
            if sheet is None:
                if len(wb.sheetnames) != 1:
                    if side == SIDE_LEGACY:
                        raise EvaluatorFault("LAYER_AMBIGUOUS", f"{len(wb.sheetnames)} sheets and none declared: "
                                                                f"{wb.sheetnames}")
                sheet = wb.sheetnames[0]
            if sheet not in wb.sheetnames:
                if side == SIDE_LEGACY:
                    raise EvaluatorFault("LAYER_AMBIGUOUS", f"declared sheet {sheet!r} not in {wb.sheetnames}")
                raise _SchemaMissing(f"sheet {sheet!r} missing in new workbook ({wb.sheetnames})")
            rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        finally:
            wb.close()
        layout = {"format": "xlsx", "sheet": sheet}
    elif p.endswith((".csv", ".tsv", ".txt")):
        import csv
        delim = cmp_.get("delimiter") or ("\t" if p.endswith(".tsv") else ",")
        try:
            with open(path, newline="", encoding=cmp_.get("encoding") or "utf-8-sig") as fh:
                rows = [list(r) for r in csv.reader(fh, delimiter=delim)]
        except (OSError, UnicodeDecodeError, csv.Error) as e:
            if side == SIDE_LEGACY:
                raise LegacyUnreadable(f"{path}: {e}") from e
            raise OutputUnreadable(f"{path}: {e}") from e
        rows = [[None if v == "" else v for v in r] for r in rows]
        layout = {"format": "csv", "delimiter": delim}
    else:
        raise EvaluatorFault("FAMILY_UNSUPPORTED", f"table format of {path} not supported")
    if len(rows) < header_row:
        if side == SIDE_LEGACY:
            raise LegacyUnreadable(f"{path}: fewer rows than header_row {header_row}")
        raise OutputUnreadable(f"{path}: fewer rows than header_row {header_row}")
    header = rows[header_row - 1]
    width = max([len(header)] + [len(r) for r in rows])
    header = list(header) + [None] * (width - len(header))
    cols = [str(h) if h is not None else f"_col{j + 1}" for j, h in enumerate(header)]
    if data_rows:
        a, b = int(data_rows[0]), int(data_rows[1])
        body_idx = [i for i in range(a, b + 1) if i <= len(rows)]
    else:
        body_idx = [i for i in range(header_row + 1, len(rows) + 1) if i != total_row]
    body = []
    for i in body_idx:
        r = list(rows[i - 1]) + [None] * (width - len(rows[i - 1]))
        if all(v is None for v in r):
            continue
        body.append(r)
    keep = [j for j in range(width) if header[j] is not None or any(r[j] is not None for r in body)]
    cols = [cols[j] for j in keep]
    data = {c: [r[j] for r in body] for c, j in zip(cols, keep)}
    if layout["format"] == "csv":
        for c in cols:                                  # one reader config: numbers where every value parses
            vals = data[c]
            parsed, ok = [], True
            for v in vals:
                if v is None:
                    parsed.append(None)
                    continue
                if str(v).strip().lower() in ("nan",):
                    parsed.append(None)
                    continue
                try:
                    parsed.append(float(v))
                except ValueError:
                    ok = False
                    break
            if ok and any(x is not None for x in parsed):
                data[c] = parsed
    tot = None
    if total_row:
        if total_row <= len(rows):
            tr = list(rows[total_row - 1]) + [None] * (width - len(rows[total_row - 1]))
            tot = {cols[k]: tr[j] for k, j in enumerate(keep)}
        else:
            tot = {}
    layout.update({"header_row": header_row, "data_rows": [body_idx[0], body_idx[-1]] if body_idx else None,
                   "total_row": total_row, "declared": bool(data_rows)})
    return Frame(columns=cols, data=data, nrows=len(body), layout=layout, total_row=tot)


class _SchemaMissing(Exception):
    pass


def _keys_for(fr: Frame, cols: list[str], strip_cols: set, case_cols: set) -> list[tuple]:
    return [tuple(_canon_key(fr.data[c][i], c in strip_cols, c in case_cols) for c in cols) for i in range(fr.nrows)]


def _bbox(g) -> list | None:
    if g is None or g.is_empty:
        return None
    return [float(x) for x in g.bounds]


def _compare_frames(ctx: "PairCtx", L: Frame, N: Frame, prefix: str, order: int) -> tuple[list[dict], dict]:
    """Shared by vectors and tables: key, counts, key sets, schema, attributes (null pattern before values)."""
    spec = ctx.spec
    cmp_ = spec.get("compare") or {}
    checks: list[dict] = []
    ignore = set(cmp_.get("ignore") or [])
    rename = dict(cmp_.get("rename") or {})
    # KD rename normalisation (schema check)
    kd_active_r, kd_cand_r = ctx.kdb.norm(f"{prefix}.schema", "rename")
    rename_kd = {}
    for kd in kd_active_r:
        rename_kd.update(kd["rule"]["rename"] or {})
    rename_cand = {}
    for kd in kd_cand_r:
        rename_cand.update(kd["rule"]["rename"] or {})

    def mapped(c, extra=None):
        m = dict(rename)
        if extra:
            m.update(extra)
        return m.get(c, c)
    lcols = [c for c in L.columns if c not in ignore]
    ncols = set(N.columns)
    # ---- schema
    sc = new_check(f"{prefix}.schema", order)
    missing_raw = [c for c in lcols if mapped(c) not in ncols]
    missing_kd = [c for c in lcols if mapped(c, rename_kd) not in ncols]
    eff = dict(rename)
    if not missing_kd and missing_raw:
        eff.update(rename_kd)
    colmap = {c: eff.get(c, c) for c in lcols if eff.get(c, c) in ncols}
    fam_bad = []
    for c, nc in colmap.items():
        fl, fn = _family_of(L.data[c]), _family_of(N.data[nc])
        if fl != fn and "empty" not in (fl, fn):
            fam_bad.append({"column": c, "legacy": fl, "new": fn})
    extra = sorted(ncols - set(colmap.values()) - {eff.get(c, c) for c in ignore})
    sc["detail"] = {"missing": sorted(missing_kd if (not missing_kd and missing_raw) else missing_raw),
                    "extra_new_columns": extra, "dtype_family_mismatch": fam_bad, "rename": eff}
    sc.update(metric="missing_columns", value=len(sc["detail"]["missing"]) + len(fam_bad))
    if fam_bad:
        set_state(sc, "FAIL", "SCHEMA_MISMATCH")
    else:
        apply_norm(sc, not missing_raw, ctx.kdb, "rename",
                   lambda kd: not [c for c in lcols if mapped(c, kd["rule"]["rename"]) not in ncols],
                   "SCHEMA_MISMATCH")
        if sc["state"] == "PASS_WITHIN_TOLERANCE":
            sc["detail"]["missing"] = []
        if sc["state"] == "NEEDS_DECISION" and not rename_cand:
            set_state(sc, "FAIL", "SCHEMA_MISMATCH")
    checks.append(sc)
    # ---- key
    key = cmp_.get("key")
    match = cmp_.get("match")
    kc = new_check(f"{prefix}.key", order + 1, gate=True)
    strip_cols, case_cols = set(), set()
    for kd in ctx.kdb.norm(f"{prefix}.key", "string_strip")[0]:
        strip_cols |= set((kd["rule"]["string_strip"] or {}).get("columns", []))
    for kd in ctx.kdb.norm(f"{prefix}.key", "casefold")[0]:
        case_cols |= set((kd["rule"]["casefold"] or {}).get("columns", []))
    if isinstance(key, str):
        key = [key]
    if key:
        miss_l = [k for k in key if k not in L.columns]
        miss_n = [k for k in key if eff.get(k, k) not in N.columns]
        if miss_l:
            raise EvaluatorFault("KEY_UNDECLARED", f"declared key columns {miss_l} are not in the legacy file")
        if miss_n:
            set_state(kc, "FAIL", "SCHEMA_MISMATCH")
            kc["detail"]["missing_key_columns_new"] = miss_n
            checks.append(kc)
            return checks, {}
        lk = _keys_for(L, key, strip_cols, case_cols)
        nk = [tuple(_canon_key(N.data[eff.get(k, k)][i], k in strip_cols, k in case_cols) for k in key)
              for i in range(N.nrows)]
        kc["detail"]["key"] = key
    elif match == "geometry" and L.geoms is not None and N.geoms is not None:
        def gk(geoms):
            seen = collections.Counter()
            out = []
            for g in geoms:
                cg = _canon_geom(g)
                h = hashlib.sha1(cg.wkb if cg is not None else b"").hexdigest()
                seen[h] += 1
                out.append((h, str(seen[h])))
            return out
        lk, nk = gk(L.geoms), gk(N.geoms)
        kc["detail"]["match"] = "geometry (normalised WKB multiset)"
    elif match == "full_row":
        def rk(fr, cmap):
            seen = collections.Counter()
            out = []
            for i in range(fr.nrows):
                t = tuple(_canon_key(fr.data[cmap[c]][i]) for c in sorted(cmap))
                seen[t] += 1
                out.append(t + (str(seen[t]),))
            return out
        lk = rk(L, {c: c for c in colmap})
        nk = rk(N, colmap)
        kc["detail"]["match"] = "full_row"
    else:
        raise EvaluatorFault("KEY_UNDECLARED", "no key declared and no match: geometry|full_row declared "
                                               "(no auto-detected key, never row order)")
    dl = [k for k, n in collections.Counter(lk).items() if n > 1]
    dn = [k for k, n in collections.Counter(nk).items() if n > 1]
    if dl or dn:
        set_state(kc, "FAIL", "KEY_DUPLICATE")
        kc["detail"].update(duplicate_legacy=[list(k) for k in dl[:SAMPLE_N]],
                            duplicate_new=[list(k) for k in dn[:SAMPLE_N]])
    else:
        set_state(kc, "PASS_EXACT")
    checks.append(kc)
    # ---- counts and key sets
    cc = new_check(f"{prefix}.{'feature_count' if prefix == 'vector' else 'row_count'}", order + 2)
    cc.update(metric="count", legacy=L.nrows, new=N.nrows, value=N.nrows - L.nrows)
    if L.nrows == N.nrows:
        set_state(cc, "PASS_EXACT")
    else:
        set_state(cc, "FAIL", "FEATURE_COUNT_MISMATCH" if prefix == "vector" else "ROW_COUNT_MISMATCH")
    if N.nrows == 0 and L.nrows > 0 and not cmp_.get("may_be_empty"):
        set_state(cc, "FAIL", "OUTPUT_EMPTY")
    checks.append(cc)
    ks = new_check(f"{prefix}.key_set", order + 3)
    li = {k: i for i, k in enumerate(lk)}
    ni = {k: i for i, k in enumerate(nk)}
    missing = sorted(set(li) - set(ni))
    extra_k = sorted(set(ni) - set(li))
    ks.update(metric="keys_missing+extra", value=len(missing) + len(extra_k))
    ks["detail"] = {"missing": len(missing), "extra": len(extra_k)}

    def ksample(keys, fr, idx):
        out = []
        for k in keys[:SAMPLE_N]:
            s = {"key": list(k)}
            if fr.geoms is not None:
                s["bbox"] = _bbox(fr.geoms[idx[k]])
            out.append(s)
        return out
    ks["samples"] = ksample(missing, L, li) + ksample(extra_k, N, ni)
    ctx.key_index_legacy, ctx.keys_missing = li, missing
    codes = (["KEY_SET_MISSING"] if missing else []) + (["KEY_SET_EXTRA"] if extra_k else [])
    set_state(ks, "FAIL" if codes else "PASS_EXACT", *codes)
    checks.append(ks)
    for k in missing:
        ctx.mismatch_rows.append({"check": ks["id"], "key": json.dumps(list(k)), "kind": "key_missing"})
    for k in extra_k:
        ctx.mismatch_rows.append({"check": ks["id"], "key": json.dumps(list(k)), "kind": "key_extra"})
    if kc["state"] != "PASS_EXACT":
        for c in lcols:
            if c in colmap and c not in (key or []):
                checks.append(gated(f"{prefix}.attr:{c}", order + 10, kc["id"]))
        return checks, {"gated_by": kc["id"]}
    matched = sorted(set(li) & set(ni))
    lidx = [li[k] for k in matched]
    nidx = [ni[k] for k in matched]
    ctx.matched_keys = matched
    # ---- attributes per column
    for c in lcols:
        if c not in colmap or c in (key or []):
            continue
        checks.append(_attr_check(ctx, prefix, c, colmap[c], L, N, lidx, nidx, matched, order + 10))
    if prefix == "table":
        checks.extend(_table_totals(ctx, L, N, colmap, order + 30))
    return checks, {"lidx": lidx, "nidx": nidx, "matched": matched}


def _attr_check(ctx, prefix, c, nc, L: Frame, N: Frame, lidx, nidx, matched, order) -> dict:
    cid = f"{prefix}.attr:{c}"
    chk = new_check(cid, order)
    lv = [L.data[c][i] for i in lidx]
    nv = [N.data[nc][i] for i in nidx]
    lnull = np.array([_is_null(v) for v in lv], dtype=bool)
    nnull = np.array([_is_null(v) for v in nv], dtype=bool)
    nullx = lnull != nnull
    both = ~lnull & ~nnull
    fam = _family_of(lv) if _family_of(lv) != "empty" else _family_of(nv)
    chk["detail"]["family"] = fam
    if nc != c:
        chk["detail"]["new_column"] = nc
    samples, rows = [], []
    null_fail = bool(nullx.any())
    if null_fail:
        for i in np.nonzero(nullx)[0]:
            if len(samples) < SAMPLE_N:
                samples.append({"key": list(matched[i]), "legacy": _jsonable(lv[i]), "new": _jsonable(nv[i]),
                                "kind": "null_pattern"})
            rows.append({"check": cid, "key": json.dumps(list(matched[i])), "legacy": str(_jsonable(lv[i])),
                         "new": str(_jsonable(nv[i])), "kind": "null_pattern"})
    chk["detail"]["null_pattern_mismatch"] = int(nullx.sum())
    bi = np.nonzero(both)[0]
    if fam == "numeric" and all(isinstance(v, (int, float)) and not isinstance(v, bool)
                                for v in [lv[i] for i in bi] + [nv[i] for i in bi]):
        l = np.array([float(lv[i]) for i in bi], dtype="float64")
        n = np.array([float(nv[i]) for i in bi], dtype="float64")
        res = ctx.book.lookup(ctx.ds_id, cid)
        rules = [res.rule] + ([res.proposed] if res.proposed is not None else [])
        overs = [r.over(l, n) for r in rules]
        d = np.abs(n - l)
        lnz = (l != 0) & (d > 0)
        chk["stats"] = {"compared": int(l.size), "count_diff": int((d > 0).sum()), "count_over": int(overs[0].sum()),
                        "max_abs": float(d.max()) if d.size else 0.0,
                        "max_rel": float((d[lnz] / np.abs(l[lnz])).max()) if lnz.any() else 0.0}
        chk.update(metric="count_over", value=int(overs[0].sum()), threshold=res.rule.describe())
        for j in np.nonzero(overs[0])[0]:
            i = bi[j]
            if len(samples) < SAMPLE_N:
                samples.append({"key": list(matched[i]), "legacy": float(l[j]), "new": float(n[j])})
            rows.append({"check": cid, "key": json.dumps(list(matched[i])), "legacy": repr(float(l[j])),
                         "new": repr(float(n[j])), "kind": "value"})
        apply_tol(chk, res, lambda r: int(overs[rules.index(r)].sum()) <= r.max_mismatch if r in rules else False,
                  bool((d > 0).any()), "ATTRIBUTE_VALUE_MISMATCH")
        if chk["state"] == "FAIL":
            _numeric_diagnostics(chk, l, n, overs[0])
        ctx.attr_values[cid] = (l, n)
    else:
        def norm(v, strip, cf):
            s = v if isinstance(v, str) else str(_jsonable(v))
            if strip:
                s = s.strip()
            if cf:
                s = s.casefold()
            return s

        def mism(strip, cf):
            return [i for i in bi if norm(lv[i], strip, cf) != norm(nv[i], strip, cf)
                    or (not isinstance(lv[i], str) and type(lv[i]) is not type(nv[i])
                        and not (isinstance(lv[i], (int, float)) and isinstance(nv[i], (int, float))))]
        raw = mism(False, False)
        chk.update(metric="count_over", value=len(raw))
        act_s, cand_s = ctx.kdb.norm(cid, "string_strip")
        act_c, cand_c = ctx.kdb.norm(cid, "casefold")

        def col_in(kd, rk):
            cols = (kd["rule"][rk] or {}).get("columns") if isinstance(kd["rule"][rk], dict) else None
            return cols is None or c in cols

        def fixed(kd):
            s = "string_strip" in kd["rule"] and col_in(kd, "string_strip")
            f = "casefold" in kd["rule"] and col_in(kd, "casefold")
            return not mism(s, f)
        if raw:
            for i in raw[:SAMPLE_N]:
                samples.append({"key": list(matched[i]), "legacy": _jsonable(lv[i]), "new": _jsonable(nv[i])})
            for i in raw:
                rows.append({"check": cid, "key": json.dumps(list(matched[i])), "legacy": str(_jsonable(lv[i])),
                             "new": str(_jsonable(nv[i])), "kind": "value"})
        active = [kd for kd in act_s + act_c if fixed(kd)]
        cands = [kd for kd in cand_s + cand_c if fixed(kd)]
        chk["tolerance_id"] = "default.attr.exact"
        if not raw:
            set_state(chk, "PASS_EXACT")
        elif active:
            set_state(chk, "PASS_WITHIN_TOLERANCE")
            chk["kd"] = sorted({kd["id"] for kd in active})
            for kd in active:
                ctx.kdb.mark(kd, cid)
            rows = [r for r in rows if r["kind"] != "value"]
        elif cands:
            set_state(chk, "NEEDS_DECISION", "KD_UNSIGNED")
            chk["proposed_kd"] = sorted({kd["id"] for kd in cands})
        else:
            set_state(chk, "FAIL", "ATTRIBUTE_VALUE_MISMATCH")
    if null_fail:
        state_codes = [x for x in chk["reason_codes"] if x not in ("TOLERANCE_UNSIGNED", "KD_UNSIGNED")]
        chk["reason_codes"] = ["NULL_PATTERN_MISMATCH"] + [x for x in state_codes if x != "NULL_PATTERN_MISMATCH"]
        chk["state"] = "FAIL"
    chk["samples"] = samples
    ctx.mismatch_rows.extend(rows[:MISMATCH_ROWS_MAX])
    return chk


def _table_totals(ctx, L: Frame, N: Frame, colmap: dict, order: int) -> list[dict]:
    out = []
    for c, nc in colmap.items():
        if _family_of(L.data[c]) != "numeric" or _family_of(N.data[nc]) not in ("numeric", "empty"):
            continue
        lvals = [float(v) for v in L.data[c] if not _is_null(v) and not isinstance(v, bool)]
        nvals = [float(v) for v in N.data[nc] if not _is_null(v) and not isinstance(v, bool)]
        tl, tn = math.fsum(lvals), math.fsum(nvals)
        cid = f"table.total:{c}"
        chk = new_check(cid, order)
        chk.update(metric="rel_drift", legacy=tl, new=tn,
                   value=(abs(tn - tl) / abs(tl)) if tl else (0.0 if tn == tl else None))
        if chk["value"] is None:
            chk.setdefault("nulls", {})["value"] = "UNDEFINED_ZERO_REFERENCE"
        res = ctx.book.lookup(ctx.ds_id, cid)
        if not res.considered and res.error is None:
            res = ctx.book.lookup(ctx.ds_id, f"table.attr:{c}")
        apply_tol(chk, res, lambda r: r.close(tl, tn), tl != tn, "TOTAL_DRIFT")
        out.append(chk)
    if L.total_row is not None:
        chk = new_check("table.total_row", order + 1)
        diffs = []
        for c, nc in colmap.items():
            lv = (L.total_row or {}).get(c)
            nv = (N.total_row or {}).get(nc) if N.total_row is not None else None
            if isinstance(lv, (int, float)) and not isinstance(lv, bool):
                res = ctx.book.lookup(ctx.ds_id, f"table.attr:{c}")
                ok = isinstance(nv, (int, float)) and (res.rule.close(float(lv), float(nv)) if res.error is None
                                                       else float(lv) == float(nv))
                if not ok:
                    diffs.append({"column": c, "legacy": _jsonable(lv), "new": _jsonable(nv)})
            elif lv != nv and not (_is_null(lv) and _is_null(nv)):
                diffs.append({"column": c, "legacy": _jsonable(lv), "new": _jsonable(nv)})
        chk.update(metric="columns_differing", value=len(diffs))
        chk["samples"] = diffs[:SAMPLE_N]
        set_state(chk, *(("PASS_EXACT",) if not diffs else ("FAIL", "TOTAL_DRIFT")))
        out.append(chk)
    return out


def _geometry_checks(ctx: "PairCtx", L: Frame, N: Frame, m: dict, gate: str | None, order: int) -> list[dict]:
    import geopandas as gpd
    import pyproj
    import shapely
    polygonal = any(_geom_family(g, True) == "Polygon" for g in L.geoms)
    out = []
    ids = ["vector.geometry_validity", "vector.hausdorff_m"] + (["vector.symdiff_ratio"] if polygonal else [])
    if gate or "lidx" not in m:
        return [gated(i, order + k, gate or m.get("gated_by")) for k, i in enumerate(ids)]
    lg = [L.geoms[i] for i in m["lidx"]]
    ng = [N.geoms[i] for i in m["nidx"]]
    matched = m["matched"]
    # validity
    va = new_check("vector.geometry_validity", order)
    inv_l = int(sum(1 for g in L.geoms if g is not None and not g.is_valid))
    inv_n = int(sum(1 for g in N.geoms if g is not None and not g.is_valid))
    va.update(metric="invalid_count", legacy=inv_l, new=inv_n, value=inv_n)
    set_state(va, *(("PASS_EXACT",) if inv_n <= inv_l else ("FAIL", "GEOMETRY_INVALID")))
    out.append(va)
    differ = []
    for i, (a, b) in enumerate(zip(lg, ng)):
        ca, cb = _canon_geom(a), _canon_geom(b)
        if ca is None and cb is None:
            continue
        if ca is None or cb is None or not shapely.equals_exact(ca, cb, 0.0):
            differ.append(i)
    hd = new_check("vector.hausdorff_m", order + 1)
    sd = new_check("vector.symdiff_ratio", order + 2) if polygonal else None
    geo_checks = [x for x in (hd, sd) if x is not None]
    if not differ:
        for c in geo_checks:
            c.update(metric=c["id"].split(".")[1], value=0.0, tolerance_id="default.geometry.exact")
            set_state(c, "PASS_EXACT")
            c["detail"]["features_differing"] = 0
        return out + geo_checks
    bad = _crs_invalid_for_measure(ctx.spec.get("measure_crs"))
    if bad:
        for c in geo_checks:
            set_state(c, "EVALUATOR_ERROR", "GEOMETRY_MEASURE_CRS_INVALID")
            c["detail"]["error"] = bad
            c["detail"]["features_differing"] = len(differ)
        return out + geo_checks
    mcrs = pyproj.CRS.from_user_input(ctx.spec["measure_crs"])
    # measurement only: both sides projected into measure_crs; the files are never changed
    gl = gpd.GeoSeries([lg[i] for i in differ], crs=L.crs).to_crs(mcrs)
    gn = gpd.GeoSeries([ng[i] for i in differ], crs=N.crs).to_crs(mcrs)
    haus, syms, ratios, area_ratio, vtx_l, vtx_n = [], [], [], [], 0, 0
    for a, b in zip(gl, gn):
        a_ok = a is not None and not a.is_empty
        b_ok = b is not None and not b.is_empty
        if not (a_ok and b_ok):
            haus.append(math.inf)
            if polygonal:                                 # one side null/empty: the whole other side differs
                sym = (a.area if a_ok else 0.0) + (b.area if b_ok else 0.0)
                syms.append(sym)
                ratios.append(math.inf if not a_ok or a.area == 0 else sym / a.area)
            continue
        haus.append(float(shapely.hausdorff_distance(a, b)))
        vtx_l += int(shapely.get_num_coordinates(a))
        vtx_n += int(shapely.get_num_coordinates(b))
        if polygonal:
            al = a.area
            sym = shapely.symmetric_difference(shapely.make_valid(a), shapely.make_valid(b)).area
            syms.append(sym)
            ratios.append(sym / al if al > 0 else (0.0 if sym == 0 else math.inf))
            if al > 0:
                area_ratio.append(b.area / al)
    thr = lambda r: (r.value if r.value is not None else r.atol)  # noqa: E731
    ctx.geom_by_key = {k: {"sym": 0.0, "haus": 0.0} for k in matched}
    for j, i in enumerate(differ):
        ctx.geom_by_key[matched[i]] = {"sym": syms[j] if polygonal else 0.0, "haus": haus[j]}
    # ---- Hausdorff: per feature, in metres (a local defect on one feature is never averaged away)
    res = ctx.book.lookup(ctx.ds_id, hd["id"])
    arr = np.array(haus, dtype="float64")
    rules = [res.rule] + ([res.proposed] if res.proposed is not None else [])
    over = [int((~(arr <= thr(r))).sum()) for r in rules]
    mx = float(arr.max()) if arr.size else 0.0
    hd.update(metric="hausdorff_m", value=mx if math.isfinite(mx) else None, threshold=thr(res.rule))
    if not math.isfinite(mx):
        hd.setdefault("nulls", {})["value"] = "EMPTY_OR_NULL_GEOMETRY"
    hd["detail"]["features_differing"] = len(differ)
    hd["stats"] = {"count_over": over[0],
                   "median": float(np.median(arr[np.isfinite(arr)])) if np.isfinite(arr).any() else None}
    for j in np.argsort(-np.nan_to_num(arr, posinf=1e300))[:SAMPLE_N]:
        i = differ[int(j)]
        if not (arr[j] <= thr(res.rule)):
            hd["samples"].append({"key": list(matched[i]), "value": _jsonable(float(arr[j])), "bbox": _bbox(lg[i])})
    for j in range(arr.size):
        if not (arr[j] <= thr(res.rule)):
            ctx.mismatch_rows.append({"check": hd["id"], "key": json.dumps(list(matched[differ[j]])), "legacy": None,
                                      "new": None, "kind": "geometry", "value": _jsonable(float(arr[j]))})
    apply_tol(hd, res, lambda r: over[rules.index(r)] <= r.max_mismatch if r in rules else False,
              bool((arr > 0).any()), "GEOMETRY_HAUSDORFF_EXCEEDED")
    if hd["state"] == "FAIL" and vtx_l and vtx_n < 0.8 * vtx_l:
        add_diag(hd, "VERTEX_COUNT_DROP", vertex_count={"legacy": vtx_l, "new": vtx_n})
    # ---- symmetric difference as a per-feature MEAN BOUNDARY DISPLACEMENT in metres: area(Li^Ni) / perimeter(Li).
    # A ratio to area(Li) explodes on slivers (T0 S1: a legitimate segment-count change gave 7.96 on a 3.87 m2 piece
    # while every mutant stayed <= 0.997: unseparable); an area-weighted ratio was only 1.8x apart from the
    # simplify-before-buffer mutant. Displacement per unit boundary keeps the per-feature guard (max over features),
    # is in metres like Hausdorff (D6-9) and separated every S1 dataset by >= 3.3x. Check id kept for consumers.
    if sd is not None:
        ml = gpd.GeoSeries(lg, crs=L.crs).to_crs(mcrs) if lg else gpd.GeoSeries([])
        areas = ml.area.fillna(0.0).tolist() if lg else []
        perims = ml.length.fillna(0.0).tolist() if lg else []
        ctx.legacy_area_by_key = dict(zip(matched, areas))
        ctx.legacy_perimeter_by_key = dict(zip(matched, perims))
        disp = []
        for j, i in enumerate(differ):
            pl = perims[i]
            disp.append(syms[j] / pl if pl > 0 else (0.0 if syms[j] == 0 else math.inf))
        for j, i in enumerate(differ):
            ctx.geom_by_key[matched[i]]["disp"] = disp[j]
        darr = np.array(disp, dtype="float64")
        res = ctx.book.lookup(ctx.ds_id, sd["id"])
        rules = [res.rule] + ([res.proposed] if res.proposed is not None else [])
        over = [int((~(darr <= thr(r))).sum()) for r in rules]
        mx = float(darr.max()) if darr.size else 0.0
        tot_l, sym_tot = math.fsum(areas), math.fsum(syms)
        rarr = np.array(ratios, dtype="float64")
        sd.update(metric="symdiff_m", value=mx if math.isfinite(mx) else None, threshold=thr(res.rule))
        if not math.isfinite(mx):
            sd.setdefault("nulls", {})["value"] = "EMPTY_OR_NULL_GEOMETRY"
        sd["detail"].update(features_differing=len(differ), symdiff_area_m2=sym_tot, legacy_area_m2=tot_l,
                            metric_definition="max over matched features of area(L xor N) / perimeter(L), metres, "
                                              "in measure_crs")
        sd["stats"] = {"count_over": over[0],
                       "area_weighted_ratio": _jsonable(sym_tot / tot_l) if tot_l > 0 else None,
                       "per_feature_max_ratio": _jsonable(float(rarr.max())) if rarr.size else 0.0}
        for j in np.argsort(-np.nan_to_num(darr, posinf=1e300))[:SAMPLE_N]:
            i = differ[int(j)]
            if not (darr[j] <= thr(res.rule)):
                sd["samples"].append({"key": list(matched[i]), "value": _jsonable(float(darr[j])),
                                      "symdiff_m2": float(syms[j]), "ratio": _jsonable(float(rarr[j])),
                                      "bbox": _bbox(lg[i])})
        for j in range(darr.size):
            if not (darr[j] <= thr(res.rule)):
                ctx.mismatch_rows.append({"check": sd["id"], "key": json.dumps(list(matched[differ[j]])),
                                          "legacy": None, "new": None, "kind": "geometry",
                                          "value": _jsonable(float(darr[j]))})
        apply_tol(sd, res, lambda r: over[rules.index(r)] <= r.max_mismatch if r in rules else False,
                  bool((darr > 0).any()), "GEOMETRY_SYMDIFF_EXCEEDED")
        if sd["state"] == "FAIL":
            if area_ratio:
                add_diag(sd, "GEOMETRY_AREA_RATIO", area_ratio={"median": float(np.median(area_ratio))})
            if vtx_l and vtx_n < 0.8 * vtx_l:
                add_diag(sd, "VERTEX_COUNT_DROP", vertex_count={"legacy": vtx_l, "new": vtx_n})
    return out + geo_checks


def _scalar_under(ctx: "PairCtx", check_id: str, value: float) -> str:
    """PASS | UNSIGNED (passes only under an unsigned proposal) | FAIL for a scalar metric with threshold rules."""
    res = ctx.book.lookup(ctx.ds_id, check_id)
    thr = lambda r: (r.value if r.value is not None else r.atol)  # noqa: E731
    ok = lambda r: math.isfinite(value) and value <= thr(r)  # noqa: E731
    if res.error is None and ok(res.rule):
        return "PASS"
    if res.error is None and res.proposed is not None and ok(res.proposed):
        return "UNSIGNED"
    return "FAIL"


def _branch_checks(ctx: "PairCtx", L: Frame, checks: list[dict], order: int) -> list[dict]:
    """compare.branches: [{id, where: {column: value}}] (a legacy code branch and the rows it produced). A branch
    the legacy data never exercised cannot be proven by parity -> NEEDS_REVIEW/BRANCH_NOT_EXERCISED; otherwise key
    set and geometry are re-checked on the branch subset (DESIGN §3: BRANCH_OUTPUT_MISMATCH -> KEY_SET_MISSING /
    GEOMETRY_SYMDIFF_EXCEEDED on the subset)."""
    out = []
    for br in (ctx.spec.get("compare") or {}).get("branches") or []:
        c = new_check(f"branch:{br.get('id')}", order)
        where = br.get("where") or {}
        c["detail"]["where"] = where
        if not where or any(col not in L.data for col in where):
            set_state(c, "EVALUATOR_ERROR", "CHECKER_EXCEPTION")
            c["detail"]["error"] = f"branch where-columns {sorted(where)} not in the legacy layer"
            out.append(c)
            continue
        li = ctx.key_index_legacy or {}
        sub = [k for k, i in li.items() if all(_canon_key(L.data[col][i]) == _canon_key(v) for col, v in where.items())]
        c.update(metric="legacy_rows_in_branch", legacy=len(sub))
        if not sub:
            set_state(c, "NEEDS_REVIEW", "BRANCH_NOT_EXERCISED")
            out.append(c)
            continue
        if any(x["id"] == "vector.key" and x["state"] != "PASS_EXACT" for x in checks):
            out.append(gated(c["id"], order, "vector.key"))
            continue
        codes, unsigned = [], False
        miss = sorted(set(sub) & set(ctx.keys_missing or []))
        if miss:
            codes.append("KEY_SET_MISSING")
            c["samples"] = [{"key": list(k)} for k in miss[:SAMPLE_N]]
        gk = ctx.geom_by_key or {}
        present = [k for k in sub if k in gk]
        detail = {"rows": len(sub), "missing": len(miss)}
        if present and ctx.legacy_area_by_key is not None:
            dmax = max(gk[k].get("disp", 0.0) for k in present)
            detail["symdiff_m_max"] = _jsonable(dmax)
            verdict_ = _scalar_under(ctx, "vector.symdiff_ratio", dmax)
            if verdict_ == "FAIL":
                codes.append("GEOMETRY_SYMDIFF_EXCEEDED")
            elif verdict_ == "UNSIGNED":
                unsigned = True
        if present:
            h = max(gk[k]["haus"] for k in present)
            detail["hausdorff_m_max"] = _jsonable(h)
            verdict_ = _scalar_under(ctx, "vector.hausdorff_m", h)
            if verdict_ == "FAIL":
                codes.append("GEOMETRY_HAUSDORFF_EXCEEDED")
            elif verdict_ == "UNSIGNED":
                unsigned = True
        c["detail"].update(detail)
        c["value"] = len(codes)
        if codes:
            set_state(c, "FAIL", *codes)
        elif unsigned:
            set_state(c, "NEEDS_DECISION", "TOLERANCE_UNSIGNED")
        else:
            set_state(c, "PASS_EXACT")
        out.append(c)
    return out


def compare_vector(ctx: "PairCtx") -> list[dict]:
    import pyproj
    ll, nl, info = _resolve_layers(ctx)
    checks: list[dict] = []
    lc = new_check("vector.layer", 1, gate=True)
    lc.update(metric="layer", legacy=ll, new=info["new_layer"])
    lc["detail"] = info
    if nl is None:
        set_state(lc, "FAIL", "SCHEMA_MISMATCH")
        lc["detail"]["error"] = f"layer {info['new_layer']!r} not in the new file ({info['new_layers']})"
        checks.append(lc)
        return checks
    set_state(lc, "PASS_EXACT")
    checks.append(lc)
    L = _read_vector(ctx.legacy, ll, SIDE_LEGACY)
    N = _read_vector(ctx.new, nl, SIDE_NEW)
    gate = None
    if L.geoms is not None:
        cc = new_check("vector.crs", 2, gate=True)
        cc["legacy"] = L.crs.to_string() if L.crs is not None else None
        cc["new"] = N.crs.to_string() if N.crs is not None else None
        if N.geoms is None:
            set_state(cc, "FAIL", "GEOMETRY_TYPE_MISMATCH")
        elif L.crs is None or N.crs is None:
            set_state(cc, "FAIL", "CRS_MISSING")
        elif pyproj.CRS(L.crs).equals(pyproj.CRS(N.crs)):
            set_state(cc, "PASS_EXACT")
        else:
            set_state(cc, "FAIL", "CRS_MISMATCH")
        checks.append(cc)
        gt = new_check("vector.geometry_type", 3, gate=True)
        if N.geoms is not None:
            fl = sorted({_geom_family(g, True) for g in L.geoms} - {"none"})
            fn = sorted({_geom_family(g, True) for g in N.geoms} - {"none"})
            gt.update(metric="geometry_families", legacy=fl, new=fn)
            gt["detail"]["normalisation"] = "multi_equivalence (default)"
            gt["detail"]["legacy_counts"] = dict(collections.Counter(_geom_family(g, False) for g in L.geoms))
            gt["detail"]["new_counts"] = dict(collections.Counter(_geom_family(g, False) for g in N.geoms))
            set_state(gt, *(("PASS_EXACT",) if (fl == fn or not fn and not fl or not fl)
                            else ("FAIL", "GEOMETRY_TYPE_MISMATCH")))
            if not fn and fl:
                set_state(gt, "PASS_EXACT") if N.nrows == 0 else set_state(gt, "FAIL", "GEOMETRY_TYPE_MISMATCH")
        else:
            set_state(gt, "FAIL", "GEOMETRY_TYPE_MISMATCH")
        checks.append(gt)
        bad = [c["id"] for c in (cc, gt) if c["state"] != "PASS_EXACT"]
        gate = bad[0] if bad else None
    fchecks, m = _compare_frames(ctx, L, N, ctx.prefix_override or "vector", 4)
    checks.extend(fchecks)
    if L.geoms is not None:
        checks.extend(_geometry_checks(ctx, L, N, m, gate, 60))
    checks.extend(_branch_checks(ctx, L, checks, 70))

    def values_for(chk: str):
        return ctx.attr_values.get(chk)
    checks.extend(kd_change_checks(ctx.kdb, 80, values_for))
    checks.extend(_invariant_checks(ctx, 90))
    return checks


def compare_table(ctx: "PairCtx") -> list[dict]:
    p = ctx.legacy.lower()
    if p.endswith(".gpkg"):
        ctx.prefix_override = "table"
        return compare_vector(ctx)
    checks: list[dict] = []
    L = _read_table(ctx, ctx.legacy, SIDE_LEGACY)
    lay = new_check("table.layout", 1, gate=True)
    lay["legacy"] = L.layout
    try:
        N = _read_table(ctx, ctx.new, SIDE_NEW)
    except _SchemaMissing as e:
        set_state(lay, "FAIL", "SCHEMA_MISMATCH")
        lay["detail"]["error"] = str(e)
        return [lay]
    lay["new"] = N.layout
    set_state(lay, "PASS_EXACT")
    checks.append(lay)
    fchecks, _ = _compare_frames(ctx, L, N, "table", 2)
    checks.extend(fchecks)
    checks.extend(kd_change_checks(ctx.kdb, 80, lambda chk: ctx.attr_values.get(chk)))
    checks.extend(_invariant_checks(ctx, 90))
    return checks


# =================================================================================================================
# Invariants (06 §4c): evaluated on legacy first
# =================================================================================================================
def _ds_col(ref: str) -> tuple[str, str | None]:
    if ":" in ref:
        a, b = ref.split(":", 1)
        return a, b
    return ref, None


def _side_values(path: str, layer: str | None, col: str | None, side: str):
    """Values of a raster (valid cells) or a vector/table column."""
    p = str(path).lower()
    if p.endswith((".tif", ".tiff", ".vrt", ".img", ".asc")) and col is None:
        return "raster", _open_raster(path, side)
    if p.endswith((".csv", ".xlsx", ".xlsm")):
        raise EvaluatorFault("CHECKER_EXCEPTION", "table invariants need a gpkg/vector column")
    fr = _read_vector(path, layer or _list_layers(path, side)[0], side)
    return "column", fr


def _invariant_eval(ctx: "PairCtx", inv: dict, side: str, rule: Rule) -> tuple[bool, float | None, dict]:
    kind = inv.get("kind")
    resolved = inv.get("resolved") or {}

    def paths_for(role: str):
        ref = inv.get(role)
        dsid, col = _ds_col(str(ref)) if ref is not None else (None, None)
        r = resolved.get(role)
        if r is None and dsid == ctx.ds_id:
            r = {"legacy": {"path": ctx.legacy, "layer": ctx.spec.get("layer")},
                 "new": {"path": ctx.new, "layer": ctx.spec.get("new_layer") or ctx.spec.get("layer")}}
        if r is None:
            raise EvaluatorFault("CHECKER_EXCEPTION", f"invariant {inv.get('id')}: {role} {ref} not resolved")
        s = r.get(side) or r
        return s["path"], s.get("layer"), col

    if kind == "min_ge":
        path, layer, col = paths_for("lhs")
        kind_, obj = _side_values(path, layer, col, side)
        bound = float(inv.get("rhs"))
        if kind_ == "raster":
            with obj:
                acc = _BandAccum()
                for win in _windows(obj.height, obj.width):
                    a, m = _read_valid(obj, 1, win)
                    acc.update(a, m)
                mn = acc.finalise()["min"]
        else:
            vals = [float(v) for v in obj.data[col] if not _is_null(v)]
            mn = min(vals) if vals else None
        ok = mn is None or mn >= bound
        return ok, (bound - mn) if (mn is not None and mn < bound) else 0.0, {"min": mn, "bound": bound}
    if kind == "total_equals":
        tots = []
        for role in ("lhs", "rhs"):
            path, layer, col = paths_for(role)
            kind_, obj = _side_values(path, layer, col, side)
            if kind_ == "raster":
                with obj:
                    tots.append(_exact_total(obj, 1)[0])
            else:
                tots.append(math.fsum(float(v) for v in obj.data[col] if not _is_null(v)))
        a, b = tots
        rel = abs(a - b) / abs(b) if b else (0.0 if a == b else math.inf)
        return rule.close(b, a), rel, {"lhs": a, "rhs": b}
    if kind == "zone_sum_equals":
        path, layer, _ = paths_for("lhs")
        # CONTRACTS §4 invariants carry no zones field: default to the lhs dataset's declared compare.zones
        zref = str(inv.get("zones") or (ctx.spec.get("compare") or {}).get("zones"))
        zds, zkey = _ds_col(zref)
        zr = resolved.get("zones") or (ctx.spec.get("zones") if (ctx.spec.get("compare") or {}).get("zones") == zref
                                       else None)
        if inv.get("rule") not in ("centre", "all_touched"):
            raise EvaluatorFault("CHECKER_EXCEPTION", f"invariant {inv.get('id')}: rule must be centre|all_touched")
        if zr is None:
            raise EvaluatorFault("CHECKER_EXCEPTION", f"invariant {inv.get('id')}: zones {zref} not resolved")
        zr = zr.get(side, zr) if isinstance(zr, dict) and side in zr else zr
        rpath, rlayer, rcol = paths_for("rhs")
        rkey = inv.get("rhs_key") or zkey
        ds = _open_raster(path, side)
        with ds:
            shapes, keys = _zones_frame({"path": zr["path"], "layer": zr.get("layer"), "key": zkey}, ds.crs)
            sums, _ = _zone_sums(ds, 1, shapes, len(keys), inv.get("rule") == "all_touched")
        fr = _read_vector(rpath, rlayer or _list_layers(rpath, side)[0], side)
        if rkey not in fr.data or rcol not in fr.data:
            raise EvaluatorFault("CHECKER_EXCEPTION", f"invariant {inv.get('id')}: {rkey}/{rcol} not in {rpath}")
        rhs: dict = {}
        for k, v in zip(fr.data[rkey], fr.data[rcol]):
            kk = _canon_key(k)
            if _is_null(v) or (kk in rhs and rhs[kk] is None):
                rhs[kk] = None                            # a null is never 0 (no fillna): the zone cannot satisfy it
            else:
                rhs[kk] = (rhs.get(kk) or 0.0) + float(v)
        worst_rel, table, ok = 0.0, [], True
        for i, k in enumerate(keys, start=1):
            kk = _canon_key(k)
            b = rhs.get(kk)
            a = sums[i]
            if b is None:                                 # missing or null on the right-hand side
                ok = False
                table.append({"zone": k, "raster": a, "table": None, "rel": None})
                continue
            rel = abs(a - b) / abs(b) if b else (0.0 if a == b else math.inf)
            worst_rel = max(worst_rel, rel)
            ok = ok and rule.close(b, a)
            table.append({"zone": k, "raster": a, "table": b, "rel": rel})
        return ok, worst_rel, {"table": table}
    raise EvaluatorFault("CHECKER_EXCEPTION", f"invariant kind {kind!r} not supported "
                                              "(zone_sum_equals | total_equals | min_ge)")


def _invariant_checks(ctx: "PairCtx", order: int) -> list[dict]:
    out = []
    for inv in ctx.spec.get("invariants") or []:
        lhs_ds = _ds_col(str(inv.get("lhs")))[0]
        if lhs_ds != ctx.ds_id:
            continue
        c = new_check(f"invariant:{inv.get('id')}", order)
        c["detail"]["invariant"] = {k: v for k, v in inv.items() if k != "resolved"}
        if inv.get("unavailable"):
            set_state(c, "NOT_COMPARED", "UPSTREAM_BLOCKED")
            c["detail"]["why"] = inv["unavailable"]
            out.append(c)
            continue
        res = ctx.book.lookup_id(inv.get("tolerance_ref"))
        if res.error:
            set_state(c, "EVALUATOR_ERROR", res.error[0])
            out.append(c)
            continue
        rules = [res.rule] + ([res.proposed] if res.proposed is not None else [])
        lres = [_invariant_eval(ctx, inv, SIDE_LEGACY, r) for r in rules]
        nres = [_invariant_eval(ctx, inv, SIDE_NEW, r) for r in rules]
        c.update(metric="max_rel" if inv.get("kind") != "min_ge" else "below_bound", value=nres[0][1],
                 legacy=lres[0][1], new=nres[0][1], tolerance_id=res.rule.id or "default.invariant.exact")
        c["detail"]["legacy"], c["detail"]["new"] = lres[0][2], nres[0][2]
        for n_ in res.notes:
            c["reason_codes"].append(n_)
        if not lres[0][0]:
            if len(rules) > 1 and lres[1][0] and (nres[1][0]):
                set_state(c, "NEEDS_DECISION", "TOLERANCE_UNSIGNED")
                c["proposed"] = res.proposed_info
            else:
                set_state(c, "NEEDS_DECISION", "LEGACY_DEFECT_SUSPECTED")
        elif not nres[0][0]:
            if len(rules) > 1 and nres[1][0]:
                set_state(c, "NEEDS_DECISION", "TOLERANCE_UNSIGNED")
                c["proposed"] = res.proposed_info
            else:
                set_state(c, "FAIL", "INVARIANT_VIOLATED")
        else:
            set_state(c, "PASS_EXACT" if (nres[0][1] in (0, 0.0, None)) else "PASS_WITHIN_TOLERANCE")
        out.append(c)
    return out


# =================================================================================================================
# One pair (shared by run mode and standalone mode)
# =================================================================================================================
@dataclass
class PairCtx:
    legacy: str
    new: str
    spec: dict
    book: ToleranceBook
    kdb: KDBook
    ds_id: str
    mismatches_path: str | None = None
    mismatch_rows: list = field(default_factory=list)
    band_stats: dict = field(default_factory=dict)
    value_arrays_hint: dict = field(default_factory=dict)
    attr_values: dict = field(default_factory=dict)
    totals: tuple | None = None
    gate_name: str | None = None
    matched_keys: list | None = None
    raster_pair_arrays: Any = None
    prefix_override: str | None = None
    key_index_legacy: dict | None = None
    keys_missing: list | None = None
    geom_by_key: dict | None = None
    legacy_area_by_key: dict | None = None
    legacy_perimeter_by_key: dict | None = None

    def legacy_integral(self, lds, band) -> bool:
        return _exact_total(lds, band)[2]


def _family(spec: dict, legacy: str) -> str:
    k = spec.get("kind") or spec.get("family")
    if k in ("raster", "vector", "table"):
        return k
    return k or "unknown"


def _load_tolerance_doc(x):
    import yaml
    if x is None:
        return None
    if isinstance(x, dict):
        return x
    p = Path(x)
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else None


def _load_kds(x) -> list[dict]:
    if x is None:
        return []
    if isinstance(x, list):
        return x
    p = Path(x)
    if p.exists():
        return parse_known_differences(p.read_text(encoding="utf-8"))[0]
    return []


def _base_result(spec: dict, legacy: str, new: str, family: str) -> dict:
    return {"schema": RESULT_SCHEMA, "comparator": comparator_block(), "run_id": None, "mode": "standalone",
            "isolated_step": None, "contract_hash": None, "dataset_id": spec.get("dataset_id"),
            "step_id": spec.get("step_id"), "family": family, "role": spec.get("role"),
            "measure": (spec.get("compare") or {}).get("measure"),
            "legacy": {"path": legacy, "layer": spec.get("layer"), "sha256": None, "content_sha256": None,
                       "pinned_sha256": (spec.get("pinned") or {}).get("sha256")},
            "new": {"path": new, "layer": spec.get("new_layer") or spec.get("layer"), "sha256": None,
                    "written_by_step": spec.get("written_by_step")},
            "pairing": spec.get("pairing") or {"method": "declared_spec", "evidence": [], "state": "standalone"},
            "verdict": None, "reason_codes": [], "diagnostic_codes": [], "checks": [], "within_tolerance": [],
            "not_checked": [], "known_differences_applied": [], "known_differences_ignored": [],
            "tolerance_proposals_ignored": [],
            "excluded": False, "refs": {"out": None, "ref": None, "tol": list(spec.get("tol_refs") or [])},
            "env": _env_block(), "plots": f"plots/{spec.get('dataset_id')}/plots.json", "mismatches": None,
            "created_at": ac.utc_now_iso()}


def finalise(result: dict, *, verdict: str | None = None, codes: Iterable[str] = ()) -> dict:
    checks = result["checks"]
    by_construction = result.pop("_by_construction", {}) or {}
    for c in checks:
        if c["id"] in by_construction:
            c["by_construction"] = True
            c["by_construction_param"] = by_construction[c["id"]]
    if verdict is None:
        verdict = verdict_of_checks(checks) or "NOT_COMPARED"
    result["verdict"] = verdict
    rc: list[str] = list(codes)
    dc: list[str] = []
    for c in sorted(checks, key=lambda c: (c.get("order", 0), c["id"])):
        if c.get("kd_check"):
            continue
        if c["state"] not in PASS and counts_for_verdict(c):
            for x in c["reason_codes"]:
                if x not in rc:
                    rc.append(x)
        else:
            for x in c["reason_codes"]:
                if x in ("TOLERANCE_UNJUSTIFIED",) and x not in rc:
                    rc.append(x)
        for x in c.get("diagnostic_codes", []):
            if x not in dc:
                dc.append(x)
    for x in result.get("diagnostic_codes", []):
        if x not in dc:
            dc.append(x)
    result["reason_codes"] = rc
    result["diagnostic_codes"] = dc
    result["within_tolerance"] = [{"check": c["id"], "metric": c.get("metric"), "value": c.get("value"),
                                   "tolerance_id": c.get("tolerance_id"), "kd": c.get("kd") or []}
                                  for c in checks if c["state"] == "PASS_WITHIN_TOLERANCE" and not c.get("kd_check")]
    result["not_checked"] = [f"{c['id']} (GATED by {c.get('gated_by')})" for c in checks
                             if c["state"] == "NOT_COMPARED" and c["reason_codes"] == ["GATED"]]
    nulls: list[str] = []
    clean = _sanitise(result, nulls)
    if nulls:
        clean["null_values"] = [{"path": p, "reason": "NOT_FINITE"} for p in sorted(set(nulls))]
    body = {k: v for k, v in clean.items() if k not in ("result_id",)}
    clean["result_id"] = ac.object_id("cmp", body)
    return clean


def compare_pair(legacy: str, new: str, spec: dict, *, mismatches_path: str | None = None,
                 book: ToleranceBook | None = None, kds: list | None = None, base: dict | None = None) -> dict:
    """Compare one legacy file with one new file under `spec` (standalone or from run mode). Never raises for a
    data problem: every exception becomes EVALUATOR_ERROR/CHECKER_EXCEPTION with the traceback."""
    family = _family(spec, legacy)
    result = base or _base_result(spec, legacy, new, family)
    result["_by_construction"] = spec.get("by_construction") or {}
    try:
        started = parse_time(spec.get("run_started_at"), end_of_day=False) or _dt.datetime.now(_dt.timezone.utc)
        if book is None:
            book = ToleranceBook(_load_tolerance_doc(spec.get("tolerances")), started, spec.get("measure_crs"),
                                 proposals=_load_tolerance_doc(spec.get("tolerances_proposed")))
        kd_list = kds if kds is not None else _load_kds(spec.get("known_differences"))
        kdb = KDBook(kd_list, started, spec.get("dataset_id") or "", spec.get("role"))
        result["tolerance_proposals_ignored"] = [n for n in book.proposal_notes
                                                 if n.get("dataset") in (None, spec.get("dataset_id"))]
        result["known_differences_ignored"] = kdb.ignored
        if spec.get("role") == "input" or spec.get("legacy_role") == "input":
            result["checks"] = []
            result["detail"] = "a legacy input is never a reference (REFERENCE_NOT_INDEPENDENT)"
            return finalise(result, verdict="EVALUATOR_ERROR", codes=["REFERENCE_NOT_INDEPENDENT"])
        lp = Path(legacy)
        if not lp.exists():
            return finalise(result, verdict="EVALUATOR_ERROR", codes=["REFERENCE_HASH_MISMATCH"]) \
                if spec.get("pinned") else finalise(result, verdict="EVALUATOR_ERROR", codes=["CHECKER_EXCEPTION"])
        result["legacy"]["sha256"] = _sha_any(lp)
        result["legacy"]["content_sha256"] = _content_sha(lp)
        result["refs"]["ref"] = "ref:" + result["legacy"]["sha256"]
        pin = spec.get("pinned")
        if pin and pin.get("sha256") and pin["sha256"] != result["legacy"]["sha256"] and not (
                pin.get("content_sha256") and pin["content_sha256"] == result["legacy"]["content_sha256"]):
            return finalise(result, verdict="EVALUATOR_ERROR", codes=["REFERENCE_HASH_MISMATCH"])
        ex = kdb.exclusion()
        if ex is not None:
            kdb.mark(ex, "*")
            result["excluded"] = True
            result["known_differences_applied"] = list(kdb.applied.values())
            return finalise(result, verdict="NOT_COMPARED", codes=[ex["rule"]["reason"]])
        np_ = Path(new)
        if not np_.exists():
            return finalise(result, verdict="FAIL", codes=["OUTPUT_MISSING"])
        result["new"]["sha256"] = _sha_any(np_)
        result["refs"]["out"] = "out:" + result["new"]["sha256"]
        identical = result["new"]["sha256"] == result["legacy"]["sha256"] or (
            result["legacy"]["content_sha256"] is not None
            and _content_sha(np_) == result["legacy"]["content_sha256"])
        if identical:
            result["diagnostic_codes"] = ["CANDIDATE_IDENTICAL_TO_REFERENCE"]
        if family not in ("raster", "vector", "table"):
            result["excluded"] = False
            return finalise(result, verdict="NOT_COMPARED", codes=["FAMILY_UNSUPPORTED"])
        ctx = PairCtx(legacy=str(legacy), new=str(new), spec=spec, book=book, kdb=kdb,
                      ds_id=spec.get("dataset_id") or "", mismatches_path=mismatches_path)
        try:
            if family == "raster":
                checks = compare_raster(ctx)
            elif family == "vector":
                checks = compare_vector(ctx)
            else:
                checks = compare_table(ctx)
        except OutputUnreadable as e:
            result["detail"] = str(e)
            return finalise(result, verdict="FAIL", codes=["OUTPUT_UNREADABLE"])
        except EvaluatorFault as e:
            result["detail"] = str(e)
            return finalise(result, verdict="EVALUATOR_ERROR", codes=[e.code])
        result["checks"] = checks
        result["known_differences_applied"] = sorted(kdb.applied.values(), key=lambda x: x["id"]) + \
            [{"id": kd["id"], "kind": kd["kind"], "checks": kd.get("checks") or [], "owner": kd.get("owner"),
              "decision_ref": kd.get("decision_ref")} for kd in kdb.preserved()]
        if ctx.band_stats:
            result["stats"] = {f"band{b}": s for b, s in ctx.band_stats.items()}
        if mismatches_path:
            mp = Path(mismatches_path)
            if mp.exists():
                mp.unlink()
            if ctx.mismatch_rows:
                _write_parquet(mp, ctx.mismatch_rows)
                result["mismatches"] = str(mismatches_path) if spec.get("_abs_paths_ok", True) else mp.name
        return finalise(result)
    except Exception as e:                                # any checker exception → EVALUATOR_ERROR, never hidden
        result["checks"] = result.get("checks") or []
        result["traceback"] = traceback.format_exc()
        result["detail"] = f"{type(e).__name__}: {e}"
        return finalise(result, verdict="EVALUATOR_ERROR", codes=["CHECKER_EXCEPTION"])


def _write_parquet(path: Path, rows: list[dict]) -> None:
    import pyarrow as pa
    import pyarrow.parquet as pq
    cols = sorted({k for r in rows for k in r})
    data = {}
    for c in cols:
        vals = [r.get(c) for r in rows]
        if all(v is None or isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals) and \
                any(isinstance(v, float) for v in vals):
            data[c] = pa.array([None if v is None else float(v) for v in vals], type=pa.float64())
        elif all(v is None or isinstance(v, int) and not isinstance(v, bool) for v in vals):
            data[c] = pa.array(vals, type=pa.int64())
        else:
            data[c] = pa.array([None if v is None else str(v) for v in vals], type=pa.string())
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    pq.write_table(pa.table(data), tmp)
    os.replace(tmp, path)


# =================================================================================================================
# Run mode: frozen contract, live checks, pairing from lineage (06 §4a; DESIGN §4.5-4.6; CONTRACTS §6-7)
# =================================================================================================================
def _read_text(p: Path) -> str | None:
    # bytes as stored: read_text() turns CRLF into LF, and the sha of that never matches the file's pin (a CRLF
    # tolerances.yaml written on Windows was reported as TOLERANCE_CHANGED_AFTER_RUN)
    return p.read_bytes().decode("utf-8") if p.exists() else None


class RunContext:
    def __init__(self, run_dir: str | os.PathLike):
        import yaml
        self.run = Path(run_dir).resolve()
        mpath = self.run / "manifest.json"
        if not mpath.exists():
            raise SetupError(f"EVALUATOR_ERROR: {mpath} missing (not a run directory)")
        self.manifest = ac.load_json(mpath)
        self.run_id = self.manifest.get("run_id") or self.run.name
        self.snap = self.run / "snapshot"
        gpath = self.snap / "graph.json"
        if not gpath.exists():
            raise SetupError(f"EVALUATOR_ERROR: {gpath} missing (the contract is frozen from the snapshot)")
        self.graph = ac.load_json(gpath)
        self.inventory = ac.load_json(self.snap / "inventory.json") if (self.snap / "inventory.json").exists() else None
        self.tol_text = _read_text(self.snap / "tolerances.yaml")
        self.kd_text = _read_text(self.snap / "known-differences.md")
        self.params_text = _read_text(self.snap / "params.yaml")
        self.tol_doc = yaml.safe_load(self.tol_text) if self.tol_text else None
        prop_text = _read_text(self.snap / "tolerances.proposed.yaml")       # never part of the contract
        self.prop_doc = yaml.safe_load(prop_text) if prop_text else None
        self.kds, self.kd_errors = parse_known_differences(self.kd_text)
        self.params = yaml.safe_load(self.params_text) if self.params_text else None
        self.chain = self.manifest.get("chain") or "chained"
        started = self.manifest.get("started_at")
        self.started_at = parse_time(started, end_of_day=False) if started else None
        mig = self.run.parent.parent
        self.migration = mig if (self.run.parent.name == "runs" and (mig / "graph.json").exists()) else None
        self.project_root = self.migration.parent if self.migration else None
        lr = (self.graph.get("project") or {}).get("legacy_root")
        self.legacy_root = (self.migration / lr).resolve() if (self.migration and lr) else None
        self.datasets = {d["id"]: d for d in self.graph.get("datasets", [])}
        self.steps = {s["id"]: s for s in self.graph.get("steps", [])}
        self.m_outputs = {o.get("dataset_id"): o for o in self.manifest.get("outputs", []) if o.get("dataset_id")}
        self.m_steps = {s.get("step_id") or s.get("id"): s for s in self.manifest.get("steps", [])}

    def param(self, pid: str):
        for p in (self.params or {}).get("params", []) or []:
            if isinstance(p, dict) and p.get("id") == pid:
                return p.get("value")
        return None

    def measure_crs(self) -> Any:
        pids = {((d.get("compare") or {}).get("measure_crs_param")) for d in self.datasets.values()} - {None}
        pid = sorted(pids)[0] if pids else "p.measure_crs"
        return self.param(pid)

    def rel_project(self, p: Path) -> str:
        p = Path(p).resolve()
        if self.project_root:
            try:
                return p.relative_to(self.project_root).as_posix()
            except ValueError:
                return p.as_posix()
        return p.as_posix()


def pairing_projection(graph: dict) -> dict:
    ds = []
    for d in sorted(graph.get("datasets", []), key=lambda d: d["id"]):
        lg = d.get("legacy") or {}
        ds.append({"id": d["id"], "kind": d.get("kind"), "role": d.get("role"),
                   "legacy": {"path": lg.get("path"), "layer": lg.get("layer"), "sheet": lg.get("sheet"),
                              "exists": lg.get("exists")},
                   "new": d.get("new"), "produced_by": (d.get("produced_by") or {}).get("step"),
                   "reference": d.get("reference"), "legacy_in_memory": d.get("legacy_in_memory", False),
                   "no_reference_reason": d.get("no_reference_reason"), "compare": d.get("compare") or {},
                   "waiver": (d.get("verify") or {}).get("waiver")})
    steps = [{"id": s["id"], "scope": s.get("scope"), "kind": s.get("kind")}
             for s in sorted(graph.get("steps", []), key=lambda s: s["id"])]
    edges = []
    for e in sorted(graph.get("edges", []), key=lambda e: (e.get("from", ""), e.get("to", ""), e.get("kind", ""))):
        if e.get("kind") in ("writes", "reads_write", "reads", "reads_grid", "equivalent_to"):
            edges.append({"from": e.get("from"), "to": e.get("to"), "kind": e.get("kind"), "state": e.get("state"),
                          "evidence": [{"state": x.get("state"), "ref": x.get("ref")} for x in e.get("evidence") or []]})
    return {"datasets": ds, "steps": steps, "edges": edges, "invariants": graph.get("invariants") or [],
            "legacy_output_dirs": graph.get("legacy_output_dirs") or [],
            "legacy_root": (graph.get("project") or {}).get("legacy_root")}


def build_contract(rc: RunContext) -> dict:
    started = rc.started_at or _dt.datetime.now(_dt.timezone.utc)
    book = ToleranceBook(rc.tol_doc, started, rc.measure_crs())
    proj = pairing_projection(rc.graph)
    refs = {}
    inv_files = {f.get("path"): f for f in (rc.inventory or {}).get("files", [])}
    wanted = set()
    for d in proj["datasets"]:
        if d["legacy"]["path"]:
            wanted.add(d["legacy"]["path"])
    for p in sorted(wanted):
        f = inv_files.get(p)
        refs[p] = {"sha256": (f or {}).get("sha256"), "content_sha256": (f or {}).get("content_sha256")}
    kd_entries = []
    for kd in rc.kds:
        k = {x: y for x, y in kd.items() if x not in ("_line", "_head")}
        kd_entries.append(_jsonable(k))
    body = {
        "schema": CONTRACT_SCHEMA, "run_id": rc.run_id, "chain": rc.chain, "comparator": comparator_block(),
        "run_started_at": started.isoformat(),
        "tolerances": {"sha256": ac.sha256_bytes(rc.tol_text.encode()) if rc.tol_text is not None else None,
                       "version": (rc.tol_doc or {}).get("version") if isinstance(rc.tol_doc, dict) else None,
                       "defaults": BUILTIN_DEFAULTS, "defaults_ignored": book.defaults_ignored,
                       "doc_errors": book.doc_errors,
                       "records": [_jsonable(dict(t["record"], status=t["status"], why=t["why"],
                                                  threshold=t["threshold"])) if isinstance(t["record"], dict)
                                   else {"status": t["status"], "why": t["why"]} for t in book.table()]},
        "known_differences": {"sha256": ac.sha256_bytes(rc.kd_text.encode()) if rc.kd_text is not None else None,
                              "entries": kd_entries, "parse_errors": rc.kd_errors},
        "pairing": proj, "references": refs, "params": {"measure_crs": _jsonable(rc.measure_crs())},
        "graph_contract_sha256": ac.graph_contract_sha256(rc.graph),
    }
    body = _sanitise(body, [])
    return dict(body, contract_hash="sha256:" + ac.sha256_bytes(ac.cjson(body)), created_at=ac.utc_now_iso())


def _contract_body_hash(c: dict) -> str:
    body = {k: v for k, v in c.items() if k not in ("contract_hash", "created_at")}
    return "sha256:" + ac.sha256_bytes(ac.cjson(body))


def global_checks(rc: RunContext, contract: dict) -> list[tuple[str, str]]:
    """Live vs snapshot, snapshot vs manifest pins, contract re-hash (D6-11). Each item = (reason code, detail)."""
    errs: list[tuple[str, str]] = []
    if rc.migration is None:
        errs.append(("CHECKER_EXCEPTION", f"{rc.run} is not under <migration>/runs/ (live contract unavailable)"))
        return errs
    if rc.started_at is None:
        errs.append(("CHECKER_EXCEPTION", "manifest.started_at missing"))
    live_tol = _read_text(rc.migration / "tolerances.yaml")
    live_kd = _read_text(rc.migration / "known-differences.md")
    if live_tol != rc.tol_text:
        errs.append(("TOLERANCE_CHANGED_AFTER_RUN", "live tolerances.yaml differs from the run snapshot"))
    if live_kd != rc.kd_text:
        errs.append(("TOLERANCE_CHANGED_AFTER_RUN", "live known-differences.md differs from the run snapshot"))
    live_graph = ac.load_json(rc.migration / "graph.json")
    if ac.cjson(pairing_projection(live_graph)) != ac.cjson(pairing_projection(rc.graph)):
        errs.append(("GRAPH_CHANGED_AFTER_RUN", "pairing projection of the live graph differs from the snapshot"))
    for key, text, code in (("tolerances", rc.tol_text, "TOLERANCE_CHANGED_AFTER_RUN"),
                            ("known_differences", rc.kd_text, "TOLERANCE_CHANGED_AFTER_RUN"),
                            ("graph", _read_text(rc.snap / "graph.json"), "GRAPH_CHANGED_AFTER_RUN")):
        pin = rc.manifest.get(key)
        if not isinstance(pin, dict) or "sha256" not in pin:
            errs.append((code, f"manifest.{key}.sha256 missing: the snapshot cannot be shown unchanged"))
            continue
        actual = ac.sha256_bytes(text.encode()) if text is not None else None
        if pin["sha256"] != actual:
            errs.append((code, f"snapshot {key} differs from the hash pinned in manifest.json"))
    cpath = rc.run / "verify" / "contract.json"
    if cpath.exists():
        old = ac.load_json(cpath)
        if old.get("comparator", {}).get("sha256") != contract["comparator"]["sha256"]:
            errs.append(("CHECKER_HASH_MISMATCH", "compare.py changed since this run's contract was frozen"))
        elif _contract_body_hash(old) != old.get("contract_hash") or old.get("contract_hash") != contract["contract_hash"]:
            errs.append(("TOLERANCE_CHANGED_AFTER_RUN", "verify/contract.json differs from the contract recomputed "
                                                        "from the snapshot"))
    return errs


def _edge_evidence(graph: dict, step: str, ds: str) -> list[str]:
    refs = []
    for e in graph.get("edges", []):
        if e.get("from") == step and e.get("to") == ds and e.get("kind") in ("writes", "reads_write"):
            if e.get("state") in ("rejected", "superseded"):
                continue
            for x in e.get("evidence") or []:
                if x.get("state") not in ("rejected", "superseded") and x.get("ref"):
                    refs.append(x["ref"])
    return sorted(set(refs))


def _upstream_steps(graph: dict) -> dict[str, set]:
    prod = {d["id"]: (d.get("produced_by") or {}).get("step") for d in graph.get("datasets", [])}
    parents: dict[str, set] = collections.defaultdict(set)
    for e in graph.get("edges", []):
        if e.get("kind") in ("reads", "reads_grid", "reads_write") and e.get("to") in {s["id"] for s in graph["steps"]}:
            p = prod.get(e.get("from"))
            if p and p != e["to"]:
                parents[e["to"]].add(p)
    return parents


def _ancestors(parents: dict, s: str) -> set:
    seen, todo = set(), [s]
    while todo:
        x = todo.pop()
        for p in parents.get(x, ()):
            if p not in seen:
                seen.add(p)
                todo.append(p)
    return seen


def _topo(steps: list[str], parents: dict) -> list[str]:
    order, done = [], set()
    pending = sorted(steps)
    while pending:
        progressed = False
        for s in list(pending):
            if all(p in done or p not in steps for p in parents.get(s, ())):
                order.append(s)
                done.add(s)
                pending.remove(s)
                progressed = True
        if not progressed:
            order.extend(sorted(pending))
            break
    return order


def _by_construction(rc: RunContext) -> dict[str, dict]:
    kind_map = {"nodata": "raster.nodata_value", "dtype": "raster.dtype", "crs": "raster.grid.crs",
                "grid": "raster.grid.transform", "transform": "raster.grid.transform", "layer": "vector.layer"}
    out: dict[str, dict] = collections.defaultdict(dict)
    for p in (rc.params or {}).get("params", []) or []:
        if not isinstance(p, dict) or p.get("source") != "reference_metadata":
            continue
        chk = kind_map.get(str(p.get("kind") or "").lower())
        if not chk:
            continue
        dsets = ((p.get("impact") or {}).get("datasets") or []) if isinstance(p.get("impact"), dict) else []
        if not dsets and p.get("node") in rc.steps:
            dsets = [d for d, x in rc.datasets.items() if (x.get("produced_by") or {}).get("step") == p["node"]]
        for d in dsets:
            out[d][chk] = p.get("id")
    return out


def _resolve_ds_paths(rc: RunContext, dsid: str) -> dict | None:
    d = rc.datasets.get(dsid)
    if d is None or rc.legacy_root is None:
        return None
    lg = d.get("legacy") or {}
    leg = {"path": str(rc.legacy_root / lg["path"]), "layer": lg.get("layer")} if lg.get("path") else None
    if d.get("role") == "input":
        return {"legacy": leg, "new": leg}
    m = rc.m_outputs.get(dsid)
    new = {"path": str(rc.run / m["path"]), "layer": (d.get("new") or {}).get("layer")} if m else None
    return {"legacy": leg, "new": new}


def _pair_one(rc: RunContext, dsid: str, contract: dict, book: ToleranceBook, gerrs: list, out_dir: Path,
              bc: dict) -> dict | None:
    d = rc.datasets[dsid]
    role = d.get("role")
    if role == "input":
        return None
    prod = (d.get("produced_by") or {}).get("step")
    step = rc.steps.get(prod) if prod else None
    if step is not None and step.get("scope") in ("drop", "keep_legacy"):
        return None
    if rc.chain == "isolated" and prod not in rc.m_steps:
        return None
    lg = d.get("legacy") or {}
    family = d.get("kind")
    spec = {"dataset_id": dsid, "kind": family, "role": role, "step_id": prod,
            "layer": lg.get("layer"), "new_layer": (d.get("new") or {}).get("layer"),
            "sheet": lg.get("sheet") or (d.get("compare") or {}).get("sheet"),
            "compare": d.get("compare") or {}, "measure_crs": rc.measure_crs(),
            "run_started_at": contract["run_started_at"], "by_construction": bc.get(dsid, {})}
    legacy_abs = str(rc.legacy_root / lg["path"]) if (rc.legacy_root and lg.get("path")) else (lg.get("path") or "")
    m = rc.m_outputs.get(dsid)
    new_abs = str(rc.run / m["path"]) if m else ""
    base = _base_result(spec, rc.rel_project(Path(legacy_abs)) if legacy_abs else None,
                        rc.rel_project(Path(new_abs)) if new_abs else None, family)
    base.update(run_id=rc.run_id, mode=rc.chain, contract_hash=contract["contract_hash"],
                isolated_step=prod if rc.chain == "isolated" else None)
    base["legacy"]["legacy_relpath"] = lg.get("path")
    base["new"]["written_by_step"] = (m or {}).get("written_by_step")
    if rc.tol_text is not None:
        base["refs"]["tol"] = ["tol:" + ac.sha256_bytes(rc.tol_text.encode())]
    base["pairing"] = {"method": (d.get("reference") or {}).get("match_method") or "lineage",
                       "evidence": _edge_evidence(rc.graph, prod, dsid) if prod else [], "state": "unconfirmed"}
    base["_by_construction"] = spec["by_construction"]

    def early(verdict, *codes, detail=None, excluded=False):
        r = dict(base)
        r["excluded"] = excluded
        if detail:
            r["detail"] = detail
        return finalise(r, verdict=verdict, codes=codes)
    if gerrs:
        return early("EVALUATOR_ERROR", *sorted({c for c, _ in gerrs}), detail="; ".join(x for _, x in gerrs))
    if step is not None and step.get("kind") == "bridge":
        return early("NOT_COMPARED", "BRIDGED", detail="bridged datasets are never compared against themselves")
    waiver = (d.get("verify") or {}).get("waiver")
    ref = d.get("reference") or {}
    L_path = lg.get("path") if lg.get("exists", bool(lg.get("path"))) else None
    if ref.get("source") == "equivalent":
        via = rc.datasets.get(ref.get("via_dataset") or "")
        if not via or not ref.get("confirmed_by"):
            return early("NOT_COMPARED", "PAIRING_UNCONFIRMED",
                         detail="reference from an equivalent step needs a confirmed via_dataset")
        vlg = via.get("legacy") or {}
        L_path = vlg.get("path") if vlg.get("exists", bool(vlg.get("path"))) else None
        if not L_path:
            return early("NOT_COMPARED", "NO_REFERENCE", detail=f"equivalent dataset {via['id']} has no legacy file")
        legacy_abs = str(rc.legacy_root / L_path)
        # the reference is the equivalent step's file: its layer/sheet, and its compare settings where this dataset
        # declares none (the target's own legacy layer names a container that never existed, e.g. a gdb item)
        vcmp = dict(via.get("compare") or {})
        for k, v in (d.get("compare") or {}).items():
            if k in ("layer", "sheet"):
                continue
            if v not in (None, [], {}, "", False):
                vcmp[k] = v
        spec["compare"] = vcmp
        spec["layer"] = vlg.get("layer")
        spec["sheet"] = vlg.get("sheet") or vcmp.get("sheet")
        base["legacy"].update(path=rc.rel_project(Path(legacy_abs)), legacy_relpath=L_path, layer=vlg.get("layer"),
                              via_dataset=via["id"])
        base["pairing"].update(method="equivalent", reference_kind="legacy_equivalent",
                               label="reference from equivalent step", via_dataset=via["id"],
                               confirmed_by=ref.get("confirmed_by"))
    elif ref.get("source") == "none":
        L_path = None
    N_decl = (d.get("new") or {}).get("path")
    if L_path and not N_decl:
        if waiver and role != "final":
            return early("NOT_COMPARED", "UNPAIRED_LEGACY", detail=f"waived: {waiver}", excluded=True)
        return early("NOT_COMPARED", "UNPAIRED_LEGACY", detail="legacy dataset has no new output in the graph")
    if N_decl and not L_path:
        return early("NOT_COMPARED", "NO_REFERENCE" if d.get("legacy_in_memory") else "UNPAIRED_NEW")
    if not L_path and not N_decl:
        return None
    ev = base["pairing"]["evidence"]
    human = ref.get("match_method") == "human" and ref.get("confirmed_by")
    if not ev and not human:
        return early("NOT_COMPARED", "PAIRING_UNCONFIRMED",
                     detail=f"no evidence that {prod} writes {L_path} (no stem/position fallback)")
    base["pairing"]["state"] = "confirmed"
    inv_inputs = {x.get("legacy_path") for x in (rc.inventory or {}).get("datasets", []) if x.get("role") == "input"}
    if L_path in inv_inputs:
        # A legacy input file may be a reference only for a ported manual step: the graph dataset is produced_by a
        # step with scope migrate (so its graph role is not input) and no graph *input* dataset owns the same file
        # (the hand-made artefact IS the reference, DESIGN §8 / 02 T-pair-2). Anything else stays refused.
        owned_as_input = [x["id"] for x in rc.graph.get("datasets", []) if x.get("role") == "input"
                          and (x.get("legacy") or {}).get("path") == L_path]
        ported_manual = step is not None and step.get("scope") == "migrate" and role != "input" and not owned_as_input
        if not ported_manual:
            return early("EVALUATOR_ERROR", "REFERENCE_NOT_INDEPENDENT",
                         detail=f"{L_path} is a legacy input; inputs are never references"
                                + (f" (owned by input dataset {owned_as_input})" if owned_as_input else ""))
        base["pairing"]["reference_kind"] = "manual_step_artefact"
    pin = contract["references"].get(L_path) or {}
    lpath = Path(legacy_abs)
    if not pin.get("sha256"):
        return early("EVALUATOR_ERROR", "REFERENCE_HASH_MISMATCH", detail=f"{L_path} has no pin in the inventory")
    if not lpath.exists():
        return early("EVALUATOR_ERROR", "REFERENCE_HASH_MISMATCH", detail=f"{L_path} missing")
    cur = _sha_any(lpath)
    if cur != pin["sha256"] and not (pin.get("content_sha256") and _content_sha(lpath) == pin["content_sha256"]):
        base["legacy"]["sha256"] = cur
        return early("EVALUATOR_ERROR", "REFERENCE_HASH_MISMATCH",
                     detail=f"{L_path} changed since the scan (pinned {pin['sha256'][:12]}…)")
    base["legacy"]["pinned_sha256"] = pin["sha256"]
    if m is None:
        ms = rc.m_steps.get(prod)
        if ms is None or ms.get("status") in ("blocked", "cancelled", "not_run", "skipped"):
            return early("BLOCKED", "UPSTREAM_BLOCKED", detail=f"step {prod} did not run")
        return early("FAIL", "OUTPUT_MISSING", detail=f"step {prod} ended {ms.get('status')} without this output")
    npath = rc.run / m["path"]
    if not npath.exists():
        return early("FAIL", "OUTPUT_MISSING", detail=f"{m['path']} listed in the manifest but absent")
    if m.get("written_by_step") != prod:
        return early("FAIL", "LINEAGE_MISMATCH",
                     detail=f"written by {m.get('written_by_step')}, graph says produced_by {prod}")
    ms = rc.m_steps.get(prod) or {}
    ledger = ms.get("writes_ledger")
    lpaths = {(x if isinstance(x, str) else (x or {}).get("path")) for x in (ledger or [])}
    if m["path"] not in lpaths:
        return early("FAIL", "OUTPUT_NOT_PRODUCED_BY_RUN",
                     detail=f"{m['path']} is not in the write ledger of {prod}")
    if m.get("sha256") and _sha_any(npath) != m["sha256"]:
        return early("FAIL", "OUTPUT_NOT_PRODUCED_BY_RUN",
                     detail=f"{m['path']} changed after the run (sha differs from manifest)")
    # ---- zones and invariants resolved from the graph (zones use the legacy zone file on both sides)
    zones = (d.get("compare") or {}).get("zones")
    if zones:
        zds, zkey = _ds_col(zones)
        zp = _resolve_ds_paths(rc, zds)
        if zp and zp.get("legacy"):
            spec["zones"] = {"path": zp["legacy"]["path"], "layer": zp["legacy"].get("layer"), "key": zkey}
    invs = []
    for inv in rc.graph.get("invariants") or []:
        if _ds_col(str(inv.get("lhs")))[0] != dsid:
            continue
        inv = dict(inv)
        resolved, missing = {}, []
        for role_ in ("lhs", "rhs"):
            ref_ = inv.get(role_)
            if not isinstance(ref_, str) or not ref_.startswith("ds."):
                continue
            rp = _resolve_ds_paths(rc, _ds_col(ref_)[0])
            if not rp or not rp.get("legacy") or not rp.get("new"):
                missing.append(ref_)
            else:
                resolved[role_] = rp
        if inv.get("zones"):
            zp = _resolve_ds_paths(rc, _ds_col(str(inv["zones"]))[0])
            if zp and zp.get("legacy"):
                resolved["zones"] = zp["legacy"]
            else:
                missing.append(inv["zones"])
        inv["resolved"] = resolved
        if missing:
            inv["unavailable"] = f"not available in this run: {missing}"
        invs.append(inv)
    spec["invariants"] = invs
    spec["pinned"] = pin
    mpath = out_dir / f"{dsid}.mismatches.parquet"
    base["new"]["sha256"] = m.get("sha256")
    r = compare_pair(legacy_abs, str(npath), spec, mismatches_path=str(mpath), book=book, kds=rc.kds, base=base)
    if r.get("mismatches"):
        r["mismatches"] = rc.rel_project(mpath)
        r["result_id"] = ac.object_id("cmp", {k: v for k, v in r.items() if k != "result_id"})
    r["pairing"]["state"] = "confirmed"
    r["result_id"] = ac.object_id("cmp", {k: v for k, v in r.items() if k != "result_id"})
    return r


def sweep(rc: RunContext) -> list[dict]:
    notes = []
    if rc.legacy_root is None:
        return notes
    mentioned = {(d.get("legacy") or {}).get("path") for d in rc.graph.get("datasets", [])}
    seen = set()
    for od in rc.graph.get("legacy_output_dirs") or []:
        base = rc.legacy_root / od
        if not base.exists():
            continue
        for p in sorted(base.rglob("*")):
            rel = p.relative_to(rc.legacy_root).as_posix()
            parts = rel.split("/")
            gdb = next((i for i, x in enumerate(parts) if x.lower().endswith(".gdb")), None)
            if gdb is not None:
                rel = "/".join(parts[: gdb + 1])
            elif not p.is_file():
                continue
            if rel in seen or ac.is_sidecar(rel) or rel in mentioned:
                continue
            seen.add(rel)
            notes.append({"code": "UNPAIRED_LEGACY", "path": rel})
    known = set(rc.datasets)
    for did, o in sorted(rc.m_outputs.items()):
        if did not in known:
            notes.append({"code": "UNPAIRED_NEW", "path": o.get("path")})
    return notes


def run_compare(run_dir: str | os.PathLike, datasets: list[str] | None = None) -> tuple[dict, int]:
    rc = RunContext(run_dir)
    if datasets:
        unknown = [d for d in datasets if d not in rc.datasets]
        if unknown:
            raise SetupError(f"unknown dataset ids {unknown}")
    contract = build_contract(rc)
    gerrs = global_checks(rc, contract)
    vdir = rc.run / "verify"
    vdir.mkdir(exist_ok=True)
    cpath = vdir / "contract.json"
    if not cpath.exists():
        ac.dump_json(cpath, contract)
    out_dir = vdir / "compare"
    out_dir.mkdir(exist_ok=True)
    book = ToleranceBook(rc.tol_doc, rc.started_at or _dt.datetime.now(_dt.timezone.utc), rc.measure_crs(),
                         proposals=rc.prop_doc)
    bc = _by_construction(rc)
    results: dict[str, dict] = {}
    pairing_rows, excluded_scope = [], []
    for dsid in sorted(datasets or rc.datasets):
        for stale in (out_dir / f"{dsid}.json", out_dir / f"{dsid}.mismatches.parquet"):
            if stale.exists():
                stale.unlink()                           # stale results never survive a rerun (D6-17)
        r = _pair_one(rc, dsid, contract, book, gerrs, out_dir, bc)
        if r is None:
            d = rc.datasets[dsid]
            excluded_scope.append({"dataset_id": dsid, "role": d.get("role"),
                                   "why": "input" if d.get("role") == "input" else "out of scope or not in this run"})
            continue
        ac.dump_json(out_dir / f"{dsid}.json", r)
        results[dsid] = r
        pairing_rows.append({"dataset_id": dsid, "verdict": r["verdict"], "reason_codes": r["reason_codes"],
                             "legacy": r["legacy"].get("legacy_relpath"), "new": r["new"].get("path"),
                             "step_id": r.get("step_id"), "family": r.get("family"), "role": r.get("role"),
                             "excluded": r.get("excluded", False)})
    ppath = vdir / "pairing.json"
    prev = ac.load_json(ppath) if (ppath.exists() and datasets) else {"pairs": []}
    keep = [p for p in prev.get("pairs", []) if p["dataset_id"] not in results]
    ac.dump_json(ppath, {"schema": "axis-migrate.verify-pairing.v1", "run_id": rc.run_id,
                         "contract_hash": contract["contract_hash"], "global_errors": [list(x) for x in gerrs],
                         "pairs": sorted(keep + pairing_rows, key=lambda p: p["dataset_id"]),
                         "not_considered": excluded_scope, "notes": sweep(rc),
                         "tolerance_proposals_ignored": book.proposal_notes})
    requested = list(datasets) if datasets else list(results)
    ok = bool(requested) and all(d in results and results[d]["verdict"] in PASS for d in requested)
    return results, (EXIT_PASS if ok else EXIT_NOT_PASS)


# =================================================================================================================
# Aggregation and blame (06 §4d-e; DESIGN §4.6). Writes verify/verdict.json with no rung.
# =================================================================================================================
def _load_results(rc: RunContext) -> dict[str, dict]:
    out = {}
    cdir = rc.run / "verify" / "compare"
    if not cdir.exists():
        return out
    for p in sorted(cdir.glob("*.json")):
        r = ac.load_json(p)
        if r.get("run_id") == rc.run_id and r.get("dataset_id"):
            out[r["dataset_id"]] = r
    return out


def _recompute_state(r: dict) -> str | None:
    checks = r.get("checks") or []
    if not checks:
        return r.get("verdict")
    return verdict_of_checks(checks)


def _visual(rc: RunContext, results: dict) -> tuple[dict, dict]:
    """Receipts of this run (lib/visual_receipt.py is the one acceptance rule): a PNG receipt must cite a plot hash
    of this run, a view receipt the GeoLibre view projects of this run with their current bytes (a project changed
    after the review does not count). Both count the same: a receipt, and a flag at most moves PASS_* down."""
    vdir = rc.run / "verify"
    vr = ac.load_json(vdir / "visual-review.json") if (vdir / "visual-review.json").exists() else {"findings": []}
    per: dict[str, dict] = collections.defaultdict(lambda: {"receipts": [], "flags": [], "notes": []})
    rejected, not_assessed = [], []
    receipts = vrc.RunReceipts(rc.run)
    receipts.run_id = rc.run_id
    for f in vr.get("findings", []) or []:
        fid, ds = f.get("finding_id"), f.get("dataset_id")
        why = receipts.problem(f, results)
        if why:
            rejected.append({"finding_id": fid, "dataset_id": ds, "why": why})
            continue
        if f["answer"] == "cannot_tell":
            not_assessed.append({"finding_id": fid, "dataset_id": ds, "question_id": f.get("question_id")})
            continue
        per[ds]["receipts"].append(fid)
        if vrc.is_view(f):
            per[ds].setdefault("view_receipts", []).append(fid)
        if f["effect"] == "flag":
            per[ds]["flags"].append({"finding_id": fid, "question_id": f.get("question_id"),
                                     "category": f.get("category"),
                                     "plot": f.get("plot") or (f.get("view") or {}).get("project")})
    return per, {"rejected": rejected, "not_visually_assessed": not_assessed,
                 "rubric_sha256": vr.get("rubric_sha256"), "findings": len(vr.get("findings", []) or [])}


def _step_verdicts(rc: RunContext, results: dict) -> dict[str, str]:
    out = {}
    for sid, s in rc.steps.items():
        ds = [r for r in results.values() if r.get("step_id") == sid and "UNPAIRED_NEW" not in r["reason_codes"]
              and not r.get("excluded")]
        if ds:
            out[sid] = ac.worst(r["verdict"] for r in ds)
    return out


def aggregate(run_dir: str | os.PathLike, isolation_runs: Iterable[str | os.PathLike] = ()) -> tuple[dict, int]:
    rc = RunContext(run_dir)
    results = _load_results(rc)
    pairing = ac.load_json(rc.run / "verify" / "pairing.json") if (rc.run / "verify" / "pairing.json").exists() \
        else {"notes": [], "contract_hash": None}
    contract_hash = pairing.get("contract_hash")
    v: dict = {"schema": VERDICT_SCHEMA, "run_id": rc.run_id, "mode": rc.chain, "contract_hash": contract_hash,
               "comparator": comparator_block(), "pipeline": {}, "steps": {}, "datasets": {},
               "first_failing_steps": [], "within_tolerance": [], "excluded": [], "not_checked": [],
               "notes": pairing.get("notes", []), "visual": {}, "isolation_runs": [], "created_at": ac.utc_now_iso()}
    if not results:
        v["pipeline"] = {"verdict": "EVALUATOR_ERROR", "reason_codes": ["CHECKER_EXCEPTION"],
                         "detail": "no compare results for this run: run `compare.py run` first"}
        ac.dump_json(rc.run / "verify" / "verdict.json", v)
        return v, EXIT_NOT_PASS
    per_vis, vis_summary = _visual(rc, results)
    v["visual"] = vis_summary
    compare_verdicts = {}
    for dsid, r in sorted(results.items()):
        cv = r["verdict"]
        codes = list(r["reason_codes"])
        rs = _recompute_state(r)
        if rs != cv and r.get("checks"):
            cv = "EVALUATOR_ERROR"
            codes = ["STATE_RECOMPUTE_MISMATCH"] + codes
        compare_verdicts[dsid] = cv
        dv = cv
        vis = per_vis.get(dsid, {"receipts": [], "flags": [], "notes": []})
        compared = r.get("family") in ("raster", "vector") and bool(r.get("checks")) and \
            cv not in ("NOT_COMPARED", "EVALUATOR_ERROR", "BLOCKED")
        if compared and not vis["receipts"]:
            if dv in PASS:
                dv = "NEEDS_REVIEW"
            codes.append("VISUAL_REVIEW_MISSING")
        if vis["flags"]:
            if dv in PASS:
                dv = "NEEDS_REVIEW"
                codes.append("VISUAL_FLAG")
            else:
                vis["notes"].append("visual flag attached as a note (state not PASS_*)")
        v["datasets"][dsid] = {"verdict": dv, "compare_verdict": cv, "reason_codes": codes,
                               "diagnostic_codes": r.get("diagnostic_codes", []), "step_id": r.get("step_id"),
                               "family": r.get("family"), "role": r.get("role"), "excluded": r.get("excluded", False),
                               "result_id": r.get("result_id"), "visual": vis}
        for w in r.get("within_tolerance", []):
            v["within_tolerance"].append(dict(w, dataset_id=dsid))
        for nc in r.get("not_checked", []):
            v["not_checked"].append(f"{dsid}: {nc}")
        if r.get("excluded") and r.get("role") != "final":
            v["excluded"].append({"dataset_id": dsid, "reason_codes": r["reason_codes"]})
    # ---- steps (chained verdicts with visual merge) and blame (compare verdicts only)
    final_ds = {d: x for d, x in v["datasets"].items()}
    parents = _upstream_steps(rc.graph)
    steps_in = [s for s, x in rc.steps.items() if x.get("scope") == "migrate"]
    step_v: dict[str, str] = {}
    step_codes: dict[str, list] = {}
    for sid in steps_in:
        ds = [d for d, x in final_ds.items() if x["step_id"] == sid and "UNPAIRED_NEW" not in x["reason_codes"]
              and not (x["excluded"] and x["role"] != "final")]
        if not ds:
            if rc.chain == "isolated" and sid not in rc.m_steps:
                continue
            step_v[sid] = "NOT_COMPARED"
            step_codes[sid] = ["NO_REFERENCE"]
            continue
        step_v[sid] = ac.worst(final_ds[d]["verdict"] for d in ds)
        codes = []
        for d in sorted(ds):
            if final_ds[d]["verdict"] not in PASS:
                for c in final_ds[d]["reason_codes"]:
                    if c not in codes:
                        codes.append(c)
        step_codes[sid] = codes
    C = {}
    for sid in steps_in:
        ds = [d for d in compare_verdicts if results[d].get("step_id") == sid
              and "UNPAIRED_NEW" not in results[d]["reason_codes"] and not results[d].get("excluded")]
        if ds:
            C[sid] = ac.worst(compare_verdicts[d] for d in ds)
    I: dict[str, str] = {}
    for iso in isolation_runs:
        irc = RunContext(iso)
        ires = _load_results(irc)
        entry = {"run_id": irc.run_id, "steps": sorted(irc.m_steps), "compared": bool(ires)}
        v["isolation_runs"].append(entry)
        if not ires:
            for s in irc.m_steps:
                I[s] = "EVALUATOR_ERROR"
            entry["error"] = "no compare results in the isolation run: run `compare.py run` on it first"
            continue
        for s, sv in _step_verdicts(irc, ires).items():
            if s in irc.m_steps:
                I[s] = sv
    blame: dict[str, str | None] = {}
    order = _topo(steps_in, parents)
    for s in order:
        c, i = C.get(s), I.get(s)
        anc = _ancestors(parents, s)
        anc_fail = [a for a in order if a in anc and C.get(a) == "FAIL"]
        b = None
        if i == "FAIL" and c in PASS:
            b = "compensating"
            step_v[s] = ac.worst([step_v.get(s, c), "NEEDS_REVIEW"])
            step_codes.setdefault(s, []).append("COMPENSATING_DEFECT")
        elif i == "FAIL":
            b = "own"
        elif c == "FAIL" and anc_fail:
            b = "inherited"
            first = anc_fail[0]
            fds = sorted(d for d in compare_verdicts if results[d].get("step_id") == first
                         and compare_verdicts[d] == "FAIL")
            step_codes.setdefault(s, []).append(f"INHERITED_FROM:{fds[0] if fds else first}")
        elif c == "FAIL" and i in PASS:
            b = "accumulated"
            step_codes.setdefault(s, []).append("ACCUMULATED_UPSTREAM_DIFFERENCE")
            step_v[s] = ac.worst([step_v.get(s, c), "NEEDS_REVIEW"])
        elif c == "FAIL":
            b = "own"
        blame[s] = b
    owns = [s for s in order if blame.get(s) == "own"]
    v["first_failing_steps"] = [s for s in owns if not (_ancestors(parents, s) & set(owns))]
    for s in order:
        if s in step_v:
            v["steps"][s] = {"verdict": step_v[s], "compare_verdict": C.get(s), "isolation_verdict": I.get(s),
                             "reason_codes": step_codes.get(s, []), "blame": blame.get(s),
                             "datasets": sorted(d for d, x in final_ds.items() if x["step_id"] == s)}
    # ---- pipeline
    pv = [x["verdict"] for x in v["steps"].values()]
    pcodes: list[str] = []
    for s, x in v["steps"].items():
        if x["verdict"] not in PASS:
            for c in x["reason_codes"]:
                if c not in pcodes:
                    pcodes.append(c)
    for n in v["notes"]:
        if n["code"] == "UNPAIRED_LEGACY":
            pv.append("NOT_COMPARED")
            if "UNPAIRED_LEGACY" not in pcodes:
                pcodes.append("UNPAIRED_LEGACY")
    # a role: final dataset is never excluded (KD exclusions and waivers do not apply to finals), so any final that
    # was not compared is already NOT_COMPARED inside its step verdict
    v["pipeline"] = {"verdict": ac.worst(pv) or "NOT_COMPARED", "reason_codes": pcodes,
                     "counts": dict(collections.Counter(x["verdict"] for x in final_ds.values()))}
    if rc.chain != "chained":
        v["pipeline"]["note"] = "isolated run: localises defects, proves nothing end to end"
    ac.dump_json(rc.run / "verify" / "verdict.json", v)
    return v, (EXIT_PASS if v["pipeline"]["verdict"] in PASS else EXIT_NOT_PASS)


# =================================================================================================================
# CLI
# =================================================================================================================
def _print_result_line(r: dict) -> str:
    codes = ",".join(r["reason_codes"]) or "-"
    diag = (" diag=" + ",".join(r["diagnostic_codes"])) if r.get("diagnostic_codes") else ""
    return f"{r['dataset_id'] or '?':40s} {r['verdict']:22s} {codes}{diag}"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="compare.py", description=__doc__.split("\n\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p_run = sub.add_parser("run", help="compare every paired dataset of a run (verify/compare/<ds>.json)")
    p_run.add_argument("--run", required=True)
    p_run.add_argument("--dataset", action="append", default=[])
    p_run.add_argument("--json", action="store_true")
    p_agg = sub.add_parser("aggregate", help="dataset/step/pipeline verdicts and blame (verify/verdict.json)")
    p_agg.add_argument("--run", required=True)
    p_agg.add_argument("--isolation-run", action="append", default=[])
    p_agg.add_argument("--json", action="store_true")
    p_pair = sub.add_parser("pair", help="standalone: one legacy file vs one new file under a spec")
    p_pair.add_argument("--legacy", required=True)
    p_pair.add_argument("--new", required=True)
    p_pair.add_argument("--spec", required=True)
    p_pair.add_argument("--out")
    p_pair.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    try:
        if a.cmd == "run":
            results, rc = run_compare(a.run, a.dataset or None)
            if a.json:
                print(json.dumps({"run": str(a.run), "exit": rc,
                                  "datasets": {k: {"verdict": r["verdict"], "reason_codes": r["reason_codes"],
                                                   "diagnostic_codes": r["diagnostic_codes"]}
                                               for k, r in sorted(results.items())}}, sort_keys=True, indent=1))
            else:
                for _, r in sorted(results.items()):
                    print(_print_result_line(r))
                print(f"exit {rc}: {'every requested dataset PASS' if rc == 0 else 'not every dataset PASS'}")
            return rc
        if a.cmd == "aggregate":
            v, rc = aggregate(a.run, a.isolation_run)
            if a.json:
                print(json.dumps(v, sort_keys=True, indent=1))
            else:
                print(f"pipeline {v['pipeline']['verdict']} {','.join(v['pipeline'].get('reason_codes', []))}")
                for s, x in v["steps"].items():
                    print(f"  {s:32s} {x['verdict']:22s} blame={x['blame']} {','.join(x['reason_codes'])}")
                print(f"first failing steps: {v['first_failing_steps']}")
            return rc
        spec = json.loads(ac.read_own_text(a.spec, "--spec"))
        mpath = None
        if a.out:
            mpath = str(Path(a.out).with_suffix("")) + ".mismatches.parquet"
            if Path(a.out).exists():
                Path(a.out).unlink()
        r = compare_pair(a.legacy, a.new, spec, mismatches_path=mpath)
        if a.out:
            ac.dump_json(a.out, r)
        if a.json:
            print(json.dumps(r, sort_keys=True, indent=1, allow_nan=False))
        else:
            print(_print_result_line(r))
        return EXIT_PASS if r["verdict"] in PASS else EXIT_NOT_PASS
    except SetupError as e:
        print(f"compare.py: {e}", file=sys.stderr)
        print(json.dumps({"verdict": "EVALUATOR_ERROR", "error": str(e)}))
        return EXIT_SETUP
    except Exception as e:                                # an evaluator fault is never a candidate verdict
        tb = traceback.format_exc()
        print(tb, file=sys.stderr)
        print(json.dumps({"verdict": "EVALUATOR_ERROR", "reason_codes": ["CHECKER_EXCEPTION"],
                          "error": f"{type(e).__name__}: {e}"}))
        return EXIT_SETUP


if __name__ == "__main__":
    sys.exit(main())

"""The redesigned report: one self-contained index.html, REPORT.md for the file pane, parity.xlsx, graph.md, and the
deliverable check (thumbnails and files hashed; a missing or changed deliverable fails)."""
from __future__ import annotations

import io
import json
import re
import subprocess
import sys
from pathlib import Path

import ladder
import report
import report_check
import synth
from report_asserts import check_report, hero_line, network_refs, parse_ok

HERE = Path(__file__).resolve().parent
GRAPHVIEW = HERE.parent / "graphview.py"


def _decision_run(proj):
    checks = synth.default_checks({"kind": "raster"})
    checks[-1].update(state="NEEDS_DECISION", reason_codes=["TOLERANCE_UNSIGNED"],
                      stats={"compared": 50, "count_diff": 7, "count_over": 7, "max_abs": 2e-9, "max_rel": 1e-12,
                             "quantiles_abs": {"p50": 0.0, "p90": 0.0, "p99": 1e-9, "p99.9": 2e-9, "max": 2e-9}},
                      samples=[{"row": 1, "col": 2, "legacy": 5.0, "new": 5.000000002}],
                      proposed={"record_id": "tol.b_exposure.v1", "source": "tolerances.proposed.yaml",
                                "rule": {"atol": 1e-8, "rtol": 0, "max_mismatch": 0, "value": None},
                                "correct_spread": 2e-9, "nearest_wrong": 0.4, "threshold": 1e-8})
    return synth.make_run(proj, 1, compares={"ds.b_exposure": ("NEEDS_DECISION", checks, ["TOLERANCE_UNSIGNED"])},
                          skip_compare=("ds.b_by_zone",))


def test_brief_assertions_on_a_mixed_run(proj):
    _decision_run(proj)
    code, rep = report.build(proj["migration"])
    lad = ladder.compute(proj["migration"])
    facts = check_report(proj["migration"], lad)
    assert facts["attention"] == {"ds.b_by_zone": "NOT_COMPARED", "ds.b_exposure": "NEEDS_DECISION"}
    page = facts["page"]
    # NOT_COMPARED before NEEDS_DECISION, both before the cards, the decision card has the proposed tolerance
    assert page.index('data-disposition="NOT_COMPARED"') < page.index('data-disposition="NEEDS_DECISION"') \
        < page.index('id="decisions"') < page.index('id="datasets"')
    dec = page[page.index('id="decisions"'):page.index('id="datasets"')]
    assert "tol.b_exposure.v1" in dec and "never signs anything" in dec and "Copy my decisions" in dec
    assert "correct spread" in dec and "0.4" in dec
    # the quantile chart with the tolerance line is drawn from the compare stats
    card = page[page.index('id="ds-ds.b_exposure"'):]
    card = card[:card.index("</article>")]
    assert 'class="dchart"' in card and "proposed 1e-08" in card
    # what the next rung needs comes from ladder.py's own blockers
    nx = page[page.index('id="next-rung"'):page.index('id="failures"')]
    for b in lad["pipeline"]["blockers"]:
        if not b.get("target"):
            assert b["code"] in nx or "DISPOSITION_" + b["code"] in nx or b["code"] in ("NOT_COMPARED",)
    assert rep["next_rung"]["next"] == "reference-verified"
    assert rep["decisions_for_you"] == ["ds.b_exposure · raster.values · TOLERANCE_UNSIGNED"] or \
        rep["decisions_for_you"][0].startswith("ds.b_exposure ·")
    md = facts["md"]
    assert md.index("## What needs attention") < md.index("## Decisions for you") < md.index("## Parity")
    assert "```mermaid" in md and "flowchart TD" in md
    assert re.search(r"\]\(assets/plots/ds\.a_mask/overview\.png\)", md)
    assert (proj["migration"] / "report" / "assets" / "plots" / "ds.a_mask" / "overview.png").is_file()


def test_thumbnails_hashed_and_report_check(proj):
    r = _decision_run(proj)
    mig = proj["migration"]
    _, rep = report.build(mig)
    thumbs = rep["thumbnails"]
    assert thumbs and all(t["source_sha256"] and t["evidence_object"] and t["sha256"] for t in thumbs.values())
    assert report_check.check(mig) == []
    assert report.check(mig) == []
    rd = mig / "report"
    # a hand edit of the summary, a deleted workbook: both caught
    (rd / "REPORT.md").write_text((rd / "REPORT.md").read_text() + "\nAll good.\n")
    (rd / "parity.xlsx").unlink()
    codes = {(x["code"], x["path"]) for x in report_check.check(mig)}
    assert ("REPORT_FILE_CHANGED", "REPORT.md") in codes and ("REPORT_FILE_MISSING", "parity.xlsx") in codes
    assert report.main([str(mig), "--check"]) == 2
    report.build(mig)
    assert report.main([str(mig), "--check"]) == 0
    # a changed source plot: the thumbnail no longer matches its sealed evidence object
    key = sorted(thumbs)[0]
    src = proj["root"] / thumbs[key]["source"]
    synth.make_writable(src.parent)
    src.write_bytes(src.read_bytes() + b"\0")
    assert ("THUMBNAIL_SOURCE_CHANGED", key) in {(x["code"], x["path"]) for x in report_check.check(mig)}
    # stdlib CLI (the Stop gate can call it without the env)
    cp = subprocess.run([sys.executable, "-I", str(HERE.parent / "report_check.py"), str(mig)], capture_output=True,
                        text=True)
    assert cp.returncode == 2 and "THUMBNAIL_SOURCE_CHANGED" in cp.stdout


def test_hand_edited_index_fails_the_render_check(sealed):
    mig = sealed["migration"]
    report.build(mig)
    p = mig / "report" / "index.html"
    rep = json.loads((mig / "report" / "REPORT.json").read_text())
    txt = p.read_text().replace("sandbox-validated", "reference-verified")
    p.write_text(txt)
    # re-hashing REPORT.json by hand does not help: a fresh render from the bundle differs
    rep["files"]["index.html"] = __import__("hashlib").sha256(p.read_bytes()).hexdigest()
    (mig / "report" / "REPORT.json").write_text(json.dumps(rep, sort_keys=True, indent=1) + "\n")
    probs = report.check(mig)
    assert any(x["code"] == "REPORT_NOT_FROM_BUNDLE" and x["path"] == "index.html" for x in probs)


def test_no_sealed_run_files(proj):
    code, rep = report.build(proj["migration"])
    assert code == 3
    md = (proj["migration"] / "report" / "REPORT.md").read_text()
    lad = ladder.compute(proj["migration"])
    assert md.splitlines()[0] == lad["status_line"] == rep["status_line"]
    page = (proj["migration"] / "report" / "index.html").read_text()
    assert hero_line(page) == lad["status_line"] and parse_ok(page) == [] and network_refs(page) == []
    assert report_check.check(proj["migration"]) == []


def test_no_plots_no_display_still_renders(proj):
    synth.make_run(proj, 1, plots=False, display=False)
    code, rep = report.build(proj["migration"])
    lad = ladder.compute(proj["migration"])
    facts = check_report(proj["migration"], lad)
    assert "swipe.html" not in rep["files"] and rep["thumbnails"] == {}
    assert "No sealed plot for this dataset" in facts["page"]


def test_parity_xlsx_values_only(sealed):
    from openpyxl import load_workbook
    report.build(sealed["migration"])
    data = (sealed["migration"] / "report" / "parity.xlsx").read_bytes()
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Status", "Parity", "Decisions", "Parameters", "Sign-off", "Runs"]
    lad = ladder.compute(sealed["migration"])
    assert wb["Status"]["B2"].value == lad["status_line"]
    ids = {row[0].value for row in wb["Parity"].iter_rows(min_row=2)}
    assert ids == set(lad["datasets"])
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), (ws.title, c.coordinate)


def test_graph_md_before_and_after_a_run(proj):
    mig = proj["migration"]

    def ask(g):
        g["questions"] = [{"id": "q07", "header": "Mask threshold", "question": "Is 60 the intended threshold?",
                           "node": "s01_mask", "answer": None}]
    synth.update_graph(mig, ask)
    cp = subprocess.run([sys.executable, str(GRAPHVIEW), "md", str(mig)], capture_output=True, text=True)
    assert cp.returncode == 0, cp.stderr
    md = (mig / "graph.md").read_text()
    assert "no sealed run yet" in md and "```mermaid" in md and "flowchart TD" in md
    assert "q07" in md and "? open: q07" in md and (mig / "graph.png").stat().st_size > 1000
    assert "![Workflow graph](graph.png)" in md
    synth.make_run(proj, 1, shift_new=True)
    cp = subprocess.run([sys.executable, str(GRAPHVIEW), "md", str(mig)], capture_output=True, text=True)
    assert cp.returncode == 0, cp.stderr
    md = (mig / "graph.md").read_text()
    lad = ladder.compute(mig)
    assert lad["status_line"] in md
    assert ":::bad" in md and "failed" in md   # s02_total fails: coloured and labelled
    assert ":::ok" in md


def test_lineage_svg_links_and_colours(proj):
    synth.make_run(proj, 1, shift_new=True)
    report.build(proj["migration"])
    page = (proj["migration"] / "report" / "index.html").read_text()
    svg = page[page.index('<svg xmlns="http://www.w3.org/2000/svg"'):]
    svg = svg[:svg.index("</svg>")]
    assert '<a href="#ds-ds.b_exposure"><g class="lin-node dataset bad"' in svg
    assert '<a href="#step-s02_total"><g class="lin-node step bad"' in svg
    assert '<a href="#ds-ds.a_mask"><g class="lin-node dataset ok"' in svg

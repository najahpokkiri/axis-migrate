#!/usr/bin/env python3
"""discover/scan.py: deterministic scanner of a read-only legacy GIS folder (axis-migrate.inventory.v1).

    scan.py <legacy_dir> --out <project>/migration [--json]
    scan.py annotate <migration> map-finding --json finding.json      (a PNG finding or a view receipt)
    scan.py annotate <migration> note --json annotation.json
    scan.py annotate <migration> png --path discover/x.png [--kind inputs_overview]   (CLI-internal: map.py calls
                                                                                       scan.annotate() itself)

The scan makes no model calls, never runs legacy code and opens nothing outside the legacy root. It hashes every file
before and after (read_only_verified), sets GDAL_PAM_ENABLED=NO and opens GeoPackages read-only/immutable.
It writes <migration>/inventory.json, <migration>/discover/profiles/<dataset>.json and <migration>/discover/scan.log.
Agent judgements enter only through `annotate`; they are merged by id and survive a rescan.
Exit codes: 0 ok, 2 usage/validation error, 3 read-only check failed (the legacy folder changed during the scan).
"""
from __future__ import annotations

import os

os.environ["GDAL_PAM_ENABLED"] = "NO"  # before any GDAL-backed import: a stats read must never write .aux.xml
os.environ.setdefault("OGR_SQLITE_JOURNAL", "OFF")

import argparse
import ast
import copy
import csv
import datetime as _dt
import difflib
import hashlib
import io
import json
import math
import re
import shlex
import sys
import tokenize
import zipfile
from collections import Counter, defaultdict
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "lib"))
import axis_common as ac  # noqa: E402

SCAN_VERSION = "0.1.0"
SCHEMA = "axis-migrate.inventory.v1"
HERE = Path(__file__).resolve().parent

RASTER_EXT = {".tif", ".tiff", ".img", ".vrt", ".asc", ".nc", ".jp2", ".bil", ".dem", ".grd", ".hdf", ".h5", ".sdat",
              ".rst", ".ecw", ".sid", ".bt", ".hgt", ".dt2", ".xyz"}
VECTOR_EXT = {".gpkg", ".shp", ".geojson", ".kml", ".kmz", ".gml", ".fgb", ".tab", ".mif", ".sqlite", ".gpx",
              ".geoparquet"}
TABLE_EXT = {".xlsx", ".xlsm", ".xls", ".xlsb", ".csv", ".tsv", ".ods", ".parquet"}
DOC_EXT = {".txt", ".md", ".rst", ".docx"}
CODE_EXT = {".py": "python", ".ipynb": "notebook", ".pyt": "python", ".r": "r", ".sql": "sql", ".sh": "shell",
            ".bat": "shell"}
PROJECT_EXT = {".qgs", ".qgz", ".model3", ".aprx", ".atbx"}
LISTED_ONLY_EXT = {".mxd", ".tbx", ".lyr", ".lyrx", ".fmw", ".pdf", ".qlr", ".mxt", ".sde"}
SHP_PARTS = {".dbf", ".shx", ".prj", ".cpg", ".sbn", ".sbx", ".qix", ".fix", ".shp.xml"}
SIDECAR_TAILS = (".aux.xml", ".ovr", ".tfw", ".tifw", ".wld", ".msk", ".qmd", "-wal", "-shm", "-journal", ".lock")
STATS_MAX_CELLS = 2 ** 26  # rasters above this many cells get sampled statistics (every k-th block)
SECRET_RE = re.compile(r"(^|[._-])(env|secret|secrets|credential|credentials|token|password|passwd|private[-_]?key|"
                       r"id_rsa)([._-]|$)|\.(pem|key|p12|pfx)$", re.I)
OUTPUT_DIR_RE = re.compile(r"(final|output|out|result|deliver|export|temp|tmp|work|scratch|intermediate)", re.I)
FINAL_DIR_RE = re.compile(r"(final|output|result|deliver|export)", re.I)
SUPERSEDED_NAME_RE = re.compile(r"(^|[_\-. ])(old|bak|backup|copy|deprecated|obsolete|unused|archive)([_\-. ]|$)",
                                re.I)
SUPERSEDED_TEXT_RE = re.compile(r"(not used|don'?t run|do not run|do not use|deprecated|obsolete|replaced by|"
                                r"superseded|old (arcgis |qgis )?version|no longer used)", re.I)
MEM_RE = re.compile(r"^(memory:.*|TEMPORARY_OUTPUT|ogr:memory.*|in_memory[\\/].*|memory[\\/].*)$", re.I)
GDAL_CLI = {"ogr2ogr": "dst_first", "gdal_translate": "last", "gdalwarp": "last", "gdal_rasterize": "last",
            "gdal_polygonize.py": "last", "gdal_polygonize": "last", "gdal_merge.py": "o_flag", "gdal_calc.py": "outfile",
            "gdal_calc": "outfile", "gdaldem": "last", "gdal_proximity.py": "last", "gdal_sieve.py": "last",
            "gdalbuildvrt": "first", "gdal_grid": "last", "ogrinfo": "read", "gdalinfo": "read",
            "gdal_contour": "last", "gdal_fillnodata.py": "last", "gdaladdo": "inplace", "gdal_edit.py": "inplace"}
LAYER_TYPES = {"source", "raster", "vector", "multilayer", "layer", "mesh", "file", "pointcloud", "vectortile"}
GRID_PARAM_MAP = {"TARGET_CRS": "outputCoordinateSystem", "CRS": "outputCoordinateSystem", "EXTENT": "extent",
                  "TARGET_EXTENT": "extent", "PROJWIN": "extent", "WIDTH": "width", "HEIGHT": "height",
                  "CELL_SIZE": "cellSize", "CELLSIZE": "cellSize", "TARGET_RESOLUTION": "cellSize",
                  "X_RESOLUTION": "cellSize", "Y_RESOLUTION": "cellSize", "PIXEL_SIZE": "cellSize",
                  "RESAMPLING": "resampling", "NODATA": "nodata", "NO_DATA": "nodata", "DATA_TYPE": "dtype",
                  "RTYPE": "dtype", "CROP_TO_CUTLINE": "crop_to_cutline", "KEEP_RESOLUTION": "keep_resolution",
                  "UNITS": "units", "ALL_TOUCHED": "all_touched", "TARGET_ALIGN": "target_aligned_pixels"}
EXTRA_FLAG_MAP = {"-at": "all_touched", "-tap": "target_aligned_pixels", "-tr": "cellSize", "-te": "extent",
                  "-ot": "dtype", "-a_nodata": "nodata", "-dstnodata": "nodata", "-r": "resampling", "-t_srs": "outputCoordinateSystem"}


# ---------------------------------------------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------------------------------------------
def posix(p: str | Path) -> str:
    return PurePosixPath(Path(p).as_posix()).as_posix()


def file_ext(name: str) -> str:
    low = name.lower()
    if low.endswith(".shp.xml"):
        return ".shp.xml"
    if low.endswith(".aux.xml"):
        return ".aux.xml"
    return os.path.splitext(low)[1]


def stem_of(name: str) -> str:
    base = name.rstrip("/").split("/")[-1]
    low = base.lower()
    for tail in (".aux.xml", ".shp.xml"):
        if low.endswith(tail):
            return base[: -len(tail)]
    return os.path.splitext(base)[0]


def iso_mtime(st) -> str:
    return _dt.datetime.fromtimestamp(st.st_mtime, _dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%fZ")


def jsonable(x: Any) -> Any:
    """Make numpy/odd scalars JSON-safe; non-finite floats become None (never NaN in an inventory)."""
    try:
        import numpy as np  # noqa: WPS433
        if isinstance(x, np.generic):
            x = x.item()
    except ImportError:  # pragma: no cover
        pass
    if isinstance(x, float) and not math.isfinite(x):
        return None
    if isinstance(x, dict):
        return {str(k): jsonable(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [jsonable(v) for v in x]
    if isinstance(x, (str, int, float, bool)) or x is None:
        return x
    return str(x)


class Log:
    def __init__(self):
        self.lines: list[str] = []

    def __call__(self, msg: str) -> None:
        self.lines.append(msg)


LOG = Log()


def load_tables() -> tuple[dict, dict]:
    qa = ac.load_json(HERE / "qgis_algs.json")
    at = ac.load_json(HERE / "arcpy_tools.json")
    return qa, at


# ---------------------------------------------------------------------------------------------------------------
# walk + hashing
# ---------------------------------------------------------------------------------------------------------------
def gdb_hash(path: Path) -> tuple[str, int]:
    h = hashlib.sha256()
    size = 0
    for dp, dn, fn in os.walk(path, followlinks=False):
        dn.sort()
        for f in sorted(fn):
            p = Path(dp) / f
            rel = p.relative_to(path).as_posix()
            size += p.stat().st_size
            h.update(rel.encode() + b"\0" + ac.sha256_file(p).encode() + b"\n")
    return h.hexdigest(), size


def walk(root: Path) -> list[dict]:
    """Every path under root, sorted. Nothing is skipped silently: VCS internals and secret-like names are listed
    with a reason and never opened; symlinks are listed and never followed; a .gdb directory is one entry."""
    out: list[dict] = []
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        dp = Path(dirpath)
        keep = []
        for d in sorted(dirnames):
            full = dp / d
            rel = posix(full.relative_to(root))
            if full.is_symlink():
                out.append({"path": rel, "kind": "symlink", "symlink_target": os.readlink(full), "dir": True})
            elif d == ".git" or d in (".svn", ".hg"):
                out.append({"path": rel, "kind": "excluded", "excluded_reason": "vcs_internals", "dir": True})
            elif d.lower().endswith(".gdb"):
                out.append({"path": rel, "kind": "gdb", "dir": True})
            elif SECRET_RE.search(d):
                out.append({"path": rel, "kind": "excluded", "excluded_reason": "secret_like_name", "dir": True})
            else:
                keep.append(d)
        dirnames[:] = keep
        for f in sorted(filenames):
            full = dp / f
            rel = posix(full.relative_to(root))
            if full.is_symlink():
                out.append({"path": rel, "kind": "symlink", "symlink_target": os.readlink(full)})
            elif SECRET_RE.search(f):
                out.append({"path": rel, "kind": "excluded", "excluded_reason": "secret_like_name",
                            "size": full.lstat().st_size})
            else:
                out.append({"path": rel, "kind": None})
    out.sort(key=lambda e: e["path"])
    return out


def hash_tree(root: Path, entries: list[dict]) -> dict[str, dict]:
    res = {}
    for e in entries:
        if e["kind"] in ("excluded", "symlink"):
            continue
        full = root / e["path"]
        if e["kind"] == "gdb":
            h, size = gdb_hash(full)
            res[e["path"]] = {"sha256": h, "size": size, "mtime": iso_mtime(full.stat())}
        else:
            st = full.stat()
            res[e["path"]] = {"sha256": ac.sha256_file(full), "size": st.st_size, "mtime": iso_mtime(st),
                              "mtime_epoch": st.st_mtime}
    return res


def classify_entries(entries: list[dict]) -> None:
    """Assign kind and sidecar_of; shapefile parts and GDAL sidecars attach to their parent."""
    paths = {e["path"] for e in entries}
    lower = {p.lower(): p for p in paths}
    for e in entries:
        if e["kind"] is not None:
            continue
        p = e["path"]
        name = p.split("/")[-1]
        low = p.lower()
        ext = file_ext(name)
        parent = None
        for tail in SIDECAR_TAILS:
            if low.endswith(tail.lower()) and len(low) > len(tail):
                cand = low[: -len(tail)]
                if cand in lower and cand != low:
                    parent = lower[cand]
                    break
        if parent is None and ext in SHP_PARTS:
            base = low[: -len(ext)]
            if base + ".shp" in lower:
                parent = lower[base + ".shp"]
        if parent is None and ext in (".tfw", ".wld", ".prj"):
            base = low[: -len(ext)]
            for rext in (".tif", ".tiff", ".jpg", ".png", ".img"):
                if base + rext in lower:
                    parent = lower[base + rext]
                    break
        if parent is not None:
            e["kind"] = "sidecar"
            e["sidecar_of"] = parent
            continue
        if ext in RASTER_EXT:
            e["kind"] = "raster"
        elif ext in VECTOR_EXT:
            e["kind"] = "vector"
        elif ext == ".json":
            e["kind"] = "json"  # sniffed later (GeoJSON or not)
        elif ext in TABLE_EXT:
            e["kind"] = "table"
        elif ext in CODE_EXT:
            e["kind"] = "notebook" if ext == ".ipynb" else "code"
        elif ext in PROJECT_EXT:
            e["kind"] = "project"
        elif ext in DOC_EXT:
            e["kind"] = "doc"
        elif ext in LISTED_ONLY_EXT:
            e["kind"] = "listed_only"
        elif ext == "":
            e["kind"] = "noext"
        else:
            e["kind"] = "other"


def sniff_json(full: Path) -> str:
    try:
        with open(full, "rb") as fh:
            head = fh.read(4096).decode("utf-8", "replace")
    except OSError:
        return "other"
    if '"FeatureCollection"' in head or '"Feature"' in head or '"geometry"' in head:
        return "vector"
    return "other"


def sniff_noext(full: Path) -> str:
    try:
        with open(full, "rb") as fh:
            head = fh.read(512)
    except OSError:
        return "other"
    if head.startswith(b"SQLite format 3"):
        return "vector"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "raster"
    if head.startswith(b"PK\x03\x04"):
        return "other"
    if head.startswith(b"#!") and b"python" in head[:80]:
        return "code"
    return "other"


# ---------------------------------------------------------------------------------------------------------------
# path / URI resolution (legacy-root relative, POSIX)
# ---------------------------------------------------------------------------------------------------------------
def parse_uri(s: str) -> dict:
    """OGR/GDAL/QGIS data source strings: p|layername=x|subset=…|layerid=n, GPKG:p:layer, /vsizip/, gdb items."""
    out = {"path": s, "layer": None, "subset": None, "layerid": None, "vsi": None}
    t = s.strip()
    m = re.match(r"^(GPKG|SQLITE):(.+):([^:\\/]+)$", t, re.I)
    if m:
        out["path"], out["layer"] = m.group(2), m.group(3)
        return out
    if "|" in t:
        head, *opts = t.split("|")
        out["path"] = head
        for o in opts:
            k, _, v = o.partition("=")
            k = k.strip().lower()
            if k == "layername":
                out["layer"] = v
            elif k == "subset":
                out["subset"] = v
            elif k == "layerid":
                out["layerid"] = v
        t = head
    m = re.match(r"^/vsi(zip|gzip|tar)/(.+?\.(zip|gz|tar|tgz))(/(.*))?$", t, re.I)
    if m:
        out["vsi"] = m.group(1)
        out["path"] = m.group(2)
        out["member"] = m.group(5)
    return out


def split_parts(raw: str) -> tuple[bool, list[str]]:
    s = raw.strip().strip('"').strip("'")
    is_abs = bool(re.match(r"^[A-Za-z]:[\\/]", s)) or s.startswith(("/", "\\\\", "~"))
    parts = [p for p in re.split(r"[\\/]+", s) if p not in ("", ".")]
    return is_abs, parts


def join_raw(parts: list[str], like: str) -> str:
    sep = "\\" if ("\\" in like and "/" not in like) else "/"
    s = sep.join(parts)
    if like.startswith("/") and not s.startswith("/"):
        s = "/" + s
    if like.startswith("\\\\"):
        s = "\\\\" + s
    return s


class Resolver:
    """Rebase code paths onto the legacy root. Never opens anything: only the walked index is consulted."""

    def __init__(self, root: Path, entries: list[dict]):
        self.root = root
        self.files_ci: dict[str, str] = {}
        self.dirs_ci: dict[str, str] = {}
        for e in entries:
            if e.get("dir") and e["kind"] != "gdb":
                self.dirs_ci[e["path"].lower()] = e["path"]
            else:
                self.files_ci[e["path"].lower()] = e["path"]
            parts = e["path"].split("/")
            for i in range(1, len(parts)):
                d = "/".join(parts[:i])
                self.dirs_ci.setdefault(d.lower(), d)
        self.prefixes: dict[str, dict] = {}  # normalised prefix -> {"value": raw prefix, "evidence": []}

    def _known(self, rel: str) -> str | None:
        r = rel.lower()
        return self.files_ci.get(r) or self.dirs_ci.get(r)

    def _gdb_split(self, rel: str) -> tuple[str, str | None]:
        parts = rel.split("/")
        for i, p in enumerate(parts):
            if p.lower().endswith(".gdb") and i < len(parts) - 1:
                return "/".join(parts[: i + 1]), parts[-1]
        return rel, None

    def learn_prefix(self, prefix_parts: list[str], like: str) -> None:
        if not prefix_parts:
            return
        key = "/".join(p.lower() for p in prefix_parts)
        if key not in self.prefixes:
            self.prefixes[key] = {"value": join_raw(prefix_parts, like), "parts": prefix_parts}

    def resolve(self, raw: str, base_dir: str = "") -> dict | None:
        """Return {rel, layer(gdb item), exists, how, prefix} or None when the path is outside the root/unknown."""
        if not isinstance(raw, str) or not raw.strip():
            return None
        if MEM_RE.match(raw.strip()):
            return {"memory": True}
        is_abs, parts = split_parts(raw)
        if not parts:
            return None
        if not is_abs:
            # relative: against the code file's directory, then against the root (".." resolved; never above it)
            cands = []
            for c in ([base_dir.split("/") + parts] if base_dir else []) + [list(parts)]:
                norm: list[str] = []
                ok = True
                for p in c:
                    if p == "..":
                        if not norm:
                            ok = False
                            break
                        norm.pop()
                    else:
                        norm.append(p)
                if ok and norm and norm not in cands:
                    cands.append(norm)
            for c in cands:
                k = self._known("/".join(c))
                if k:
                    return self._finish(k, "relative", None)
            # not on disk (an output not written yet): the candidate whose folder exists, else the first
            for c in cands:
                parent = "/".join(c[:-1])
                if not parent or parent.lower() in self.dirs_ci:
                    return self._finish("/".join(c), "relative_missing", None, exists=False)
            if cands:
                return self._finish("/".join(cands[0]), "relative_missing", None, exists=False)
            return None
        # absolute: longest suffix (>= 2 components) that exists under the root
        for k in range(1, len(parts) - 1):
            suffix = "/".join(parts[k:])
            hit = self._known(suffix)
            if hit:
                self.learn_prefix(parts[:k], raw)
                return self._finish(hit, "rebased", join_raw(parts[:k], raw))
        # a known prefix learned from other paths
        for key, pre in sorted(self.prefixes.items(), key=lambda kv: -len(kv[1]["parts"])):
            n = len(pre["parts"])
            if len(parts) > n and [p.lower() for p in parts[:n]] == [p.lower() for p in pre["parts"]]:
                rel = "/".join(parts[n:])
                hit = self._known(rel)
                return self._finish(hit or rel, "known_prefix", pre["value"], exists=bool(hit))
        # parent directory suffix exists (file not there): rebase by directory
        for k in range(1, len(parts) - 1):
            dsuf = "/".join(parts[k:-1])
            hit = self.dirs_ci.get(dsuf.lower())
            if hit:
                self.learn_prefix(parts[:k], raw)
                return self._finish(hit + "/" + parts[-1], "rebased_by_dir", join_raw(parts[:k], raw), exists=False)
        return None

    def _finish(self, rel: str, how: str, prefix: str | None, exists: bool | None = None) -> dict:
        container, item = self._gdb_split(rel)
        if item is not None:
            cont = self._known(container)
            return {"rel": cont or container, "layer": item, "gdb_item": True, "exists": False if exists is False
                    else None, "how": how, "prefix": prefix}
        real = self._known(rel)
        return {"rel": real or rel, "layer": None, "exists": bool(real) if exists is None else (exists and bool(real)),
                "how": how, "prefix": prefix, "is_dir": bool(real and real.lower() in self.dirs_ci
                                                          and real.lower() not in self.files_ci)}


# ---------------------------------------------------------------------------------------------------------------
# abstract values for the flow-sensitive Python resolver (01 §4.2)
# ---------------------------------------------------------------------------------------------------------------
class V:
    """Abstract value. k: const | list | dict | mod | func | mem | layer | toolres | derived | obj | phi | unknown."""
    __slots__ = ("k", "v", "chain", "x")

    def __init__(self, k: str, v: Any = None, chain: tuple = (), **x):
        self.k, self.v, self.chain, self.x = k, v, tuple(chain), x

    def with_chain(self, *lines) -> "V":
        n = V(self.k, self.v, merge_chain(tuple(lines), self.chain))
        n.x = self.x
        return n

    def __repr__(self):  # pragma: no cover - debugging aid
        return f"V({self.k},{self.v!r},{self.chain},{self.x})"


def merge_chain(*chains) -> tuple:
    out: dict = {}
    for c in chains:
        for i in c:
            out[i] = None
    return tuple(out)


UNK = V("unknown", "unknown")


def is_str(v: V) -> bool:
    return v.k == "const" and isinstance(v.v, str)


def is_num(v: V) -> bool:
    return v.k == "const" and isinstance(v.v, (int, float)) and not isinstance(v.v, bool)


def first_fmt(s: str) -> str | None:
    m = re.search(r"%[-+ #0]*\d*(?:\.\d+)?[a-zA-Z]", s)
    return m.group(0) if m else None


def looks_like_path(s: str) -> bool:
    if not isinstance(s, str) or not s or len(s) > 400 or "\n" in s:
        return False
    if MEM_RE.match(s.strip()):
        return True
    if "/" in s or "\\" in s or "|layername=" in s:
        return True
    return file_ext(s.split("|")[0]) in (RASTER_EXT | VECTOR_EXT | TABLE_EXT | {".json", ".txt", ".gdb"})


class Ctx:
    """Shared state of one scan: resolver, tables, the op registry and the records the analysers emit."""

    def __init__(self, root: Path, resolver: Resolver, qgis: dict, arcpy: dict):
        self.root, self.res, self.qgis, self.arcpy = root, resolver, qgis, arcpy
        self.algs = qgis.get("algorithms", {})
        self.tools = arcpy.get("tools", {})
        self.reset()

    def reset(self):
        self.ops: dict[tuple, dict] = {}
        self.refs: dict[str, dict] = {}
        self.unresolved: list[dict] = []
        self.settings: list[dict] = []
        self.params: list[dict] = []
        self.literals: list[dict] = []
        self.code: dict[str, dict] = {}
        self.local_modules: dict[str, str] = {}

    def ref(self, rel: str, layer: str | None, res: dict, raw: str) -> str:
        key = rel + ("|" + layer if layer else "")
        r = self.refs.setdefault(key, {"rel": rel, "layer": layer, "raws": [], "how": res.get("how"),
                                       "gdb_item": bool(res.get("gdb_item")), "exists": res.get("exists")})
        if raw not in r["raws"]:
            r["raws"].append(raw)
        return key


class FoldFrame:
    def __init__(self):
        self.reads: list[dict] = []
        self.writes: list[dict] = []
        self.reads_grid: list[dict] = []
        self.nested: list[str] = []
        self.lines: list[int] = []


class PyAnalyzer:
    """Flow-sensitive abstract interpretation of one Python module (or notebook cell sequence).

    One op per I/O call site (visited once, merged on loops and repeated inlining); I/O calls nested inside another
    tool call's arguments are folded into the outer op; edits through an opened layer/cursor handle are folded into
    the op that opened the handle."""

    def __init__(self, ctx: Ctx, rel: str, code_id: str, base_dir: str, star_env: dict | None = None):
        self.ctx, self.rel, self.code_id = ctx, rel, code_id
        self.base_dir = base_dir
        self.env: dict[str, V] = {}
        self.star_mods: list[str] = []
        self.cell: int | None = None
        self.fold: list[FoldFrame] = []
        self.depth = 0
        self.ret: V | None = None
        self.funcs: dict[str, dict] = {}
        self.workspace: V | None = None
        self.env_settings: dict[str, dict] = {}
        self.imports: list[dict] = []
        self.flags: list[dict] = []
        self.src_lines: dict[Any, list[str]] = {}
        self.param_lines: dict[tuple, str] = {}
        self.top_level = True
        self.exists_checks: list[dict] = []
        self.branches: dict[tuple, dict] = {}
        self.loop_stack: list[tuple] = []
        self.last_arms = ([], [])
        self.cwd: str | None = None  # os.chdir(<const>) base for later relative paths (M4)
        self.callsite: list[tuple] = []  # (call node, helper ref) while inlining a helper from another module (M5)
        self.env["__file__"] = V("const", rel)  # paths built from __file__ stay legacy-root relative
        if star_env:
            self.env.update(star_env)

    # ---------------------------------------------------------------- driving
    def run_units(self, units: list[tuple[int | None, str]]) -> None:
        for cell, src in units:
            self.cell = cell
            self.src_lines[cell] = src.splitlines()
            try:
                tree = ast.parse(src)
            except SyntaxError as exc:
                self.flags.append({"kind": "parse_error", "cell": cell, "line": exc.lineno, "msg": str(exc.msg)})
                continue
            self.collect_literal_params(tree)
            self.flags += containment_flags(tree, cell)
            self.exec_body(tree.body, module_level=True)
        for name, f in sorted(self.funcs.items()):
            if not f["called"]:
                self.flags.append({"kind": "uncalled_function", "name": name, "line": f["node"].lineno,
                                   "cell": f["cell"]})
                saved = self.env
                self.env = dict(saved)
                self.cell = f["cell"]
                for a in f["node"].args.args:
                    self.env[a.arg] = V("unknown", "function_param")
                self.depth += 1
                self.exec_body(f["node"].body)
                self.depth -= 1
                self.env = saved
                for op in self.ctx.ops.values():
                    if op["file"] == self.rel and f["node"].lineno <= op["line"] <= (f["node"].end_lineno or 0):
                        if "in_uncalled_function" not in op["flags"]:
                            op["flags"].append("in_uncalled_function")

    def seg(self, node) -> str:
        lines = self.src_lines.get(self.cell) or []
        try:
            s = ast.get_source_segment("\n".join(lines), node)
        except Exception:  # noqa: BLE001
            s = None
        if s is None:
            return ""
        s = re.sub(r"\s+", " ", s)
        return s if len(s) <= 160 else s[:157] + "..."

    # ---------------------------------------------------------------- literal parameters (def-use)
    def collect_literal_params(self, tree: ast.Module) -> None:
        def lit(node):
            if isinstance(node, ast.Constant) and not isinstance(node.value, bytes):
                return True, node.value
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub) and isinstance(node.operand, ast.Constant) \
                    and isinstance(node.operand.value, (int, float)):
                return True, -node.operand.value
            if isinstance(node, (ast.List, ast.Tuple)) and all(isinstance(e, ast.Constant) for e in node.elts) and node.elts:
                return True, [e.value for e in node.elts]
            return False, None

        def visit(body, scope):
            for st in body:
                if isinstance(st, (ast.Assign, ast.AnnAssign)):
                    targets = st.targets if isinstance(st, ast.Assign) else [st.target]
                    if st.value is None:
                        continue
                    ok, val = lit(st.value)
                    if ok:
                        for t in targets:
                            if isinstance(t, ast.Name):
                                self.ctx.params.append({"name": t.id, "value": val, "file": self.rel, "line": st.lineno,
                                                        "cell": self.cell, "scope": scope, "code": self.code_id})
                                self.param_lines[(self.cell, st.lineno)] = t.id
                elif isinstance(st, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    visit(st.body, st.name)
                elif isinstance(st, ast.If) and scope is None:
                    visit(st.body, None)
                    visit(st.orelse, None)
        visit(tree.body, None)
        ctx_of = {ast.Mult: "mult", ast.Div: "div", ast.FloorDiv: "div", ast.Pow: "pow", ast.Add: "add",
                  ast.Sub: "sub", ast.Mod: "mod"}
        for node in ast.walk(tree):
            # every numeric operand of an arithmetic expression (M7): `(30 ** 2)/10000` gives 30 and 10000
            if isinstance(node, ast.BinOp) and type(node.op) in ctx_of:
                for side in (node.left, node.right):
                    if isinstance(side, ast.Constant) and isinstance(side.value, (int, float)) \
                            and not isinstance(side.value, bool) and side.value not in (0, 1, 2, -1):
                        rec = {"value": side.value, "line": side.lineno, "cell": self.cell,
                               "context": ctx_of[type(node.op)], "file": self.rel, "code": self.code_id}
                        cs = cell_area_literal(node)
                        if cs:
                            rec["cell_size_area"] = cs
                        self.ctx.literals.append(rec)
            # dtype casts (M7): .astype(int), .astype({'col': int}), int(<expr>), float(<expr>)
            if isinstance(node, ast.Call):
                f = node.func
                if isinstance(f, ast.Attribute) and f.attr == "astype" and node.args:
                    a = node.args[0]
                    pairs = [(k.value if isinstance(k, ast.Constant) else None, v) for k, v in zip(a.keys, a.values)] \
                        if isinstance(a, ast.Dict) else [(None, a)]
                    for col, v in pairs:
                        t = v.id if isinstance(v, ast.Name) else (v.value if isinstance(v, ast.Constant) else
                                                                  ast.unparse(v))
                        self.ctx.literals.append({"value": str(t), "line": node.lineno, "cell": self.cell,
                                                  "context": "dtype", "column": col, "via": "astype",
                                                  "file": self.rel, "code": self.code_id})
                elif isinstance(f, ast.Name) and f.id in ("int", "float") and len(node.args) == 1 and \
                        not isinstance(node.args[0], ast.Constant) and self.env.get(f.id) is None:
                    self.ctx.literals.append({"value": f.id, "line": node.lineno, "cell": self.cell,
                                              "context": "dtype", "column": None, "via": f"{f.id}()",
                                              "file": self.rel, "code": self.code_id})

    # ---------------------------------------------------------------- statements
    def exec_body(self, body, module_level=False):
        for i, st in enumerate(body):
            self.exec_stmt(st, module_level and i == 0)

    def exec_stmt(self, st, first=False):
        m = getattr(self, "st_" + type(st).__name__, None)
        if m:
            m(st, first) if type(st).__name__ == "Expr" else m(st)
        else:
            for ch in ast.iter_child_nodes(st):
                if isinstance(ch, ast.expr):
                    self.ev(ch)

    def st_Import(self, st):
        for a in st.names:
            name = a.asname or a.name.split(".")[0]
            target = a.name if a.asname else a.name.split(".")[0]
            self.env[name] = V("mod", target)
            self.imports.append({"module": a.name, "line": st.lineno, "cell": self.cell, "as": a.asname})

    def st_ImportFrom(self, st):
        mod = ("." * (st.level or 0)) + (st.module or "")
        for a in st.names:
            if a.name == "*":
                self.star_mods.append(st.module or "")
                self.imports.append({"module": mod, "names": "*", "line": st.lineno, "cell": self.cell})
                local = self.ctx.local_modules.get((st.module or "").split(".")[-1])
                if local and local != self.rel:
                    for k, v in self.ctx.code.get(local, {}).get("_exports", {}).items():
                        self.env.setdefault(k, v)
                continue
            local = self.ctx.local_modules.get((st.module or "").split(".")[-1])
            fn = self.ctx.code.get(local, {}).get("_exports", {}).get(a.name) if local and local != self.rel else None
            self.env[a.asname or a.name] = fn if fn is not None and fn.k == "func" else \
                V("mod", f"{st.module}.{a.name}" if st.module else a.name)
            self.imports.append({"module": mod, "names": a.name, "line": st.lineno, "cell": self.cell,
                                 "as": a.asname})

    def st_FunctionDef(self, st):
        self.env[st.name] = V("func", st)
        self.funcs.setdefault(st.name, {"node": st, "called": False, "cell": self.cell})

    st_AsyncFunctionDef = st_FunctionDef

    def st_ClassDef(self, st):
        for b in st.body:
            if isinstance(b, (ast.FunctionDef, ast.AsyncFunctionDef)):
                self.funcs.setdefault(f"{st.name}.{b.name}", {"node": b, "called": False, "cell": self.cell})

    def st_Return(self, st):
        self.ret = self.ev(st.value) if st.value is not None else V("const", None)

    def st_Delete(self, st):
        for t in st.targets:
            if isinstance(t, ast.Name):
                self.env.pop(t.id, None)

    def st_Expr(self, st, first=False):
        if isinstance(st.value, ast.Constant) and isinstance(st.value.value, str):
            if not first and "\n" in st.value.value:
                try:
                    t = ast.parse(st.value.value.strip())
                    if any(isinstance(n, (ast.Call, ast.Assign)) for n in ast.walk(t)) and len(t.body) >= 1:
                        self.flags.append({"kind": "disabled_by_docstring", "line": st.lineno, "cell": self.cell})
                except SyntaxError:
                    pass
            return
        self.ev(st.value)

    def assign_to(self, t, v: V, line: int):
        if isinstance(t, ast.Name):
            self.env[t.id] = v.with_chain(line)
        elif isinstance(t, (ast.Tuple, ast.List)):
            vals = v.v if v.k == "list" and len(v.v) == len(t.elts) else [V("unknown", "unpack")] * len(t.elts)
            for tt, vv in zip(t.elts, vals):
                self.assign_to(tt, vv, line)
        elif isinstance(t, ast.Attribute):
            dotted = self.dotted(t)
            if dotted and dotted.startswith("arcpy.env."):
                self.set_env(dotted.split(".", 2)[2], v, t.lineno, "assignment")
            else:
                base = self.ev(t.value)
                if base.k == "obj" and base.v == "sheet" and t.attr == "title" and is_str(v):
                    base.x["title"] = v.v
                    wb = base.x.get("wb")
                    if wb is not None:
                        wb.x.setdefault("sheets", []).append(v.v)

    def st_Assign(self, st):
        v = self.ev(st.value)
        for t in st.targets:
            self.assign_to(t, v, st.lineno)

    def st_AnnAssign(self, st):
        if st.value is not None:
            self.assign_to(st.target, self.ev(st.value), st.lineno)

    def st_AugAssign(self, st):
        if isinstance(st.target, ast.Name):
            cur = self.env.get(st.target.id, V("unknown", "undefined"))
            self.env[st.target.id] = self.binop(cur, st.op, self.ev(st.value)).with_chain(st.lineno)
        else:
            self.ev(st.value)

    def st_NamedExpr(self, st):  # pragma: no cover
        self.ev(st)

    def branch(self, bodies):
        before = dict(self.env)
        envs, arm_ops = [], []
        for b in bodies:
            self.env = dict(before)
            k0 = set(self.ctx.ops)
            self.exec_body(b)
            envs.append(self.env)
            arm_ops.append([k for k in self.ctx.ops if k not in k0])
        self.last_arms = [{k: v for k, v in e.items() if before.get(k) is not v} for e in envs], arm_ops
        merged = {}
        for k in set().union(*[set(e) for e in envs]):
            vals = [e.get(k) for e in envs]
            if all(v is vals[0] for v in vals):
                merged[k] = vals[0]
            else:
                alts = [v for v in vals if v is not None]
                same = all(a.k == alts[0].k and a.v == alts[0].v for a in alts if a.k == "const")
                merged[k] = alts[0] if (same and all(a.k == "const" for a in alts)) else \
                    V("phi", alts, merge_chain(*[a.chain for a in alts]))
        self.env = merged

    def st_If(self, st):
        t = st.test
        if isinstance(t, ast.Compare) and isinstance(t.left, ast.Name) and t.left.id == "__name__":
            self.exec_body(st.body)
            return
        n_checks = len(self.exists_checks)
        tv = self.ev(t)
        checks = self.exists_checks[n_checks:]
        before_ops = set(self.ctx.ops)
        self.branch([st.body, st.orelse])
        assigns, arm_ops = self.last_arms
        rec = self.branches.setdefault((self.cell, st.lineno), {
            "line": st.lineno, "cell": self.cell, "end_line": getattr(st, "end_lineno", st.lineno),
            "condition": self.seg(t), "arms": [
                {"arm": "if", "lines": [st.body[0].lineno, max(getattr(b, "end_lineno", b.lineno) for b in st.body)],
                 "assign_lines": {}, "values": {}, "op_keys": []},
                {"arm": "elif" if len(st.orelse) == 1 and isinstance(st.orelse[0], ast.If) else "else",
                 "lines": [st.orelse[0].lineno, max(getattr(b, "end_lineno", b.lineno) for b in st.orelse)]
                 if st.orelse else None, "assign_lines": {}, "values": {}, "op_keys": []}],
            "visits": []})
        for arm, asg, ks in zip(rec["arms"], assigns, arm_ops):
            for name, v in asg.items():
                if v.k in ("func", "mod"):
                    continue
                if v.chain:
                    arm["assign_lines"].setdefault(name, v.chain[0])
                val = self.vjson(v)
                lst = arm["values"].setdefault(name, [])
                if val not in lst:
                    lst.append(val)
            for k in ks:
                if k not in arm["op_keys"]:
                    arm["op_keys"].append(k)
        visit = {"loop": [{"var": n, "line": ln, "value": self.vjson(val)} for n, ln, val in self.loop_stack],
                 "taken": ("if" if tv.v else rec["arms"][1]["arm"]) if tv.k == "const" else None}
        if visit not in rec["visits"]:
            rec["visits"].append(visit)
        new_ops = [self.ctx.ops[k] for k in self.ctx.ops if k not in before_ops]
        for op in new_ops:
            if "conditional" not in op["flags"]:
                op["flags"].append("conditional")
        for c in checks:
            written = {w.get("ref") for op in new_ops for w in op["writes"]}
            if c.get("ref") and c["ref"] in written:
                self.flags.append({"kind": "exists_guard", "line": st.lineno, "cell": self.cell,
                                   "checked_path": c["ref"], "written_path": c["ref"]})
                for op in new_ops:
                    if "exists_guard" not in op["flags"]:
                        op["flags"].append("exists_guard")

    def loop_items(self, it: V) -> list[V] | None:
        if it.k == "list" and len(it.v) <= 50:
            return it.v
        if it.k == "const" and isinstance(it.v, str) and len(it.v) <= 50 and False:
            return None
        return None

    def st_For(self, st):
        it = self.ev(st.iter)
        if it.k == "obj" and it.v in ("features", "cursor"):
            self.assign_to(st.target, V("obj", "feature", layer=it.x.get("layer", it)), st.lineno)
            self.exec_body(st.body)
            self.exec_body(st.orelse)
            return
        items = self.loop_items(it)
        if items:
            before = dict(self.env)
            for item in items:
                self.assign_to(st.target, item, st.lineno)
                name = st.target.id if isinstance(st.target, ast.Name) else self.seg(st.target)
                self.loop_stack.append((name, st.lineno, item))
                self.exec_body(st.body)
                self.loop_stack.pop()
        else:
            reason = it.x.get("reason") if it.k == "unknown" else None
            self.assign_to(st.target, V("unknown", reason or "loop_variable"), st.lineno)
            self.exec_body(st.body)
        self.exec_body(st.orelse)

    st_AsyncFor = st_For

    def st_While(self, st):
        self.ev(st.test)
        self.exec_body(st.body)
        self.exec_body(st.orelse)

    def st_With(self, st):
        mgrs = []
        for item in st.items:
            v = self.ev(item.context_expr)
            if v.k == "obj" and v.v == "envmanager":
                mgrs.append(v)
            if item.optional_vars is not None:
                self.assign_to(item.optional_vars, v, st.lineno)
        if mgrs and st.body:
            a, b = st.body[0].lineno, max((getattr(s, "end_lineno", s.lineno) or s.lineno) for s in st.body)
            for m in mgrs:
                for name, val in m.x.get("kw", {}).items():
                    self.set_env(name, val, st.lineno, "env_manager", scope=(a, b))
        self.exec_body(st.body)
        for m in mgrs:
            for name in m.x.get("kw", {}):
                prev = m.x.get("prev", {}).get(name)
                if prev is None:
                    self.env_settings.pop(name, None)
                else:
                    self.env_settings[name] = prev

    st_AsyncWith = st_With

    def st_Try(self, st):
        self.exec_body(st.body)
        for h in st.handlers:
            if h.name:
                self.env[h.name] = V("unknown", "exception")
            raises = any(isinstance(n, ast.Raise) for n in ast.walk(h))
            exits = any(isinstance(n, ast.Call) and isinstance(n.func, (ast.Attribute, ast.Name))
                        and (getattr(n.func, "attr", None) == "exit" or getattr(n.func, "id", None) == "exit")
                        for n in ast.walk(h))
            if not raises and not exits:
                self.flags.append({"kind": "swallow_except", "line": h.lineno, "cell": self.cell,
                                   "except": self.seg(h.type) if h.type is not None else "bare"})
            self.exec_body(h.body)
        self.exec_body(st.orelse)
        self.exec_body(st.finalbody)

    st_TryStar = st_Try

    # ---------------------------------------------------------------- arcpy env settings
    def set_env(self, name: str, v: V, line: int, set_by: str, scope=None):
        prev = self.env_settings.get(name)
        rec = {"name": name, "family": "arcpy_env", "set_by": set_by, "file": self.rel, "line": line,
               "cell": self.cell, "code": self.code_id, "value_v": v,
               "scope": {"from_line": scope[0] if scope else line, "to_line": scope[1] if scope else None},
               "applies_to": []}
        if set_by == "assignment" and prev is not None and prev["set_by"] == "assignment" \
                and prev["scope"]["to_line"] is None:
            prev["scope"]["to_line"] = line - 1
        self.ctx.settings.append(rec)
        if set_by == "env_manager":
            # remember what was active so the with-exit restores it
            for st in reversed(self.ctx.settings[:-1]):
                pass
        self.env_settings[name] = rec
        if name == "workspace":
            self.workspace = v

    # ---------------------------------------------------------------- expressions
    def dotted(self, node) -> str | None:
        if isinstance(node, ast.Name):
            v = self.env.get(node.id)
            if v is not None and v.k == "mod":
                return v.v
            star = self.star_lookup(node.id)
            if star:
                return star
            return None if v is not None else node.id
        if isinstance(node, ast.Attribute):
            b = self.dotted(node.value)
            return f"{b}.{node.attr}" if b else None
        return None

    def star_lookup(self, name: str) -> str | None:
        for m in self.star_mods:
            if m in ("arcpy.sa", "arcpy.ia"):
                if f"sa.{name}" in self.ctx.tools or name in ("Raster",):
                    return f"arcpy.sa.{name}"
            if m == "arcpy":
                if re.match(r"^\w+_\w+$", name) or name in self.ctx.arcpy.get("objects", {}):
                    return f"arcpy.{name}"
            if m.startswith("qgis") and name.startswith("Qgs"):
                return f"{m}.{name}"
        return None

    def ev(self, node) -> V:
        if node is None:
            return V("const", None)
        m = getattr(self, "ev_" + type(node).__name__, None)
        if m:
            return m(node)
        for ch in ast.iter_child_nodes(node):
            if isinstance(ch, ast.expr):
                self.ev(ch)
        return V("unknown", "expr:" + type(node).__name__)

    def ev_Constant(self, node):
        if isinstance(node.value, bytes):
            return V("unknown", "bytes")
        return V("const", node.value)

    def ev_JoinedStr(self, node):
        parts, chain, known = [], (), True
        for p in node.values:
            if isinstance(p, ast.Constant):
                parts.append(str(p.value))
            else:
                v = self.ev(p.value)
                chain = merge_chain(chain, v.chain)
                if v.k == "const" and v.v is not None:
                    try:
                        spec = self.ev(p.format_spec).v if p.format_spec is not None else ""
                        parts.append(format(v.v, spec or ""))
                    except Exception:  # noqa: BLE001
                        parts.append(str(v.v))
                else:
                    known = False
                    parts.append("{?}")
        if known:
            return V("const", "".join(parts), chain)
        return V("unknown", "fstring_unknown", chain, template="".join(parts), reason="fstring_unknown")

    def ev_Name(self, node):
        if node.id in self.env:
            return self.env[node.id]
        star = self.star_lookup(node.id)
        if star:
            return V("mod", star)
        if node.id in ("True", "False", "None"):
            return V("const", {"True": True, "False": False, "None": None}[node.id])
        import builtins
        if hasattr(builtins, node.id):
            return V("mod", "builtins." + node.id)
        return V("unknown", "undefined:" + node.id, reason="undefined")

    def ev_List(self, node):
        return V("list", [self.ev(e) for e in node.elts])

    ev_Tuple = ev_List
    ev_Set = ev_List

    def ev_Dict(self, node):
        d, lines, keys = {}, {}, {}
        for k, v in zip(node.keys, node.values):
            kv = self.ev(k) if k is not None else V("unknown", "splat")
            vv = self.ev(v)
            if kv.k == "const":
                d[str(kv.v)] = vv
                lines[str(kv.v)] = v.lineno
                keys[str(kv.v)] = v
        return V("dict", d, lines=lines, nodes=keys)

    def ev_IfExp(self, node):
        self.ev(node.test)
        return V("phi", [self.ev(node.body), self.ev(node.orelse)])

    def ev_Starred(self, node):
        return self.ev(node.value)

    def ev_NamedExpr(self, node):
        v = self.ev(node.value)
        self.assign_to(node.target, v, node.lineno)
        return v

    def ev_UnaryOp(self, node):
        v = self.ev(node.operand)
        if is_num(v) and isinstance(node.op, ast.USub):
            return V("const", -v.v, v.chain)
        if v.k == "obj" and v.v == "arcraster":
            return V("obj", "arcraster", mem=None, sources=v.x.get("sources", []))
        return V("unknown", "unaryop")

    def ev_BoolOp(self, node):
        vals = [self.ev(v) for v in node.values]
        return self._raster_algebra(vals) or V("unknown", "boolop")

    def ev_Compare(self, node):
        vals = [self.ev(node.left)] + [self.ev(c) for c in node.comparators]
        ra = self._raster_algebra(vals)
        if ra is not None:
            return ra
        if all(v.k == "const" or (v.k == "list" and all(i.k == "const" for i in v.v)) for v in vals):
            import operator
            fns = {ast.Eq: operator.eq, ast.NotEq: operator.ne, ast.Lt: operator.lt, ast.LtE: operator.le,
                   ast.Gt: operator.gt, ast.GtE: operator.ge, ast.In: lambda a, b: a in b,
                   ast.NotIn: lambda a, b: a not in b, ast.Is: operator.is_, ast.IsNot: operator.is_not}
            py = [v.v if v.k == "const" else [i.v for i in v.v] for v in vals]
            try:
                ok = all(fns[type(o)](py[i], py[i + 1]) for i, o in enumerate(node.ops))
                return V("const", bool(ok), merge_chain(*[v.chain for v in vals]))
            except (KeyError, TypeError):
                pass
        return V("unknown", "compare")

    def _raster_algebra(self, vals):
        rs = [v for v in vals if v.k == "obj" and v.v == "arcraster"]
        if rs:
            src = []
            for r in rs:
                src.extend(r.x.get("sources") or ([r.x["mem"]] if r.x.get("mem") else [])
                           or ([{"ref": r.x["ref"]}] if r.x.get("ref") else []))
            return V("obj", "arcraster", mem=None, sources=src)
        return None

    def ev_BinOp(self, node):
        return self.binop(self.ev(node.left), node.op, self.ev(node.right))

    def binop(self, l: V, op, r: V) -> V:
        chain = merge_chain(l.chain, r.chain)
        ra = self._raster_algebra([l, r])
        if ra is not None:
            return ra
        if isinstance(op, ast.Add):
            if is_str(l) and is_str(r):
                return V("const", l.v + r.v, chain)
            if is_num(l) and is_num(r):
                return V("const", l.v + r.v, chain)
            if l.k == "list" and r.k == "list":
                return V("list", l.v + r.v, chain)
            if is_str(l) or is_str(r):
                tl = l.v if is_str(l) else l.x.get("template", "{?}")
                tr = r.v if is_str(r) else r.x.get("template", "{?}")
                reason = (l.x.get("reason") if l.k == "unknown" else None) or \
                         (r.x.get("reason") if r.k == "unknown" else None) or "concat_unknown"
                return V("unknown", reason, chain, template=tl + tr, reason=reason)
        if isinstance(op, ast.Mod) and is_str(l):
            derived = None
            items = r.v if r.k == "list" else [r]
            for it in items:
                if it.k == "derived":
                    derived = it
                    break
            if derived is not None:
                return V("derived", derived.v.split(".")[0], chain, ref=derived.x.get("ref"),
                         fmt=first_fmt(l.v), template=l.v)
            if all(i.k == "const" for i in items):
                try:
                    return V("const", l.v % tuple(i.v for i in items) if r.k == "list" else l.v % r.v, chain)
                except (TypeError, ValueError):
                    pass
            return V("unknown", "format_unknown", chain, template=l.v, reason="fstring_unknown")
        if isinstance(op, ast.Div) and is_str(l) and is_str(r):
            return V("const", l.v.rstrip("/\\") + "/" + r.v, chain)
        if is_num(l) and is_num(r):
            try:
                val = {ast.Mult: lambda a, b: a * b, ast.Sub: lambda a, b: a - b, ast.Div: lambda a, b: a / b,
                       ast.FloorDiv: lambda a, b: a // b, ast.Pow: lambda a, b: a ** b}[type(op)](l.v, r.v)
                return V("const", val, chain)
            except Exception:  # noqa: BLE001
                pass
        return V("unknown", "binop", chain)

    def ev_Attribute(self, node):
        base = self.ev(node.value)
        if base.k == "mod":
            dotted = f"{base.v}.{node.attr}"
            if dotted.startswith("arcpy.env."):
                s = self.env_settings.get(node.attr)
                return s["value_v"] if s else V("unknown", "env_unset")
            if dotted == "os.sep":
                return V("const", "/")
            return V("mod", dotted)
        if is_str(base):
            p = base.v.replace("\\", "/").rstrip("/")
            if node.attr == "parent":
                return V("const", p.rsplit("/", 1)[0] if "/" in p else ".", base.chain)
            if node.attr == "name":
                return V("const", p.rsplit("/", 1)[-1], base.chain)
            if node.attr == "stem":
                return V("const", os.path.splitext(p.rsplit("/", 1)[-1])[0], base.chain)
            return V("unknown", "attr")
        if base.k == "obj":
            if base.v == "describe":
                return V("derived", f"{node.attr}_of", ref=base.x.get("ref"))
            if base.v == "workbook" and node.attr == "active":
                return V("obj", "sheet", ref=base.x.get("ref"), op=base.x.get("op"), sheet="<active>", wb=base)
            if base.v == "newworkbook" and node.attr == "active":
                return V("obj", "sheet", wb=base, new=True)
        if base.k == "layer" and node.attr in ("source",):
            return V("unknown", "layer_source")
        return V("unknown", "attr:" + node.attr)

    def ev_Subscript(self, node):
        base = self.ev(node.value)
        idx = self.ev(node.slice)
        if base.k == "obj" and base.v == "feature" and is_str(idx):
            self.field_read(base.x.get("layer"), idx.v, node.lineno, "feature[]")
            return V("unknown", "field_value")
        if base.k == "toolres" and idx.k == "const":
            out = base.x.get("outputs", {}).get(str(idx.v))
            return out if out is not None else V("unknown", "tool_output")
        if base.k == "dict" and idx.k == "const":
            return base.v.get(str(idx.v), V("unknown", "key"))
        if base.k == "list" and is_num(idx) and isinstance(idx.v, int) and -len(base.v) <= idx.v < len(base.v):
            return base.v[idx.v]
        if base.k == "obj" and base.v == "workbook" and is_str(idx):
            op = self.ctx.ops.get(base.x.get("op"))
            if op is not None:
                op["params"].setdefault("sheet", {"value": idx.v, "line": node.lineno})
            return V("obj", "sheet", ref=base.x.get("ref"), op=base.x.get("op"), sheet=idx.v, wb=base)
        if base.k == "mod" and base.v in ("sys.argv",):
            self.unresolved(node, "argv")
            return V("unknown", "argv", reason="argv")
        if base.k == "mod" and base.v == "os.environ":
            self.unresolved(node, "env_var")
            return V("unknown", "env_var", reason="env_var")
        return V("unknown", "subscript")

    def ev_Lambda(self, node):
        return V("unknown", "lambda")

    def ev_ListComp(self, node):
        for g in node.generators:
            self.ev(g.iter)
        return V("unknown", "comprehension")

    ev_SetComp = ev_ListComp
    ev_GeneratorExp = ev_ListComp
    ev_DictComp = ev_ListComp

    def unresolved(self, node, reason, op=None, expr=None):
        if self.callsite:
            node, expr = self.callsite[0][0], expr or f"{self.seg(self.callsite[0][0])} ({self.callsite[0][1][11:]})"
        self.ctx.unresolved.append({"file": self.rel, "line": node.lineno, "cell": self.cell,
                                    "expr": expr or self.seg(node), "reason": reason, "op": op, "code": self.code_id})

    # ---------------------------------------------------------------- calls
    def ev_Call(self, node):
        func = node.func
        base, method, dotted = None, None, None
        if isinstance(func, ast.Name):
            v = self.env.get(func.id)
            if v is not None and v.k == "mod":
                dotted = v.v
            elif v is not None and v.k == "func":
                return self.inline(v.v, node, src=v.x.get("src"))
            elif v is not None:
                base = v
            else:
                dotted = self.star_lookup(func.id) or ("builtins." + func.id)
        elif isinstance(func, ast.Attribute):
            b = self.ev(func.value)
            if b.k == "mod":
                dotted = f"{b.v}.{func.attr}"
            else:
                base, method = b, func.attr
        else:
            base = self.ev(func)
        if dotted and "." in dotted:
            local = self.ctx.local_modules.get(dotted.rsplit(".", 1)[0].split(".")[-1])
            fn = self.ctx.code.get(local, {}).get("_exports", {}).get(dotted.rsplit(".", 1)[1]) \
                if local and local != self.rel else None
            if fn is not None and fn.k == "func":
                return self.inline(fn.v, node, src=fn.x.get("src"))
        h = self.handler_for(dotted, base, method)
        if h is not None:
            return h(node, dotted, base, method)
        # unknown call: evaluate args (nested I/O calls are still found), then name heuristic
        args = [self.ev(a) for a in node.args]
        kws = {k.arg: self.ev(k.value) for k in node.keywords if k.arg}
        if dotted and dotted.startswith("builtins."):
            name = dotted.split(".", 1)[1]
            if name in ("str", "repr") and args:
                return args[0]
            if name == "float" and args and is_num(args[0]):
                return args[0]
            return V("unknown", "call:" + name)
        if method in ("replace",) and is_str(base) and len(args) == 2 and all(is_str(a) for a in args):
            return V("const", base.v.replace(args[0].v, args[1].v), base.chain)
        if method in ("format",) and is_str(base):
            if all(a.k == "const" for a in args) and all(v.k == "const" for v in kws.values()):
                try:
                    return V("const", base.v.format(*[a.v for a in args], **{k: v.v for k, v in kws.items()}),
                             merge_chain(base.chain, *[a.chain for a in args]))
                except (IndexError, KeyError, ValueError):
                    pass
            return V("unknown", "fstring_unknown", template=base.v, reason="fstring_unknown")
        if method in ("partition", "rpartition") and is_str(base) and len(args) == 1 and is_str(args[0]):
            return V("list", [V("const", x, base.chain) for x in getattr(base.v, method)(args[0].v)], base.chain)
        if method in ("strip", "lower", "upper", "rstrip", "lstrip") and is_str(base):
            return V("const", getattr(base.v, method)(*[a.v for a in args if a.k == "const"]), base.chain)
        if method in ("join",) and is_str(base) and args and args[0].k == "list" and all(is_str(x) for x in args[0].v):
            return V("const", base.v.join(x.v for x in args[0].v), base.chain)
        return self.heuristic(node, dotted or (method and f"?.{method}") or "?", args, kws)

    def inline(self, fn: ast.FunctionDef, node, src: str | None = None) -> V:
        """Run a function body with the call-site arguments bound (one level for helpers of another module, M5).
        I/O inside a helper from another file becomes an op of THIS file at the call site, flagged
        `via_helper:<file>:<def line>:<name>`; its own module stays a helper."""
        foreign = bool(src) and src != self.rel
        rec = self.funcs.get(fn.name) if not foreign else None
        if rec is not None:
            rec["called"] = True
        args = [self.ev(a) for a in node.args]
        kws = {k.arg: self.ev(k.value) for k in node.keywords if k.arg}
        if self.depth >= 2 or (foreign and self.callsite):
            return V("unknown", "call_depth")
        if foreign:
            self.callsite.append((node, f"via_helper:{src}:{fn.lineno}:{fn.name}"))
            try:
                return self._inline_body(fn, args, kws, src)
            finally:
                self.callsite.pop()
        return self._inline_body(fn, args, kws)

    def _inline_body(self, fn: ast.FunctionDef, args: list, kws: dict, src: str | None = None) -> V:
        saved, saved_ret = self.env, self.ret
        self.env = dict(saved)
        if src:  # the helper's own module imports (rio, gpd, ...) as its body sees them
            self.env.update(self.ctx.code.get(src, {}).get("_globals", {}))
        params = [a.arg for a in fn.args.args]
        defaults = fn.args.defaults
        for i, p in enumerate(params):
            if i < len(args):
                self.env[p] = args[i]
            elif p in kws:
                self.env[p] = kws[p]
            else:
                di = i - (len(params) - len(defaults))
                self.env[p] = self.ev(defaults[di]) if di >= 0 else V("unknown", "function_param",
                                                                        reason="function_param")
        self.ret = None
        self.depth += 1
        self.exec_body(fn.body)
        self.depth -= 1
        ret = self.ret or V("const", None)
        self.env, self.ret = saved, saved_ret
        return ret

    def bind(self, node, names: list[str], skip: int = 0) -> dict[str, tuple[V, int, Any]]:
        """Bind positional and keyword args to parameter names: name -> (value, line, node)."""
        out = {}
        pos = node.args[skip:]
        for i, a in enumerate(pos):
            if isinstance(a, ast.Starred):
                continue
            nm = names[i] if i < len(names) else f"arg{i + skip}"
            out[nm] = (self.ev(a), a.lineno, a)
        for k in node.keywords:
            if k.arg:
                out[k.arg] = (self.ev(k.value), k.value.lineno, k.value)
            else:
                v = self.ev(k.value)
                if v.k == "dict":
                    for kk, vv in v.v.items():
                        out[kk] = (vv, k.value.lineno, k.value)
        return out

    def handler_for(self, dotted, base, method):
        if dotted:
            d = dotted
            if d.startswith("qgis.") or d.startswith("processing."):
                last = d.split(".")[-1]
                if d in ("processing.run", "processing.runAndLoadResults", "processing.runAlgorithm",
                         "processing.core.Processing.Processing.runAlgorithm"):
                    return self.h_processing
                if last in ("QgsRasterLayer", "QgsVectorLayer"):
                    return self.h_qgs_layer
                if last.startswith("QgsVectorFileWriter"):
                    return self.h_qgs_writer
                if last == "QgsRasterFileWriter":
                    return self.h_qgs_rwriter
                return None
            if d.startswith("arcpy"):
                return self.h_arcpy
            if d in ("os.path.join", "posixpath.join", "ntpath.join"):
                return self.h_join
            if d in ("os.path.dirname", "os.path.basename", "os.path.abspath", "os.path.normpath",
                     "os.path.realpath", "os.path.expanduser", "pathlib.Path", "pathlib.PurePath",
                     "pathlib.PureWindowsPath", "pathlib.WindowsPath", "pathlib.PosixPath", "builtins.Path"):
                return self.h_pathfn
            if d in ("os.path.exists", "os.path.isfile", "os.path.isdir", "os.access"):
                return self.h_exists
            if d == "os.chdir":
                return self.h_chdir
            if d == "os.getcwd":
                return self.h_getcwd
            if d in ("os.environ.get", "os.getenv"):
                return self.h_getenv
            if d in ("builtins.input",):
                return self.h_input
            if d in ("glob.glob", "glob.iglob", "os.listdir", "os.scandir"):
                return self.h_glob
            if d in ("builtins.open", "io.open", "codecs.open"):
                return self.h_open
            if d.startswith(("subprocess.", "os.system", "os.popen")):
                return self.h_subprocess
            if d.startswith("shutil.") or d in ("os.remove", "os.unlink", "os.rename", "os.replace", "os.rmdir"):
                return self.h_shutil
            last = d.split(".")[-1]
            root_mod = d.split(".")[0]
            if last in ("load_workbook",):
                return self.h_load_workbook
            if d.endswith("openpyxl.Workbook") or d == "openpyxl.Workbook" or last == "Workbook" and "openpyxl" in d:
                return self.h_new_workbook
            if root_mod in ("pd", "pandas", "gpd", "geopandas", "pyogrio", "xr", "xarray", "np", "numpy",
                            "polars", "pl", "fiona") and (last.startswith("read") or last in ("load", "open_dataset",
                                                                                           "open_rasterio")):
                return self.h_reader
            if root_mod in ("pyogrio", "fiona") and last.startswith("write"):
                return self.h_writer_fn
            if d in ("rasterio.open",) or (root_mod == "rasterio" and last == "open"):
                return self.h_rasterio_open
            if root_mod in ("gdal", "osgeo", "ogr") and last in ("Open", "OpenEx", "Translate", "Warp", "Rasterize",
                                                                   "VectorTranslate", "BuildVRT", "Polygonize",
                                                                   "DEMProcessing", "Grid", "Info", "VectorInfo"):
                return self.h_gdal
            if d in ("pd.ExcelWriter", "pandas.ExcelWriter"):
                return self.h_excelwriter
            return None
        if base is not None and method:
            if base.k == "obj" and base.v == "feature" and method == "attribute":
                return self.h_feature_attr
            if base.k == "layer":
                return self.h_layer_method
            if base.k == "obj" and base.v in ("provider", "cursor", "editlayer"):
                return self.h_fold_method
            if base.k == "derived" or (base.k == "obj" and base.v == "extent"):
                return self.h_derived_method
            if base.k == "obj" and base.v == "arcraster" and method == "save":
                return self.h_arcraster_save
            if base.k == "obj" and base.v in ("workbook", "newworkbook") and method == "save":
                return self.h_wb_save
            if base.k == "obj" and base.v == "sheet" and method in ("iter_rows", "cell", "iter_cols"):
                return self.h_sheet_rows
            if base.k == "obj" and base.v in ("workbook",) and method in ("get_sheet_by_name",):
                return self.h_wb_sheet
            if base.k == "obj" and base.v == "newworkbook" and method in ("create_sheet",):
                return self.h_wb_sheet
            if method in ("to_csv", "to_excel", "to_file", "to_parquet", "to_json", "to_feather", "to_raster",
                          "to_netcdf", "to_sql", "to_geojson") and base.k != "mod":
                return self.h_df_write
            if is_str(base) and method in ("joinpath",):
                return self.h_join
            if base.k == "obj" and base.v == "gdaldriver" and method in ("Create", "CreateCopy", "CreateDataSource"):
                return self.h_gdal_create
        return None

    # ---------------------------------------------------------------- op emission
    def emit(self, node, api, tool, family, io_source, reads=(), writes=(), reads_grid=(), params=None,
             state="directly_observed", flags=(), mem_out=None, end_line=None) -> tuple:
        if self.callsite:  # inside a helper of another module: the op belongs to the call site (M5)
            cs, via = self.callsite[0]
            reads = [dict(r, line=cs.lineno) for r in reads]
            writes = [dict(w, line=cs.lineno) for w in writes]
            reads_grid = [dict(g, line=cs.lineno) for g in reads_grid]
            params = {k: (dict(v, arg_line=cs.lineno) if isinstance(v, dict) and "arg_line" in v else v)
                      for k, v in (params or {}).items()}
            flags = list(flags) + [via]
            node, end_line = cs, None
        if self.fold:
            fr = self.fold[-1]
            fr.reads.extend(reads)
            fr.writes.extend(writes)
            fr.reads_grid.extend(reads_grid)
            fr.nested.append(api)
            fr.lines.append(node.lineno)
            return None
        key = (self.rel, self.cell if self.cell is not None else -1, node.lineno, node.col_offset,
               getattr(node, "end_lineno", None) or node.lineno, getattr(node, "end_col_offset", None) or 0)
        op = self.ctx.ops.get(key)
        if op is None:
            op = {"key": key, "code": self.code_id, "file": self.rel, "cell": self.cell, "line": node.lineno,
                  "col": node.col_offset, "end_line": end_line or getattr(node, "end_lineno", node.lineno),
                  "api": api, "tool": tool, "family": family, "io_source": io_source, "params": {},
                  "reads": [], "writes": [], "reads_grid": [], "flags": [], "state": state, "nested": [],
                  "memory_out": None, "visits": 0}
            self.ctx.ops[key] = op
        op["visits"] += 1

        def add(lst, items):
            for it in items:
                if it not in lst:
                    lst.append(it)
        add(op["reads"], reads)
        add(op["writes"], writes)
        add(op["reads_grid"], reads_grid)
        for f in flags:
            if f not in op["flags"]:
                op["flags"].append(f)
        if params:
            for k, v in params.items():
                op["params"].setdefault(k, v)
        if mem_out:
            op["memory_out"] = mem_out
        if self.depth and "in_function" not in op["flags"]:
            op["flags"].append("in_function")
        return key

    def mem_id(self, line: int) -> str:
        return f"mem:{self.code_id}:{self.cell if self.cell is not None else ''}:{line}"

    def path_refs(self, v: V, line: int, param: str, *, force_path=False, workspace: V | None = None) -> list[dict]:
        """Turn an abstract value used as an I/O argument into dataset refs / memory refs / unresolved records."""
        out: list[dict] = []
        if v.k == "const" and isinstance(v.v, str):
            s = v.v
            if not (force_path or looks_like_path(s)):
                return out
            if workspace is not None and is_str(workspace) and not re.search(r"[\\/]", s) and not MEM_RE.match(s) \
                    and not split_parts(s)[0]:
                s = workspace.v.rstrip("\\/") + ("\\" if "\\" in workspace.v else "/") + s
            if self.cwd and not split_parts(s)[0] and not MEM_RE.match(s.strip()) and "://" not in s:
                s = self.cwd.rstrip("\\/") + ("\\" if "\\" in self.cwd and "/" not in self.cwd else "/") + s
            uri = parse_uri(s)
            res = self.ctx.res.resolve(uri["path"], self.base_dir)
            if res is None:
                out.append({"unresolved": s, "param": param, "line": line, "reason": "outside_root"})
                if split_parts(uri["path"])[0]:
                    self.flag_foreign(uri["path"].replace("\\", "/").rsplit("/", 1)[0], line, "path")
            elif res.get("memory"):
                out.append({"memory": f"mem:{s.strip()}:{self.code_id}:{line}", "param": param, "line": line,
                            "literal_memory": s.strip()})
            elif res.get("is_dir") and not force_path:
                return out
            else:
                layer = uri["layer"] or res.get("layer")
                key = self.ctx.ref(res["rel"], layer, res, s)
                d = {"ref": key, "param": param, "line": line}
                if uri.get("subset"):
                    d["subset"] = uri["subset"]
                out.append(d)
        elif v.k == "list":
            for it in v.v:
                out.extend(self.path_refs(it, line, param, force_path=force_path, workspace=workspace))
        elif v.k == "phi":
            for it in v.v:
                for r in self.path_refs(it, line, param, force_path=force_path, workspace=workspace):
                    r["phi"] = True
                    out.append(r)
        elif v.k == "layer":
            if v.x.get("ref"):
                out.append({"ref": v.x["ref"], "param": param, "line": line, "via": "layer"})
            elif v.x.get("mem"):
                out.append({"memory": v.x["mem"], "param": param, "line": line})
        elif v.k == "mem":
            out.append({"memory": v.v, "param": param, "line": line})
        elif v.k == "obj" and v.v == "arcraster":
            if v.x.get("ref"):
                out.append({"ref": v.x["ref"], "param": param, "line": line, "via": "raster_object"})
            elif v.x.get("mem"):
                out.append({"memory": v.x["mem"], "param": param, "line": line})
            for s in v.x.get("sources") or []:
                if isinstance(s, str):
                    out.append({"memory": s, "param": param, "line": line})
                elif isinstance(s, dict) and s.get("ref"):
                    out.append({"ref": s["ref"], "param": param, "line": line, "via": "raster_object"})
        elif v.k == "toolres":
            o = v.x.get("outputs", {}).get("OUTPUT")
            if o is not None:
                out.extend(self.path_refs(o, line, param, force_path=force_path))
        elif v.k == "unknown" and force_path:
            reason = v.x.get("reason") or v.v
            tmpl = v.x.get("template") or ""
            head = tmpl.split("{?}")[0].replace("\\", "/")
            if self.cwd and tmpl and not split_parts(tmpl)[0]:
                head = self.cwd.rstrip("/\\") + "/" + head
            if head and split_parts(head)[0] and "/" in head.rstrip("/") and \
                    self.ctx.res.resolve(head.rsplit("/", 1)[0] + "/__probe__") is None:
                reason = "outside_root"  # a partly unknown path under a directory of another machine (M4)
                self.flag_foreign(head.rsplit("/", 1)[0], line, "path")
            out.append({"unresolved": v.x.get("template") or v.v, "param": param, "line": line, "reason": reason})
            self.ctx.unresolved.append({"file": self.rel, "line": line, "cell": self.cell,
                                        "expr": v.x.get("template") or str(v.v), "reason": reason, "op": None,
                                        "code": self.code_id})
        return out

    def vjson(self, v: V, depth=0) -> Any:
        if depth > 6:
            return {"unknown": "deep"}
        if v.k == "const":
            return v.v
        if v.k == "list":
            return [self.vjson(i, depth + 1) for i in v.v]
        if v.k == "dict":
            return {k: self.vjson(i, depth + 1) for k, i in v.v.items()}
        if v.k == "layer":
            return {"layer_of": v.x.get("ref")} if v.x.get("ref") else {"memory": v.x.get("mem")}
        if v.k == "mem":
            return {"memory": v.v}
        if v.k == "derived":
            d = {"derived": v.v, "ref": v.x.get("ref")}
            if v.x.get("fmt"):
                d["format"] = v.x["fmt"]
            return d
        if v.k == "phi":
            return {"phi": [self.vjson(i, depth + 1) for i in v.v]}
        if v.k == "toolres":
            return {"tool_result": "op"}
        if v.k == "obj":
            if v.v == "arcraster" and v.x.get("ref"):
                return {"raster_of": v.x["ref"]}
            if v.v == "arcraster":
                return {"memory": v.x.get("mem") or "raster_expression"}
            return {"object": v.v}
        if v.k == "mod":
            return {"module": v.v}
        if v.k == "unknown":
            d = {"unknown": v.x.get("reason") or v.v}
            if v.x.get("template"):
                d["template"] = v.x["template"]
            return d
        return {"unknown": v.k}

    def pjson(self, v: V, line: int, node, refs: list[dict] | None = None) -> dict:
        d = {"expr": self.seg(node) if node is not None else "", "value": self.vjson(v), "arg_line": line}
        if refs:
            ds = [r.get("ref") or r.get("memory") for r in refs if r.get("ref") or r.get("memory")]
            if len(ds) == 1:
                d["value"] = {"ref": ds[0]} if refs[0].get("ref") else {"memory": ds[0]}
            elif ds:
                d["value"] = [{"ref": x} if x and not x.startswith("mem:") else {"memory": x} for x in ds]
        ch = [c for c in v.chain if isinstance(c, int)]
        if ch:
            d["chain"] = ch
        return d

    def grid_refs(self, v: V, param: str, line: int) -> list[dict]:
        out = []
        if v.k == "derived" and v.x.get("ref"):
            out.append({"ref": v.x["ref"], "param": param, "line": line, "derived": v.v})
        elif v.k == "list":
            for i in v.v:
                out.extend(self.grid_refs(i, param, line))
        return out

    # ---------------------------------------------------------------- handlers: generic python
    def h_join(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        if base is not None:
            args = [base] + args
        if args and all(is_str(a) for a in args):
            parts = [a.v for a in args]
            s = parts[0]
            for p in parts[1:]:
                if split_parts(p)[0]:
                    s = p
                else:
                    sep = "\\" if ("\\" in s and "/" not in s) else "/"
                    s = s.rstrip("\\/") + sep + p.lstrip("\\/")
            return V("const", s, merge_chain(*[a.chain for a in args]))
        tmpl = "/".join(a.v if is_str(a) else a.x.get("template", "{?}") for a in args)
        reason = next((a.x.get("reason") for a in args if a.k == "unknown" and a.x.get("reason")), "join_unknown")
        return V("unknown", reason, merge_chain(*[a.chain for a in args]), template=tmpl, reason=reason)

    def h_pathfn(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        if not args:
            return V("unknown", "cwd", reason="cwd")
        a = args[0]
        last = dotted.split(".")[-1]
        if not is_str(a):
            if a.k == "unknown" and a.x.get("reason") == "__file__":
                return V("const", self.base_dir or ".", a.chain)
            return a
        p = a.v
        if last == "dirname":
            q = p.replace("\\", "/").rstrip("/")
            return V("const", q.rsplit("/", 1)[0] if "/" in q else "", a.chain)
        if last == "basename":
            return V("const", p.replace("\\", "/").rstrip("/").rsplit("/", 1)[-1], a.chain)
        if len(args) > 1 and all(is_str(x) for x in args):
            return self.h_join(node, dotted, None, None)
        return V("const", p, a.chain)

    def h_exists(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        if args:
            refs = self.path_refs(args[0], node.lineno, "path", force_path=False)
            self.exists_checks.append({"ref": refs[0].get("ref") if refs else None, "line": node.lineno})
        return V("unknown", "exists")

    def h_chdir(self, node, dotted, base, method):
        """os.chdir(<const>) is the base of every later relative path (M4); a directory outside the legacy root is a
        `foreign_host_path` (another machine or study area): paths under it are never looked for in the legacy folder."""
        args = [self.ev(a) for a in node.args]
        if not args or not is_str(args[0]):
            self.unresolved(node, "chdir_unknown")
            self.cwd = None
            return V("const", None)
        p = args[0].v
        if not split_parts(p)[0]:
            p = (self.cwd or self.base_dir or "").rstrip("/\\") + ("/" if (self.cwd or self.base_dir) else "") + p
        self.cwd = p
        if split_parts(p)[0] and self.ctx.res.resolve(p.rstrip("/\\") + "/__probe__") is None:
            self.flag_foreign(p, node.lineno, "os.chdir")
        return V("const", None)

    def h_getcwd(self, node, dotted, base, method):
        return V("const", self.cwd) if self.cwd else V("unknown", "cwd", reason="cwd")

    def flag_foreign(self, path: str, line: int, via: str) -> None:
        d = path.rstrip("/\\")
        if not any(f["kind"] == "foreign_host_path" and f["value"] == d for f in self.flags):
            self.flags.append({"kind": "foreign_host_path", "line": self.callsite[0][0].lineno if self.callsite else line,
                               "cell": self.cell, "value": d, "via": via})

    def h_getenv(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        kws = {k.arg: self.ev(k.value) for k in node.keywords if k.arg}
        default = args[1] if len(args) > 1 else kws.get("default")
        self.unresolved(node, "env_var")
        if default is not None and default.k == "const":
            return V("const", default.v, default.chain, env_var=True)
        return V("unknown", "env_var", reason="env_var")

    def h_input(self, node, dotted, base, method):
        for a in node.args:
            self.ev(a)
        self.flags.append({"kind": "input_call", "line": node.lineno, "cell": self.cell})
        self.unresolved(node, "input_call")
        return V("unknown", "input_call", reason="input_call")

    def h_glob(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        if not args or not is_str(args[0]):
            self.unresolved(node, "glob")
            return V("unknown", "glob", reason="glob")
        pat = args[0].v
        if self.cwd and not split_parts(pat)[0]:
            pat = self.cwd.rstrip("/\\") + "/" + pat
        last = dotted.split(".")[-1]
        if last in ("listdir", "scandir"):
            pat = pat.rstrip("/\\") + "/*"
        uri_dir, _, tail = pat.replace("\\", "/").rpartition("/")
        res = self.ctx.res.resolve(uri_dir + "/x", self.base_dir) if uri_dir else None
        if not res or "rel" not in res:
            if uri_dir and split_parts(uri_dir)[0]:
                self.flag_foreign(uri_dir, node.lineno, dotted)
                self.unresolved(node, "outside_root")
                return V("unknown", "outside_root", reason="outside_root", template=pat)
            self.unresolved(node, "glob")
            return V("unknown", "glob", reason="glob", template=pat)
        rel_dir = res["rel"].rsplit("/", 1)[0] if "/" in res["rel"] else ""
        import fnmatch
        names = sorted(p for p in list(self.ctx.res.files_ci.values()) + list(self.ctx.res.dirs_ci.values())
                       if (p.rsplit("/", 1)[0] if "/" in p else "") == rel_dir
                       and fnmatch.fnmatch(p.rsplit("/", 1)[-1], tail or "*"))
        names = list(dict.fromkeys(names))
        prefix = pat.replace("\\", "/").rpartition("/")[0]
        vals = [V("const", (uri_dir + "/" + n.rsplit("/", 1)[-1]) if last == "glob" or last == "iglob"
                  else n.rsplit("/", 1)[-1], args[0].chain) for n in names]
        self.flags.append({"kind": "listing_expanded", "line": node.lineno, "cell": self.cell, "pattern": pat,
                           "matches": [n for n in names], "order": "sorted"})
        if len(vals) > 50:
            self.unresolved(node, "loop_over_listdir")
            return V("unknown", "loop_over_listdir", reason="loop_over_listdir", template=prefix + "/*")
        return V("list", vals)

    def h_open(self, node, dotted, base, method):
        b = self.bind(node, ["file", "mode"])
        fv = b.get("file", (UNK, node.lineno, None))
        mode = b.get("mode", (V("const", "r"), node.lineno, None))[0]
        m = mode.v if is_str(mode) else "r"
        refs = self.path_refs(fv[0], fv[1], "file", force_path=True)
        reads = [r for r in refs if "ref" in r] if ("r" in m and "+" not in m) else []
        wmode = "append" if "a" in m else ("in_place" if "+" in m else "create")
        writes = [dict(r, mode=wmode) for r in refs if "ref" in r] if any(c in m for c in "wax+") else []
        for r in refs:
            if "unresolved" in r:
                self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
        key = self.emit(node, "open", "open", "python", "python_io", reads=reads, writes=writes,
                        params={"file": self.pjson(fv[0], fv[1], fv[2], refs), "mode": {"value": m}})
        return V("obj", "file", ref=(refs[0].get("ref") if refs else None), op=key)

    def h_subprocess(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        for k in node.keywords:
            self.ev(k.value)
        if not args:
            return V("unknown", "subprocess")
        a = args[0]
        if a.k == "list" and all(is_str(x) for x in a.v):
            argv = [x.v for x in a.v]
        elif is_str(a):
            try:
                argv = shlex.split(a.v)
            except ValueError:
                argv = a.v.split()
        else:
            self.unresolved(node, "subprocess_args")
            return V("unknown", "subprocess")
        return self.cli_op(node, argv, api=dotted)

    def cli_op(self, node, argv: list[str], api: str, line_override=None) -> V:
        if not argv:
            return V("unknown", "cli")
        prog = argv[0].replace("\\", "/").split("/")[-1]
        if prog in ("python", "python3") and len(argv) > 1:
            argv = argv[1:]
            prog = argv[0].replace("\\", "/").split("/")[-1]
        if prog == "pip" or (len(argv) > 2 and argv[1] == "-m" and argv[2] == "pip") or "pip" in argv[:3] and "install" in argv:
            self.flags.append({"kind": "pip_install", "line": node.lineno, "cell": self.cell})
            return V("unknown", "pip")
        rule = GDAL_CLI.get(prog)
        if rule is None:
            if prog in ("rm", "mv", "cp", "del", "copy", "move"):
                self.flags.append({"kind": "mutates_files", "line": node.lineno, "cell": self.cell, "cmd": prog})
            return V("unknown", "cli")
        pos, flags_with_val, opts = [], set(), {}
        i = 1
        takes_val = {"-f", "-of", "-t_srs", "-s_srs", "-a_srs", "-tr", "-te", "-ot", "-r", "-a_nodata", "-dstnodata",
                     "-srcnodata", "-where", "-sql", "-nln", "-nlt", "-burn", "-a", "-l", "-co", "-lco", "-dsco",
                     "-b", "-outsize", "-projwin", "-clipsrc", "-cutline", "-cl", "-o", "--outfile", "-A", "-B",
                     "-C", "--calc", "--type", "--NoDataValue", "-tap"}
        multi = {"-tr": 2, "-te": 4, "-outsize": 2, "-projwin": 4}
        while i < len(argv):
            t = argv[i]
            if t.startswith("--") and "=" in t:
                k, _, v = t.partition("=")
                opts[k] = v
            elif t.startswith("-") and len(t) > 1 and not re.match(r"^-\d", t):
                n = multi.get(t, 1 if t in takes_val and t != "-tap" else 0)
                opts[t] = argv[i + 1: i + 1 + n] if n > 1 else (argv[i + 1] if n == 1 and i + 1 < len(argv) else True)
                i += n
            else:
                pos.append(t)
            i += 1
        reads, writes = [], []
        srcs, dsts = [], []
        if rule == "dst_first" and len(pos) >= 2:
            dsts, srcs = [pos[0]], pos[1:]
        elif rule == "last" and len(pos) >= 2:
            srcs, dsts = pos[:-1], [pos[-1]]
        elif rule == "first" and len(pos) >= 2:
            dsts, srcs = [pos[0]], pos[1:]
        elif rule == "o_flag":
            srcs, dsts = pos, [opts.get("-o")] if opts.get("-o") else []
        elif rule == "outfile":
            dsts = [opts.get("--outfile")] if opts.get("--outfile") else []
            srcs = [v for k, v in opts.items() if re.match(r"^-[A-Z]$", k) and isinstance(v, str)]
        elif rule == "read":
            srcs = pos
        elif rule == "inplace":
            srcs = pos[:1]
            dsts = pos[:1]
        for s in srcs:
            reads += [r for r in self.path_refs(V("const", s), node.lineno, "src", force_path=True) if "ref" in r]
        for d in dsts:
            mode = "in_place" if rule == "inplace" else "create"
            writes += [dict(r, mode=mode) for r in self.path_refs(V("const", d), node.lineno, "dst", force_path=True)
                       if "ref" in r]
        self.emit(node, f"cli:{prog}", prog, "cli", "io_inferred_by_cli", reads=reads, writes=writes,
                  params={"argv": {"value": argv}}, state="structurally_inferred", flags=["io_inferred_by_cli"])
        return V("unknown", "cli")

    def h_shutil(self, node, dotted, base, method):
        b = self.bind(node, ["src", "dst"])
        last = dotted.split(".")[-1]
        src = b.get("src", (UNK, node.lineno, None))
        dst = b.get("dst")
        reads, writes, flags = [], [], []
        sref = self.path_refs(src[0], src[1], "src", force_path=True)
        if last in ("remove", "unlink", "rmtree", "rmdir"):
            writes = [dict(r, mode="delete") for r in sref if "ref" in r]
            flags.append("mutates_files")
        else:
            reads = [r for r in sref if "ref" in r]
            if dst is not None:
                writes = [dict(r, mode="create") for r in self.path_refs(dst[0], dst[1], "dst", force_path=True)
                          if "ref" in r]
            if last in ("move", "rename", "replace"):
                flags.append("mutates_files")
        if flags:
            self.flags.append({"kind": "mutates_files", "line": node.lineno, "cell": self.cell, "call": dotted})
        self.emit(node, dotted, last, "python", "python_io", reads=reads, writes=writes, flags=flags,
                  params={k: self.pjson(v[0], v[1], v[2]) for k, v in b.items()})
        return V("unknown", "shutil")

    def h_load_workbook(self, node, dotted, base, method):
        b = self.bind(node, ["filename", "read_only", "keep_vba", "data_only", "keep_links"])
        f = b.get("filename", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "filename", force_path=True)
        params = {"filename": self.pjson(f[0], f[1], f[2], refs)}
        for k in ("data_only", "read_only"):
            if k in b:
                params[k] = {"value": self.vjson(b[k][0]), "arg_line": b[k][1]}
        key = self.emit(node, "openpyxl.load_workbook", "load_workbook", "python", "python_io",
                        reads=[r for r in refs if "ref" in r], params=params)
        return V("obj", "workbook", ref=refs[0].get("ref") if refs else None, op=key)

    def h_new_workbook(self, node, dotted, base, method):
        for a in node.args:
            self.ev(a)
        return V("obj", "newworkbook", sheets=[])

    def h_wb_sheet(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        name = args[0].v if args and is_str(args[0]) else None
        if base.v == "newworkbook":
            if name:
                base.x.setdefault("sheets", []).append(name)
            return V("obj", "sheet", wb=base, new=True, sheet=name)
        op = self.ctx.ops.get(base.x.get("op"))
        if op is not None and name:
            op["params"].setdefault("sheet", {"value": name, "line": node.lineno})
        return V("obj", "sheet", ref=base.x.get("ref"), op=base.x.get("op"), sheet=name, wb=base)

    def h_sheet_rows(self, node, dotted, base, method):
        b = self.bind(node, ["min_row", "max_row", "min_col", "max_col", "values_only"])
        op = self.ctx.ops.get(base.x.get("op")) if base.x.get("op") else None
        if op is not None:
            rows = {k: self.vjson(v[0]) for k, v in b.items()}
            rows.update({"sheet": base.x.get("sheet"), "line": node.lineno, "call": method})
            lst = op["params"].setdefault("rows", {"value": []})["value"]
            if rows not in lst:
                lst.append(rows)
        return V("unknown", "rows")

    def h_wb_save(self, node, dotted, base, method):
        b = self.bind(node, ["filename"])
        f = b.get("filename", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "filename", force_path=True)
        mode = "in_place" if base.v == "workbook" and refs and refs[0].get("ref") == base.x.get("ref") else "create"
        params = {"filename": self.pjson(f[0], f[1], f[2], refs)}
        if base.x.get("sheets"):
            params["sheets"] = {"value": list(base.x["sheets"])}
        self.emit(node, "openpyxl.Workbook.save", "save", "python", "python_io",
                  writes=[dict(r, mode=mode) for r in refs if "ref" in r], params=params)
        return V("const", None)

    def h_excelwriter(self, node, dotted, base, method):
        b = self.bind(node, ["path"])
        f = b.get("path", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "path", force_path=True)
        self.emit(node, dotted, "ExcelWriter", "python", "python_io",
                  writes=[dict(r, mode="create") for r in refs if "ref" in r],
                  params={"path": self.pjson(f[0], f[1], f[2], refs)})
        return V("obj", "excelwriter")

    def h_reader(self, node, dotted, base, method):
        b = self.bind(node, ["path"])
        f = b.get("path", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "path", force_path=True)
        params = {"path": self.pjson(f[0], f[1], f[2], refs)}
        for k, v in b.items():
            if k != "path":
                params[k] = {"value": self.vjson(v[0]), "arg_line": v[1]}
        for r in refs:
            if "unresolved" in r:
                self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
            if r.get("ref") and "layer" in b and is_str(b["layer"][0]):
                rr = self.ctx.refs[r["ref"]]
                if not rr["layer"]:
                    r["ref"] = self.ctx.ref(rr["rel"], b["layer"][0].v, {"how": rr["how"], "exists": rr["exists"]},
                                            rr["raws"][0] if rr["raws"] else "")
        self.emit(node, dotted, dotted.split(".")[-1], "python", "python_io",
                  reads=[r for r in refs if "ref" in r], params=params)
        return V("obj", "frame")

    def h_writer_fn(self, node, dotted, base, method):
        b = self.bind(node, ["df", "path"])
        f = b.get("path", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "path", force_path=True)
        self.emit(node, dotted, dotted.split(".")[-1], "python", "python_io",
                  writes=[dict(r, mode="create") for r in refs if "ref" in r],
                  params={"path": self.pjson(f[0], f[1], f[2], refs)})
        return V("const", None)

    def h_df_write(self, node, dotted, base, method):
        b = self.bind(node, ["path"])
        f = b.get("path") or b.get("path_or_buf") or b.get("filename") or b.get("excel_writer") or (UNK, node.lineno, None)
        refs = self.path_refs(f[0], f[1], "path", force_path=True)
        params = {"path": self.pjson(f[0], f[1], f[2], refs)}
        for k, v in b.items():
            if k not in ("path", "path_or_buf", "filename"):
                params[k] = {"value": self.vjson(v[0]), "arg_line": v[1]}
        mode = "append" if (b.get("mode") and is_str(b["mode"][0]) and b["mode"][0].v == "a") else "create"
        if refs and refs[0].get("ref") and "layer" in b and is_str(b["layer"][0]):
            rr = self.ctx.refs[refs[0]["ref"]]
            refs[0]["ref"] = self.ctx.ref(rr["rel"], b["layer"][0].v, {"how": rr["how"], "exists": rr["exists"]},
                                          rr["raws"][0] if rr["raws"] else "")
        self.emit(node, f"DataFrame.{method}", method, "python", "python_io",
                  writes=[dict(r, mode=mode) for r in refs if "ref" in r], params=params)
        return V("const", None)

    def h_rasterio_open(self, node, dotted, base, method):
        b = self.bind(node, ["fp", "mode", "driver", "width", "height", "count", "crs", "transform", "dtype",
                             "nodata"])
        f = b.get("fp", (UNK, node.lineno, None))
        mode = b.get("mode", (V("const", "r"), node.lineno, None))[0]
        m = mode.v if is_str(mode) else "r"
        refs = self.path_refs(f[0], f[1], "fp", force_path=True)
        reads = [r for r in refs if "ref" in r] if m in ("r", "r+") else []
        writes = [dict(r, mode="in_place" if m == "r+" else "create") for r in refs if "ref" in r] \
            if m in ("w", "w+", "r+") else []
        params = {k: {"value": self.vjson(v[0]), "arg_line": v[1]} for k, v in b.items() if k != "fp"}
        params["fp"] = self.pjson(f[0], f[1], f[2], refs)
        key = self.emit(node, "rasterio.open", "open", "python", "python_io", reads=reads, writes=writes,
                        params=params)
        return V("obj", "rasterio_ds", ref=refs[0].get("ref") if refs else None, op=key)

    def h_gdal(self, node, dotted, base, method):
        last = dotted.split(".")[-1]
        args = [self.ev(a) for a in node.args]
        for k in node.keywords:
            self.ev(k.value)
        reads, writes = [], []
        if last in ("Open", "OpenEx", "Info", "VectorInfo"):
            update = len(args) > 1 and args[1].k in ("const", "mod") and (args[1].v in (1, True) or "Update" in str(args[1].v))
            refs = self.path_refs(args[0], node.lineno, "path", force_path=True) if args else []
            reads = [r for r in refs if "ref" in r]
            if update:
                writes = [dict(r, mode="in_place") for r in reads]
        elif args:
            dst = self.path_refs(args[0], node.lineno, "destName", force_path=True)
            writes = [dict(r, mode="create") for r in dst if "ref" in r]
            for a in args[1:2]:
                reads += [r for r in self.path_refs(a, node.lineno, "srcDS", force_path=True) if "ref" in r]
        key = self.emit(node, dotted, last, "gdal", "python_io", reads=reads, writes=writes)
        return V("obj", "gdal_ds", ref=(reads or writes or [{}])[0].get("ref"), op=key)

    def h_gdal_create(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        refs = self.path_refs(args[0], node.lineno, "path", force_path=True) if args else []
        self.emit(node, f"gdal.Driver.{method}", method, "gdal", "python_io",
                  writes=[dict(r, mode="create") for r in refs if "ref" in r])
        return V("obj", "gdal_ds")

    def heuristic(self, node, name: str, args: list[V], kws: dict[str, V]) -> V:
        """A call not in any table: I/O only if an argument is a path under the legacy root (name heuristic)."""
        if name.startswith(("builtins.", "os.", "sys.", "re.", "math.", "json.dumps", "print", "logging.")):
            return V("unknown", "call")
        reads, writes = [], []
        writeish = re.search(r"(write|save|export|to_|dump|create|copy)", name.split(".")[-1], re.I)
        for i, a in enumerate(args):
            for r in self.path_refs(a, node.lineno, f"arg{i}"):
                if "ref" not in r:
                    continue
                rr = self.ctx.refs[r["ref"]]
                if writeish and (i > 0 or not rr.get("exists")):
                    writes.append(dict(r, mode="create"))
                else:
                    reads.append(r)
        for k, a in kws.items():
            for r in self.path_refs(a, node.lineno, k):
                if "ref" not in r:
                    continue
                if re.match(r"^(out|dst|dest|output|target)", k, re.I):
                    writes.append(dict(r, mode="create"))
                else:
                    reads.append(r)
        if reads or writes:
            self.emit(node, name, name.split(".")[-1], "python", "name_heuristic", reads=reads, writes=writes,
                      state="structurally_inferred", flags=["name_heuristic"])
        return V("unknown", "call")

    # ---------------------------------------------------------------- handlers: PyQGIS
    def h_processing(self, node, dotted, base, method):
        fr = FoldFrame()
        self.fold.append(fr)
        try:
            b = self.bind(node, ["algOrName", "parameters"])
        finally:
            self.fold.pop()
        algv = b.get("algOrName", (UNK, node.lineno, None))[0]
        alg = algv.v if is_str(algv) else None
        pv, pline, pnode = b.get("parameters", (V("dict", {}), node.lineno, None))
        spec = alg_spec(self.ctx.algs, alg) or None
        io_source = "qgis_algs.json" if spec else "name_rule"
        reads, writes, grid, params, outputs = list(fr.reads), list(fr.writes), list(fr.reads_grid), {}, {}
        fields: list[dict] = []
        mem_out = None
        items = pv.v if pv.k == "dict" else {}
        lines = pv.x.get("lines", {}) if pv.k == "dict" else {}
        nodes = pv.x.get("nodes", {}) if pv.k == "dict" else {}
        if pv.k != "dict":
            self.unresolved(node, "dynamic_attr", expr="processing parameters not a literal dict")
        for pname, val in items.items():
            line = lines.get(pname, node.lineno)
            ps = (spec or {}).get("params", {}).get(pname)
            is_dest = ps["is_destination"] if ps else pname.upper().startswith("OUTPUT")
            ptype = ps["type"] if ps else None
            layerish = ptype in LAYER_TYPES if ps else (False if is_expression_param(pname, {}) else
                                                        looks_like_path(val.v) if is_str(val) else val.k in (
                                                            "layer", "mem", "toolres", "list"))
            refs = []
            if is_dest:
                refs = self.path_refs(val, line, pname, force_path=True)
                for r in refs:
                    if "ref" in r:
                        writes.append(dict(r, mode="create"))
                    elif "memory" in r:
                        mem_out = self.mem_id(line)
                        r["memory"] = mem_out
                    elif "unresolved" in r:
                        self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
                outputs[pname] = (V("mem", mem_out) if any("memory" in r for r in refs)
                                  else val)
            elif layerish:
                refs = self.path_refs(val, line, pname, force_path=bool(ps))
                for r in refs:
                    if "ref" in r or "memory" in r:
                        reads.append(r)
                    elif "unresolved" in r:
                        self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
            g = self.grid_refs(val, pname, line)
            grid.extend(g)
            if ptype == "field":
                vals = [val] if is_str(val) else (val.v if val.k == "list" else [])
                for fv in vals:
                    if is_str(fv) and fv.v:
                        for f in fv.v.split(";"):
                            fields.append({"field": f, "param": pname, "line": line,
                                           "of_param": (ps or {}).get("parent") or "INPUT"})
            params[pname] = self.pjson(val, line, nodes.get(pname), refs if (is_dest or layerish) else None)
            if is_dest and mem_out and any("memory" in r for r in refs):
                params[pname]["value"] = {"memory": mem_out}
        key = self.emit(node, dotted.split(".", 1)[-1] if dotted.startswith("processing") else dotted, alg or "?",
                        "pyqgis", io_source, reads=reads, writes=writes, reads_grid=grid,
                        params=params, mem_out=mem_out, flags=([] if spec else ["alg_not_in_table"]))
        if key and fr.nested:
            op = self.ctx.ops[key]
            for n in fr.nested:
                if n not in op["nested"]:
                    op["nested"].append(n)
        if key and fields:
            lst = self.ctx.ops[key].setdefault("fields_read", [])
            for f in fields:  # the layer the field belongs to (qgis_algs.json `parent`), never a guess
                src = next((r for r in reads if r.get("param") == f["of_param"] and r.get("ref")), None)
                if src is None:
                    continue  # parent is a memory layer or unresolved: nothing on disk to check against
                rec = {"ref": src["ref"], "field": f["field"], "line": f["line"], "param": f["param"],
                       "of_param": f["of_param"]}
                if rec not in lst:
                    lst.append(rec)
        out = {}
        for k, v in outputs.items():
            out[k] = v
        res = V("toolres", None, outputs=out, op=key)
        return res

    def h_qgs_layer(self, node, dotted, base, method):
        cls = dotted.split(".")[-1]
        b = self.bind(node, ["path", "baseName", "providerLib"])
        f = b.get("path", (UNK, node.lineno, None))
        prov = b.get("providerLib", (V("const", None), 0, None))[0]
        if is_str(prov) and prov.v in ("memory", "virtual"):
            return V("layer", None, mem=self.mem_id(node.lineno), cls=cls)
        refs = self.path_refs(f[0], f[1], "path", force_path=True)
        for r in refs:
            if "unresolved" in r:
                self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
        key = self.emit(node, cls, cls, "pyqgis", "pyqgis_api", reads=[r for r in refs if "ref" in r],
                        params={"path": self.pjson(f[0], f[1], f[2], refs)})
        ref = refs[0].get("ref") if refs else None
        return V("layer", None, ref=ref, op=key, cls=cls)

    def h_layer_method(self, node, dotted, base, method):
        for a in node.args:
            self.ev(a)
        ref = base.x.get("ref")
        if method == "extent":
            return V("derived", "extent_of", ref=ref)
        if method in ("width", "height"):
            return V("derived", f"{method}_of", ref=ref)
        if method in ("crs",):
            return V("derived", "crs_of", ref=ref)
        if method in ("rasterUnitsPerPixelX", "rasterUnitsPerPixelY"):
            return V("derived", "res_of", ref=ref)
        if method == "dataProvider":
            return V("obj", "provider", layer=base)
        if method in ("getFeatures", "selectedFeatures", "getFeature"):
            return V("obj", "features", layer=base)
        if method in ("startEditing", "commitChanges", "changeAttributeValue", "addFeature", "deleteFeature",
                      "updateFeature", "addAttribute", "deleteAttribute", "changeGeometry"):
            self.fold_into(base, node, method)
            return V("unknown", "edit")
        return V("unknown", "layer." + method)

    def field_read(self, layer_v: V | None, field: str, line: int, how: str) -> None:
        if layer_v is None or field.endswith("@"):
            return
        op = self.ctx.ops.get(layer_v.x.get("op")) if layer_v.x.get("op") else None
        if op is None or not layer_v.x.get("ref"):
            return
        rec = {"ref": layer_v.x["ref"], "field": field, "line": line, "param": how}
        if rec not in op.setdefault("fields_read", []):
            op["fields_read"].append(rec)

    def fold_into(self, layer_v: V, node, method):
        key = layer_v.x.get("op")
        op = self.ctx.ops.get(key) if key else None
        if op is None or not layer_v.x.get("ref"):
            return
        w = {"ref": layer_v.x["ref"], "param": method, "line": node.lineno, "mode": "in_place"}
        if w not in op["writes"]:
            op["writes"].append(w)
        op["end_line"] = max(op["end_line"], getattr(node, "end_lineno", node.lineno))
        if method not in op["nested"]:
            op["nested"].append(method)

    def h_feature_attr(self, node, dotted, base, method):
        args = [self.ev(a) for a in node.args]
        if args and is_str(args[0]):
            self.field_read(base.x.get("layer"), args[0].v, node.lineno, "feature.attribute")
        return V("unknown", "field_value")

    def h_fold_method(self, node, dotted, base, method):
        for a in node.args:
            self.ev(a)
        for k in node.keywords:
            self.ev(k.value)
        if base.v == "provider":
            if method in ("addAttributes", "changeAttributeValues", "addFeatures", "deleteFeatures",
                          "changeGeometryValues", "deleteAttributes", "truncate", "renameAttributes"):
                self.fold_into(base.x["layer"], node, method)
            return V("unknown", "provider." + method)
        if base.v == "cursor":
            if method in ("updateRow", "insertRow", "deleteRow"):
                self.fold_into(base, node, method)
            return V("unknown", "cursor")
        return V("unknown", "fold")

    def h_derived_method(self, node, dotted, base, method):
        for a in node.args:
            self.ev(a)
        if base.k == "derived":
            return V("derived", f"{base.v}.{method}", ref=base.x.get("ref"))
        return V("unknown", "derived")

    def h_qgs_writer(self, node, dotted, base, method):
        b = self.bind(node, ["layer", "fileName"])
        lay = b.get("layer", (UNK, node.lineno, None))
        f = b.get("fileName", (UNK, node.lineno, None))
        reads = [r for r in self.path_refs(lay[0], lay[1], "layer") if "ref" in r or "memory" in r]
        refs = self.path_refs(f[0], f[1], "fileName", force_path=True)
        self.emit(node, dotted.split(".", 2)[-1], "QgsVectorFileWriter", "pyqgis", "pyqgis_api", reads=reads,
                  writes=[dict(r, mode="create") for r in refs if "ref" in r],
                  params={"fileName": self.pjson(f[0], f[1], f[2], refs)})
        return V("unknown", "writer")

    def h_qgs_rwriter(self, node, dotted, base, method):
        b = self.bind(node, ["outputUrl"])
        f = b.get("outputUrl", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "outputUrl", force_path=True)
        self.emit(node, "QgsRasterFileWriter", "QgsRasterFileWriter", "pyqgis", "pyqgis_api",
                  writes=[dict(r, mode="create") for r in refs if "ref" in r])
        return V("obj", "rwriter")

    # ---------------------------------------------------------------- handlers: ArcPy
    def arcpy_name(self, dotted: str) -> tuple[str | None, str]:
        parts = dotted.split(".")
        if parts[0] != "arcpy":
            return None, dotted
        rest = parts[1:]
        if rest and rest[0] == "gp":
            rest = rest[1:]
        tbs = {"analysis", "management", "conversion", "sa", "ia", "ddd", "cartography", "stats", "edit", "da",
               "mapping", "md", "na", "lr", "geocoding", "ra"}
        alias = {"ia": "sa", "ra": "ra"}
        if len(rest) == 2 and rest[0] in tbs:
            tb, tool = alias.get(rest[0], rest[0]), rest[1]
            return f"{tb}.{tool}", f"arcpy.{rest[0]}.{tool}"
        if len(rest) == 1:
            m = re.match(r"^(\w+?)_(analysis|management|conversion|sa|ia|ddd|cartography|stats|edit|md|na)$", rest[0])
            if m:
                tb = alias.get(m.group(2), m.group(2))
                return f"{tb}.{m.group(1)}", f"arcpy.{m.group(1)}_{m.group(2)}"
            return rest[0], f"arcpy.{rest[0]}"
        return ".".join(rest), dotted

    def h_arcpy(self, node, dotted, base, method):
        name, api = self.arcpy_name(dotted)
        objs = self.ctx.arcpy.get("objects", {})
        if dotted.startswith("arcpy.env.") or name in ("CheckOutExtension", "CheckInExtension", "AddMessage",
                                                       "AddError", "AddWarning", "GetMessages", "SetProgressor",
                                                       "GetParameterAsText", "GetParameter", "SetParameter",
                                                       "SetParameterAsText", "SpatialReference", "Extent", "Point",
                                                       "Array", "Polygon", "Polyline", "PointGeometry", "Usage",
                                                       "ResetEnvironments", "ClearEnvironment"):
            for a in node.args:
                self.ev(a)
            for k in node.keywords:
                self.ev(k.value)
            if name in ("GetParameterAsText", "GetParameter"):
                self.unresolved(node, "argv")
                return V("unknown", "argv", reason="argv")
            return V("unknown", "arcpy." + str(name))
        if name in ("Raster", "sa.Raster"):
            b = self.bind(node, ["inRaster"])
            f = b.get("inRaster", (UNK, node.lineno, None))
            refs = self.path_refs(f[0], f[1], "inRaster", force_path=True, workspace=self.workspace)
            self.emit(node, "arcpy.Raster", "Raster", "arcpy", "arcpy_tools.json", reads=[r for r in refs if "ref" in r],
                      params={"inRaster": self.pjson(f[0], f[1], f[2], refs)})
            return V("obj", "arcraster", ref=refs[0].get("ref") if refs and refs[0].get("ref") else None,
                     mem=refs[0].get("memory") if refs else None)
        if name == "Describe" or name == "da.Describe":
            b = self.bind(node, ["value", "datatype"])
            f = b.get("value", (UNK, node.lineno, None))
            refs = self.path_refs(f[0], f[1], "value", force_path=True, workspace=self.workspace)
            self.emit(node, "arcpy.Describe", "Describe", "arcpy", "arcpy_tools.json",
                      reads=[r for r in refs if "ref" in r], params={"value": self.pjson(f[0], f[1], f[2], refs)},
                      flags=["metadata_read"])
            return V("obj", "describe", ref=refs[0].get("ref") if refs else None)
        if name in ("Exists",):
            args = [self.ev(a) for a in node.args]
            refs = self.path_refs(args[0], node.lineno, "dataset", force_path=True, workspace=self.workspace) \
                if args else []
            self.exists_checks.append({"ref": refs[0].get("ref") if refs else None, "line": node.lineno})
            return V("unknown", "exists")
        if name == "EnvManager":
            kw = {k.arg: self.ev(k.value) for k in node.keywords if k.arg}
            prev = {k: self.env_settings.get(k) for k in kw}
            return V("obj", "envmanager", kw=kw, prev=prev)
        if name in ("da.SearchCursor", "da.UpdateCursor", "da.InsertCursor", "SearchCursor", "UpdateCursor",
                    "InsertCursor"):
            spec = objs.get(name, {})
            b = self.bind(node, spec.get("params", ["in_table", "field_names"]))
            first = spec.get("params", ["in_table"])[0]
            f = b.get(first, (UNK, node.lineno, None))
            refs = self.path_refs(f[0], f[1], first, force_path=True, workspace=self.workspace)
            kind = spec.get("kind", "read")
            reads = [r for r in refs if "ref" in r]
            writes = [dict(r, mode="in_place" if kind == "in_place" else "append") for r in reads] \
                if kind in ("in_place", "append") else []
            params = {k: self.pjson(v[0], v[1], v[2], refs if k == first else None) for k, v in b.items()}
            key = self.emit(node, f"arcpy.{name}", name, "arcpy", "arcpy_tools.json", reads=reads, writes=writes,
                            params=params)
            fn = b.get("field_names") or b.get("fields")
            if key and fn and reads:
                vals = fn[0].v if fn[0].k == "list" else [fn[0]]
                for fv in vals:
                    if is_str(fv) and not fv.v.endswith("@") and fv.v != "*":
                        rec = {"ref": reads[0]["ref"], "field": fv.v, "line": fn[1], "param": "field_names"}
                        if rec not in self.ctx.ops[key].setdefault("fields_read", []):
                            self.ctx.ops[key]["fields_read"].append(rec)
            return V("obj", "cursor", ref=refs[0].get("ref") if refs else None, op=key)
        if name.startswith(("List",)):
            for a in node.args:
                self.ev(a)
            self.unresolved(node, "loop_over_listdir")
            return V("unknown", "listing", reason="loop_over_listdir")
        spec = self.ctx.tools.get(name)
        if spec is None:
            # unique bare tool name across toolboxes (arcpy.Buffer is not valid, but be lenient)
            cands = [k for k in self.ctx.tools if k.split(".")[-1] == name]
            if len(cands) == 1:
                name, spec = cands[0], self.ctx.tools[cands[0]]
        fr = FoldFrame()
        self.fold.append(fr)
        try:
            b = self.bind(node, spec["params"] if spec else [])
        finally:
            self.fold.pop()
        if spec is None:
            args = [v for v in b.values()]
            reads, writes = [], []
            for k, (v, line, n) in b.items():
                for r in self.path_refs(v, line, k, workspace=self.workspace):
                    if "ref" in r:
                        (writes if re.match(r"^(out|arg1$)", k) else reads).append(
                            dict(r, mode="create") if re.match(r"^(out|arg1$)", k) else r)
            self.emit(node, api, name, "arcpy", "name_heuristic", reads=reads + fr.reads,
                      writes=[w for w in writes] + fr.writes, state="structurally_inferred",
                      flags=["tool_not_in_table"],
                      params={k: self.pjson(v[0], v[1], v[2]) for k, v in b.items()})
            return V("unknown", "arcpy_call")
        reads, writes = list(fr.reads), list(fr.writes)
        params = {}
        for pname, (v, line, n) in b.items():
            refs = None
            if pname in spec.get("inputs", []) or pname in spec.get("in_place", []):
                refs = self.path_refs(v, line, pname, force_path=v.k != "const" or isinstance(v.v, str),
                                      workspace=self.workspace)
                if pname in spec.get("multi_inputs", []) and is_str(v) and ";" in v.v:
                    refs = []
                    for part in v.v.split(";"):
                        refs += self.path_refs(V("const", part.strip()), line, pname, force_path=True,
                                               workspace=self.workspace)
                for r in refs:
                    if "ref" in r or "memory" in r:
                        reads.append(r)
                        if pname in spec.get("in_place", []) and "ref" in r:
                            writes.append(dict(r, mode="in_place"))
                    elif "unresolved" in r:
                        self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
            elif pname in spec.get("outputs", []):
                refs = self.path_refs(v, line, pname, force_path=True, workspace=self.workspace)
                for r in refs:
                    if "ref" in r:
                        writes.append(dict(r, mode="create"))
                    elif "unresolved" in r:
                        self.unresolved(node, r["reason"], expr=str(r["unresolved"]))
            params[pname] = self.pjson(v, line, n, refs)
        if "out_path" in b and "out_name" in b:
            op_, on = b["out_path"][0], b["out_name"][0]
            if is_str(op_) and is_str(on):
                sep = "\\" if "\\" in op_.v else "/"
                for r in self.path_refs(V("const", op_.v.rstrip("\\/") + sep + on.v), node.lineno, "out_name",
                                        force_path=True):
                    if "ref" in r:
                        writes.append(dict(r, mode="create"))
        mem_out = f"{self.mem_id(node.lineno)}:{node.col_offset}" if spec.get("returns") == "Raster" else None
        key = self.emit(node, api, name, "arcpy", "arcpy_tools.json", reads=reads, writes=writes,
                        params=params, mem_out=mem_out)
        creates_fields = name in ("management.AddField", "management.AddFields", "management.CalculateField",
                                  "management.CalculateFields", "management.AlterField", "management.DeleteField")
        first_in = next((r for r in reads if r.get("ref")), None)
        if key and first_in and not creates_fields:
            for pname, (v, line, n) in b.items():
                if re.search(r"(^|_)fields?$", pname) and is_str(v) and v.v:
                    for f in v.v.split(";"):
                        rec = {"ref": first_in["ref"], "field": f.strip(), "line": line, "param": pname}
                        if rec not in self.ctx.ops[key].setdefault("fields_read", []):
                            self.ctx.ops[key]["fields_read"].append(rec)
        if key and fr.nested:
            op = self.ctx.ops[key]
            for n in fr.nested:
                if n not in op["nested"]:
                    op["nested"].append(n)
        if spec.get("returns") == "Raster":
            return V("obj", "arcraster", mem=mem_out if key else None, op=key)
        return V("unknown", "arcpy_result")

    def h_arcraster_save(self, node, dotted, base, method):
        b = self.bind(node, ["path"])
        f = b.get("path", (UNK, node.lineno, None))
        refs = self.path_refs(f[0], f[1], "path", force_path=True, workspace=self.workspace)
        reads = []
        if base.x.get("mem"):
            reads.append({"memory": base.x["mem"], "param": "raster", "line": node.lineno})
        for s in base.x.get("sources") or []:
            if isinstance(s, str):
                reads.append({"memory": s, "param": "raster", "line": node.lineno})
        self.emit(node, "arcpy.Raster.save", "save", "arcpy", "arcpy_tools.json", reads=reads,
                  writes=[dict(r, mode="create") for r in refs if "ref" in r],
                  params={"path": self.pjson(f[0], f[1], f[2], refs)})
        if refs and refs[0].get("ref"):
            base.x["ref"] = refs[0]["ref"]
            base.x["mem"] = None
            base.x["sources"] = []
            if isinstance(node.func, ast.Attribute) and isinstance(node.func.value, ast.Name):
                self.env[node.func.value.id] = V("obj", "arcraster", self.env.get(node.func.value.id, base).chain,
                                                 ref=refs[0]["ref"], mem=None)
        return V("const", None)


# ---------------------------------------------------------------------------------------------------------------
# comments (tokenize) and notebooks
# ---------------------------------------------------------------------------------------------------------------
def read_comments(src: str) -> list[dict]:
    out = []
    try:
        toks = list(tokenize.generate_tokens(io.StringIO(src).readline))
    except (tokenize.TokenError, IndentationError, SyntaxError):
        return out
    code_lines = set()
    for t in toks:
        if t.type not in (tokenize.COMMENT, tokenize.NL, tokenize.NEWLINE, tokenize.INDENT, tokenize.DEDENT,
                          tokenize.ENDMARKER, tokenize.ENCODING):
            for ln in range(t.start[0], t.end[0] + 1):
                code_lines.add(ln)
    for t in toks:
        if t.type == tokenize.COMMENT:
            text = t.string[1:].strip()
            out.append({"line": t.start[0], "text": text, "own_line": t.start[0] not in code_lines})
    return out


def commented_code(comments: list[dict]) -> list[dict]:
    out = []
    for c in comments:
        t = c["text"].strip()
        if not t or t.startswith(("!", "%", "-*-")) or len(t) > 300:
            continue
        try:
            tree = ast.parse(t)
        except SyntaxError:
            continue
        if len(tree.body) == 1 and isinstance(tree.body[0], (ast.Assign, ast.AugAssign, ast.Import, ast.ImportFrom)) \
                or (len(tree.body) == 1 and isinstance(tree.body[0], ast.Expr) and isinstance(tree.body[0].value,
                                                                                               ast.Call)):
            rec = {"line": c["line"], "text": t}
            st = tree.body[0]
            if isinstance(st, ast.Assign) and isinstance(st.value, ast.Constant) and isinstance(st.value.value, str) \
                    and len(st.targets) == 1 and isinstance(st.targets[0], ast.Name):
                rec["assigns"] = {"name": st.targets[0].id, "value": st.value.value}
            out.append(rec)
    return out


def stmt_starts(src: str) -> list[int]:
    try:
        tree = ast.parse(src)
    except SyntaxError:
        return []
    return sorted({n.lineno for n in ast.walk(tree) if isinstance(n, ast.stmt)})


def notebook_units(full: Path) -> tuple[list[dict], dict]:
    """Code cells with magics blanked in place (line numbers kept) + notebook facts."""
    try:
        nb = json.loads(full.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        return [], {"parse_error": str(exc)}
    cells = nb.get("cells", [])
    units, magics, outputs_text, errors = [], [], [], []
    for i, c in enumerate(cells):
        if c.get("cell_type") != "code":
            continue
        src = c.get("source", "")
        src = "".join(src) if isinstance(src, list) else str(src)
        lines = src.split("\n")
        for j, ln in enumerate(lines):
            s = ln.lstrip()
            if s.startswith(("%", "!")) or re.match(r"^\w+\s*=\s*!", s):
                kind = "shell" if s.startswith("!") or "=" in s.split("!")[0] else (
                    "cell_magic" if s.startswith("%%") else "line_magic")
                magics.append({"cell": i, "line": j + 1, "text": s, "kind": kind})
                lines[j] = ""
            elif "get_ipython()" in s:
                magics.append({"cell": i, "line": j + 1, "text": s, "kind": "get_ipython"})
                lines[j] = ""
        units.append({"cell": i, "src": "\n".join(lines), "exec": c.get("execution_count"),
                      "has_code": bool(src.strip()), "raw": src})
        for o in c.get("outputs", []) or []:
            if o.get("output_type") == "error":
                errors.append({"cell": i, "ename": o.get("ename")})
            txt = o.get("text") or (o.get("data") or {}).get("text/plain")
            if txt:
                outputs_text.append({"cell": i, "text": "".join(txt) if isinstance(txt, list) else str(txt)})
    execs = [u["exec"] for u in units if u["exec"] is not None]
    facts = {"exec_counts": [u["exec"] for u in units], "magics": magics, "errors": errors,
             "monotonic": execs == sorted(execs) and len(set(execs)) == len(execs),
             "duplicates": sorted(k for k, n in Counter(execs).items() if n > 1),
             "gaps": bool(execs) and sorted(execs) != list(range(min(execs), min(execs) + len(execs))),
             "unexecuted_code_cells": [u["cell"] for u in units if u["exec"] is None and u["has_code"]],
             "outputs_text": outputs_text[:50]}
    return units, facts


# ---------------------------------------------------------------------------------------------------------------
# documents
# ---------------------------------------------------------------------------------------------------------------
def doc_text(full: Path) -> list[str]:
    if full.suffix.lower() == ".docx":
        try:
            with zipfile.ZipFile(full) as z:
                xml = z.read("word/document.xml").decode("utf-8", "replace")
            paras = re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)
            return [re.sub(r"<[^>]+>", "", p) for p in paras]
        except (OSError, KeyError, zipfile.BadZipFile):
            return []
    raw = full.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc).splitlines()
        except UnicodeDecodeError:
            continue
    return []


def mine_doc(rel: str, lines: list[str], all_paths: list[str]) -> dict:
    by_base: dict[str, list[str]] = defaultdict(list)
    for p in all_paths:
        by_base[p.split("/")[-1].lower()].append(p)
    path_re = re.compile(r"[\w\-./\\:]+\.(py|ipynb|tif|tiff|gpkg|geojson|shp|xlsx|xls|csv|qgz|qgs|gdb|json|txt|model3|"
                         r"aprx|vrt|img|nc|parquet)\b", re.I)
    mentions, run_order, dont_run, constants, mults, cell_sizes, manual, outputs = [], [], [], [], [], [], [], []
    in_run = False
    section = None
    for i, raw in enumerate(lines, start=1):
        line = raw.rstrip()
        low = line.lower()
        if re.search(r"run order|order of (the )?scripts|how to run|steps:", low):
            in_run = True
            section = "run"
        elif not line.strip():
            in_run = False
            section = None
        elif re.match(r"^\S", line) and section == "run" and not path_re.search(line):
            in_run = False
            section = None
        if re.search(r"\bmanual\b|by hand|in the attribute table", low) and section != "manual":
            section = "manual"
            manual.append({"line": i, "text": line.strip(), "header": True})
        elif section == "manual" and line.strip():
            manual.append({"line": i, "text": line.strip()})
        hits = []
        for m in path_re.finditer(line):
            tok = m.group(0).strip("./").replace("\\", "/")
            cand = [p for p in all_paths if p.lower() == tok.lower() or p.lower().endswith("/" + tok.lower())]
            if not cand:
                cand = by_base.get(tok.split("/")[-1].lower(), [])
            hits.append({"line": i, "text": m.group(0), "resolved": cand[0] if len(cand) == 1 else None,
                         "candidates": cand if len(cand) > 1 else []})
        mentions.extend(hits)
        if in_run:
            for h in hits:
                if h["resolved"] and h["resolved"].lower().endswith((".py", ".ipynb", ".model3")) and \
                        not SUPERSEDED_TEXT_RE.search(low):
                    run_order.append({"path": h["resolved"], "line": i})
        if SUPERSEDED_TEXT_RE.search(low):
            for h in hits:
                if h["resolved"]:
                    dont_run.append({"path": h["resolved"], "line": i, "text": line.strip()})
        for m in re.finditer(r"\b([A-Z][A-Z0-9_]{2,})(?:\s*[:=]\s*|\s+)(-?\d+(?:\.\d+)?)\b(?:\s*([A-Za-z%²/]+\d?))?", line):
            constants.append({"name": m.group(1), "value": float(m.group(2)) if "." in m.group(2) else int(m.group(2)),
                              "unit": m.group(3), "line": i, "text": line.strip()})
        for m in re.finditer(r"(?:^|[\s(])[x×*]\s?(\d+(?:\.\d+)?)\b", line):
            v = float(m.group(1))
            mults.append({"value": int(v) if v.is_integer() else v, "line": i, "text": line.strip()})
        for m in re.finditer(r"(\d+(?:\.\d+)?)\s?(m|km|metre|meter|metres|meters)\b(?![²2])", line):
            ctx_words = re.search(r"cell|grid|resolution|pixel|res\b|raster", low)
            if ctx_words:
                v = float(m.group(1)) * (1000.0 if m.group(2) == "km" else 1.0)
                cell_sizes.append({"value_m": v, "line": i, "text": line.strip()})
        if re.search(r"^\s*outputs?\b|\boutputs?\s*:", low) and hits:
            for h in hits:
                if h["resolved"]:
                    outputs.append({"path": h["resolved"], "line": i})
        elif outputs and hits and outputs[-1]["line"] == i - 1:
            for h in hits:
                if h["resolved"]:
                    outputs.append({"path": h["resolved"], "line": i})
    # expand "x.gpkg/.csv/.xlsx" shorthand in output lines
    for i, raw in enumerate(lines, start=1):
        for m in re.finditer(r"([\w\-]+)\.(\w+)((?:/\.\w+)+)", raw):
            stem = m.group(1)
            exts = [m.group(2)] + [e.lstrip("/.") for e in m.group(3).split("/") if e.strip(".")]
            for e in exts:
                for p in by_base.get(f"{stem}.{e}".lower(), []):
                    if any(o["line"] in (i, i - 1) for o in outputs) and not any(o["path"] == p for o in outputs):
                        outputs.append({"path": p, "line": i})
                    if not any(mm["resolved"] == p and mm["line"] == i for mm in mentions):
                        mentions.append({"line": i, "text": f"{stem}.{e}", "resolved": p, "candidates": []})
    return {"path": rel, "kind": "notes", "lines": len(lines), "mentions": mentions, "run_order": run_order,
            "dont_run": dont_run, "constants": constants, "multipliers": mults, "cell_sizes": cell_sizes,
            "manual": manual, "outputs": outputs}


# ---------------------------------------------------------------------------------------------------------------
# QGIS / ArcGIS Pro projects and models
# ---------------------------------------------------------------------------------------------------------------
def redact(s: str) -> str:
    s = re.sub(r"(password|pwd|passwd)\s*=\s*('[^']*'|\"[^\"]*\"|\S+)", r"\1=<redacted:password>", s, flags=re.I)
    return re.sub(r"(user(name)?)\s*=\s*('[^']*'|\"[^\"]*\"|\S+)", r"\1=<redacted:user>", s, flags=re.I)


def parse_subset(s: str | None) -> list[dict]:
    if not s:
        return []
    out = []
    for m in re.finditer(r'"?(\w+)"?\s*(=|!=|<>|>=|<=|>|<|\bin\b|\blike\b)\s*(\'[^\']*\'|\([^)]*\)|-?[\d.]+)', s, re.I):
        out.append({"field": m.group(1), "op": m.group(2).lower(), "value": m.group(3).strip("'")})
    return out


def parse_qgis_project(rel: str, data: bytes, res: Resolver) -> dict:
    from defusedxml import ElementTree as DET
    root = DET.fromstring(data)
    proj_dir = rel.rsplit("/", 1)[0] if "/" in rel else ""
    crs = root.find("./projectCrs/spatialrefsys")
    pcrs = {"authid": (crs.findtext("authid") if crs is not None else None),
            "description": (crs.findtext("description") if crs is not None else None)}
    absolute = (root.findtext("./properties/Paths/Absolute") or "").strip().lower() in ("true", "1")
    groups: dict[str, dict] = {}

    def walk_tree(el, path):
        for ch in el:
            if ch.tag == "layer-tree-group":
                walk_tree(ch, path + [ch.get("name")])
            elif ch.tag == "layer-tree-layer":
                groups[ch.get("id")] = {"group": "/".join(p for p in path if p) or None,
                                        "visible": ch.get("checked") == "Qt::Checked"}
    lt = root.find("./layer-tree-group")
    if lt is not None:
        walk_tree(lt, [])
    layers = []
    for ml in root.iter("maplayer"):
        ltype = ml.get("type")
        name = ml.findtext("layername")
        ds = ml.findtext("datasource") or ""
        if ltype == "annotation" or not ds.strip():
            continue
        lid = ml.findtext("id")
        provider = ml.findtext("provider")
        uri = parse_uri(ds)
        rec = {"name": name, "type": ltype, "provider": provider, "datasource": redact(ds),
               "layer": uri["layer"], "subset": uri["subset"], "subset_parsed": parse_subset(uri["subset"]),
               "crs": ml.findtext("./srs/spatialrefsys/authid"), **groups.get(lid, {"group": None, "visible": None})}
        if provider in ("postgres", "oracle", "mssql", "wms", "wfs", "xyz", "arcgisfeatureserver"):
            rec["remote"] = True
        else:
            r = res.resolve(uri["path"], proj_dir)
            if r and r.get("rel"):
                rec["path"] = r["rel"]
                rec["exists"] = bool(r.get("exists"))
                if r.get("layer"):
                    rec["layer"] = r["layer"]
            else:
                rec["path"] = None
                rec["unresolved"] = True
        layers.append(rec)
    models = []
    for m in root.iter("Option"):
        if m.get("name") == "alg_id":
            models.append(m.get("value"))
    layers.sort(key=lambda l: (l.get("group") or "", l.get("name") or ""))
    return {"path": rel, "kind": "qgis_project", "crs": pcrs, "absolute_paths": absolute, "layers": layers,
            "groups": sorted({l["group"] for l in layers if l.get("group")}), "embedded_alg_ids": sorted(set(models)),
            "annotations_ignored": sum(1 for ml in root.iter("maplayer") if ml.get("type") == "annotation")}


EXPR_PARAM_NAMES = {"FORMULA", "EXPRESSION", "GROUP_BY", "FILTER", "WHERE", "SQL", "QUERY", "AGGREGATES",
                    "FIELDS_MAPPING"}
NON_PATH_TYPES = {"expression", "raster_calc_expression", "string", "number", "enum", "boolean", "field", "aggregates",
                  "fields_mapping", "crs", "extent", "distance", "band", "matrix", "point", "scale", "duration",
                  "color", "area"}


def alg_spec(algs: dict, alg_id: str | None) -> dict:
    """qgis_algs.json entry; a QGIS 3 `qgis:` id that QGIS 4 moved to `native:` resolves to the native entry."""
    if not alg_id:
        return {}
    if alg_id in algs:
        return algs[alg_id]
    if ":" in alg_id:
        return algs.get("native:" + alg_id.split(":", 1)[1], {})
    return {}


def is_expression_param(pname: str, ps: dict) -> bool:
    """FORMULA / EXPRESSION / string-like parameter values are never datasets, whatever they contain (B5)."""
    return (ps.get("type") in NON_PATH_TYPES) if ps else pname.upper() in EXPR_PARAM_NAMES


def name_tokens(text: str) -> set[str]:
    """Lower-case word tokens of a name: snake/kebab/space and camelCase split, digits dropped, plus the acronym of
    a camelCase name (OpenStreetMap -> osm)."""
    parts = re.findall(r"[A-Z]+(?=[A-Z][a-z])|[A-Z]?[a-z]+|[A-Z]+", text or "")
    toks = {p.lower() for p in parts if len(p) > 1}
    for w in re.split(r"[^A-Za-z0-9]+", text or ""):
        if w and not w.isdigit():
            toks.add(re.sub(r"\d+", "", w.lower()))
            caps = re.findall(r"[A-Z][a-z]+", w)
            if len(caps) >= 3:
                toks.add("".join(c[0] for c in caps).lower())
    return {t for t in toks if len(t) > 1}


def token_near(a: str, b: str) -> bool:
    """Weak token match: one is a prefix of the other, or the first three letters of one occur in the other."""
    if min(len(a), len(b)) < 3:
        return False
    return a.startswith(b) or b.startswith(a) or a[:3] in b or b[:3] in a


def parse_tbx_strings(full: Path, rel: str, tools: dict) -> dict:
    """ArcGIS .tbx (OLE compound file, no reader in the env): a partial parse of its UTF-16 strings (M6). Keeps the
    tool display names that match arcpy_tools.json (`management.CalculateField` -> "Calculate Field"), path-like
    strings and every distinct short identifier-like string (field names), as evidence only."""
    try:
        data = full.read_bytes()
    except OSError as exc:
        return {"path": rel, "kind": "arcgis_toolbox", "parse": "failed", "error": str(exc)[:200]}
    strs = [m.group(0).decode("utf-16-le") for m in re.finditer(rb"(?:[\x20-\x7e]\x00){4,}", data)]
    display = {}
    for k in tools:
        name = k.split(".")[-1]
        display[" ".join(re.findall(r"[A-Z]+(?![a-z])|[A-Z][a-z0-9]*|[a-z0-9]+", name))] = k
    found = Counter(x for x in strs if x in display)
    paths = sorted({x for x in strs if ("\\" in x or x.lower().endswith((".gdb", ".shp", ".tif"))) and len(x) < 200})
    idents = Counter(x for x in strs if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{2,40}", x))
    return {"path": rel, "kind": "arcgis_toolbox", "parse": "partial", "method": "utf16_strings",
            "strings": len(strs), "tools": [{"display": d, "tool": display[d], "count": n} for d, n in sorted(found.items())],
            "paths": paths[:60], "paths_total": len(paths), "identifiers": sorted(idents)}


def containment_flags(tree: ast.Module, cell) -> list[dict]:
    """`reversed_containment` (M8): inside a loop over a constant list of strings, `item.__contains__(arg)` or
    `arg in item` tests whether the ARGUMENT is a substring of each list item. An empty argument matches the first
    item of the first such loop, so a missing value falls into that loop's group (recorded as `empty_matches`)."""
    out = []
    scopes = [n for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))] + [tree]
    for sc in scopes:
        lists, first = {}, None
        body = sc.body
        for st in ast.walk(sc) if sc is not tree else body:
            if isinstance(st, ast.Assign) and isinstance(st.value, (ast.List, ast.Tuple)) and st.value.elts and \
                    all(isinstance(e, ast.Constant) and isinstance(e.value, str) for e in st.value.elts):
                for t in st.targets:
                    if isinstance(t, ast.Name):
                        lists[t.id] = [e.value for e in st.value.elts]
        loops = [n for n in (ast.walk(sc) if sc is not tree else body) if isinstance(n, ast.For)
                 and isinstance(n.target, ast.Name)]
        for lp in sorted(loops, key=lambda n: n.lineno):
            items = lists.get(lp.iter.id) if isinstance(lp.iter, ast.Name) else (
                [e.value for e in lp.iter.elts] if isinstance(lp.iter, (ast.List, ast.Tuple)) and lp.iter.elts and all(
                    isinstance(e, ast.Constant) and isinstance(e.value, str) for e in lp.iter.elts) else None)
            if not items:
                continue
            var = lp.target.id
            for n in ast.walk(lp):
                arg = None
                if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr == "__contains__" \
                        and isinstance(n.func.value, ast.Name) and n.func.value.id == var and n.args:
                    arg = n.args[0]
                elif isinstance(n, ast.Compare) and len(n.ops) == 1 and isinstance(n.ops[0], ast.In) and \
                        isinstance(n.comparators[0], ast.Name) and n.comparators[0].id == var:
                    arg = n.left
                if arg is None or isinstance(arg, ast.Constant):
                    continue
                ret = next((r.value.value for r in ast.walk(lp) if isinstance(r, ast.Return)
                            and isinstance(r.value, ast.Constant)), None)
                rec = {"kind": "reversed_containment", "line": n.lineno, "cell": cell, "loop_line": lp.lineno,
                       "list": lp.iter.id if isinstance(lp.iter, ast.Name) else None, "items": len(items),
                       "first_item": items[0], "arg": ast.unparse(arg), "returns": ret,
                       "function": sc.name if sc is not tree else None,
                       "note": f"tests whether {ast.unparse(arg)} is a substring of each item; an empty value matches "
                               f"{items[0]!r}"}
                if first is None:
                    first = rec
                    rec["empty_matches"] = ret if ret is not None else items[0]
                out.append(rec)
                break
    return out


def expression_functions(expr: str) -> list[str]:
    """Function names called in a QGIS expression (outside quoted strings and "field" names)."""
    bare = re.sub(r"'(?:[^']|'')*'|\"(?:[^\"]|\"\")*\"", " ", expr or "")
    return list(dict.fromkeys(m.group(1) for m in re.finditer(r"(?<![\w.])([A-Za-z_]\w*)\s*\(", bare)))


AREA_DIVISORS = {10000: "ha", 1000000: "km2"}


def cell_area_literal(node) -> dict | None:
    """`(c ** 2) / 10000`, `c * c / 10000` (also / 1e6): the area of a nominal c-metre cell in ha (km2)."""
    if not (isinstance(node, ast.BinOp) and isinstance(node.op, ast.Div) and isinstance(node.right, ast.Constant)
            and node.right.value in AREA_DIVISORS):
        return None
    sq = node.left
    if isinstance(sq, ast.BinOp) and isinstance(sq.left, ast.Constant) and isinstance(sq.left.value, (int, float)):
        c = sq.left.value
        if (isinstance(sq.op, ast.Pow) and isinstance(sq.right, ast.Constant) and sq.right.value == 2) or \
                (isinstance(sq.op, ast.Mult) and isinstance(sq.right, ast.Constant) and sq.right.value == c):
            return {"cell_m": c, "unit": AREA_DIVISORS[node.right.value], "area": c * c / node.right.value}
    return None


def cell_area_expression(expr: str) -> dict | None:
    """The same pattern in a QGIS expression string (`(30*30)/10000`, `30^2/10000`)."""
    m = re.search(r"\(?\s*(\d+(?:\.\d+)?)\s*(?:\*\s*(\d+(?:\.\d+)?)|(?:\^|\*\*)\s*2)\s*\)?\s*/\s*(10000|1000000)\b",
                  expr or "")
    if not m or (m.group(2) and float(m.group(2)) != float(m.group(1))):
        return None
    c = float(m.group(1))
    c = int(c) if c.is_integer() else c
    d = int(m.group(3))
    return {"cell_m": c, "unit": AREA_DIVISORS[d], "area": c * c / d}


def expression_numbers(expr: str) -> list[int | float]:
    """Numeric literals of a QGIS expression, in order, outside quoted strings and field names."""
    bare = re.sub(r"'(?:[^']|'')*'|\"(?:[^\"]|\"\")*\"", " ", expr or "")
    out = []
    for m in re.finditer(r"(?<![\w.])(\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?(?![\w.])", bare):
        t = m.group(0)
        out.append(float(t) if any(c in t for c in ".eE") else int(t))
    return out


def xml_with_lines(data: bytes):
    """defusedxml parse that records each element's start line as the attribute `__line` (no lxml in the env)."""
    import xml.etree.ElementTree as ET
    from defusedxml.ElementTree import DefusedXMLParser

    class LineParser(DefusedXMLParser):
        def _start(self, tag, attrs):
            el = super()._start(tag, attrs)
            el.set("__line", str(self.parser.CurrentLineNumber))
            return el
    p = LineParser(target=ET.TreeBuilder())
    p.feed(data)
    return p.close()


def _opt(el, name: str):
    return next((c for c in el if c.tag == "Option" and c.get("name") == name), None) if el is not None else None


def _line(el) -> int | None:
    return int(el.get("__line")) if el is not None and el.get("__line") else None


MODEL_DEST_TYPES = {"sink", "vectorDestination", "rasterDestination", "fileDestination", "folderDestination",
                    "vectorTileDestination", "pointCloudDestination"}


def parse_model3(rel: str, data: bytes) -> dict:
    """QGIS model (.model3, Option XML): child algorithms with static parameter values (and their XML lines), model
    input parameters (type, geometry types) and model outputs (child, output name)."""
    root = xml_with_lines(data)

    def opt_value(el):
        t = el.get("type")
        if t == "Map":
            return {c.get("name"): opt_value(c) for c in el if c.tag == "Option"}
        if t in ("List", "StringList"):
            return [opt_value(c) for c in el if c.tag == "Option"]
        return el.get("value")
    top = root.find("Option") if root.tag != "Option" else root
    doc = opt_value(top) if top is not None else {}
    children, outputs = [], []
    ch_el = _opt(top, "children")
    for cid, ch in sorted((doc.get("children") or {}).items()):
        if not isinstance(ch, dict):
            continue
        cel = _opt(ch_el, cid)
        pel = _opt(cel, "params")
        params = {}
        for pname, pv in (ch.get("params") or {}).items():
            vals = pv if isinstance(pv, list) else [pv]
            pe = _opt(pel, pname)
            items = [c for c in pe if c.tag == "Option"] if pe is not None and pe.get("type") == "List" else [pe]
            for v, ve in zip(vals, items + [None] * len(vals)):
                if isinstance(v, dict):
                    src = v.get("source")
                    line = _line(_opt(ve, "static_value")) or _line(_opt(ve, "parameter_name")) or _line(ve) \
                        or _line(pe)
                    params[pname] = {"static_value": v.get("static_value"), "source": src,
                                     "parameter_name": v.get("parameter_name"), "child_id": v.get("child_id"),
                                     "output_name": v.get("output_name"), "line": line}
        outs = {}
        oel = _opt(cel, "outputs")
        for oname, ov in (ch.get("outputs") or {}).items():
            if isinstance(ov, dict):
                outs[oname] = {"name": ov.get("name"), "default_value": ov.get("default_value"),
                               "output_name": ov.get("output_name"), "line": _line(_opt(oel, oname))}
                outputs.append({"name": ov.get("name") or oname, "child": cid,
                                "output_name": ov.get("output_name") or "OUTPUT", "line": _line(_opt(oel, oname))})
        lines = [int(x.get("__line")) for x in cel.iter() if x.get("__line")] if cel is not None else []
        children.append({"id": cid, "alg_id": ch.get("alg_id"), "params": params, "outputs": outs,
                         "dependencies": ch.get("dependencies"), "description": ch.get("component_description"),
                         "active": ch.get("active") != "false",
                         "line": _line(cel), "end_line": max(lines) if lines else _line(cel)})
    defs = doc.get("parameterDefinitions") or {}
    dfel = _opt(top, "parameterDefinitions")
    input_defs = {}
    for name, d in sorted(defs.items()):
        if not isinstance(d, dict) or d.get("parameter_type") in MODEL_DEST_TYPES:
            continue
        dt = d.get("data_types")
        input_defs[name] = {"type": d.get("parameter_type"), "description": d.get("description"),
                            "data_types": [int(x) for x in dt if str(x).lstrip("-").isdigit()]
                            if isinstance(dt, list) else None, "line": _line(_opt(dfel, name))}
    inputs = sorted((doc.get("parameterDefinitions") or doc.get("parameters") or {}).keys())
    return {"path": rel, "kind": "qgis_model", "name": doc.get("model_name"), "group": doc.get("model_group"),
            "children": children, "inputs": inputs, "input_defs": input_defs,
            "outputs": sorted(outputs, key=lambda o: (o["child"], o["name"]))}


def parse_aprx(rel: str, full: Path, res: Resolver) -> dict:
    """ArcGIS Pro .aprx/.atbx (zip of CIM JSON), best effort: GISProject.json, follow CIMPATH= to map and layer JSON."""
    out = {"path": rel, "kind": "arcgis_pro_project" if rel.lower().endswith(".aprx") else "arcgis_toolbox",
           "maps": [], "layers": [], "tools": [], "unparsed": []}
    proj_dir = rel.rsplit("/", 1)[0] if "/" in rel else ""
    try:
        z = zipfile.ZipFile(full)
    except zipfile.BadZipFile:
        out["unparsed"].append("not_a_zip (pre-2.x binary or corrupt)")
        return out
    names = set(z.namelist())

    def load(n):
        try:
            return json.loads(z.read(n).decode("utf-8-sig"))
        except (KeyError, ValueError):
            return None
    if out["kind"] == "arcgis_toolbox":
        tb = load("toolbox.content") or {}
        out["tools"] = sorted((tb.get("toolsets") or {}).keys()) + sorted(tb.get("tools", []) if isinstance(
            tb.get("tools"), list) else [])
        for n in sorted(names):
            if n.endswith("tool.content"):
                out["tools"].append(n.split("/")[0])
            if n.endswith(".py") or n.endswith("tool.script.execute.py"):
                out.setdefault("scripts", []).append(n)
        return out
    proj = load("GISProject.json")
    if proj is None:
        out["unparsed"].append("GISProject.json missing")
        return out

    def cimpaths(obj):
        if isinstance(obj, str) and obj.startswith("CIMPATH="):
            yield obj[len("CIMPATH="):]
        elif isinstance(obj, dict):
            for v in obj.values():
                yield from cimpaths(v)
        elif isinstance(obj, list):
            for v in obj:
                yield from cimpaths(v)
    seen = set()
    queue = sorted(set(cimpaths(proj)))
    while queue:
        p = queue.pop(0)
        if p in seen:
            continue
        seen.add(p)
        if p not in names:
            out["unparsed"].append(f"CIMPATH {p} not in archive")
            continue
        doc = load(p)
        if not isinstance(doc, dict):
            out["unparsed"].append(f"{p} not JSON")
            continue
        t = doc.get("type", "")
        if t == "CIMMapDocument" or "mapDefinition" in doc:
            md = doc.get("mapDefinition", doc)
            sr = (md.get("spatialReference") or {})
            out["maps"].append({"cim": p, "name": md.get("name"), "wkid": sr.get("latestWkid") or sr.get("wkid")})
        dc = doc.get("featureTable", {}).get("dataConnection") if isinstance(doc.get("featureTable"), dict) else None
        dc = dc or doc.get("dataConnection")
        if dc:
            wcs = dc.get("workspaceConnectionString", "")
            m = re.search(r"DATABASE=([^;]+)", wcs)
            wpath = m.group(1) if m else None
            lay = {"cim": p, "name": doc.get("name"), "workspace": redact(wcs), "factory": dc.get("workspaceFactory"),
                   "dataset": dc.get("dataset"), "dataset_type": dc.get("datasetType")}
            sr = (doc.get("spatialReference") or {})
            if sr:
                lay["wkid"] = sr.get("latestWkid") or sr.get("wkid")
            if wpath and dc.get("dataset"):
                sep = "\\" if "\\" in wpath else "/"
                target = wpath.rstrip("\\/") + sep + dc["dataset"]
                if dc.get("workspaceFactory") in ("Shapefile", "Raster") and not re.search(r"\.\w+$", dc["dataset"]):
                    target += ".shp" if dc.get("workspaceFactory") == "Shapefile" else ""
                r = res.resolve(target, proj_dir)
                if r and r.get("rel"):
                    lay["path"] = r["rel"]
                    lay["layer"] = r.get("layer")
            out["layers"].append(lay)
        for q in sorted(set(cimpaths(doc))):
            if q not in seen:
                queue.append(q)
    return out


# ---------------------------------------------------------------------------------------------------------------
# data profiles
# ---------------------------------------------------------------------------------------------------------------
def laea_for(lon: float, lat: float):
    from pyproj import CRS
    return CRS.from_proj4(f"+proj=laea +lat_0={lat:.6f} +lon_0={lon:.6f} +x_0=0 +y_0=0 +datum=WGS84 +units=m +no_defs")


def aux_xml_stats(full: Path) -> dict | None:
    side = Path(str(full) + ".aux.xml")
    if not side.exists():
        return None
    try:
        from defusedxml import ElementTree as DET
        root = DET.fromstring(side.read_bytes())
    except Exception:  # noqa: BLE001
        return {"parse_error": True}
    out = {}
    for mdi in root.iter("MDI"):
        k = mdi.get("key") or ""
        if k.startswith("STATISTICS_"):
            v = (mdi.text or "").strip()
            try:
                out[k[len("STATISTICS_"):].lower()] = float(v)
            except ValueError:
                out[k[len("STATISTICS_"):].lower()] = v
    if "approximate" in out:
        out["approximate"] = str(out["approximate"]).upper() == "YES"
    out["source"] = "aux.xml"
    return out


def profile_raster(full: Path, stats: bool = True) -> dict:
    import numpy as np
    import rasterio
    from pyproj import CRS, Geod
    with rasterio.open(full) as src:
        crs = src.crs
        t = src.transform
        prof: dict[str, Any] = {"driver": src.driver, "bands": src.count, "shape": [src.height, src.width],
                                "transform": [t.a, t.b, t.c, t.d, t.e, t.f], "res": [t.a, t.e],
                                "origin": [t.c, t.f], "bounds": list(src.bounds), "dtype": src.dtypes[0],
                                "dtypes": list(src.dtypes), "nodata": src.nodata,
                                "block": list(src.block_shapes[0]) if src.block_shapes else None,
                                "compress": (src.compression.value if src.compression else None),
                                "overviews": src.overviews(1) if src.count else [],
                                "scale_offset": [list(src.scales), list(src.offsets)],
                                "colortable": False}
        try:
            src.colormap(1)
            prof["colortable"] = True
        except ValueError:
            pass
        if crs:
            pc = CRS.from_wkt(crs.to_wkt())
            prof["crs"] = {"wkt": crs.to_wkt(), "epsg": pc.to_epsg(), "is_geographic": pc.is_geographic,
                           "units": pc.axis_info[0].unit_name if pc.axis_info else None}
        else:
            prof["crs"] = None
        dt = np.dtype(src.dtypes[0])
        nd = src.nodata
        if nd is not None:
            try:
                prof["nodata_in_domain"] = bool(np.array([nd]).astype(dt)[0] == nd) or (
                    np.issubdtype(dt, np.floating) and math.isnan(nd))
            except (OverflowError, ValueError):
                prof["nodata_in_domain"] = False
        cx = (prof["bounds"][0] + prof["bounds"][2]) / 2
        cy = (prof["bounds"][1] + prof["bounds"][3]) / 2
        if crs and prof["crs"]["is_geographic"]:
            g = Geod(ellps="WGS84")
            _, _, dx = g.inv(cx, cy, cx + abs(t.a), cy)
            _, _, dy = g.inv(cx, cy, cx, cy + abs(t.e))
            prof["res_m_at_centre"] = [round(dx, 3), round(dy, 3)]
        elif crs:
            prof["res_m_at_centre"] = [abs(t.a), abs(t.e)]
        tags = src.tags(1) if src.count else {}
        emb = {k[len("STATISTICS_"):].lower(): v for k, v in tags.items() if k.startswith("STATISTICS_")}
        for k, v in list(emb.items()):
            try:
                emb[k] = float(v)
            except ValueError:
                pass
        if "approximate" in emb:
            emb["approximate"] = str(emb["approximate"]).upper() == "YES"
        sidecar = aux_xml_stats(full)
        prof["embedded_stats"] = {"tags": emb or None, "aux_xml": sidecar,
                                  "approximate": bool((emb or {}).get("approximate") or (sidecar or {}).get("approximate")),
                                  "trusted": False, "note": "recorded only; every statistic used is computed"}
        if not stats:
            return prof
        # computed statistics (float64 accumulators over block windows; strided above 2 GB or STATS_MAX_CELLS cells:
        # every k-th block, a systematic sample; counts and sums are then those of the sample)
        cells = src.width * src.height
        big = full.stat().st_size > 2 * 1024 ** 3 or cells > STATS_MAX_CELLS
        windows = [w for _, w in src.block_windows(1)] if src.count else []
        if big:
            step = max(1, len(windows) // 2000, math.ceil(cells / STATS_MAX_CELLS))
            windows = windows[::step]
            prof["stats_sample"] = {"every_nth_block": step, "blocks": len(windows), "cells_total": cells}
        n = valid = nodata_n = nan_n = zero_n = 0
        s = 0.0
        mean = m2 = 0.0
        vmin, vmax = math.inf, -math.inf
        isint = np.issubdtype(dt, np.integer)
        uniq: Counter = Counter()
        for w in windows:
            a = src.read(1, window=w)
            n += a.size
            mask = np.ones(a.shape, bool)
            if np.issubdtype(a.dtype, np.floating):
                nanm = np.isnan(a)
                nan_n += int(nanm.sum())
                mask &= ~nanm
            if nd is not None and not (isinstance(nd, float) and math.isnan(nd)):
                ndm = a == np.array(nd).astype(a.dtype)
                nodata_n += int(ndm.sum())
                mask &= ~ndm
            v = a[mask].astype("float64")
            if v.size:
                k = v.size
                bm = float(v.mean())
                bm2 = float(((v - bm) ** 2).sum())
                delta = bm - mean
                tot = valid + k
                mean += delta * k / tot
                m2 += bm2 + delta * delta * valid * k / tot
                valid = tot
                s += float(v.sum())
                zero_n += int((v == 0).sum())
                vmin = min(vmin, float(v.min()))
                vmax = max(vmax, float(v.max()))
                if isint and len(uniq) <= 256:
                    vals, cnts = np.unique(a[mask], return_counts=True)
                    for vv, cc in zip(vals.tolist(), cnts.tolist()):
                        uniq[vv] += cc
        comp = {"cells": n, "valid": valid, "nodata": nodata_n, "nan": nan_n, "zero": zero_n,
                "min": vmin if valid else None, "max": vmax if valid else None, "mean": mean if valid else None,
                "std": math.sqrt(m2 / valid) if valid else None, "sum": s if valid else None}
        prof["computed"] = comp
        prof["stats_mode"] = "strided" if big else "full"
        if isint and 0 < len(uniq) <= 256:
            prof["unique"] = sorted(uniq)
            prof["unique_counts"] = {str(k): v for k, v in sorted(uniq.items())}
        if valid and vmax > vmin:
            hist = np.zeros(64, dtype="int64")
            for w in windows:
                a = src.read(1, window=w).astype("float64")
                m = ~np.isnan(a)
                if nd is not None and not (isinstance(nd, float) and math.isnan(nd)):
                    m &= a != float(np.array(nd).astype(dt))
                h, _ = np.histogram(a[m], bins=64, range=(vmin, vmax))
                hist += h
            prof["hist64"] = {"range": [vmin, vmax], "counts": hist.tolist()}
        # embedded vs computed
        diffs = {}
        for srcname, st in (("tags", emb), ("aux_xml", sidecar or {})):
            if st and isinstance(st.get("mean"), float) and comp["mean"] is not None:
                diffs[srcname] = {"mean_diff": st["mean"] - comp["mean"],
                                  "approximate": st.get("approximate")}
        prof["embedded_vs_computed"] = diffs or None
    return prof


def vrt_alias_source(full: Path, rel: str, known: set[str]) -> str | None:
    """A .vrt with one band whose only source is a sibling file in the legacy tree, read whole and placed whole
    (SrcRect = DstRect = the full raster): an alias of that file (m1). Returns the source's relative path."""
    from defusedxml import ElementTree as DET
    try:
        root = DET.fromstring(full.read_bytes())
    except Exception:  # noqa: BLE001
        return None
    if root.tag != "VRTDataset":
        return None
    bands = root.findall("VRTRasterBand")
    srcs = [x for b in bands for x in b if x.tag in ("SimpleSource", "ComplexSource")]
    if len(bands) != 1 or len(srcs) != 1 or bands[0].get("subClass"):
        return None
    fn = srcs[0].find("SourceFilename")
    if fn is None or fn.get("relativeToVRT") != "1" or not fn.text or "/" in fn.text or "\\" in fn.text:
        return None
    sib = (rel.rsplit("/", 1)[0] + "/" if "/" in rel else "") + fn.text.strip()
    if sib not in known:
        return None
    w, h = root.get("rasterXSize"), root.get("rasterYSize")
    for tag in ("SrcRect", "DstRect"):
        r = srcs[0].find(tag)
        if r is not None and (r.get("xOff", "0") != "0" or r.get("yOff", "0") != "0" or r.get("xSize") != w
                              or r.get("ySize") != h):
            return None
    return sib


def read_layer(full: Path, layer: str | None):
    import pyogrio
    kw = {"IMMUTABLE": "YES"} if full.suffix.lower() in (".gpkg", ".sqlite") else {}
    return pyogrio.read_dataframe(full, layer=layer, **kw)


def profile_vector(full: Path, keep_values: dict, key: str) -> dict:
    import numpy as np
    import pyogrio
    import shapely
    from pyproj import CRS, Transformer
    immutable = full.suffix.lower() in (".gpkg", ".sqlite")
    if full.suffix.lower() == ".gpkg":
        # list layers through SQLite opened immutable: GDAL never opens a legacy GeoPackage writable
        import sqlite3
        con = sqlite3.connect("file:" + str(full.resolve()) + "?mode=ro&immutable=1", uri=True)
        try:
            gcols = dict(con.execute("SELECT table_name, geometry_type_name FROM gpkg_geometry_columns").fetchall())
            layers = [(str(t), str(gcols.get(t, "None"))) for (t,) in con.execute(
                "SELECT table_name FROM gpkg_contents WHERE data_type IN ('features', 'attributes') ORDER BY table_name")]
        finally:
            con.close()
    else:
        layers = [(str(n), str(g)) for n, g in pyogrio.list_layers(full)]
    out = {"driver": None, "layers": []}
    for lname, gtype in layers:
        rec: dict[str, Any] = {"name": lname, "declared_geometry": gtype}
        try:
            info = pyogrio.read_info(full, layer=lname, **({"IMMUTABLE": "YES"} if immutable else {}))
            out["driver"] = info.get("driver")
            df = read_layer(full, lname)
        except Exception as exc:  # noqa: BLE001
            rec["error"] = str(exc)[:200]
            out["layers"].append(rec)
            continue
        rec["count"] = int(len(df))
        crs = getattr(df, "crs", None)
        if crs is not None:
            rec["crs"] = {"wkt": crs.to_wkt(), "epsg": crs.to_epsg(), "is_geographic": crs.is_geographic,
                          "name": crs.name}
        else:
            rec["crs"] = None
        rec["schema"] = [{"name": c, "dtype": str(df[c].dtype)} for c in df.columns if c != "geometry"]
        geoms = df.geometry.values if "geometry" in df.columns else None
        if geoms is not None and len(df):
            arr = np.asarray(geoms)
            types = Counter(str(t) for t in shapely.get_type_id(arr).tolist())
            names = {"-1": "None", "0": "Point", "1": "LineString", "2": "LinearRing", "3": "Polygon",
                     "4": "MultiPoint", "5": "MultiLineString", "6": "MultiPolygon", "7": "GeometryCollection"}
            rec["geom_types"] = {names.get(k, k): v for k, v in sorted(types.items())}
            nulls = shapely.is_missing(arr)
            rec["null_geom"] = int(nulls.sum())
            rec["empty_geom"] = int(shapely.is_empty(arr[~nulls]).sum()) if (~nulls).any() else 0
            valid = shapely.is_valid(arr)
            inv = (~valid) & ~nulls
            reasons = Counter(re.sub(r"\[.*", "", r).strip() for r in shapely.is_valid_reason(arr[inv]).tolist())
            rec["invalid"] = {"count": int(inv.sum()), "top_reasons": dict(reasons.most_common(5))}
            b = shapely.total_bounds(arr[~nulls]) if (~nulls).any() else None
            rec["bounds"] = [float(x) for x in b] if b is not None else None
            polys = shapely.get_type_id(arr) % 3 == 0
            polys &= shapely.get_type_id(arr) > 0
            if crs is not None and b is not None and polys.any():
                to4326 = Transformer.from_crs(crs, 4326, always_xy=True)
                cx, cy = to4326.transform((b[0] + b[2]) / 2, (b[1] + b[3]) / 2)
                tr = Transformer.from_crs(crs, laea_for(cx, cy), always_xy=True)
                g = shapely.transform(arr[polys & valid], lambda xy: np.column_stack(tr.transform(xy[:, 0], xy[:, 1])))
                rec["area_km2"] = float(shapely.area(g).sum() / 1e6)
                rec["area_note"] = "valid polygons only, LAEA centred on the bbox"
            wkb = [hashlib.sha256(shapely.to_wkb(gg, hex=False)).hexdigest() if gg is not None else "null"
                   for gg in arr]
            rec["geometry_multiset_sha256"] = hashlib.sha256("\n".join(sorted(wkb)).encode()).hexdigest()
            keep_values[(key, lname)] = {"wkb": wkb, "cols": {}}
        keys, domains = [], {}
        for c in df.columns:
            if c == "geometry":
                continue
            col = df[c]
            nonnull = int(col.notna().sum())
            uniq = int(col.nunique(dropna=True))
            entry = {"col": c, "unique": uniq == len(df) and len(df) > 0, "nonnull": nonnull == len(df)}
            if col.dtype == object or str(col.dtype).startswith("str"):
                vals = [v for v in col.tolist()]
                sc = {str(v).strip().casefold() for v in vals if v is not None}
                entry["unique_after_strip_casefold"] = len(sc) == len(df)
                if uniq <= 50:
                    domains[c] = sorted(str(v) for v in set(vals) if v is not None)
                if (key, lname) in keep_values and len(df) <= 20000:
                    keep_values[(key, lname)]["cols"][c] = [None if v is None else str(v) for v in vals]
            if entry["unique"] and entry["nonnull"]:
                keys.append(entry)
        rec["key_candidates"] = keys
        rec["domains"] = domains
        out["layers"].append(rec)
    return out


def profile_xlsx(full: Path) -> dict:
    import openpyxl
    wbf = openpyxl.load_workbook(full, data_only=False)
    wbv = openpyxl.load_workbook(full, data_only=True)
    sheets = []
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        rec: dict[str, Any] = {"name": ws.title, "hidden": ws.sheet_state != "visible", "dims": ws.dimensions,
                               "merged": sorted(str(r) for r in ws.merged_cells.ranges)}
        formula_cells, no_cache, units, strip_issues = [], [], [], []
        rows: dict[int, dict[int, Any]] = defaultdict(dict)
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                cached = wv[c.coordinate].value
                rows[c.row][c.column] = cached if not (isinstance(c.value, str) and c.value.startswith("=")) else (
                    "=f" if cached is None else cached)
                if isinstance(c.value, str) and c.value.startswith("="):
                    formula_cells.append({"cell": c.coordinate, "formula": c.value, "cached": jsonable(cached)})
                    if cached is None:
                        no_cache.append(c.coordinate)
                if isinstance(c.value, str):
                    if c.value != c.value.strip():
                        strip_issues.append({"cell": c.coordinate, "value": c.value})
                    if re.search(r"\bunits?\b", c.value, re.I):
                        nxt = ws.cell(row=c.row, column=c.column + 1).value
                        units.append({"cell": c.coordinate, "label": c.value, "value": jsonable(nxt),
                                      "value_cell": ws.cell(row=c.row, column=c.column + 1).coordinate})
        # header candidates: rows of >= 2 text cells followed by rows typed like data
        cands = []
        maxr = max(rows) if rows else 0
        for r in sorted(rows)[:40]:
            cells = rows[r]
            texts = [v for v in cells.values() if isinstance(v, str) and v != "=f"]
            if len(cells) < 2 or len(texts) != len(cells):
                continue
            below, cnum = 0, 0
            for rr in range(r + 1, maxr + 1):
                nxt = rows.get(rr)
                if not nxt:
                    break
                if any(isinstance(nxt.get(col), (int, float)) and not isinstance(nxt.get(col), bool) for col in cells):
                    cnum += 1
                below += 1
            score = len(cells) + 2 * cnum + (1 if below >= 3 else 0)
            reason = f"{len(cells)} text cells, {cnum} typed data rows below"
            if cnum:
                cands.append({"row": r, "score": score, "reason": reason, "data_rows": [r + 1, r + below],
                              "columns": {openpyxl.utils.get_column_letter(k): v for k, v in sorted(cells.items())}})
        cands.sort(key=lambda c: (-c["score"], c["row"]))
        rec.update({"header_candidates": cands[:3], "header_row": None, "header_state": "unresolved",
                    "data_rows": None, "formula_cells": formula_cells[:200], "formula_count": len(formula_cells),
                    "formula_no_cache": no_cache, "units_cells": units, "strip_issues": strip_issues,
                    "max_row": ws.max_row, "max_column": ws.max_column})
        rec["_values"] = {r: {openpyxl.utils.get_column_letter(k): v for k, v in cells.items()}
                          for r, cells in rows.items()}
        sheets.append(rec)
    named = []
    try:
        for name, dn in wbf.defined_names.items():
            named.append({"name": name, "value": str(dn.attr_text)})
    except AttributeError:  # pragma: no cover - older openpyxl
        pass
    ext = [str(getattr(l, "file_link", None) and l.file_link.Target) for l in getattr(wbf, "_external_links", [])]
    return {"format": "xlsx", "sheets": sheets, "named_ranges": named, "external_links": ext}


def profile_csv(full: Path) -> dict:
    raw = full.read_bytes()
    bom = raw.startswith(b"\xef\xbb\xbf")
    enc = "utf-8-sig" if bom else "utf-8"
    try:
        text = raw.decode(enc)
    except UnicodeDecodeError:
        enc = "cp1252"
        text = raw.decode(enc, "replace")
    sample = text[:8192]
    try:
        sep = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        sep = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=sep))
    header = rows[0] if rows else []
    body = rows[1:]
    decimal = "," if sep == ";" and re.search(r"\d,\d", sample) else "."
    code_like = []
    for j, h in enumerate(header):
        vals = [r[j] for r in body if j < len(r)]
        if vals and all(re.match(r"^0\d+$", v) for v in vals if v):
            code_like.append(h)
    return {"format": "csv", "encoding": "utf-8" if enc.startswith("utf-8") else enc, "bom": bom, "sep": sep,
            "decimal": decimal, "header": header, "rows": len(body), "columns": len(header), "code_like_cols": code_like,
            "_rows": body[:20000]}


# ---------------------------------------------------------------------------------------------------------------
# the scan
# ---------------------------------------------------------------------------------------------------------------
def step_ids_for(code_paths: list[str]) -> dict[str, str]:
    out, used = {}, set()
    numbered, rest = [], []
    for p in sorted(code_paths):
        st = stem_of(p)
        m = re.match(r"^(\d+)([a-zA-Z]?)(?:[_\- .]+(.*))?$", st)
        (numbered if m else rest).append((p, m))
    maxn = 0
    for p, m in numbered:
        n = int(m.group(1))
        maxn = max(maxn, n)
        tail = ac.slug(m.group(3) or "") or ac.slug(stem_of(p)) or "step"
        rid = f"s{n:02d}{m.group(2).lower()}_{tail}"
        if rid in used:
            rid += "_" + ac.slug(p.rsplit("/", 1)[0] if "/" in p else "root")
        used.add(rid)
        out[p] = rid
    for i, (p, _) in enumerate(rest, start=1):
        rid = f"s{maxn + i:02d}_{ac.slug(stem_of(p)) or 'code'}"
        if rid in used:
            rid += "_" + ac.slug(p.rsplit("/", 1)[0] if "/" in p else "root")
        used.add(rid)
        out[p] = rid
    return out


def short_step(step_id: str) -> str:
    return step_id.split("_", 1)[0]


class Scan:
    def __init__(self, legacy: Path, out: Path):
        self.root = legacy.resolve()
        self.out = out.resolve()
        self.log = LOG

    # ------------------------------------------------------------------ driver
    def run(self) -> tuple[dict, int]:
        root = self.root
        started = ac.utc_now_iso()
        self.log(f"scan {SCAN_VERSION} of {root.name}")
        entries = walk(root)
        classify_entries(entries)
        for e in entries:
            if e["kind"] == "json":
                e["kind"] = sniff_json(root / e["path"])
            elif e["kind"] == "noext":
                e["kind"] = sniff_noext(root / e["path"])
        before = hash_tree(root, entries)
        self.entries = entries
        self.hashes = before
        self.res = Resolver(root, entries)
        qa, at = load_tables()
        self.ctx = Ctx(root, self.res, qa, at)
        code_entries = [e for e in entries if e["kind"] in ("code", "notebook") and
                        file_ext(e["path"]) in (".py", ".ipynb", ".pyt")]
        other_code = [e for e in entries if e["kind"] == "code" and file_ext(e["path"]) not in (".py", ".ipynb", ".pyt")]
        self.step_of = step_ids_for([e["path"] for e in code_entries] +
                                    [e["path"] for e in entries if file_ext(e["path"]) == ".model3"])
        for e in code_entries:
            if e["path"].endswith(".py"):
                self.ctx.local_modules[stem_of(e["path"])] = e["path"]
        exports: dict[str, dict] = {}
        globals_: dict[str, dict] = {}
        self.nb_facts: dict[str, dict] = {}
        self.analysers: dict[str, PyAnalyzer] = {}
        for pass_no in (1, 2):
            self.ctx.reset()
            for e in code_entries:
                if e["path"].endswith(".py"):
                    self.ctx.local_modules[stem_of(e["path"])] = e["path"]
            for rel, ex in exports.items():
                self.ctx.code[rel] = {"_exports": ex, "_globals": globals_.get(rel, {})}
            for e in code_entries:
                self.analyse_code(e)
            for rel, an in self.analysers.items():
                exports[rel] = {k: v for k, v in an.env.items() if v.k == "const" and k.isupper()}
                # module-level functions of a local module: inlined at the importer's call sites (M5)
                exports[rel].update({k: V("func", v.v, src=v.x.get("src") or rel) for k, v in an.env.items()
                                     if v.k == "func" and (v.x.get("src") or rel) == rel})
                globals_[rel] = {k: v for k, v in an.env.items() if v.k == "mod"}
        self.docs = [mine_doc(e["path"], doc_text(root / e["path"]), [x["path"] for x in entries])
                     for e in entries if e["kind"] == "doc"]
        self.projects = self.parse_projects(entries)
        self.profiles, self.keep_values = self.profile_all(entries)
        self.bind_model_inputs()
        self.build_datasets()
        self.build_ops()
        self.code_records(code_entries, other_code)
        self.step_order()
        self.classify_roles()
        self.collect_branches()
        self.finish_settings()
        self.finish_params()
        self.grids()
        self.derived_copies()
        self.excel_corroboration()
        self.join_coverage()
        self.extent_relations()
        self.toolboxes = [parse_tbx_strings(self.root / e["path"], e["path"], self.ctx.tools) for e in self.entries
                          if file_ext(e["path"]) == ".tbx"]
        self.implementation_groups()
        self.open_items()
        after_entries = walk(root)
        classify_entries(after_entries)
        for e in after_entries:
            if e["kind"] == "json":
                e["kind"] = sniff_json(root / e["path"])
            elif e["kind"] == "noext":
                e["kind"] = sniff_noext(root / e["path"])
        after = hash_tree(root, after_entries)
        ro_ok = after == before and [e["path"] for e in after_entries] == [e["path"] for e in entries]
        changed = sorted({p for p in set(before) | set(after) if before.get(p) != after.get(p)})
        inv = self.assemble(started, ro_ok, changed)
        self.log(f"files {len(inv['files'])}, datasets {len(inv['datasets'])}, ops {len(inv['ops'])}, "
                 f"open items {len(inv['open_items'])}, unresolved {len(inv['unresolved'])}")
        for u in inv["unresolved"]:
            self.log(f"unresolved {u['file']}:{u['line']} {u['reason']}: {u['expr']}")
        self.log(f"read-only verified: {ro_ok}" + (f" (changed: {changed})" if changed else ""))
        return inv, (0 if ro_ok else 3)

    # ------------------------------------------------------------------ code
    def analyse_code(self, e: dict) -> None:
        rel = e["path"]
        sid = self.step_of[rel]
        base_dir = rel.rsplit("/", 1)[0] if "/" in rel else ""
        if base_dir.split("/")[-1] == ".ipynb_checkpoints":  # a checkpoint runs in its notebook's folder
            base_dir = base_dir.rsplit("/", 1)[0] if "/" in base_dir else ""
        an = PyAnalyzer(self.ctx, rel, "code." + sid, base_dir)
        full = self.root / rel
        if rel.endswith(".ipynb"):
            units, facts = notebook_units(full)
            self.nb_facts[rel] = facts
            an.run_units([(u["cell"], u["src"]) for u in units])
            # magics: shell commands become ops, %run a code dependency, %cd changes the base directory
            for mg in facts.get("magics", []):
                txt = mg["text"]
                if mg["kind"] == "shell":
                    cmd = txt.split("!", 1)[1]
                    try:
                        argv = shlex.split(cmd)
                    except ValueError:
                        argv = cmd.split()
                    node = ast.Constant(value=None)
                    node.lineno, node.col_offset, node.end_lineno = mg["line"], 0, mg["line"]
                    an.cell = mg["cell"]
                    an.cli_op(node, argv, api="!shell")
                elif txt.startswith("%run"):
                    an.imports.append({"module": txt.split()[1] if len(txt.split()) > 1 else "?", "line": mg["line"],
                                       "cell": mg["cell"], "via": "%run"})
            # exec-order replay: paths that differ between cell order and exec order are flagged
            order = sorted((u for u in units if u["exec"] is not None), key=lambda u: u["exec"]) + \
                [u for u in units if u["exec"] is None]
            if [u["cell"] for u in order] != [u["cell"] for u in units]:
                saved_ops = dict(self.ctx.ops)
                saved_refs = dict(self.ctx.refs)
                saved_lists = (list(self.ctx.unresolved), list(self.ctx.settings), list(self.ctx.params),
                               list(self.ctx.literals))
                self.ctx.ops = {}
                an2 = PyAnalyzer(self.ctx, rel, "code." + sid, base_dir)
                an2.run_units([(u["cell"], u["src"]) for u in order])
                exec_ops = self.ctx.ops
                self.ctx.ops = saved_ops
                self.ctx.refs.update({k: v for k, v in saved_refs.items()})
                (self.ctx.unresolved, self.ctx.settings, self.ctx.params, self.ctx.literals) = saved_lists
                diffs = []
                for k, op in exec_ops.items():
                    base = self.ctx.ops.get(k)
                    if base is None:
                        continue
                    a = sorted(r.get("ref") or "" for r in base["reads"] + base["writes"])
                    b = sorted(r.get("ref") or "" for r in op["reads"] + op["writes"])
                    if a != b:
                        diffs.append({"cell": op["cell"], "line": op["line"], "cell_order": a, "exec_order": b})
                        if "order_dependent_path" not in base["flags"]:
                            base["flags"].append("order_dependent_path")
                facts["order_dependent_paths"] = diffs
            src_for_comments = {u["cell"]: u["raw"] for u in units}
        else:
            try:
                src = full.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                src = full.read_text(encoding="cp1252", errors="replace")
            an.run_units([(None, src)])
            src_for_comments = {None: src}
        self.analysers[rel] = an
        an.comments = {}
        an.stmt_starts = {}
        for cell, src in src_for_comments.items():
            an.comments[cell] = read_comments(src)
            an.stmt_starts[cell] = stmt_starts(src if cell is None else
                                               "\n".join(("" if l.lstrip().startswith(("%", "!")) else l)
                                                         for l in src.split("\n")))

    def parse_projects(self, entries) -> list[dict]:
        out = []
        for e in entries:
            if e["kind"] != "project":
                continue
            rel, full = e["path"], self.root / e["path"]
            ext = file_ext(rel)
            try:
                if ext == ".qgz":
                    with zipfile.ZipFile(full) as z:
                        qgs = [n for n in z.namelist() if n.lower().endswith(".qgs")]
                        rec = parse_qgis_project(rel, z.read(qgs[0]), self.res) if qgs else {
                            "path": rel, "kind": "qgis_project", "unparsed": ["no .qgs inside"]}
                        rec["archive_members"] = sorted(z.namelist())
                elif ext == ".qgs":
                    rec = parse_qgis_project(rel, full.read_bytes(), self.res)
                elif ext == ".model3":
                    rec = parse_model3(rel, full.read_bytes())
                    self.model_ops(rec)
                else:
                    rec = parse_aprx(rel, full, self.res)
            except Exception as exc:  # noqa: BLE001 - recorded as unparsed, never silently dropped
                rec = {"path": rel, "kind": "project", "unparsed": [f"{type(exc).__name__}: {str(exc)[:200]}"]}
            out.append(rec)
        return out

    def model_ops(self, rec: dict) -> None:
        """Child algorithms of a .model3 become ops (family qgis_model) at their XML lines; static paths are resolved;
        expression values (FORMULA, EXPRESSION, ...) are never paths (B5). Model outputs are bound to a file only
        through a folder named after the model holding `<output name, spaces→_>.<ext>` (B3); model input parameters
        are bound later (bind_model_inputs, needs the profiles)."""
        rel = rec["path"]
        sid = self.step_of[rel]
        an = PyAnalyzer(self.ctx, rel, "code." + sid, rel.rsplit("/", 1)[0] if "/" in rel else "")
        an.cell = None
        out_files = self.model_output_files(rel, rec)
        rec["pending_inputs"] = []
        for i, ch in enumerate(rec["children"], start=1):
            spec = alg_spec(self.ctx.algs, ch["alg_id"])
            line = ch.get("line") or i
            node = ast.Constant(value=None)
            node.lineno, node.col_offset, node.end_lineno = line, 0, ch.get("end_line") or line
            reads, writes, params = [], [], {}
            for pname, pv in ch["params"].items():
                ps = spec.get("params", {}).get(pname, {})
                sv = pv.get("static_value")
                params[pname] = {"value": sv, "source": pv.get("source"), "from_child": pv.get("child_id"),
                                 "from_input": pv.get("parameter_name"), "arg_line": pv.get("line")}
                if pv.get("parameter_name"):
                    rec["pending_inputs"].append({"child": ch["id"], "param": pname, "input": pv["parameter_name"],
                                                  "line": pv.get("line") or line, "key": None})
                if isinstance(sv, str) and not is_expression_param(pname, ps) and \
                        (ps.get("type") in LAYER_TYPES or (not ps and looks_like_path(sv))):
                    reads += [r for r in an.path_refs(V("const", sv), pv.get("line") or line, pname, force_path=True)
                              if "ref" in r]
            for oname, ov in ch["outputs"].items():
                dv = ov.get("default_value")
                if isinstance(dv, str) and dv:
                    writes += [dict(r, mode="create") for r in an.path_refs(V("const", dv), ov.get("line") or line,
                                                                            oname, force_path=True) if "ref" in r]
                bound = out_files.get((ch["id"], ov.get("name") or oname))
                if bound and not (isinstance(dv, str) and dv):
                    writes += [dict(r, mode="create", binding="model_output_folder")
                               for r in an.path_refs(V("const", bound), ov.get("line") or line,
                                                     ov.get("output_name") or "OUTPUT", force_path=True) if "ref" in r]
            key = an.emit(node, "model3.child", ch["alg_id"] or "?", "qgis_model",
                          "qgis_algs.json" if spec else "name_rule", reads=reads, writes=writes, params=params,
                          flags=["model_child:" + ch["id"]] + ([] if ch.get("active", True) else ["inactive_child"]))
            for pi in rec["pending_inputs"]:
                if pi["child"] == ch["id"]:
                    pi["key"] = key
        rec["output_bindings"] = [{"child": c, "output": o, "path": p, "rule": "model_output_folder"}
                                  for (c, o), p in sorted(out_files.items())]
        self.analysers[rel] = an
        an.comments, an.stmt_starts = {None: []}, {None: []}

    def model_output_files(self, rel: str, rec: dict) -> dict:
        """(child, output name) -> legacy file: a folder named exactly like the model (stem) holding exactly one data
        file whose stem is the output name with spaces as underscores. Folder + exact name is the evidence; a
        name-only match elsewhere is never used."""
        stem = stem_of(rel).lower()
        data = [e["path"] for e in self.entries if e["kind"] in ("vector", "raster", "table")]
        out = {}
        for o in rec.get("outputs", []):
            want = re.sub(r"\s+", "_", (o["name"] or "").strip()).lower()
            hits = [p for p in data if "/" in p and p.rsplit("/", 2)[-2].lower() == stem
                    and stem_of(p).lower() == want]
            if len(hits) == 1:
                out[(o["child"], o["name"])] = hits[0]
        return out

    def bind_model_inputs(self) -> None:
        """Model input parameters -> files on disk, as agent guesses (B3): by type (raster / vector, point / line /
        polygon from the profile), then name tokens (parameter name and description vs file stem and folders), then
        elimination. The reads join the child ops with `guess: model_input_binding`; one open item asks to confirm."""
        written = {self.ctx.refs[w["ref"]]["rel"] for o in self.ctx.ops.values() for w in o["writes"] if w.get("ref")}
        alias = set(getattr(self, "vrt_alias", {}))
        geom = {}
        for e in self.entries:
            prof = self.profiles.get(e["path"], {})
            if e["kind"] == "vector":
                gts = set()
                for l in prof.get("layers", []):
                    gts |= {re.sub(r"^Multi|[ZM]+$", "", g).lower() for g in (l.get("geom_types") or {})}
                    if l.get("declared_geometry"):
                        gts.add(re.sub(r"^Multi|[ZM ]+$", "", l["declared_geometry"]).lower())
                geom[e["path"]] = gts
        cands = [e for e in self.entries if e["kind"] in ("raster", "vector") and e["path"] not in written
                 and e["path"] not in alias and "error" not in self.profiles.get(e["path"], {})]
        qtype = {0: "point", 1: "linestring", 2: "polygon"}
        self.model_bindings = []
        for rec in self.projects:
            if rec.get("kind") != "qgis_model" or not rec.get("input_defs"):
                continue
            bound_out = {b["path"] for b in rec.get("output_bindings", [])}
            pool = [e for e in cands if e["path"] not in bound_out]
            scores = []
            for name, d in rec["input_defs"].items():
                kind = {"raster": "raster", "vector": "vector", "source": "vector", "layer": None,
                        "multilayer": None, "file": None}.get(d.get("type"), "skip")
                if kind == "skip":
                    continue
                want = {qtype[t] for t in (d.get("data_types") or []) if t in qtype}
                ptoks = name_tokens(name) | name_tokens(d.get("description") or "")
                for e in pool:
                    if kind and e["kind"] != kind:
                        continue
                    if want and e["kind"] == "vector" and geom.get(e["path"]) and not geom[e["path"]] & want:
                        continue
                    stoks = name_tokens(stem_of(e["path"]))
                    ftoks = set().union(*[name_tokens(x) for x in e["path"].split("/")[:-1]]) if "/" in e["path"] \
                        else set()
                    sc, why = 0, []
                    for t in sorted(ptoks):
                        if t in stoks:
                            sc, why = sc + 3, why + [f"name token {t!r}"]
                        elif t in ftoks:
                            sc, why = sc + 2, why + [f"folder token {t!r}"]
                        elif any(token_near(t, x) for x in stoks | ftoks):
                            sc, why = sc + 1, why + [f"near token {t!r}"]
                    scores.append((sc, name, e["path"], why, kind, want))
            used_p, used_f, out = set(), set(), []
            for sc, name, path, why, kind, want in sorted(scores, key=lambda x: (-x[0], x[1], x[2])):
                if sc <= 0 or name in used_p or path in used_f:
                    continue
                rivals = [x for x in scores if x[1] == name and x[0] == sc and x[2] not in used_f and x[2] != path]
                if rivals:
                    continue  # a tie is not a guess
                used_p.add(name)
                used_f.add(path)
                out.append({"input": name, "path": path, "score": sc, "why": why, "rule": "name_tokens"})
            for name, d in sorted(rec["input_defs"].items()):  # elimination: one type-compatible file left
                if name in used_p:
                    continue
                left = sorted({x[2] for x in scores if x[1] == name and x[2] not in used_f})
                if len(left) == 1:
                    used_p.add(name)
                    used_f.add(left[0])
                    out.append({"input": name, "path": left[0], "score": 0, "why": ["only type-compatible file left"],
                                "rule": "elimination"})
            for b in out:
                d = rec["input_defs"][b["input"]]
                b.update({"model": rec["path"], "type": d.get("type"), "data_types": d.get("data_types"),
                          "line": d.get("line"), "state": "agent_guess"})
                res = self.res.resolve(b["path"])
                key = self.ctx.ref(res["rel"], None, res, b["path"])
                for pi in rec.get("pending_inputs", []):
                    op = self.ctx.ops.get(pi["key"]) if pi["key"] else None
                    if op is not None and pi["input"] == b["input"]:
                        r = {"ref": key, "param": pi["param"], "line": pi["line"], "guess": "model_input_binding"}
                        if r not in op["reads"]:
                            op["reads"].append(r)
            rec["input_bindings"] = sorted(out, key=lambda b: b["input"])
            rec["unbound_inputs"] = sorted(set(rec["input_defs"]) - {b["input"] for b in out})
            self.model_bindings.append({"model": rec["path"], "bindings": rec["input_bindings"],
                                        "unbound": rec["unbound_inputs"]})

    # ------------------------------------------------------------------ profiles
    def profile_all(self, entries):
        profiles: dict[str, dict] = {}
        keep: dict = {}
        rasters = {e["path"] for e in entries if e["kind"] == "raster"}
        self.vrt_alias: dict[str, str] = {}
        for e in sorted(entries, key=lambda e: file_ext(e["path"]) == ".vrt"):  # sources before their VRTs
            rel, kind = e["path"], e["kind"]
            full = self.root / rel
            try:
                sib = vrt_alias_source(full, rel, rasters) if kind == "raster" and file_ext(rel) == ".vrt" else None
                if sib and "transform" in profiles.get(sib, {}):
                    p = profile_raster(full, stats=False)
                    src = profiles[sib]
                    same = {k: p.get(k) == src.get(k) for k in ("shape", "transform", "dtype")}
                    if all(same.values()):
                        diff = {k: [p.get(k), src.get(k)] for k in ("nodata", "embedded_stats") if p.get(k) != src.get(k)}
                        profiles[rel] = dict(src, driver=p.get("driver"), nodata=p.get("nodata"),
                                             embedded_stats=p.get("embedded_stats"), alias_of=sib,
                                             alias_differs=sorted(diff), computed_from=sib)
                        self.vrt_alias[rel] = sib
                        continue
                if kind == "raster":
                    profiles[rel] = profile_raster(full)
                elif kind in ("vector", "gdb"):
                    profiles[rel] = profile_vector(full, keep, rel)
                elif kind == "table":
                    ext = file_ext(rel)
                    if ext in (".xlsx", ".xlsm"):
                        profiles[rel] = profile_xlsx(full)
                    elif ext in (".csv", ".tsv"):
                        profiles[rel] = profile_csv(full)
                    else:
                        profiles[rel] = {"format": ext.lstrip("."), "unparsed": "no reader in the pinned env"}
            except Exception as exc:  # noqa: BLE001
                profiles[rel] = {"error": f"{type(exc).__name__}: {str(exc)[:300]}"}
                self.log(f"profile error {rel}: {exc}")
        return profiles, keep

    # ------------------------------------------------------------------ datasets
    def build_datasets(self) -> None:
        ents: dict[tuple, dict] = {}
        file_kind = {e["path"]: e["kind"] for e in self.entries}
        layers_of: dict[str, list] = {}
        for e in self.entries:
            if e["kind"] not in ("raster", "vector", "table", "gdb"):
                continue
            rel = e["path"]
            prof = self.profiles.get(rel, {})
            if e["kind"] in ("vector", "gdb"):
                lays = [l["name"] for l in prof.get("layers", [])] or [None]
                layers_of[rel] = lays
                multi = len(lays) > 1 or e["kind"] == "gdb"
                for l in lays:
                    ents[(rel, l)] = {"rel": rel, "layer": l, "exists": True, "kind": "vector", "multi": multi,
                                      "file_kind": e["kind"]}
            else:
                ents[(rel, None)] = {"rel": rel, "layer": None, "exists": True, "kind": e["kind"], "multi": False,
                                     "file_kind": e["kind"]}
        self.ref_to_ent: dict[str, tuple] = {}
        self.layer_mismatch: list[dict] = []
        for key, r in sorted(self.ctx.refs.items()):
            rel, layer = r["rel"], r["layer"]
            if r.get("gdb_item"):
                k = (rel, layer)
                if k not in ents:
                    ents[k] = {"rel": rel, "layer": layer, "exists": False, "kind": None, "multi": True,
                               "gdb_item": True, "file_kind": "gdb"}
                self.ref_to_ent[key] = k
                continue
            if rel in layers_of:
                lays = layers_of[rel]
                if len(lays) == 1 and not ents[(rel, lays[0])].get("gdb_item") and file_kind.get(rel) != "gdb":
                    self.ref_to_ent[key] = (rel, lays[0])
                    if layer and lays[0] is not None and layer != lays[0]:
                        self.layer_mismatch.append({"ref": key, "rel": rel, "layer": layer, "actual": lays})
                elif layer in lays:
                    self.ref_to_ent[key] = (rel, layer)
                elif layer is None:
                    st = stem_of(rel)
                    pick = st if st in lays else sorted(l for l in lays if l is not None)[0]
                    self.ref_to_ent[key] = (rel, pick)
                    self.layer_mismatch.append({"ref": key, "rel": rel, "layer": None, "actual": lays,
                                                "reason": "layer_not_given"})
                else:
                    k = (rel, layer)
                    ents.setdefault(k, {"rel": rel, "layer": layer, "exists": False, "kind": "vector", "multi": True,
                                        "file_kind": "vector", "layer_missing": True})
                    self.ref_to_ent[key] = k
                    self.layer_mismatch.append({"ref": key, "rel": rel, "layer": layer, "actual": lays})
                continue
            k = (rel, None) if (rel, None) in ents else (rel, layer)
            if k not in ents:
                fk = file_kind.get(rel)
                ext = file_ext(rel)
                kind = ("raster" if ext in RASTER_EXT else "vector" if ext in VECTOR_EXT or ext in (".json", ".geojson")
                        else "table" if ext in TABLE_EXT else "document")
                ents[k] = {"rel": rel, "layer": layer if ext in VECTOR_EXT else None, "exists": fk is not None,
                           "kind": kind, "multi": False, "file_kind": fk}
                k = (rel, ents[k]["layer"]) if (rel, ents[k]["layer"]) in ents else k
            self.ref_to_ent[key] = k
        # project and doc references to data files that code never names attach to the same entities
        # ids
        stems: dict[str, set] = defaultdict(set)
        for (rel, layer), en in ents.items():
            if not en.get("gdb_item") and en.get("file_kind") != "gdb":
                stems[ac.slug(stem_of(rel))].add(rel)
        ids: dict[tuple, str] = {}
        used: dict[str, tuple] = {}
        for k in sorted(ents, key=lambda kk: (kk[0], kk[1] or "")):
            en = ents[k]
            rel, layer = k
            st = ac.slug(stem_of(rel))
            st = st or "file"
            if en.get("gdb_item") or en.get("file_kind") == "gdb":
                did = f"ds.{st}__{ac.slug(layer or 'item') or 'item'}"
            else:
                did = f"ds.{st}" if len(stems[st]) <= 1 else f"ds.{st}_{ac.slug(file_ext(rel).lstrip('.'))}"
                if en.get("multi") and layer and ac.slug(layer) != st:
                    did += f"__{ac.slug(layer)}"
            if did in used:
                did += "_" + ac.slug(rel.rsplit("/", 1)[0] if "/" in rel else "root")
            used[did] = k
            ids[k] = did
        self.ents = ents
        self.ent_id = ids
        self.id_ent = {v: k for k, v in ids.items()}
        for key, k in self.ref_to_ent.items():
            self.ctx.refs[key]["dataset"] = ids[k]

    def ds_of_ref(self, key: str | None) -> str | None:
        if not key:
            return None
        k = self.ref_to_ent.get(key)
        return self.ent_id.get(k) if k else None

    # ------------------------------------------------------------------ ops
    def build_ops(self) -> None:
        by_file: dict[str, list[dict]] = defaultdict(list)
        for op in self.ctx.ops.values():
            by_file[op["file"]].append(op)
        self.ops: list[dict] = []
        self.op_by_key: dict[tuple, dict] = {}
        mem_final: dict[str, str] = {}
        used_mem = set()
        for rel in sorted(by_file):
            sid = self.step_of.get(rel, "s99_" + ac.slug(rel))
            raw = sorted(by_file[rel], key=lambda o: (o["cell"] if o["cell"] is not None else -1, o["line"], o["col"],
                                                      o["key"][4], o["key"][5]))
            for i, op in enumerate(raw, start=1):
                op["id"] = f"{sid}.op{i}"
                self.op_by_key[op["key"]] = op
                if op["memory_out"]:
                    base = f"ds.mem.{short_step(sid)}" + (f"_c{op['cell']}" if op["cell"] is not None else "") + \
                           f"_l{op['memory_out'].split(':')[3]}"
                    mid = base if base not in used_mem else f"{base}_c{op['col']}"
                    used_mem.add(mid)
                    mem_final[op["memory_out"]] = mid
        self.mem_final = mem_final

        def conv(x):
            if isinstance(x, dict):
                if set(x) == {"ref"}:
                    return {"dataset": self.ds_of_ref(x["ref"])}
                if "layer_of" in x and len(x) == 1:
                    return {"dataset": self.ds_of_ref(x["layer_of"]), "via": "layer"}
                if "raster_of" in x and len(x) == 1:
                    return {"dataset": self.ds_of_ref(x["raster_of"]), "via": "raster_object"}
                if "derived" in x and "ref" in x:
                    d = {"derived": x["derived"], "dataset": self.ds_of_ref(x["ref"])}
                    if x.get("format"):
                        d["format"] = x["format"]
                    return d
                if set(x) == {"memory"}:
                    return {"memory": mem_final.get(x["memory"], x["memory"])}
                return {k: conv(v) for k, v in x.items()}
            if isinstance(x, list):
                return [conv(v) for v in x]
            return x

        self.mem_records: dict[str, dict] = {}
        for rel in sorted(by_file):
            an = self.analysers.get(rel)
            for op in sorted(by_file[rel], key=lambda o: o["id"]):
                reads, mreads, writes, grid = [], [], [], []
                for r in op["reads"]:
                    if r.get("ref"):
                        d = {"dataset": self.ds_of_ref(r["ref"]), "param": r["param"], "line": r["line"]}
                        for k in ("via", "subset", "phi", "guess"):
                            if r.get(k):
                                d[k] = r[k]
                        if d not in reads:
                            reads.append(d)
                    elif r.get("memory"):
                        mid = mem_final.get(r["memory"], r["memory"])
                        d = {"memory": mid, "param": r["param"], "line": r["line"]}
                        if d not in mreads and not mid.startswith("mem:"):
                            mreads.append(d)
                for w in op["writes"]:
                    if w.get("ref"):
                        d = {"dataset": self.ds_of_ref(w["ref"]), "param": w["param"], "line": w["line"],
                             "mode": w.get("mode", "create")}
                        if w.get("binding"):
                            d["binding"] = w["binding"]
                        if d not in writes:
                            writes.append(d)
                for g in op["reads_grid"]:
                    d = {"dataset": self.ds_of_ref(g["ref"]), "param": g["param"], "line": g["line"],
                         "derived": g["derived"]}
                    if d not in grid:
                        grid.append(d)
                params = conv(op["params"])
                fields_read = []
                for f in op.get("fields_read", []):
                    d = {"dataset": self.ds_of_ref(f["ref"]), "field": f["field"], "line": f["line"], "param": f["param"]}
                    if f.get("of_param"):
                        d["of_param"] = f["of_param"]
                    if d not in fields_read:
                        fields_read.append(d)
                comments = []
                if an is not None:
                    cell = op["cell"]
                    starts = an.stmt_starts.get(cell, [])
                    for c in an.comments.get(cell, []):
                        tgt = c["line"]
                        if c["own_line"]:
                            nxt = [s for s in starts if s > c["line"]]
                            tgt = nxt[0] if nxt else c["line"]
                        if op["line"] <= tgt <= op["end_line"] or any(tgt == w.get("line") for w in op["writes"]):
                            comments.append({"line": c["line"], "text": c["text"]})
                rec = {"id": op["id"], "code": op["code"], "file": op["file"], "line": op["line"],
                       "end_line": op["end_line"], "api": op["api"], "tool": op["tool"], "family": op["family"],
                       "io_source": op["io_source"], "params": params, "reads": reads, "writes": writes,
                       "reads_grid": grid, "memory_reads": mreads, "fields_read": fields_read,
                       "memory_out": mem_final.get(op["memory_out"]) if op["memory_out"] else None,
                       "comments": comments, "flags": sorted(set(op["flags"])), "nested": op["nested"],
                       "state": op["state"]}
                if op["cell"] is not None:
                    rec["cell"] = op["cell"]
                    rec["ref"] = f"{op['file']}#c{op['cell']}:L{op['line']}"
                else:
                    rec["ref"] = f"{op['file']}:{op['line']}" + (f"-{op['end_line']}" if op["end_line"] != op["line"]
                                                                 else "")
                if op["visits"] > 1:
                    rec["visits"] = op["visits"]
                self.ops.append(rec)
                if rec["memory_out"]:
                    self.mem_records[rec["memory_out"]] = {"id": rec["memory_out"], "op": rec["id"], "file": rec["file"],
                                                           "line": int(op["memory_out"].split(":")[3]),
                                                           "consumed_by": []}
        for rec in self.ops:
            for m in rec["memory_reads"]:
                if m["memory"] in self.mem_records:
                    mr = self.mem_records[m["memory"]]
                    if rec["id"] not in mr["consumed_by"]:
                        mr["consumed_by"].append(rec["id"])
        self.ops.sort(key=lambda o: (o["file"], o.get("cell", -1), o["line"], o["id"]))
        self.op_by_id = {o["id"]: o for o in self.ops}

    # ------------------------------------------------------------------ code records, status, twins
    def code_records(self, code_entries, other_code) -> None:
        self.codes: dict[str, dict] = {}
        imported_by: dict[str, list] = defaultdict(list)
        for rel, an in self.analysers.items():
            for imp in an.imports:
                mod = (imp.get("module") or "").lstrip(".").split(".")[-1]
                if mod in self.ctx.local_modules and self.ctx.local_modules[mod] != rel:
                    imported_by[self.ctx.local_modules[mod]].append(rel)
        dont_run = defaultdict(list)
        for d in self.docs:
            for x in d["dont_run"]:
                dont_run[x["path"]].append({"kind": "doc_mention", "file": d["path"], "line": x["line"],
                                            "text": x["text"]})
        for e in code_entries + [{"path": p["path"]} for p in self.projects if p.get("kind") == "qgis_model"]:
            rel = e["path"]
            sid = self.step_of[rel]
            an = self.analysers.get(rel)
            ops = [o["id"] for o in self.ops if o["file"] == rel]
            runtime = "python"
            mods = {(i.get("module") or "").split(".")[0] for i in (an.imports if an else [])}
            if "arcpy" in mods:
                runtime = "arcpy"
            elif mods & {"qgis", "processing"}:
                runtime = "pyqgis"
            if rel.endswith(".model3"):
                runtime = "qgis_model"
            comments = []
            if an:
                for cell, cs in an.comments.items():
                    for c in cs:
                        comments.append(dict(c, cell=cell) if cell is not None else c)
            status_ev = []
            st = stem_of(rel)
            m = SUPERSEDED_NAME_RE.search(st)
            if m:
                status_ev.append({"kind": "name", "value": m.group(2)})
            head = [c for c in comments if c["line"] <= 15 and (c.get("cell") in (None, 0))]
            for c in head:
                if SUPERSEDED_TEXT_RE.search(c["text"]):
                    status_ev.append({"kind": "comment", "file": rel, "line": c["line"], "text": c["text"]})
            status_ev += dont_run.get(rel, [])
            twin_of = None
            status = "unknown"
            if "/.ipynb_checkpoints/" in "/" + rel or rel.startswith(".ipynb_checkpoints/"):
                status = "twin"
                orig = re.sub(r"(^|/)\.ipynb_checkpoints/", r"\1", rel).replace("-checkpoint.ipynb", ".ipynb")
                if orig in self.step_of:
                    twin_of = "code." + self.step_of[orig]
                status_ev.append({"kind": "checkpoint_copy", "value": ".ipynb_checkpoints"})
            elif status_ev and any(x["kind"] in ("comment", "doc_mention") for x in status_ev) or \
                    (status_ev and not ops):
                status = "superseded"
            elif status_ev:
                status = "superseded"
            elif ops and not (rel in imported_by and all(
                    "in_uncalled_function" in self.op_by_id[o]["flags"] for o in ops)):
                status = "live"
            elif rel in imported_by:
                status = "helper"  # its I/O sits in functions the importers call (inlined there, M5)
            self.codes[rel] = {"id": "code." + sid, "step_id": sid, "path": rel,
                               "language": "notebook" if rel.endswith(".ipynb") else (
                                   "qgis_model" if rel.endswith(".model3") else "python"),
                               "runtime": runtime, "status": status, "status_evidence": status_ev, "twin_of": twin_of,
                               "imports": sorted((an.imports if an else []), key=lambda i: (i.get("cell") or 0,
                                                                                         i["line"])),
                               "imported_by": sorted(imported_by.get(rel, [])),
                               "entry": "module", "functions": sorted(
                                   [{"name": k, "line": f["node"].lineno, "called": f["called"]}
                                    for k, f in (an.funcs.items() if an else [])], key=lambda f: f["line"]),
                               "flags": sorted((an.flags if an else []), key=lambda f: (f.get("cell") or 0,
                                                                                         f.get("line") or 0,
                                                                                         f["kind"])),
                               "comments": len(comments),
                               "commented_code": commented_code([c for c in comments if c.get("cell") is None]),
                               "parse_error": next((f["msg"] for f in (an.flags if an else [])
                                                    if f["kind"] == "parse_error"), None),
                               "ops": ops}
            if rel in self.nb_facts:
                f = self.nb_facts[rel]
                self.codes[rel]["notebook"] = {k: v for k, v in f.items() if k != "outputs_text"}
                self.codes[rel]["notebook"]["output_paths"] = sorted({
                    m.group(0) for o in f.get("outputs_text", [])
                    for m in re.finditer(r"[\w./\\-]+\.(tif|gpkg|shp|csv|xlsx|geojson)", o["text"])})
        for e in other_code:
            self.codes[e["path"]] = {"id": "code." + ac.slug(e["path"]), "path": e["path"], "step_id": None,
                                     "language": CODE_EXT.get(file_ext(e["path"]), "other"), "status": "unknown",
                                     "status_evidence": [], "twin_of": None, "ops": [], "flags": [],
                                     "parse_error": "not analysed (listed only)"}
        self.twins = []
        self.alternative_implementations()
        # twins of superseded code: replacement named in a comment, the same numeric prefix, or similar code
        live = {r: c for r, c in self.codes.items() if c["status"] == "live"}
        for rel, c in sorted(self.codes.items()):
            if c["status"] not in ("superseded", "twin") or c.get("twin_of"):
                if c["status"] == "twin" and c.get("twin_of"):
                    self.twins.append(self.twin_record(rel, c))
                continue
            target, why = None, None
            for x in c["status_evidence"]:
                for r2 in live:
                    if x.get("text") and stem_of(r2).lower() in x["text"].lower():
                        target, why = r2, {"kind": x["kind"], "line": x.get("line"), "file": x.get("file")}
                        break
                if target:
                    break
            if not target:
                m = re.match(r"^(\d+)[a-zA-Z]", stem_of(rel))
                if m:
                    cands = [r2 for r2 in live if re.match(rf"^0*{int(m.group(1))}([^0-9a-zA-Z]|$)", stem_of(r2))]
                    if len(cands) == 1:
                        target, why = cands[0], {"kind": "name_prefix", "value": m.group(1)}
            if not target and live:
                best = max(live, key=lambda r2: self.similarity(rel, r2))
                if self.similarity(rel, best) >= 0.6:
                    target, why = best, {"kind": "code_similarity", "value": round(self.similarity(rel, best), 3)}
            if target:
                c["twin_of"] = live[target]["id"]
                c["twin_evidence"] = why
                self.twins.append(self.twin_record(rel, c))

    def implementation_groups(self) -> None:
        """Alternative implementations of one workflow across tools (M3/M6): QGIS models, Python scripts and ArcGIS
        toolboxes that name the same output fields. Field names are the models' FIELD_NAME/NEW_NAME values; a
        script or toolbox joins when at least two of them occur in it as exact string constants (case-insensitive);
        two models join on two shared field names; superseded twins follow their canonical. The member whose
        outputs exist in the legacy folder is the guessed parity reference; the choice is one question (lineage)."""
        fields: dict[str, set] = {}
        for rel, c in self.codes.items():
            if not rel.endswith(".model3") or not c.get("ops"):
                continue
            f = set()
            for oid in c["ops"]:
                for pname, pv in self.op_by_id[oid]["params"].items():
                    v = pv.get("value") if isinstance(pv, dict) else None
                    if pname in ("FIELD_NAME", "NEW_NAME", "OUTPUT_FIELD", "FIELD") and isinstance(v, str) and \
                            re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{2,40}", v) and v.lower() not in ("value", "id", "fid"):
                        f.add(v)
            if f:
                fields[rel] = f
        consts: dict[str, set] = {}
        for rel, an in self.analysers.items():
            if rel.endswith(".model3"):
                continue
            try:
                src = "\n".join("\n".join(v) for v in an.src_lines.values())
                consts[rel] = {n.value.lower() for n in ast.walk(ast.parse(src)) if isinstance(n, ast.Constant)
                               and isinstance(n.value, str) and len(n.value) <= 40}
            except SyntaxError:
                consts[rel] = set()
        for tb in self.toolboxes:
            consts[tb["path"]] = {x.lower() for x in tb.get("identifiers", [])}
        nodes = sorted(set(fields) | {r for r in consts if r in self.codes and self.codes[r].get("ops")
                                      and self.codes[r]["status"] in ("live", "superseded")}
                       | {tb["path"] for tb in self.toolboxes})
        parent = {n: n for n in nodes}

        def find(n):
            while parent[n] != n:
                n = parent[n]
            return n
        evidence = defaultdict(list)
        for m in sorted(fields):
            low = {f.lower(): f for f in fields[m]}
            for o in nodes:
                if o == m:
                    continue
                shared = sorted(low[k] for k in low if k in (consts.get(o) or {f.lower() for f in fields.get(o, ())}))
                if len(shared) >= 2:
                    parent[find(o)] = find(m)
                    evidence[(m, o)] = shared
        for rel, c in self.codes.items():
            if c.get("twin_of") and rel in parent:
                canon = next((r for r, x in self.codes.items() if x["id"] == c["twin_of"]), None)
                if canon in parent:
                    parent[find(rel)] = find(canon)
        groups = defaultdict(list)
        for n in nodes:
            groups[find(n)].append(n)
        out = []
        for members in groups.values():
            if len(members) < 2 or not any(m in fields for m in members):
                continue
            recs = []
            for m in sorted(members):
                c = self.codes.get(m)
                kind = "qgis_model" if m.endswith(".model3") else "arcgis_toolbox" if m.endswith(".tbx") else "python"
                outs = sorted({w["dataset"] for o in (c or {}).get("ops", []) for w in self.op_by_id[o]["writes"]
                               if w["mode"] in ("create", "append") and self.datasets.get(w["dataset"], {}).get("exists")})
                foreign = sorted({f["value"] for f in (c or {}).get("flags", []) if f["kind"] == "foreign_host_path"})
                recs.append({"source": m, "code": c["id"] if c else None, "step": c["step_id"] if c else None,
                             "kind": kind, "status": c["status"] if c else "listed_only",
                             "twin_of": c.get("twin_of") if c else None, "existing_outputs": outs,
                             "foreign_host_paths": foreign[:6]})
            ref = sorted([r for r in recs if r["existing_outputs"]], key=lambda r: (-len(r["existing_outputs"]),
                                                                                    r["source"]))
            ev = [{"kind": "shared_field_names", "a": a, "b": b, "fields": f} for (a, b), f in sorted(evidence.items())
                  if a in members and b in members]
            gid = "impl." + ac.slug(os.path.commonprefix([stem_of(m).lower() for m in members]).strip("_-") or
                                    stem_of(sorted(members)[0]))
            out.append({"id": gid, "members": recs, "evidence": ev,
                        "reference_guess": ref[0]["code"] if len(ref) == 1 or (ref and len(ref[0]["existing_outputs"])
                                                                              > len(ref[1]["existing_outputs"]))
                        else None,
                        "reference_why": (f"{ref[0]['source']} is the only implementation whose outputs exist in the "
                                          f"legacy folder ({len(ref[0]['existing_outputs'])})" if len(ref) == 1 else
                                          "no member's outputs exist" if not ref else
                                          f"{ref[0]['source']} has the most existing outputs")})
        self.impl_groups = sorted(out, key=lambda g: g["id"])

    def alternative_implementations(self) -> None:
        """Live scripts whose write sets overlap are alternative implementations of one workflow (B1): one stays
        canonical (its stem is the others' common prefix, else the most similar to the rest, else the first path);
        the others become `superseded` with `twin_of` the canonical and twin kind `alternative_impl`."""
        writes = {}
        for rel, c in self.codes.items():
            if c["status"] != "live":
                continue
            w = {x["dataset"] for o in self.ops if o["file"] == rel for x in o["writes"]
                 if x["mode"] in ("create", "append") and x.get("dataset")}
            if w:
                writes[rel] = w
        parent = {r: r for r in writes}

        def find(r):
            while parent[r] != r:
                r = parent[r]
            return r
        rels = sorted(writes)
        for i, a in enumerate(rels):
            for b in rels[i + 1:]:
                if writes[a] & writes[b]:
                    parent[find(b)] = find(a)
        groups = defaultdict(list)
        for r in rels:
            groups[find(r)].append(r)
        for members in groups.values():
            if len(members) < 2:
                continue
            stems = {r: stem_of(r).lower() for r in members}
            prefix = [r for r in members if all(stems[o].startswith(stems[r]) for o in members)]
            sim = {r: sum(self.similarity(r, o) for o in members if o != r) for r in members}
            canon = prefix[0] if len(prefix) == 1 else max(sorted(members), key=lambda r: sim[r])
            why = "stem_prefix" if len(prefix) == 1 else "most_similar"
            for r in sorted(members):
                if r == canon:
                    continue
                c = self.codes[r]
                shared = sorted(writes[r] & writes[canon]) or sorted(writes[r] & set().union(
                    *[writes[o] for o in members if o != r]))
                c["status"] = "superseded"
                c["status_evidence"].append({"kind": "alternative_impl", "value": stem_of(canon),
                                             "shared_writes": shared})
                c["twin_of"] = self.codes[canon]["id"]
                c["twin_evidence"] = {"kind": "shared_writes", "value": shared, "canonical_by": why,
                                      "similarity": round(self.similarity(r, canon), 3)}
                rec = self.twin_record(r, c)
                rec["kind"] = "alternative_impl"
                self.twins.append(rec)

    def similarity(self, a: str, b: str) -> float:
        def norm(rel):
            try:
                src = (self.root / rel).read_text(encoding="utf-8", errors="replace")
            except OSError:
                return []
            if rel.endswith(".ipynb"):
                units, _ = notebook_units(self.root / rel)
                src = "\n".join(u["src"] for u in units)
            return [re.sub(r"\s+", " ", l.split("#")[0]).strip() for l in src.splitlines() if l.split("#")[0].strip()]
        return difflib.SequenceMatcher(None, norm(a), norm(b)).ratio()

    def twin_record(self, rel: str, c: dict) -> dict:
        other = next((r for r, x in self.codes.items() if x["id"] == c["twin_of"]), None)
        rec = {"a": c["id"], "b": c["twin_of"], "kind": "checkpoint" if c["status"] == "twin" else
               "superseded_alternative", "evidence": c.get("twin_evidence") or c["status_evidence"][:1]}
        if other:
            ta = Counter(o["tool"] for o in self.ops if o["file"] == rel)
            tb = Counter(o["tool"] for o in self.ops if o["file"] == other)
            rec["tools_only_in_a"] = sorted((ta - tb).elements())
            rec["tools_only_in_b"] = sorted((tb - ta).elements())
            rec["similarity"] = round(self.similarity(rel, other), 3)
            if rel.endswith(".ipynb") and other.endswith(".ipynb"):
                ua, _ = notebook_units(self.root / rel)
                ub, _ = notebook_units(self.root / other)
                da = {u["cell"]: u["raw"].strip() for u in ua}
                db = {u["cell"]: u["raw"].strip() for u in ub}
                rec["cells_differ"] = sorted(k for k in set(da) | set(db) if da.get(k) != db.get(k))
        return rec

    # ------------------------------------------------------------------ order of live steps
    def step_order(self) -> None:
        live = sorted(r for r, c in self.codes.items() if c["status"] == "live")
        writes_of = defaultdict(set)
        reads_of = defaultdict(set)
        for o in self.ops:
            if o["file"] not in live:
                continue
            for w in o["writes"]:
                if w["mode"] in ("create", "append"):
                    writes_of[w["dataset"]].add(o["file"])
            for r in o["reads"] + o["reads_grid"]:
                reads_of[r["dataset"]].add(o["file"])
        deps = defaultdict(set)
        for ds, ws in writes_of.items():
            for r in reads_of.get(ds, ()):
                for w in ws:
                    if w != r:
                        deps[r].add(w)
        doc_idx = {}
        doc_ev = {}
        for d in self.docs:
            for i, x in enumerate(d["run_order"]):
                if x["path"] not in doc_idx:
                    doc_idx[x["path"]] = i
                    doc_ev[x["path"]] = {"kind": "doc_run_order", "file": d["path"], "line": x["line"]}

        def prefix(rel):
            m = re.match(r"^(\d+)", stem_of(rel))
            return int(m.group(1)) if m else 10 ** 6
        order, done = [], set()
        remaining = set(live)
        while remaining:
            ready = [r for r in remaining if deps[r] <= done]
            if not ready:  # cycle: fall back to documented / prefix order and record it
                ready = list(remaining)
            ready.sort(key=lambda r: (doc_idx.get(r, 10 ** 6), prefix(r), r))
            order.append(ready[0])
            done.add(ready[0])
            remaining.discard(ready[0])
        self.live_order = order
        self.order_conflicts = []
        for i, rel in enumerate(order):
            ev = []
            if rel in doc_ev:
                ev.append(doc_ev[rel])
            if deps[rel]:
                ev.append({"kind": "data_dependency", "after": sorted("code." + self.step_of[d] for d in deps[rel])})
            if prefix(rel) < 10 ** 6:
                ev.append({"kind": "name_prefix", "value": re.match(r"^(\d+)", stem_of(rel)).group(1)})
            self.codes[rel]["order"] = i + 1
            self.codes[rel]["order_evidence"] = ev
        doc_sorted = sorted([r for r in order if r in doc_idx], key=lambda r: doc_idx[r])
        if doc_sorted and doc_sorted != [r for r in order if r in doc_idx]:
            self.order_conflicts.append({"kind": "doc_vs_data", "doc": doc_sorted, "data": order})
        pre_sorted = sorted(order, key=lambda r: (prefix(r), r))
        if pre_sorted != order:
            self.order_conflicts.append({"kind": "prefix_vs_data", "prefix": pre_sorted, "data": order})

    # ------------------------------------------------------------------ roles
    def classify_roles(self) -> None:
        live_idx = {rel: i for i, rel in enumerate(self.live_order)}
        status = {rel: c["status"] for rel, c in self.codes.items()}
        writers, readers, all_readers, all_writers = (defaultdict(list) for _ in range(4))
        for o in self.ops:
            live = status.get(o["file"]) == "live"
            for w in o["writes"]:
                all_writers[w["dataset"]].append((o, w))
                if live:
                    writers[w["dataset"]].append((o, w))
            for r in o["reads"] + o["reads_grid"]:
                all_readers[r["dataset"]].append((o, r))
                if live:
                    readers[r["dataset"]].append((o, r))
        live_write_dirs = {self.id_ent[w["dataset"]][0].rsplit("/", 1)[0]
                           for ds, lst in writers.items() for _, w in lst
                           if w["dataset"] in self.id_ent and "/" in self.id_ent[w["dataset"]][0]}
        self.legacy_output_dirs = sorted(d for d in live_write_dirs if d)
        doc_outputs = {x["path"]: (d["path"], x["line"]) for d in self.docs for x in d["outputs"]}
        doc_mentions = defaultdict(list)
        for d in self.docs:
            for m in d["mentions"]:
                if m["resolved"]:
                    doc_mentions[m["resolved"]].append({"kind": "doc_mention", "file": d["path"], "line": m["line"]})
        proj_refs = defaultdict(list)
        for p in self.projects:
            for l in p.get("layers", []):
                if l.get("path"):
                    proj_refs[(l["path"], l.get("layer"))].append({"kind": "project_layer", "file": p["path"],
                                                                   "layer": l.get("name"), "group": l.get("group")})
        self.datasets: dict[str, dict] = {}
        for k, did in sorted(self.ent_id.items(), key=lambda kv: kv[1]):
            en = self.ents[k]
            rel, layer = k

            def pos(o):
                return (live_idx.get(o["file"], -1), o["cell"] if o.get("cell") is not None else -1, o["line"])
            ev, role, unused_kind, conflict = [], None, None, False
            W, R = writers.get(did, []), readers.get(did, [])
            for o, w in W:
                ev.append({"kind": "code_write", "op": o["id"], "file": o["file"], "line": o["line"],
                           "arg_line": w["line"], "mode": w["mode"]})
                if w.get("binding") == "model_output_folder":
                    ev[-1]["binding"] = w["binding"]
                    ev[-1]["rule"] = "folder named after the model + file named after the model output"
            for o, r in R:
                e = {"kind": "code_read", "op": o["id"], "file": o["file"], "line": r["line"]}
                if r.get("guess"):
                    e["guess"] = r["guess"]
                if e not in ev:
                    ev.append(e)
            top = rel.split("/")[0] if "/" in rel else ""
            if top:
                ev.append({"kind": "folder", "value": top})
            ev += doc_mentions.get(rel, [])[:3]
            pr = proj_refs.get((rel, layer)) or proj_refs.get((rel, None)) or []
            if not pr and en.get("layer"):
                pr = [x for (p, l), lst in proj_refs.items() if p == rel for x in lst]
            ev += pr
            created = [(o, w) for o, w in W if w["mode"] in ("create", "append")]
            final_like = bool(FINAL_DIR_RE.search(top)) or rel in doc_outputs
            if created:
                wpos = min(pos(o) for o, _ in created)
                wstep = min(live_idx.get(o["file"], 0) for o, _ in created)
                later_other = any(live_idx.get(o["file"], -1) > wstep for o, _ in R)
                later_same = any(live_idx.get(o["file"], -1) == wstep and pos(o) > wpos for o, _ in R)
                if later_other:
                    role = "intermediate"
                elif later_same and not final_like:
                    role = "intermediate"
                else:
                    role = "final"
                if rel in doc_outputs:
                    ev.append({"kind": "doc_output", "file": doc_outputs[rel][0], "line": doc_outputs[rel][1]})
            elif R:
                role = "input"
                if W:
                    ev.append({"kind": "modified_in_place", "ops": sorted({o["id"] for o, _ in W})})
                d = rel.rsplit("/", 1)[0] if "/" in rel else ""
                if d in live_write_dirs and en["exists"]:
                    role, conflict = "unknown", True
                    ev.append({"kind": "conflict", "value": "read-only by code but sits in a folder the code writes to"})
            elif all_readers.get(did) or all_writers.get(did):
                role, unused_kind = "unused", "superseded_only"
            elif pr:
                role, unused_kind = "unused", "display_only"
            elif doc_mentions.get(rel):
                role, unused_kind = "unused", "doc_only"
            else:
                d = rel.rsplit("/", 1)[0] if "/" in rel else ""
                role = "unknown"
                if en["exists"] and (d in live_write_dirs):
                    unused_kind = "orphan_output"
            exists = bool(en["exists"])
            if en.get("gdb_item") and exists is False:
                exists = False
            ds = {"id": did, "legacy_path": rel, "layer": layer, "kind": en["kind"] or "vector",
                  "format": self.format_of(rel, en), "exists": exists, "role": role,
                  "role_state": "structurally_inferred" if any(x["kind"] in ("code_write", "code_read", "project_layer")
                                                               and not x.get("guess") for x in ev) else
                  "ai_inferred" if any(x.get("guess") for x in ev) else (
                      "documented" if doc_mentions.get(rel) else "ai_inferred"),
                  "role_evidence": ev, "unused_kind": unused_kind, "conflict": conflict,
                  "produced_by": sorted({o["id"] for o, _ in W}),
                  "consumed_by": sorted({o["id"] for o, _ in R}),
                  "referenced_by": sorted({o["id"] for o, _ in all_readers.get(did, []) + all_writers.get(did, [])}),
                  "grid_id": None, "derived_from": None, "profile_ref": None, "sha256": None, "content_sha256": None,
                  "persistence": "gdb" if en.get("gdb_item") or en.get("file_kind") == "gdb" else "file"}
            if en.get("layer_missing"):
                ds["layer_missing"] = True
            if exists and rel in self.hashes:
                ds["sha256"] = self.hashes[rel]["sha256"]
            self.datasets[did] = ds
        # kinds of missing gdb items from how code uses them
        for did, ds in self.datasets.items():
            if ds["exists"] or ds["kind"] not in (None, "vector") or ds["persistence"] != "gdb":
                continue
            prod = [(o, x) for o in self.ops for x in o["writes"] if x["dataset"] == did and x["mode"] == "create"]
            uses = prod + [(o, x) for o in self.ops for x in o["reads"] if x["dataset"] == did]
            for o, x in uses:
                pn = (x.get("param") or "").lower()
                if o["api"] == "arcpy.Raster.save" or "raster" in pn:
                    ds["kind"] = "raster"
                    break
                if "feature" in pn or pn.endswith("_fc"):
                    ds["kind"] = "vector"
                    break
                if "table" in pn and (o, x) in prod:
                    ds["kind"] = "table"
                    break
        self.all_readers, self.all_writers = all_readers, all_writers
        for lm in self.layer_mismatch:
            did = self.ds_of_ref(lm["ref"])
            if did in self.datasets and lm.get("reason") != "layer_not_given":
                rec = {"requested": lm["layer"], "actual": lm["actual"],
                       "ops": sorted(raw["id"] for raw in self.ctx.ops.values() if "id" in raw and any(
                           x.get("ref") == lm["ref"] for x in raw["reads"] + raw["writes"])),
                       "raw": self.ctx.refs.get(lm["ref"], {}).get("raws", [])[:3]}
                if rec not in self.datasets[did].setdefault("layer_mismatch", []):
                    self.datasets[did]["layer_mismatch"].append(rec)

    def format_of(self, rel: str, en: dict) -> str:
        prof = self.profiles.get(rel, {})
        if prof.get("driver"):
            return prof["driver"]
        ext = file_ext(rel)
        return {".tif": "GTiff", ".tiff": "GTiff", ".gpkg": "GPKG", ".geojson": "GeoJSON", ".json": "GeoJSON",
                ".shp": "ESRI Shapefile", ".xlsx": "XLSX", ".xlsm": "XLSX", ".xls": "XLS", ".csv": "CSV",
                ".gdb": "OpenFileGDB", ".parquet": "Parquet", ".vrt": "VRT", ".nc": "netCDF"}.get(ext, ext.lstrip(".")
                                                                                                  or "unknown")

    # ------------------------------------------------------------------ branches (10 §10 branches[])
    def collect_branches(self) -> None:
        """Conditionals whose arms run different I/O ops or assign values that reach an op parameter."""
        out = []
        for rel, an in sorted(self.analysers.items()):
            c = self.codes.get(rel)
            if not c or not getattr(an, "branches", None):
                continue
            file_ops = [o for o in self.ops if o["file"] == rel]
            for (cell, line), rec in sorted(an.branches.items(), key=lambda kv: (kv[0][0] or -1, kv[0][1])):
                arms, affected = [], set()
                for arm in rec["arms"]:
                    ids = sorted(self.op_by_key[k]["id"] for k in arm["op_keys"] if k in self.op_by_key)
                    lines = set(arm["assign_lines"].values())
                    used = {}
                    for o in file_ops:
                        if o.get("cell") != cell and cell is not None:
                            continue
                        for pname, pv in o["params"].items():
                            hit = lines & set(pv.get("chain") or []) if isinstance(pv, dict) else set()
                            for name, ln in arm["assign_lines"].items():
                                if ln in hit:
                                    used.setdefault(name, set()).add(o["id"])
                    affected |= set(ids) | {x for v in used.values() for x in v}
                    arms.append({"arm": arm["arm"], "lines": arm["lines"], "ops": ids,
                                 "values": {k: v for k, v in sorted(arm["values"].items()) if k in used},
                                 "values_reach": {k: sorted(v) for k, v in sorted(used.items())}})
                if not affected:
                    continue
                sid = c.get("step_id") or "s00_x"
                bid = f"br.{short_step(sid)}." + (f"c{cell}." if cell is not None else "") + f"L{line}"
                writes = sorted({w["dataset"] for o in self.ops if o["id"] in affected for w in o["writes"]})
                out.append({"id": bid, "code": c["id"], "step": sid, "file": rel, "line": line,
                            "end_line": rec["end_line"], "cell": cell, "condition": rec["condition"], "arms": arms,
                            "affects_ops": sorted(affected), "affects_datasets": writes, "visits": rec["visits"],
                            "code_status": c["status"], "status": "unknown", "state": "directly_observed",
                            "ref": f"{rel}:{line}-{rec['end_line']}" if cell is None else f"{rel}#c{cell}:L{line}"})
        self.branch_list = out

    # ------------------------------------------------------------------ settings and params
    def finish_settings(self) -> None:
        raster_like = re.compile(r"^(sa\.|conversion\.\w*ToRaster|management\.(ProjectRaster|Resample|Clip|CopyRaster|"
                                 r"MosaicToNewRaster)|Raster$)")
        out = []
        status = {c["id"]: c["status"] for c in self.codes.values()}
        for s in self.ctx.settings:
            v = s.get("value_v")
            an = self.analysers.get(s["file"])
            val = None
            if v is not None and an is not None:
                refs = an.path_refs(v, s["line"], s["name"]) if v.k in ("const", "list", "obj", "layer") else []
                ds = [self.ds_of_ref(r["ref"]) for r in refs if r.get("ref")]
                if ds and ds[0]:
                    val = {"dataset": ds[0]}
                elif refs and refs[0].get("ref"):
                    val = {"path": self.ctx.refs[refs[0]["ref"]]["rel"]}
                elif v.k == "derived":
                    val = {"derived": v.v, "dataset": self.ds_of_ref(v.x.get("ref"))}
                else:
                    val = an.vjson(v)
            a, b = s["scope"]["from_line"], s["scope"]["to_line"]
            applies = [o["id"] for o in self.ops if o["file"] == s["file"] and o["family"] == "arcpy"
                       and o["line"] >= a and (b is None or o["line"] <= b)
                       and (s["name"] in ("workspace", "scratchWorkspace", "overwriteOutput")
                            or raster_like.search(o["tool"]) or o["tool"].startswith("sa."))]
            rec = {"name": s["name"], "family": "arcpy_env", "set_by": s["set_by"], "file": s["file"],
                   "line": s["line"], "code": s["code"], "value": val, "scope": s["scope"], "applies_to": applies,
                   "from_superseded": status.get(s["code"]) in ("superseded", "twin")}
            out.append(rec)
        # arcpy defaults never set for raster-producing ops (FH-13)
        report = self.ctx.arcpy.get("env", {}).get("default_unset_report", [])
        by_code = defaultdict(list)
        for o in self.ops:
            if o["family"] == "arcpy" and (o["tool"].startswith("sa.") or raster_like.search(o["tool"])) \
                    and o["api"] != "arcpy.Raster":
                by_code[o["code"]].append(o)
        for code, ops in sorted(by_code.items()):
            for name in report:
                unset = [o["id"] for o in ops if not any(
                    r["code"] == code and r["name"] == name and r["set_by"] != "default_unset"
                    and r["scope"]["from_line"] <= o["line"] and (r["scope"]["to_line"] is None
                                                                  or o["line"] <= r["scope"]["to_line"])
                    for r in out)]
                if unset:
                    f = self.op_by_id[unset[0]]["file"]
                    out.append({"name": name, "family": "arcpy_env", "set_by": "default_unset", "file": f, "line": None,
                                "code": code, "value": None, "scope": None, "applies_to": unset,
                                "from_superseded": status.get(code) in ("superseded", "twin")})
        # QGIS per-call grid parameters of raster-producing algorithms
        for o in self.ops:
            if o["family"] not in ("pyqgis", "qgis_model") or not o["tool"]:
                continue
            spec = self.ctx.algs.get(o["tool"], {})
            if not any(p.get("type") == "rasterDestination" for p in spec.get("params", {}).values()):
                continue
            for pname, pv in sorted(o["params"].items()):
                canon = GRID_PARAM_MAP.get(pname)
                if canon is None and pname != "EXTRA":
                    continue
                val = pv.get("value") if isinstance(pv, dict) else pv
                line = pv.get("arg_line") if isinstance(pv, dict) else o["line"]
                if pname == "EXTRA":
                    if isinstance(val, str):
                        toks = val.split()
                        for i, t in enumerate(toks):
                            if t in EXTRA_FLAG_MAP:
                                arg = toks[i + 1] if i + 1 < len(toks) and not toks[i + 1].startswith("-") else True
                                out.append(self.qsetting(o, EXTRA_FLAG_MAP[t], "tool_param", line,
                                                         True if arg is True else arg, via="EXTRA " + t))
                    continue
                if val is None:
                    sb = "default_unset"
                elif isinstance(val, dict) and val.get("derived"):
                    sb = "derived_from_dataset"
                else:
                    sb = "tool_param"
                out.append(self.qsetting(o, canon, sb, line, val, via=pname))
            if o["tool"] in ("gdal:rasterize", "gdal:rasterize_over", "gdal:rasterize_over_fixed_value"):
                extra = o["params"].get("EXTRA", {}).get("value") if isinstance(o["params"].get("EXTRA"), dict) else None
                if not (isinstance(extra, str) and "-at" in extra.split()):
                    out.append(self.qsetting(o, "all_touched", "default_unset",
                                             (o["params"].get("EXTRA") or {}).get("arg_line"), False, via="EXTRA"))
        for i, s in enumerate(out):
            st = self.step_of.get(s["file"], "x")
            s["id"] = f"set.{short_step(st)}." + (s.get("op_id", "").split(".")[-1] + "." if s.get("op_id") else "") + \
                s["name"] + (f".L{s['line']}" if s.get("line") and s["set_by"] in ("assignment", "env_manager") else "")
        seen = Counter()
        for s in sorted(out, key=lambda s: (s["id"], s.get("line") or 0)):
            seen[s["id"]] += 1
            if seen[s["id"]] > 1:
                s["id"] += f".{seen[s['id']]}"
        self.settings = sorted(out, key=lambda s: s["id"])

    def qsetting(self, o, name, set_by, line, value, via):
        status = {c["id"]: c["status"] for c in self.codes.values()}
        rec = {"name": name, "family": "qgis_param", "set_by": set_by, "file": o["file"], "line": line,
               "code": o["code"], "op_id": o["id"], "value": value, "via": via, "scope": {"op": o["id"]},
               "applies_to": [o["id"]], "from_superseded": status.get(o["code"]) in ("superseded", "twin")}
        opts = alg_spec(self.ctx.algs, o["tool"]).get("params", {}).get(via, {}).get("options") or {}
        if opts and str(value) in opts:  # decoded enum label (DATA_TYPE 5 -> Int32), M7
            rec["label"] = opts[str(value)]
        return rec

    def finish_params(self) -> None:
        status = {c["id"]: c["status"] for c in self.codes.values()}
        out = []
        loads = {}
        for rel, an in self.analysers.items():
            try:
                src = (self.root / rel).read_text(encoding="utf-8", errors="replace") if rel.endswith(".py") else \
                    "\n".join(u["src"] for u in notebook_units(self.root / rel)[0]) if rel.endswith(".ipynb") else ""
                loads[rel] = Counter(n.id for n in ast.walk(ast.parse(src)) if isinstance(n, ast.Name)
                                     and isinstance(n.ctx, ast.Load))
            except SyntaxError:
                loads[rel] = Counter()
        comments_at = {}
        for rel, an in self.analysers.items():
            for cell, cs in getattr(an, "comments", {}).items():
                for c in cs:
                    if not c["own_line"]:
                        comments_at[(rel, cell, c["line"])] = c["text"]
        defs = defaultdict(list)
        for p in self.ctx.params:
            defs[(p["file"], p["name"], p["scope"])].append(p)
        for (rel, name, scope), lst in sorted(defs.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2] or "")):
            lst.sort(key=lambda p: (p["cell"] or 0, p["line"]))
            sid = self.step_of.get(rel, "x")
            for i, p in enumerate(lst):
                pid = f"p.{short_step(sid)}." + (f"{scope}." if scope else "") + name + (f".L{p['line']}" if i else "")
                used_at = []
                for o in self.ops:
                    if o["file"] != rel:
                        continue
                    for pname, pv in o["params"].items():
                        if isinstance(pv, dict) and p["line"] in (pv.get("chain") or []):
                            nxt = lst[i + 1]["line"] if i + 1 < len(lst) else None
                            if nxt is None or o["line"] < nxt:
                                used_at.append({"op": o["id"], "param": pname, "line": pv.get("arg_line")})
                rec = {"id": pid, "name": name, "value": jsonable(p["value"]), "file": rel, "line": p["line"],
                       "comment": comments_at.get((rel, p["cell"], p["line"])), "used_at": used_at,
                       "doc_mentions": [], "shadowed_by": (f"{rel}:{lst[i + 1]['line']}" if i + 1 < len(lst) else None),
                       "flags": [], "kind": "assignment", "code": p["code"],
                       "from_superseded": status.get(p["code"]) in ("superseded", "twin")}
                if p["cell"] is not None:
                    rec["cell"] = p["cell"]
                if loads.get(rel, Counter())[name] == 0 and isinstance(p["value"], (int, float)) \
                        and not isinstance(p["value"], bool):
                    rec["flags"].append("dead_knob")
                out.append(rec)
        # anonymous literals in tool parameters
        for o in self.ops:
            if o["family"] not in ("pyqgis", "arcpy", "qgis_model"):
                continue
            ospec = alg_spec(self.ctx.algs, o["tool"]).get("params", {}) if o["family"] != "arcpy" else {}
            for pname, pv in sorted(o["params"].items()):
                val = pv.get("value") if isinstance(pv, dict) else None
                expr = o["family"] != "arcpy" and is_expression_param(pname, ospec.get(pname, {}))
                if isinstance(val, (int, float, bool)) or (isinstance(val, str) and (expr or not looks_like_path(val))
                                                            and len(val) < 200):
                    if isinstance(pv, dict) and pv.get("chain"):
                        continue  # a named constant: reported under its name
                    out.append({"id": f"p.{short_step(o['id'])}.{o['id'].split('.')[-1]}.{pname}", "name": pname,
                                "value": val, "file": o["file"], "line": pv.get("arg_line") or o["line"],
                                "kind": "tool_literal", "op": o["id"], "used_at": [{"op": o["id"], "param": pname,
                                                                                    "line": pv.get("arg_line")}],
                                "doc_mentions": [], "shadowed_by": None, "flags": [], "comment": None,
                                "code": o["code"], "from_superseded": status.get(o["code"]) in ("superseded", "twin")})
                    opts = ospec.get(pname, {}).get("options") or {}
                    if opts and str(val) in opts:  # decoded enum label next to the raw value (M7)
                        out[-1]["label"] = opts[str(val)]
                    cs = cell_area_expression(val) if expr and isinstance(val, str) else None
                    if cs:
                        out[-1]["cell_size_area"] = cs
        # numeric literals inside expression parameters (FORMULA `(30*30)/10000`): each number is a parameter (B5)
        for o in self.ops:
            if o["family"] not in ("pyqgis", "qgis_model"):
                continue
            spec = alg_spec(self.ctx.algs, o["tool"])
            for pname, pv in sorted(o["params"].items()):
                val = pv.get("value") if isinstance(pv, dict) else None
                ptype = spec.get("params", {}).get(pname, {}).get("type")
                if not isinstance(val, str) or not (ptype in ("expression", "raster_calc_expression") or
                                                    (ptype is None and pname.upper() in EXPR_PARAM_NAMES)):
                    continue
                for i, num in enumerate(expression_numbers(val), start=1):
                    out.append({"id": f"p.{short_step(o['id'])}.{o['id'].split('.')[-1]}.{pname.lower()}.n{i}",
                                "name": f"{pname}#{i}", "value": num, "file": o["file"],
                                "line": pv.get("arg_line") or o["line"], "kind": "expression_literal",
                                "expression": val, "op": o["id"],
                                "used_at": [{"op": o["id"], "param": pname, "line": pv.get("arg_line")}],
                                "doc_mentions": [], "shadowed_by": None, "flags": [], "comment": None,
                                "code": o["code"], "from_superseded": status.get(o["code"]) in ("superseded", "twin")})
        for l in self.ctx.literals:
            sid = self.step_of.get(l["file"], "x")
            rec = {"id": f"p.{short_step(sid)}.L{l['line']}.{l['context']}", "name": l.get("column"),
                   "value": l["value"], "file": l["file"], "line": l["line"],
                   "kind": "dtype" if l["context"] == "dtype" else "inline_literal", "context": l["context"],
                   "used_at": [], "doc_mentions": [], "shadowed_by": None, "flags": [],
                   "comment": comments_at.get((l["file"], l["cell"], l["line"])), "code": l["code"],
                   "from_superseded": status.get(l["code"]) in ("superseded", "twin")}
            for k in ("via", "cell_size_area"):
                if l.get(k):
                    rec[k] = l[k]
            out.append(rec)
        # documented constants vs code
        for d in self.docs:
            for c in d["constants"]:
                for p in out:
                    if p["name"] == c["name"]:
                        agrees = isinstance(p["value"], (int, float)) and float(p["value"]) == float(c["value"])
                        p["doc_mentions"].append({"file": d["path"], "line": c["line"], "value": c["value"],
                                                  "unit": c["unit"], "agrees": agrees})
            for m in d["multipliers"]:
                for p in out:
                    if p["kind"] == "inline_literal" and p.get("context") == "mult" and \
                            float(p["value"]) == float(m["value"]):
                        p["doc_mentions"].append({"file": d["path"], "line": m["line"], "value": m["value"],
                                                  "agrees": True, "text": m["text"]})
        seen = Counter()
        for p in sorted(out, key=lambda p: (p["id"], p["line"] or 0)):
            seen[p["id"]] += 1
            if seen[p["id"]] > 1:
                p["id"] += f".{seen[p['id']]}"
        self.params = sorted(out, key=lambda p: p["id"])

    # ------------------------------------------------------------------ grids
    def grids(self) -> None:
        rasters = [d for d in self.datasets.values() if d["kind"] == "raster" and d["exists"]
                   and "transform" in self.profiles.get(d["legacy_path"], {})]
        groups: dict[tuple, list] = defaultdict(list)
        for d in rasters:
            p = self.profiles[d["legacy_path"]]
            crs = (p.get("crs") or {}).get("wkt")
            t = p["transform"]
            key = (hashlib.sha256((crs or "").encode()).hexdigest()[:16], round(t[0], 12), round(t[4], 12))
            groups[key].append(d)
        self.grid_list = []
        for key, members in sorted(groups.items(), key=lambda kv: sorted(d["id"] for d in kv[1])[0]):
            members.sort(key=lambda d: d["id"])

            def offsets(ref):
                t0 = self.profiles[ref["legacy_path"]]["transform"]
                o = {}
                for d in members:
                    t = self.profiles[d["legacy_path"]]["transform"]
                    o[d["id"]] = ((t[2] - t0[2]) / t0[0], (t[5] - t0[5]) / t0[4])
                return o

            def aligned(off):
                return all(abs(v - round(v)) < 1e-6 for v in off)
            pool = members[:]
            while pool:
                cands = sorted(pool, key=lambda d: (d["role"] != "input", -sum(aligned(v) for v in offsets(d).values()
                                                                               if True), d["id"]))
                ref = cands[0]
                off = offsets(ref)
                fam = [d for d in pool if aligned(off[d["id"]])]
                gid = "grid." + ref["id"][3:]
                p0 = self.profiles[ref["legacy_path"]]
                rec = {"id": gid, "reference": ref["id"], "crs_epsg": (p0.get("crs") or {}).get("epsg"),
                       "res": [p0["transform"][0], p0["transform"][4]], "origin": p0["origin"],
                       "members": [{"dataset": d["id"], "offset_px": [round(off[d["id"]][0], 9) + 0.0,
                                                                       round(off[d["id"]][1], 9) + 0.0],
                                    "shape": self.profiles[d["legacy_path"]]["shape"],
                                    "same_transform": self.profiles[d["legacy_path"]]["transform"] == p0["transform"]}
                                   for d in fam],
                       "misaligned": []}
                for d in fam:
                    d["grid_id"] = gid
                    pool.remove(d)
                for d in pool:
                    rec["misaligned"].append({"dataset": d["id"], "offset_px": [round(off[d["id"]][0], 6),
                                                                                round(off[d["id"]][1], 6)]})
                self.grid_list.append(rec)

    # ------------------------------------------------------------------ derived copies (manual steps)
    def derived_copies(self) -> None:
        import shapely  # noqa: F401
        vecs = []
        for did, d in self.datasets.items():
            if d["kind"] != "vector" or not d["exists"]:
                continue
            kv = self.keep_values.get((d["legacy_path"], d["layer"]))
            if kv:
                vecs.append((d, kv))
        written = {did for did in self.datasets if self.all_writers.get(did)}
        self.dcopies, self.manual = [], []
        for i, (a, ka) in enumerate(vecs):
            for b, kb in vecs[i + 1:]:
                if a["id"] in written or b["id"] in written:
                    continue
                if len(ka["wkb"]) != len(kb["wkb"]) or sorted(ka["wkb"]) != sorted(kb["wkb"]) or not ka["wkb"]:
                    continue
                a_read = bool(self.all_readers.get(a["id"]))
                b_read = bool(self.all_readers.get(b["id"]))
                if a_read and not b_read:
                    src, dst, ks, kd = b, a, kb, ka
                elif b_read and not a_read:
                    src, dst, ks, kd = a, b, ka, kb
                else:
                    src, dst, ks, kd = (a, b, ka, kb) if len(kb["cols"]) >= len(ka["cols"]) else (b, a, kb, ka)
                self.dcopies.append(self.dc_record(src, dst, ks, kd))
        for n, dc in enumerate(sorted(self.dcopies, key=lambda d: (d["target"], d["source"])), start=1):
            dc["id"] = f"dc.{n:02d}"
        self.dcopies.sort(key=lambda d: d["id"])
        for n, dc in enumerate(self.dcopies, start=1):
            tgt = self.datasets[dc["target"]]
            src = self.datasets[dc["source"]]
            mid = f"m{n:02d}_{ac.slug(stem_of(tgt['legacy_path']))}"
            dc["manual_step"] = mid
            tgt["derived_from"] = src["id"]
            tgt["role_evidence"].append({"kind": "derived_copy", "ref": dc["id"]})
            if src["role"] in ("unused", "unknown"):
                src["role"] = "input"
                src["unused_kind"] = None
                src["role_state"] = "structurally_inferred"
            src["role_evidence"].append({"kind": "derived_copy_source", "ref": dc["id"], "manual_step": mid})
            src["input_of_manual_step"] = mid
            self.manual.append({"id": mid, "inputs": [src["id"]], "outputs": [tgt["id"]], "derived_copy": dc["id"],
                                "evidence": dc["evidence"], "kind": "manual"})

    def dc_record(self, src, dst, ks, kd) -> dict:
        sc, dcl = set(ks["cols"]), set(kd["cols"])
        sp = self.profiles.get(src["legacy_path"], {})
        dp = self.profiles.get(dst["legacy_path"], {})

        def schema(p, layer):
            for l in p.get("layers", []):
                if l["name"] == layer:
                    return [c["name"] for c in l.get("schema", [])]
            return []
        s_cols, d_cols = schema(sp, src["layer"]), schema(dp, dst["layer"])
        added = sorted(set(d_cols) - set(s_cols))
        removed = sorted(set(s_cols) - set(d_cols))
        # pair features by WKB
        pos_s = defaultdict(list)
        for i, w in enumerate(ks["wkb"]):
            pos_s[w].append(i)
        pairs = []
        for j, w in enumerate(kd["wkb"]):
            if pos_s.get(w):
                pairs.append((pos_s[w].pop(0), j))
        value_maps = []
        changed = sorted(c for c in sc & dcl if any(ks["cols"][c][i] != kd["cols"][c][j] for i, j in pairs))
        for col in added + changed:
            if col not in kd["cols"]:
                continue
            best = None
            for scol in sorted(sc):
                mapping = [(ks["cols"][scol][i], kd["cols"][col][j]) for i, j in pairs]
                rule = self.infer_rule(mapping)
                if best is None or rule["rule_count"] > best["rule_count"]:
                    best = dict(rule, from_column=scol)
            if best:
                best["column"] = col
                value_maps.append(best)
        evidence = [{"kind": "data_diff", "summary": f"{len(pairs)}/{len(kd['wkb'])} features with equal WKB"}]
        names = {src["legacy_path"].split("/")[-1].lower(), dst["legacy_path"].split("/")[-1].lower()}
        for d in self.docs:
            lines = [m for m in d["manual"] if not m.get("header")]
            if lines and any(any(n in m["text"].lower() for n in names) for m in d["manual"]):
                head = next((m for m in d["manual"] if m.get("header") and m["line"] < lines[0]["line"]), lines[0])
                evidence.append({"kind": "doc_mention", "file": d["path"], "line": head["line"],
                                 "end_line": lines[-1]["line"]})
        return {"source": src["id"], "target": dst["id"], "count": len(kd["wkb"]), "wkb_equal": len(pairs),
                "schema": {"added": added, "removed": removed, "kept": sorted(set(s_cols) & set(d_cols))},
                "value_maps": value_maps, "evidence": evidence, "state": "directly_observed"}

    @staticmethod
    def infer_rule(mapping: list[tuple]) -> dict:
        pre, suf = Counter(), Counter()
        ident = 0
        for s, t in mapping:
            if s is None or t is None:
                continue
            if s == t:
                ident += 1
            if s.endswith(t) and len(s) > len(t):
                pre[s[: len(s) - len(t)]] += 1
            if s.startswith(t) and len(s) > len(t):
                suf[s[len(t):]] += 1
        cands = [("identity", None, ident)]
        if pre:
            p, n = pre.most_common(1)[0]
            cands.append(("strip_prefix", p, n))
        if suf:
            p, n = suf.most_common(1)[0]
            cands.append(("strip_suffix", p, n))
        cf = sum(1 for s, t in mapping if s is not None and t is not None and s.strip().casefold() == t.casefold()
                 and s != t)
        cands.append(("casefold_strip", None, cf + ident))
        kind, val, n = max(cands, key=lambda c: (c[2], c[0] != "identity"))

        def apply(s):
            if s is None:
                return None
            if kind == "strip_prefix":
                return s[len(val):] if s.startswith(val) else s
            if kind == "strip_suffix":
                return s[: -len(val)] if s.endswith(val) else s
            if kind == "casefold_strip":
                return s.strip().casefold()
            return s
        exceptions, n_exc = {}, 0
        for s, t in mapping:
            r = apply(s)
            if (r.casefold() if (kind == "casefold_strip" and r is not None) else r) != \
                    (t.casefold() if (kind == "casefold_strip" and t is not None) else t):
                n_exc += 1
                exceptions.setdefault(str(s), t)
        consistent = len(exceptions) == n_exc  # a many-to-one source column cannot explain the target
        return {"rule": {"kind": kind, "value": val}, "rule_count": len(mapping) - n_exc if consistent else 0,
                "exceptions": dict(sorted(exceptions.items())), "pairs": len(mapping)}

    # ------------------------------------------------------------------ Excel header corroboration
    def excel_corroboration(self) -> None:
        for did, d in self.datasets.items():
            prof = self.profiles.get(d["legacy_path"], {})
            if prof.get("format") != "xlsx":
                continue
            for o in self.ops:
                if not any(r["dataset"] == did for r in o["reads"]):
                    continue
                sheet = (o["params"].get("sheet") or {}).get("value")
                rows = (o["params"].get("rows") or {}).get("value") or []
                for rr in rows:
                    sh = rr.get("sheet") if rr.get("sheet") not in (None, "<active>") else sheet
                    for s in prof["sheets"]:
                        if (sh and s["name"] != sh) or (not sh and s is not prof["sheets"][0]):
                            continue
                        mn = rr.get("min_row")
                        if isinstance(mn, int) and mn > 1:
                            hdr = mn - 1
                            cand = next((c for c in s["header_candidates"] if c["row"] == hdr), None)
                            s["header_row"] = hdr
                            s["header_state"] = "resolved" if cand else "code_only"
                            s["data_rows"] = [mn, rr.get("max_row")]
                            s["header_corroboration"] = {"op": o["id"], "file": o["file"], "line": rr.get("line"),
                                                         "min_row": mn, "max_row": rr.get("max_row"),
                                                         "candidate": bool(cand)}
                for pname in ("header", "skiprows"):
                    pv = o["params"].get(pname)
                    if isinstance(pv, dict) and isinstance(pv.get("value"), int):
                        hdr = pv["value"] + 1 if pname == "header" else pv["value"] + 1
                        for s in prof["sheets"][:1]:
                            s["header_row"] = hdr
                            s["header_state"] = "resolved" if any(c["row"] == hdr for c in s["header_candidates"]) \
                                else "code_only"
                            s["header_corroboration"] = {"op": o["id"], "file": o["file"], "line": pv.get("arg_line")}

    # ------------------------------------------------------------------ join coverage
    def join_coverage(self) -> None:
        self.joins = []
        tables = []
        for did, d in self.datasets.items():
            prof = self.profiles.get(d["legacy_path"], {})
            if prof.get("format") == "xlsx":
                for s in prof["sheets"]:
                    if s.get("header_row") and s.get("data_rows"):
                        a, b = s["data_rows"][0], s["data_rows"][1] or s["max_row"]
                        hdr = s["_values"].get(s["header_row"], {})
                        for col, name in hdr.items():
                            vals = [s["_values"].get(r, {}).get(col) for r in range(a, b + 1)]
                            vals = [v for v in vals if isinstance(v, str)]
                            if vals:
                                tables.append((did, s["name"], f"{col} ({name})", vals))
        if not tables:
            return
        step_reads = defaultdict(set)
        for o in self.ops:
            for r in o["reads"]:
                step_reads[o["code"]].add(r["dataset"])
        for tdid, sheet, col, vals in tables:
            same_step = {x for c, s in step_reads.items() if tdid in s for x in s}
            for (rel, layer), kv in sorted(self.keep_values.items()):
                vdid = self.ent_id.get((rel, layer))
                if not vdid:
                    continue
                rows = []
                for vcol, vvals in sorted(kv["cols"].items()):
                    vset = {v for v in vvals if v is not None}
                    vset_n = {v.strip().casefold() for v in vset}
                    exact = sum(1 for v in vals if v in vset)
                    strip = sum(1 for v in vals if v.strip().casefold() in vset_n)
                    tnorm = {v.strip().casefold() for v in vals}
                    rows.append({"vector_column": vcol, "table_values": len(vals), "matched_exact": exact,
                                 "matched_after_strip": strip, "vector_values": len(vset),
                                 "vector_matched": sum(1 for v in vset_n if v in tnorm)})
                vd = self.datasets.get(vdid, {})
                relevant = vd.get("role") == "input" or (vdid in same_step and not any(
                    self.op_by_id[o]["code"] in {self.op_by_id[c]["code"] for c in self.datasets[tdid]["consumed_by"]}
                    for o in vd.get("produced_by", [])))
                if relevant and any(r["matched_after_strip"] >= 0.8 * len(vals) for r in rows):
                    for r in rows:
                        self.joins.append({"table": tdid, "sheet": sheet, "column": col, "vector": vdid, **r})

    # ------------------------------------------------------------------ extent relations
    def extent_relations(self) -> None:
        import numpy as np
        import shapely
        from pyproj import CRS, Transformer
        self.extents = []
        self.aoi = None
        inputs = [d for d in self.datasets.values() if d["role"] == "input" and d["exists"]]
        polys = []
        for d in inputs:
            if d["kind"] != "vector":
                continue
            prof = next((l for l in self.profiles.get(d["legacy_path"], {}).get("layers", [])
                         if l["name"] == d["layer"]), None)
            if not prof or not prof.get("geom_types"):
                continue
            if any(k in ("Polygon", "MultiPolygon") for k in prof["geom_types"]):
                polys.append((d, prof))
        if not polys:
            return

        def score(dp):
            d, prof = dp
            tools = " ".join(self.op_by_id[o]["tool"] for o in d["consumed_by"] if o in self.op_by_id).lower()
            s = sum(2 for w in ("dissolve", "clip", "mask", "zonal", "extractbymask") if w in tools)
            return (-s, prof.get("count", 0), d["id"])
        aoi, aoi_prof = sorted(polys, key=score)[0]
        self.aoi = aoi["id"]
        try:
            ga = read_layer(self.root / aoi["legacy_path"], aoi["layer"])
        except Exception as exc:  # noqa: BLE001
            self.log(f"extent relations skipped: {exc}")
            return
        crs_a = ga.crs
        u = shapely.union_all(shapely.make_valid(ga.geometry.values))
        to4326 = Transformer.from_crs(crs_a, 4326, always_xy=True)
        c = u.centroid
        lon, lat = to4326.transform(c.x, c.y)
        laea = laea_for(lon, lat)
        tr_a = Transformer.from_crs(crs_a, laea, always_xy=True)

        def to_m(g, tr):
            return shapely.transform(g, lambda xy: np.column_stack(tr.transform(xy[:, 0], xy[:, 1])))
        u_m = to_m(u, tr_a)
        aoi_area = u_m.area / 1e6
        measured = f"LAEA lat_0={lat:.4f} lon_0={lon:.4f}"
        for d in sorted(inputs, key=lambda d: d["id"]):
            if d["id"] == aoi["id"]:
                continue
            try:
                if d["kind"] == "raster":
                    p = self.profiles[d["legacy_path"]]
                    rc = CRS.from_wkt(p["crs"]["wkt"])
                    b = p["bounds"]
                    box = shapely.box(*b)
                    tr_r = Transformer.from_crs(rc, laea, always_xy=True)
                    box_m = to_m(shapely.segmentize(box, max(abs(b[2] - b[0]), abs(b[3] - b[1])) / 200), tr_r)
                    out_m = u_m.difference(box_m)
                    self.extents.append({"a": aoi["id"], "b": d["id"], "relation": "a_outside_b_extent",
                                         "within": out_m.is_empty or out_m.area < 1.0,
                                         "area_outside_km2": round(out_m.area / 1e6, 4),
                                         "a_area_km2": round(aoi_area, 4), "measured_in": measured})
                    nd = self.nodata_inside(d, u, crs_a)
                    if nd is not None:
                        self.extents.append({"a": d["id"], "b": aoi["id"], "relation": "nodata_inside_aoi",
                                             **nd, "measured_in": "cell areas on the ellipsoid (LAEA per cell row)"})
                elif d["kind"] == "vector":
                    g = read_layer(self.root / d["legacy_path"], d["layer"])
                    tr = Transformer.from_crs(g.crs, crs_a, always_xy=True) if g.crs and crs_a and \
                        not CRS(g.crs).equals(CRS(crs_a)) else None
                    geoms = g.geometry.values
                    if tr is not None:
                        geoms = shapely.transform(geoms, lambda xy: np.column_stack(tr.transform(xy[:, 0], xy[:, 1])))
                    inter = shapely.intersects(geoms, u)
                    within = shapely.within(geoms, u)
                    rec = {"a": d["id"], "b": aoi["id"], "relation": "features_vs_aoi", "count": int(len(g)),
                           "features_outside": int((~inter).sum()), "features_not_within": int((~within).sum()),
                           "within": bool(within.all()), "measured_in": measured}
                    by_value = {}
                    prof = next((l for l in self.profiles.get(d["legacy_path"], {}).get("layers", [])
                                 if l["name"] == d["layer"]), {})
                    for col in sorted(prof.get("domains", {})):
                        if len(prof["domains"][col]) <= 20 and col in g.columns:
                            vc = Counter(str(v) for v in g.loc[~inter, col].tolist())
                            if vc:
                                by_value[col] = dict(sorted(vc.items()))
                    rec["features_outside_by_value"] = by_value
                    gm = to_m(shapely.union_all(shapely.make_valid(geoms)), Transformer.from_crs(crs_a, laea,
                                                                                                always_xy=True))
                    if gm.area > 0:
                        rec["area_km2"] = round(gm.area / 1e6, 4)
                        rec["area_outside_km2"] = round(gm.difference(u_m).area / 1e6, 4)
                        rec["overlap_pct"] = round(100 * gm.intersection(u_m).area / gm.area, 3)
                    self.extents.append(rec)
            except Exception as exc:  # noqa: BLE001
                self.log(f"extent relation {d['id']} failed: {exc}")

    def nodata_inside(self, d, aoi_geom, aoi_crs) -> dict | None:
        import numpy as np
        import rasterio
        from rasterio import features
        from pyproj import CRS, Geod, Transformer
        import shapely
        with rasterio.open(self.root / d["legacy_path"]) as src:
            if src.width * src.height > 50_000_000:
                return None
            rc = CRS.from_wkt(src.crs.to_wkt())
            g = aoi_geom
            if not rc.equals(CRS(aoi_crs)):
                tr = Transformer.from_crs(aoi_crs, rc, always_xy=True)
                g = shapely.transform(g, lambda xy: np.column_stack(tr.transform(xy[:, 0], xy[:, 1])))
            inside = features.geometry_mask([g], out_shape=(src.height, src.width), transform=src.transform,
                                            invert=True)
            a = src.read(1)
            nd = src.nodata
            bad = np.isnan(a) if np.issubdtype(a.dtype, np.floating) else np.zeros(a.shape, bool)
            if nd is not None:
                bad |= a == np.array(nd).astype(a.dtype)
            cells = inside & bad
            t = src.transform
            if rc.is_geographic:
                geod = Geod(ellps="WGS84")
                rows = np.arange(src.height)
                lat0 = t.f + rows * t.e
                area_row = np.array([abs(geod.polygon_area_perimeter(
                    [t.c, t.c + t.a, t.c + t.a, t.c], [la, la, la + t.e, la + t.e])[0]) for la in lat0])
                km2 = float((cells.sum(axis=1) * area_row).sum() / 1e6)
            else:
                km2 = float(cells.sum() * abs(t.a * t.e) / 1e6)
            return {"cells_inside_aoi": int(inside.sum()), "nodata_cells_inside_aoi": int(cells.sum()),
                    "nodata_inside_aoi_km2": round(km2, 4), "rule": "cell centre"}

    # ------------------------------------------------------------------ open items
    def open_items(self) -> None:
        items = []

        def add(kind, severity, about, evidence, state="inferred", **kw):
            items.append({"kind": kind, "severity": severity, "about": about, "evidence": evidence, "state": state,
                          "agent_guess": None, **kw})
        live_codes = {c["id"] for c in self.codes.values() if c["status"] == "live"}
        for did, d in sorted(self.datasets.items()):
            if d["role"] == "input" and not d["exists"]:
                cons = sorted(set(d["consumed_by"]) | {o for o in d["referenced_by"]})
                add("missing_input", "high", [did], [{"kind": "code_read", "op": o} for o in cons], blocking=True,
                    consumers=cons)
            elif not d["exists"] and d["role"] in ("intermediate", "final"):
                add("missing_output", "medium", [did] + d["produced_by"],
                    [{"kind": "code_write", "op": o} for o in d["produced_by"]])
            if d["role"] == "unknown" and d["exists"]:
                if d.get("unused_kind") == "orphan_output":
                    add("orphan_output", "medium", [did], [{"kind": "folder", "value": d["legacy_path"].split("/")[0]}])
                else:
                    add("unknown_role", "medium" if d.get("conflict") else "low", [did], d["role_evidence"][:3])
            if d.get("layer_missing"):
                add("layer_missing", "high", [did], [{"kind": "code_read", "op": o} for o in d["referenced_by"]])
        for m in self.manual:
            add("manual_step", "high", [m["id"]] + m["outputs"], m["evidence"] + [{"kind": "data_diff",
                                                                                  "ref": m["derived_copy"]}])
        for c in self.codes.values():
            nb = c.get("notebook")
            if nb and (not nb["monotonic"] or nb.get("duplicates") or nb.get("unexecuted_code_cells")) \
                    and c["status"] != "twin":
                add("untrusted_notebook_outputs", "medium", [c["id"]],
                    [{"kind": "exec_counts", "value": nb["exec_counts"]}])
        for t in self.twins:
            if t["kind"] == "checkpoint" and t.get("cells_differ"):
                add("twin_conflict", "medium", [t["a"], t["b"]], [{"kind": "cells_differ", "cells": t["cells_differ"]}],
                    cells=t["cells_differ"])
        for p in self.projects:
            for u in p.get("unparsed", []):
                add("unparsed_format", "low", [p["path"]], [{"kind": "parser", "text": u}])
        for rel, prof in sorted(self.profiles.items()):
            if prof.get("unparsed") or prof.get("error"):
                add("unparsed_format", "low", [rel], [{"kind": "profile", "text": prof.get("unparsed") or prof["error"]}])
        for e in self.entries:
            if e["kind"] == "listed_only":
                add("unparsed_format", "low", [e["path"]], [{"kind": "format", "value": file_ext(e["path"])}])
        for g in self.grid_list:
            for m in g["misaligned"]:
                add("misaligned_grid", "high", [m["dataset"], g["id"]], [{"kind": "grid_offset", "value": m["offset_px"]}])
        for did, d in sorted(self.datasets.items()):
            prof = self.profiles.get(d["legacy_path"], {})
            for s in prof.get("sheets", []) if prof.get("format") == "xlsx" else []:
                if s.get("header_state") != "resolved" and (d["role"] == "input" or d["consumed_by"]):
                    add("header_unresolved", "medium", [did], [{"kind": "header_candidates",
                                                                "value": [c["row"] for c in s["header_candidates"]]}],
                        sheet=s["name"])
                if s.get("formula_no_cache"):
                    add("formula_no_cache", "low", [did], [{"kind": "cells", "sheet": s["name"],
                                                            "cells": s["formula_no_cache"]}], sheet=s["name"])
        for p in self.params:
            if "dead_knob" in p["flags"] and not p.get("from_superseded"):
                add("dead_knob", "low", [p["id"]], [{"kind": "code", "file": p["file"], "line": p["line"]}])
            for dm in p["doc_mentions"]:
                if not dm["agrees"]:
                    add("doc_code_conflict", "medium", [p["id"]], [{"kind": "doc_mention", "file": dm["file"],
                                                                     "line": dm["line"], "value": dm["value"]},
                                                                    {"kind": "code", "file": p["file"],
                                                                     "line": p["line"], "value": p["value"]}])
        # documented nominal cell size vs the real raster grid
        ins = [d for d in self.datasets.values() if d["kind"] == "raster" and d["exists"] and d["role"] == "input"]
        for doc in self.docs:
            for cs in doc["cell_sizes"]:
                for d in ins:
                    rm = self.profiles[d["legacy_path"]].get("res_m_at_centre")
                    if not rm:
                        continue
                    if max(abs(r - cs["value_m"]) for r in rm) / cs["value_m"] > 0.2:
                        p = self.profiles[d["legacy_path"]]
                        add("doc_code_conflict", "info", [d["id"]],
                            [{"kind": "doc_mention", "file": doc["path"], "line": cs["line"], "text": cs["text"]},
                             {"kind": "raster_profile", "dataset": d["id"], "res": p["res"],
                              "res_m_at_centre": rm}],
                            detail=f"documented {cs['value_m']:g} m vs grid {abs(p['res'][0]):.10g} x "
                                   f"{abs(p['res'][1]):.10g} {(p.get('crs') or {}).get('units') or ''} "
                                   f"= {rm[0]:.1f} x {rm[1]:.1f} m at the centre")
        # filters that use only part of a categorical domain (possible missing branch)
        for o in self.ops:
            if self.op_by_id[o["id"]]["code"] not in live_codes:
                continue
            for pname, pv in o["params"].items():
                val = pv.get("value") if isinstance(pv, dict) else None
                if not isinstance(val, str) or pname not in ("EXPRESSION", "where_clause", "SQL", "query", "FILTER"):
                    continue
                conds = parse_subset(val)
                for src in self.origin_datasets(o):
                    d = self.datasets.get(src)
                    if not d:
                        continue
                    prof = next((l for l in self.profiles.get(d["legacy_path"], {}).get("layers", [])
                                 if l["name"] == d["layer"]), {})
                    for c in conds:
                        dom = prof.get("domains", {}).get(c["field"])
                        if c["op"] == "=" and dom and c["value"] in dom and len(dom) > 1:
                            unused = [v for v in dom if v != c["value"]]
                            ev = [{"kind": "code", "op": o["id"], "file": o["file"],
                                   "line": pv.get("arg_line") or o["line"]}]
                            for doc in self.docs:
                                for i, line in enumerate(doc_text(self.root / doc["path"]), start=1):
                                    if sum(1 for v in unused if v.lower() in line.lower()) >= max(1, len(unused) - 1):
                                        ev.append({"kind": "doc_mention", "file": doc["path"], "line": i})
                            add("unused_branch", "info", [src, o["id"]], ev, field=c["field"], used=[c["value"]],
                                unused=unused)
        for oc in self.order_conflicts:
            add("order_conflict", "medium", [], [oc])
        for c in sorted(self.codes.values(), key=lambda c: c["path"]):
            for f in c.get("flags", []):
                if f["kind"] == "reversed_containment" and f.get("empty_matches") is not None:
                    add("reversed_containment", "medium", [c["id"]],
                        [{"kind": "code", "file": c["path"], "line": f["line"], "loop_line": f["loop_line"]}],
                        function=f.get("function"), arg=f["arg"], empty_matches=f["empty_matches"],
                        detail=f"{f.get('function') or 'module'}(): {f['note']}; an empty or missing value is "
                               f"classified as {f['empty_matches']!r}")
        # a nominal cell size used for areas next to a geographic grid (M7): `(30*30)/10000` ha per cell vs the real
        # cell of the rasters the same code reads (else the input rasters), at the AOI's latitude when there is one
        nominal = defaultdict(list)
        for p in self.params:
            if p.get("cell_size_area") and not p.get("from_superseded"):
                nominal[p["code"]].append(p)
        aoi_lat = None
        if self.aoi:
            try:
                b = self.profiles[self.datasets[self.aoi]["legacy_path"]]
                lay = next((l for l in b.get("layers", []) if l["name"] == self.datasets[self.aoi]["layer"]), {})
                if lay.get("bounds") and (lay.get("crs") or {}).get("is_geographic"):
                    aoi_lat = (lay["bounds"][1] + lay["bounds"][3]) / 2
            except (KeyError, TypeError):
                aoi_lat = None
        for cid, ps in sorted(nominal.items()):
            own = sorted({r["dataset"] for o in self.ops if o["code"] == cid for r in o["reads"]
                          if self.datasets.get(r["dataset"], {}).get("kind") == "raster"
                          and self.datasets[r["dataset"]]["exists"]})
            pool = own or sorted(d for d, x in self.datasets.items() if x["kind"] == "raster" and x["exists"]
                                 and x["role"] == "input" and not x.get("alias_of"))
            rasters = []
            for did in pool:
                pr = self.profiles.get(self.datasets[did]["legacy_path"], {})
                if not (pr.get("crs") or {}).get("is_geographic") or not pr.get("res"):
                    continue
                from pyproj import Geod
                lat = aoi_lat if aoi_lat is not None else (pr["bounds"][1] + pr["bounds"][3]) / 2
                g = Geod(ellps="WGS84")
                _, _, dx = g.inv(0, lat, abs(pr["res"][0]), lat)
                _, _, dy = g.inv(0, lat, 0, lat + abs(pr["res"][1]))
                rasters.append({"dataset": did, "res_deg": [abs(pr["res"][0]), abs(pr["res"][1])],
                                "res_m": [round(dx, 3), round(dy, 3)], "at_lat": round(lat, 4),
                                "lat_from": "aoi" if aoi_lat is not None else "raster_centre",
                                "cell_area_m2": round(dx * dy, 2)})
            if not rasters:
                continue
            p0 = ps[0]
            cs = p0["cell_size_area"]
            add("nominal_vs_real_cell", "medium", [cid] + [r["dataset"] for r in rasters],
                [{"kind": "code", "file": p["file"], "line": p["line"], "param": p["id"]} for p in ps]
                + [{"kind": "raster_profile", **r} for r in rasters], nominal_m=cs["cell_m"], unit=cs["unit"],
                area_per_cell=cs["area"], rasters_from="same_code" if own else "inventory_inputs",
                detail=f"nominal {cs['cell_m']:g} m cell = {cs['area']:g} {cs['unit']} per cell; real cells "
                       + "; ".join(f"{r['dataset']} {r['res_deg'][0]:.6g}° = {r['res_m'][0]:.2f} x "
                                   f"{r['res_m'][1]:.2f} m at {r['at_lat']:.2f}° ({r['cell_area_m2'] / (1e4 if cs['unit'] == 'ha' else 1e6):.5g} {cs['unit']})"
                                   for r in rasters))
        # custom expression functions (M1): a function the expression calls that is not a QGIS built-in and is not
        # shipped with the model; a Python def of the same name is cross-referenced as a guess
        builtin = {n.lower() for n in (self.ctx.qgis.get("expression_functions") or {}).get("names", [])}
        qv = str((self.ctx.qgis.get("expression_functions") or {}).get("qgis_version") or "").split("-")[0]
        defs = defaultdict(list)
        for rel, an in sorted(self.analysers.items()):
            for name, f in an.funcs.items():
                deco = [ast.unparse(d) for d in getattr(f["node"], "decorator_list", [])]
                defs[name.lower()].append({"file": rel, "line": f["node"].lineno, "name": name,
                                           "qgsfunction": any("qgsfunction" in d for d in deco)})
        rev = {}
        for c in self.codes.values():
            for f in c.get("flags", []):
                if f["kind"] == "reversed_containment" and f.get("empty_matches") is not None:
                    rev[(c["path"], f.get("function"))] = f
        seen = {}
        for o in self.ops:
            if o["family"] not in ("qgis_model", "pyqgis") or not builtin:
                continue
            spec = alg_spec(self.ctx.algs, o["tool"]).get("params", {})
            for pname, pv in sorted(o["params"].items()):
                val = pv.get("value") if isinstance(pv, dict) else None
                ptype = spec.get(pname, {}).get("type")
                if not isinstance(val, str) or not (ptype in ("expression",) or (ptype is None and pname.upper() in
                                                                                 EXPR_PARAM_NAMES)):
                    continue
                for fn in expression_functions(val):
                    if fn.lower() in builtin:
                        continue
                    ev = {"kind": "expression", "op": o["id"], "file": o["file"], "line": pv.get("arg_line") or o["line"],
                          "param": pname, "text": val.strip()}
                    if (o["code"], fn) in seen:
                        seen[(o["code"], fn)]["evidence"].append(ev)
                        continue
                    cands = defs.get(fn.lower(), [])
                    add("custom_expression_function", "high", [o["code"], o["id"]], [ev],
                        state="agent_guess" if cands else "inferred", function=fn,
                        builtin_table=qv, python_defs=cands,
                        detail=f"{fn}() is not a QGIS {qv} built-in expression function and is not shipped with {o['file']}; the model fails "
                               f"with 'Function {fn} is not known' unless it is registered")
                    it = items[-1]
                    if cands:
                        c0 = sorted(cands, key=lambda c: (not c["qgsfunction"], c["file"]))[0]
                        r = rev.get((c0["file"], c0["name"]))
                        it["agent_guess"] = {"python_def": f"{c0['file']}:{c0['line']}",
                                             "why": "same name" + ("; decorated with @qgsfunction" if c0["qgsfunction"]
                                                                   else "; a plain Python function, not registered "
                                                                        "as a QGIS function: same logic is a guess"),
                                             "reversed_containment": ({"line": r["line"], "empty_matches":
                                                                       r["empty_matches"]} if r else None)}
                    seen[(o["code"], fn)] = it
        # model input parameters bound to files by guess (B3): one confirmation, never silent
        for mb in getattr(self, "model_bindings", []):
            if not mb["bindings"] and not mb["unbound"]:
                continue
            code = self.codes.get(mb["model"], {})
            ids = {b["input"]: next((d for d, x in self.datasets.items() if x["legacy_path"] == b["path"]), None)
                   for b in mb["bindings"]}
            add("model_input_binding", "medium", [code.get("id")] + sorted(i for i in ids.values() if i),
                [{"kind": "model_parameter", "file": mb["model"], "line": b["line"], "input": b["input"],
                  "type": b["type"], "data_types": b["data_types"], "dataset": ids[b["input"]], "rule": b["rule"],
                  "why": b["why"]} for b in mb["bindings"]], state="agent_guess", model=mb["model"],
                unbound=mb["unbound"])
            items[-1]["agent_guess"] = {b["input"]: ids[b["input"]] for b in mb["bindings"]}
        for lm in self.layer_mismatch:
            if lm.get("reason") == "layer_not_given":
                continue
            ds = self.ds_of_ref(lm["ref"])
            if ds and not self.datasets.get(ds, {}).get("layer_missing"):
                add("layer_missing", "high", [ds], [{"kind": "code_ref", "layer": lm["layer"], "actual": lm["actual"]}])
        for u in self.ctx.unresolved:
            if u["reason"] in ("outside_root",):
                continue
        # a project layer's subset filter that equals a code filter corroborates it (01 AT-14)
        for p in self.projects:
            for l in p.get("layers", []):
                if not l.get("subset_parsed") or not l.get("path"):
                    continue
                did = next((d for d, x in self.datasets.items() if x["legacy_path"] == l["path"]
                            and (l.get("layer") in (None, x["layer"]))), None)
                hits = []
                for o in self.ops:
                    for pname, pv in o["params"].items():
                        val = pv.get("value") if isinstance(pv, dict) else None
                        if isinstance(val, str) and pname in ("EXPRESSION", "where_clause", "SQL", "query") and \
                                did in self.origin_datasets(o) and all(c in parse_subset(val) for c in l["subset_parsed"]):
                            hits.append({"op": o["id"], "file": o["file"], "line": pv.get("arg_line") or o["line"]})
                l["corroborates"] = hits
        items.sort(key=lambda i: (i["kind"], json.dumps(i["about"]), json.dumps(i["evidence"], sort_keys=True)))
        for n, it in enumerate(items, start=1):
            it["id"] = f"OI-{n:02d}"
        self.items = items

    def origin_datasets(self, o: dict, depth=0) -> list[str]:
        out = [r["dataset"] for r in o["reads"]]
        if depth > 10:
            return out
        for m in o["memory_reads"]:
            mr = self.mem_records.get(m["memory"])
            if mr:
                out += self.origin_datasets(self.op_by_id[mr["op"]], depth + 1)
        return sorted(set(out))

    # ------------------------------------------------------------------ assemble
    def assemble(self, started: str, ro_ok: bool, changed: list[str]) -> dict:
        root = self.root
        files = []
        ds_by_rel = defaultdict(list)
        for did, d in self.datasets.items():
            if d["exists"]:
                ds_by_rel[d["legacy_path"]].append(did)
        sidecars = defaultdict(list)
        for e in self.entries:
            if e["kind"] == "sidecar":
                sidecars[e["sidecar_of"]].append(e["path"])
        for e in self.entries:
            h = self.hashes.get(e["path"], {})
            rec = {"path": e["path"], "kind": e["kind"]}
            if e["kind"] == "excluded":
                rec["excluded_reason"] = e["excluded_reason"]
                if "size" in e:
                    rec["size"] = e["size"]
            elif e["kind"] == "symlink":
                rec["symlink_target"] = e["symlink_target"]
            else:
                rec.update({"size": h.get("size"), "sha256": h.get("sha256"), "mtime": h.get("mtime"),
                            "content_sha256": None, "dataset_ids": sorted(ds_by_rel.get(e["path"], [])),
                            "sidecars": sorted(sidecars.get(e["path"], []))})
                if e["kind"] == "sidecar":
                    rec["sidecar_of"] = e["sidecar_of"]
                    rec["stats_untrusted"] = True
                if e["kind"] in ("vector", "table", "gdb") and file_ext(e["path"]) in (".gpkg", ".xlsx", ".xlsm"):
                    try:
                        rec["content_sha256"] = ac.content_sha256(root / e["path"])
                    except Exception as exc:  # noqa: BLE001
                        rec["content_sha256_error"] = str(exc)[:200]
            files.append(rec)
        fmap = {f["path"]: f for f in files}
        mt = [h["mtime_epoch"] for h in self.hashes.values() if "mtime_epoch" in h]
        spread = (max(mt) - min(mt)) if mt else 0
        n_files = len([f for f in files if f["kind"] not in ("excluded", "symlink")])
        mtime_reliable = not (mt and spread <= 3600)  # one copy/checkout session, not a work history
        manifest = "sha256:" + ac.sha256_bytes(ac.cjson([{"path": f["path"], "sha256": f.get("sha256"),
                                                           "size": f.get("size")} for f in files]))
        # legacy roots (absolute prefixes the code used; commented alternatives included)
        roots = {}
        for key, pre in self.res.prefixes.items():
            roots[key] = {"value": pre["value"], "evidence": []}
        for rel, an in sorted(self.analysers.items()):
            for cell, src in an.src_lines.items():
                for i, line in enumerate(src, start=1):
                    for key, r in roots.items():
                        norm = re.sub(r"[\\/]+", "/", line).lower()
                        if key in norm and not line.lstrip().startswith("#"):
                            if not any(e["file"] == rel for e in r["evidence"]):
                                r["evidence"].append({"kind": "code_literal", "file": rel, "line": i})
            for cc in self.codes.get(rel, {}).get("commented_code", []):
                a = cc.get("assigns")
                if not a or not split_parts(a["value"])[0]:
                    continue
                live_val = an.env.get(a["name"])
                if live_val is not None and is_str(live_val):
                    k2 = "/".join(p.lower() for p in split_parts(live_val.v)[1])
                    if k2 in roots or any(k2.startswith(k) for k in roots):
                        key = "/".join(p.lower() for p in split_parts(a["value"])[1])
                        r = roots.setdefault(key, {"value": a["value"], "evidence": []})
                        if not any(e["file"] == rel for e in r["evidence"]):
                            r["evidence"].append({"kind": "commented_code", "file": rel, "line": cc["line"]})
        legacy_roots = sorted(({"value": r["value"], "evidence": sorted(r["evidence"], key=lambda e: (e["file"],
                                                                                                         e["line"]))}
                               for r in roots.values()), key=lambda r: r["value"])
        # datasets: profiles + hints
        profiles_out = {}
        for did, d in self.datasets.items():
            rel = d["legacy_path"]
            prof = self.profiles.get(rel)
            f = fmap.get(rel)
            if f:
                d["sha256"] = f.get("sha256")
                d["content_sha256"] = f.get("content_sha256")
                d["sidecars"] = f.get("sidecars", [])
            if d["exists"] and prof:
                if d["kind"] == "vector":
                    lay = next((l for l in prof.get("layers", []) if l["name"] == d["layer"]), None)
                    p = {"driver": prof.get("driver"), "layers": [l["name"] for l in prof.get("layers", [])],
                         **(lay or {})}
                elif prof.get("format") == "xlsx":
                    p = {k: v for k, v in prof.items() if k != "sheets"}
                    p["sheets"] = [{k: v for k, v in s.items() if not k.startswith("_")} for s in prof["sheets"]]
                elif prof.get("format") == "csv":
                    p = {k: v for k, v in prof.items() if not k.startswith("_")}
                else:
                    p = prof
                p = jsonable(p)
                profiles_out[did] = dict(p, dataset=did, legacy_path=rel, layer=d["layer"])
                d["profile_ref"] = f"discover/profiles/{did}.json"
            hints = {"kind": d["kind"], "volatile_bytes": []}
            if rel.lower().endswith(".gpkg"):
                hints["volatile_bytes"] = ["gpkg_contents.last_change"]
            if rel.lower().endswith((".xlsx", ".xlsm")):
                hints["volatile_bytes"] = ["docProps/core.xml"]
            if d["kind"] == "raster" and prof:
                hints["categorical"] = bool(prof.get("unique") is not None and len(prof["unique"]) <= 64)
                hints["dtype"] = prof.get("dtype")
                hints["nodata"] = prof.get("nodata")
            elif d["kind"] == "vector" and did in profiles_out:
                hints["categorical"] = False
                hints["key_candidates"] = [k["col"] for k in profiles_out[did].get("key_candidates", [])]
                hints["layer"] = d["layer"]
            elif prof and prof.get("format") == "xlsx":
                s = prof["sheets"][0] if prof["sheets"] else {}
                res_s = next((x for x in prof["sheets"] if x.get("header_state") == "resolved"), s)
                hints.update({"sheet": res_s.get("name"), "header_row": res_s.get("header_row"),
                              "data_rows": res_s.get("data_rows"), "categorical": False})
            elif prof and prof.get("format") == "csv":
                hints.update({"header_row": 1, "sep": prof.get("sep"), "categorical": False})
            if rel in self.vrt_alias:
                d["alias_of"] = next((x for x, y in self.datasets.items() if y["legacy_path"] == self.vrt_alias[rel]),
                                     None)
            d["compare_hints"] = hints
            d["reference_candidate"] = bool(d["exists"] and d["role"] in ("intermediate", "final"))
        step_candidates = []
        for rel, c in sorted(self.codes.items()):
            if not c.get("ops"):
                continue
            step_candidates.append({"id": c["step_id"], "code": c["id"], "ops": c["ops"], "status": c["status"],
                                    "order": c.get("order"), "order_evidence": c.get("order_evidence", []),
                                    "kind": "notebook" if rel.endswith(".ipynb") else (
                                        "qgis_model" if rel.endswith(".model3") else "script"),
                                    "runtime": c.get("runtime"), "source": rel,
                                    "lines": [[min(self.op_by_id[o]["line"] for o in c["ops"]),
                                               max(self.op_by_id[o]["end_line"] for o in c["ops"])]]})
        step_candidates.sort(key=lambda s: (s["order"] is None, s["order"] or 0, s["id"]))
        unresolved = sorted({json.dumps({k: (self.step_of.get(u["file"]) and f"code.{self.step_of[u['file']]}"
                                             if k == "code" else v) for k, v in u.items()}, sort_keys=True)
                             for u in self.ctx.unresolved})
        prev = self.previous()
        inv = {
            "schema": SCHEMA,
            "legacy_root": posix(os.path.relpath(root, self.out)),
            "scan": {"version": SCAN_VERSION, "started_at": started, "libs": self.libs(),
                     "manifest_sha256": manifest, "read_only_verified": ro_ok, "changed_during_scan": changed,
                     "mtime_reliable": mtime_reliable,
                     "mtime_reason": (f"{n_files}/{n_files} files within {int(math.ceil(spread))} s (copy or checkout "
                                      "artefact; corroboration only)" if not mtime_reliable
                                      else f"mtimes spread over {int(spread)} s"),
                     "legacy_roots": legacy_roots, "gdal_pam_enabled": os.environ.get("GDAL_PAM_ENABLED"),
                     "tables": {"qgis_algs": self.ctx.qgis.get("source", {}).get("qgis_version"),
                                "qgis_algs_count": self.ctx.qgis.get("count"),
                                "arcpy_tools_count": len(self.ctx.tools)}},
            "files": files,
            "datasets": [self.datasets[k] for k in sorted(self.datasets)],
            "ops": self.ops,
            "code": [self.codes[k] for k in sorted(self.codes)],
            "settings": self.settings,
            "params": self.params,
            "unresolved": [json.loads(u) for u in unresolved],
            "memory_datasets": sorted(self.mem_records.values(), key=lambda m: m["id"]),
            "step_candidates": step_candidates,
            "manual_step_candidates": self.manual,
            "derived_copies": self.dcopies,
            "twins": sorted(self.twins, key=lambda t: (t["a"], t["b"])),
            "implementations": self.impl_groups,
            "toolboxes": self.toolboxes,
            "projects": sorted(self.projects, key=lambda p: p["path"]),
            "docs": self.docs,
            "grids": self.grid_list,
            "extent_relations": self.extents,
            "aoi": self.aoi,
            "join_coverage": self.joins,
            "branches": self.branch_list,
            "legacy_output_dirs": self.legacy_output_dirs,
            "open_items": self.items,
            "map": prev.get("map") or {"pngs": [], "findings": []},
            "annotations": prev.get("annotations") or {},
        }
        inv = jsonable(inv)
        self.profiles_out = profiles_out
        return inv

    def previous(self) -> dict:
        p = self.out / "inventory.json"
        if p.exists():
            try:
                return ac.load_json(p)
            except ValueError:
                return {}
        return {}

    @staticmethod
    def libs() -> dict:
        out = {}
        try:
            import rasterio
            out["rasterio"] = rasterio.__version__
            out["gdal"] = rasterio.__gdal_version__
        except ImportError:  # pragma: no cover
            pass
        try:
            import pyogrio
            out["pyogrio"] = pyogrio.__version__
            out["pyogrio_gdal"] = pyogrio.__gdal_version_string__
        except ImportError:  # pragma: no cover
            pass
        try:
            import shapely
            import pyproj
            out["shapely"] = shapely.__version__
            out["pyproj"] = pyproj.__version__
            out["proj"] = pyproj.proj_version_str
        except ImportError:  # pragma: no cover
            pass
        try:
            import openpyxl
            out["openpyxl"] = openpyxl.__version__
        except ImportError:  # pragma: no cover
            pass
        return out

    def write(self, inv: dict) -> None:
        out = self.out
        (out / "discover" / "profiles").mkdir(parents=True, exist_ok=True)
        keep = set()
        for did, p in sorted(self.profiles_out.items()):
            path = out / "discover" / "profiles" / f"{did}.json"
            ac.dump_json(path, p)
            keep.add(path.name)
        for old in (out / "discover" / "profiles").glob("*.json"):
            if old.name not in keep:
                old.unlink()
        ac.dump_json(out / "inventory.json", inv)
        ac.atomic_write_text(out / "discover" / "scan.log", "\n".join(self.log.lines) + "\n")


# ---------------------------------------------------------------------------------------------------------------
# annotate (agent findings; merged by id; survive a rescan)
# ---------------------------------------------------------------------------------------------------------------
def annotate(migration: Path, what: str, payload: dict | None, path: str | None, kind: str | None) -> tuple[int, str]:
    inv_p = migration / "inventory.json"
    if not inv_p.exists():
        return 2, "inventory.json not found: run scan.py first"
    inv = ac.load_json(inv_p)
    inv.setdefault("map", {"pngs": [], "findings": []})
    inv.setdefault("annotations", {})
    if what == "png":
        p = migration / path
        if not p.is_file():
            return 2, f"PNG not found: {path}"
        rec = {"path": posix(path), "kind": kind or stem_of(path), "sha256": ac.sha256_file(p)}
        inv["map"]["pngs"] = sorted([x for x in inv["map"]["pngs"] if x["path"] != rec["path"]] + [rec],
                                    key=lambda x: x["path"])
        ac.dump_json(inv_p, inv)
        return 0, f"registered {rec['path']}"
    if payload is None:
        return 2, "--json is required"
    if what == "map-finding":
        f = dict(payload)
        is_view = "view" in f or f.get("receipt") == "view"
        if is_view and "png" in f:
            return 2, "a finding cites a png or a view, not both"
        for k in ("view" if is_view else "png", "observation", "severity"):
            if not f.get(k):
                return 2, f"finding needs {k}"
        if f["severity"] not in ("info", "watch", "question"):
            return 2, "severity must be info | watch | question"
        if len(f["observation"]) > 400:
            return 2, "observation longer than 400 characters"
        m = f.get("metric")
        has_metric = isinstance(m, dict) and isinstance(m.get("value"), (int, float)) and not isinstance(
            m.get("value"), bool) and m.get("name") and m.get("computed_by")
        if m is not None and not has_metric:
            return 2, "metric must be {name, value (number), computed_by}"
        if not has_metric and f["severity"] != "question":
            return 2, "a finding without a computed metric can only raise a question (severity: question)"
        f.setdefault("about", [])
        if is_view:
            import visual_receipt as vr       # lib/: the one rule for view receipts (verify and discover)
            extra = sorted(set(f) - {"id", "view", "views_seen", "engine", "cells_read", "features_read", "about",
                                     "observation", "severity", "metric", "location", "receipt", "model"})
            if extra:
                return 2, f"unknown field(s) in a view finding: {', '.join(extra)}"
            probs = vr.check_inputs_on_disk(migration, f)
            if probs:
                return 2, "view receipt refused:\n  - " + "\n  - ".join(probs)
            f["receipt"] = "view"
            f["about"] = sorted(set(f["about"]) | {x["dataset"] for x in (f.get("cells_read") or [])
                                                   + (f.get("features_read") or [])})
            cite = f["view"]["project"]
        else:
            if not (migration / f["png"]).is_file():
                return 2, f"png {f['png']} not found under the migration folder"
            f["png_sha256"] = ac.sha256_file(migration / f["png"])
            cite = f["png"]
        f["state"] = "ai_inferred"
        f["raises_question"] = not has_metric or f["severity"] == "question"
        f["id"] = f.get("id") or "MF-" + hashlib.sha256(f"{cite}|{f['observation']}".encode()).hexdigest()[:8]
        inv["map"]["findings"] = sorted([x for x in inv["map"]["findings"] if x["id"] != f["id"]] + [f],
                                        key=lambda x: x["id"])
        ac.dump_json(inv_p, inv)
        return 0, f"finding {f['id']} recorded"
    if what == "note":
        a = dict(payload)
        if not a.get("text"):
            return 2, "annotation needs text"
        a["id"] = a.get("id") or "AN-" + hashlib.sha256(a["text"].encode()).hexdigest()[:8]
        a["state"] = "ai_inferred"
        a.setdefault("about", [])
        inv["annotations"][a["id"]] = a
        inv["annotations"] = dict(sorted(inv["annotations"].items()))
        ac.dump_json(inv_p, inv)
        return 0, f"annotation {a['id']} recorded"
    return 2, f"unknown annotate target {what}"


# ---------------------------------------------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------------------------------------------
def summary(inv: dict) -> str:
    roles = Counter(d["role"] for d in inv["datasets"])
    files = [f for f in inv["files"]]
    lines = [f"axis-migrate scan: {len(files)} files ({sum(1 for f in files if f['kind'] == 'sidecar')} sidecars), "
             f"{len(inv['datasets'])} datasets, {len(inv['ops'])} ops",
             "roles: " + ", ".join(f"{k} {v}" for k, v in sorted(roles.items())),
             f"reference candidates: {sum(1 for d in inv['datasets'] if d['reference_candidate'])}",
             f"grids: " + ", ".join(f"{g['id']} ({len(g['members'])} rasters"
                                    + (f", {len(g['misaligned'])} misaligned" if g["misaligned"] else "") + ")"
                                    for g in inv["grids"]),
             "code: " + ", ".join(f"{c['path']} {c['status']}" for c in inv["code"]),
             f"open items: {len(inv['open_items'])} (" + ", ".join(
                 f"{k} {v}" for k, v in sorted(Counter(i['kind'] for i in inv['open_items']).items())) + ")",
             f"read-only verified: {inv['scan']['read_only_verified']}"]
    return "\n".join(lines)


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv and argv[0] == "annotate":
        ap = argparse.ArgumentParser(prog="scan.py annotate", description="Record an agent finding or annotation.")
        ap.add_argument("migration", type=Path)
        ap.add_argument("what", choices=["map-finding", "note", "png"])
        ap.add_argument("--json", dest="json_path", help="JSON file with the finding / annotation ('-' = stdin)")
        ap.add_argument("--path", help="png path relative to the migration folder (what=png; CLI-internal, "
                                       "map.py registers its PNGs through scan.annotate())")
        ap.add_argument("--kind", help="png kind (what=png; CLI-internal)")
        a = ap.parse_args(argv[1:])
        payload = None
        if a.json_path:
            payload = json.loads(sys.stdin.read() if a.json_path == "-" else Path(a.json_path).read_text())
        code, msg = annotate(a.migration.resolve(), a.what, payload, a.path, a.kind)
        print(msg, file=sys.stderr if code else sys.stdout)
        return code
    ap = argparse.ArgumentParser(prog="scan.py", description=__doc__.split("\n\n")[0])
    ap.add_argument("legacy", type=Path, help="read-only legacy folder")
    ap.add_argument("--out", type=Path, required=True, help="the migration/ folder to write")
    ap.add_argument("--json", action="store_true", help="print the inventory summary as JSON")
    a = ap.parse_args(argv)
    if not a.legacy.is_dir():
        print(f"legacy folder not found: {a.legacy}", file=sys.stderr)
        return 2
    legacy = a.legacy.resolve()
    out = a.out.resolve()
    if out == legacy or legacy in out.parents:
        print("--out must be outside the legacy folder (the legacy folder is never written)", file=sys.stderr)
        return 2
    s = Scan(legacy, out)
    inv, code = s.run()
    s.write(inv)
    if a.json:
        print(json.dumps({"files": len(inv["files"]), "datasets": len(inv["datasets"]), "ops": len(inv["ops"]),
                          "open_items": len(inv["open_items"]), "read_only_verified": inv["scan"]["read_only_verified"]},
                         sort_keys=True))
    else:
        print(summary(inv))
    if code:
        print("READ-ONLY CHECK FAILED: the legacy folder changed during the scan: "
              + ", ".join(inv["scan"]["changed_during_scan"][:10]), file=sys.stderr)
    return code


if __name__ == "__main__":
    sys.exit(main())

"""Synthetic legacy folders for the discover scan/map/lineage tests (generic names, no fixture facts).

`build_project(root)` writes a small PyQGIS-style project:
  legacy/in/{raw.tif (+ .aux.xml), zones.gpkg|zones, zones_src.geojson, rates.xlsx, extra_display.geojson}
  legacy/scripts/{01_clip.py, 02_mask.py, 03_table.py, 03b_mask_arcpy_OLD.py}
  legacy/tmp/{01_outline.gpkg|01_outline, 01_clip.tif, 02_mask.tif}   legacy/out/{result.csv}
  legacy/NOTES.txt   legacy/project/map.qgs
`build_arcpy(root)`, `build_notebook(root)`, `build_aprx(root)`, `build_model3(root)` write single-purpose folders.
Each helper returns the legacy folder path. Tests copy nothing from the real fixtures.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
DISCOVER = HERE.parent
PLUGIN = DISCOVER.parents[1]
AXPY = PLUGIN / "scripts" / "axpy"
AXPY_CMD = ([shutil.which("sh") or "sh"] if os.name == "nt" else []) + [str(AXPY)]  # Windows cannot exec a shell script
SCAN = DISCOVER / "scan.py"
MAP = DISCOVER / "map.py"
LINEAGE = DISCOVER / "lineage.py"
GRAPH = DISCOVER / "graph.py"
for p in (str(DISCOVER), str(PLUGIN / "lib")):
    if p not in sys.path:
        sys.path.insert(0, p)

os.environ["GDAL_PAM_ENABLED"] = "NO"

RES = 0.01
ORIGIN = (10.0, 50.2)
W, H = 30, 20


def _transform(dx=0.0, dy=0.0):
    from rasterio.transform import from_origin
    return from_origin(ORIGIN[0] + dx * RES, ORIGIN[1] - dy * RES, RES, RES)


def write_raster(path: Path, arr, *, dtype="float32", nodata=-1.0, dx=0.0, dy=0.0):
    import rasterio
    path.parent.mkdir(parents=True, exist_ok=True)
    with rasterio.open(path, "w", driver="GTiff", width=arr.shape[1], height=arr.shape[0], count=1, dtype=dtype,
                       crs="EPSG:4326", transform=_transform(dx, dy), nodata=nodata) as dst:
        dst.write(arr.astype(dtype), 1)


def zones_gdf():
    import geopandas as gpd
    from shapely.geometry import box
    x0, y1 = ORIGIN
    geoms = [box(x0 + 0.02, y1 - 0.18, x0 + 0.10, y1 - 0.02), box(x0 + 0.10, y1 - 0.18, x0 + 0.20, y1 - 0.02),
             box(x0 + 0.20, y1 - 0.18, x0 + 0.28, y1 - 0.02)]
    return gpd.GeoDataFrame({"name": ["Zone North", "Zone Middle", "Zone South-East"], "code": ["N", "M", "SE"]},
                            geometry=geoms, crs="EPSG:4326")


def build_project(root: Path, *, shift_mask: bool = False) -> Path:
    import numpy as np
    import openpyxl
    import pyogrio
    leg = root / "legacy"
    (leg / "in").mkdir(parents=True, exist_ok=True)
    (leg / "tmp").mkdir(parents=True, exist_ok=True)
    (leg / "out").mkdir(parents=True, exist_ok=True)
    (leg / "scripts").mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(7)
    raw = rng.random((H, W)) * 10
    raw[0, 0] = -1.0
    write_raster(leg / "in" / "raw.tif", raw)
    (leg / "in" / "raw.tif.aux.xml").write_text(
        '<PAMDataset><PAMRasterBand band="1"><Metadata><MDI key="STATISTICS_APPROXIMATE">YES</MDI>'
        '<MDI key="STATISTICS_MEAN">4.9</MDI></Metadata></PAMRasterBand></PAMDataset>')
    src = zones_gdf()
    pyogrio.write_dataframe(src, leg / "in" / "zones_src.geojson", driver="GeoJSON")
    edit = src.copy()
    edit["zone"] = ["North", "Middle", "SE"]  # prefix "Zone " dropped twice, one exception
    edit = edit.drop(columns=["code"])
    pyogrio.write_dataframe(edit, leg / "in" / "zones.gpkg", layer="zones", driver="GPKG")
    disp = src.iloc[:1].copy()
    pyogrio.write_dataframe(disp, leg / "in" / "extra_display.geojson", driver="GeoJSON")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "rates"
    ws["A1"] = "Rates by zone (working copy)"
    ws["A2"] = "units:"
    ws["B2"] = "k per item"
    ws["A4"], ws["B4"] = "Zone", "Rate"
    for i, (z, r) in enumerate([("North", 1.5), ("Middle", 2.0), ("SE ", 2.5)], start=5):
        ws.cell(row=i, column=1, value=z)
        ws.cell(row=i, column=2, value=r)
    ws["A9"] = "mean"
    ws["B9"] = "=AVERAGE(B5:B7)"
    wb.save(leg / "in" / "rates.xlsx")
    # legacy outputs
    clip = raw.copy()
    clip[:, :2] = -1.0
    write_raster(leg / "tmp" / "01_clip.tif", clip)
    mask = (raw > 5).astype("uint8")
    write_raster(leg / "tmp" / "02_mask.tif", mask, dtype="uint8", nodata=0, dx=0.5 if shift_mask else 0.0)
    out = src.dissolve()
    pyogrio.write_dataframe(out[["geometry"]], leg / "tmp" / "01_outline.gpkg", layer="01_outline", driver="GPKG")
    (leg / "out" / "result.csv").write_text("zone,value\nNorth,1\n")
    (leg / "scripts" / "01_clip.py").write_text('''# 01 - clip the raw raster to the zones
import os
import processing

#BASE = r"E:\\\\work\\\\demo"
BASE = "/srv/old/demo"
IN_RAW = BASE + "/in/raw.tif"
ZONES = BASE + "/in/zones.gpkg|layername=zones"
TMP = BASE + "/tmp"
NODATA = -1

outline = processing.run("native:dissolve", {
    'INPUT': ZONES,
    'FIELD': [],
    'OUTPUT': TMP + "/01_outline.gpkg"})['OUTPUT']

# keep the source grid
processing.run("gdal:cliprasterbymasklayer", {
    'INPUT': IN_RAW,
    'MASK': outline,
    'NODATA': NODATA,
    'CROP_TO_CUTLINE': False,
    'KEEP_RESOLUTION': False,
    'OUTPUT': TMP + "/01_clip.tif"})
''')
    (leg / "scripts" / "02_mask.py").write_text('''# 02 - mask of high values on the clip grid
import processing
from qgis.core import QgsRasterLayer

BASE = "/srv/old/demo"
TMP = BASE + "/tmp"
THRESH = 5   # value units

poly = processing.run("native:pixelstopolygons", {
    'INPUT_RASTER': TMP + "/01_clip.tif", 'RASTER_BAND': 1, 'FIELD_NAME': 'v',
    'OUTPUT': 'memory:'})['OUTPUT']
poly = processing.run("native:extractbyexpression", {
    'INPUT': poly, 'EXPRESSION': '"v" > %d' % THRESH, 'OUTPUT': 'memory:'})['OUTPUT']
ref = QgsRasterLayer(TMP + "/01_clip.tif", "clip")
ext = ref.extent()
extent = "%.10f,%.10f,%.10f,%.10f" % (ext.xMinimum(), ext.xMaximum(), ext.yMinimum(), ext.yMaximum())
processing.run("gdal:rasterize", {
    'INPUT': poly,
    'BURN': 1,
    'UNITS': 0,
    'WIDTH': ref.width(),
    'HEIGHT': ref.height(),
    'EXTENT': extent,
    'NODATA': None,
    'DATA_TYPE': 0,
    'EXTRA': '-at',   # all touched
    'OUTPUT': TMP + "/02_mask.tif"})
''')
    (leg / "scripts" / "03_table.py").write_text('''# 03 - rate table to csv
import csv
import openpyxl
from qgis.core import QgsVectorLayer

BASE = "/srv/old/demo"
wb = openpyxl.load_workbook(BASE + "/in/rates.xlsx", data_only=True)
ws = wb["rates"]
rates = {}
for row in ws.iter_rows(min_row=5, max_row=7, values_only=True):
    rates[row[0].strip()] = row[1] * 1000
lyr = QgsVectorLayer(BASE + "/in/zones.gpkg|layername=zones", "z", "ogr")
rows = []
for f in lyr.getFeatures():
    rows.append([f["zone"], rates.get(f["zone"])])
with open(BASE + "/out/result.csv", "w") as fh:
    w = csv.writer(fh)
    w.writerows(rows)
''')
    (leg / "scripts" / "03b_mask_arcpy_OLD.py").write_text('''# OLD - mask step from the ArcGIS days
# replaced by 02_mask.py. NOT USED.
import arcpy
from arcpy.sa import *
ws = r"Q:\\\\demo"
arcpy.env.workspace = ws + r"\\\\old.gdb"
arcpy.env.snapRaster = ws + r"\\\\tmp\\\\01_clip.tif"
m = Con(Raster(ws + r"\\\\tmp\\\\01_clip.tif") > 5, 1)
m.save(ws + r"\\\\old.gdb\\\\mask")
''')
    (leg / "NOTES.txt").write_text('''demo notes

run order:
  scripts/01_clip.py
  scripts/02_mask.py
  scripts/03_table.py
  03b_mask_arcpy_OLD.py is the old version of 02, don't run.

THRESH 5 (02)
rates are in k -> x1000 in 03.
grid is 500m cells.

manual step:
  zone names fixed by hand: zones_src.geojson -> dropped "Zone " in the attribute table, saved as zones.gpkg.

outputs: out/result.csv
''')
    (leg / "project").mkdir(exist_ok=True)
    (leg / "project" / "map.qgs").write_text('''<!DOCTYPE qgis><qgis version="3.40">
<projectCrs><spatialrefsys><authid>EPSG:4326</authid></spatialrefsys></projectCrs>
<layer-tree-group><layer-tree-group name="input" checked="Qt::Checked">
<layer-tree-layer id="l1" name="zones" checked="Qt::Checked"/>
<layer-tree-layer id="l2" name="extra" checked="Qt::Unchecked"/></layer-tree-group></layer-tree-group>
<projectlayers>
<maplayer type="vector"><id>l1</id><layername>zones</layername><provider>ogr</provider>
<datasource>../in/zones.gpkg|layername=zones|subset="zone" = 'North'</datasource>
<srs><spatialrefsys><authid>EPSG:4326</authid></spatialrefsys></srs></maplayer>
<maplayer type="vector"><id>l2</id><layername>extra</layername><provider>ogr</provider>
<datasource>../in/extra_display.geojson</datasource></maplayer>
<maplayer type="vector"><id>l3</id><layername>db</layername><provider>postgres</provider>
<datasource>dbname='gis' host=db.example user='bob' password='hunter2' table="x"</datasource></maplayer>
<maplayer type="annotation"><id>a</id><layername>Annotations</layername><datasource></datasource></maplayer>
</projectlayers></qgis>
''')
    return leg


AK31 = '''import arcpy
import arcpy as ap
from arcpy.sa import *

src = r"C:\\\\proj\\\\in\\\\roads.shp"
arcpy.analysis.Buffer(src, r"C:\\\\proj\\\\tmp\\\\buf_modern.shp", "10 Meters")
arcpy.Buffer_analysis(src, r"C:\\\\proj\\\\tmp\\\\buf_legacy.shp", "10 Meters")
arcpy.analysis.Clip(in_features=src, clip_features=r"C:\\\\proj\\\\in\\\\aoi.shp", out_feature_class=r"C:\\\\proj\\\\tmp\\\\clip_kw.shp")
r = Raster(r"C:\\\\proj\\\\in\\\\dem.tif")
with arcpy.EnvManager(snapRaster=r, cellSize=10):
    slope = Slope(r, "DEGREE")
    slope.save(r"C:\\\\proj\\\\tmp\\\\slope.tif")
ap.management.CopyFeatures(src, r"C:\\\\proj\\\\tmp\\\\copy_alias.shp")
arcpy.gp.ExtractByMask_sa(r"C:\\\\proj\\\\in\\\\dem.tif", r"C:\\\\proj\\\\in\\\\aoi.shp", r"C:\\\\proj\\\\tmp\\\\ebm_gp.tif")
'''


def build_arcpy(root: Path) -> Path:
    import numpy as np
    leg = root / "legacy_arcpy"
    (leg / "in").mkdir(parents=True, exist_ok=True)
    (leg / "scripts").mkdir(parents=True, exist_ok=True)
    (leg / "tmp").mkdir(parents=True, exist_ok=True)
    write_raster(leg / "in" / "dem.tif", np.arange(W * H, dtype="float32").reshape(H, W))
    import pyogrio
    z = zones_gdf()
    for n in ("roads", "aoi"):
        pyogrio.write_dataframe(z, leg / "in" / f"{n}.shp", driver="ESRI Shapefile")
    (leg / "scripts" / "01_styles.py").write_text(AK31)
    return leg


def nb_cell(src: str, n, outputs=None) -> dict:
    return {"cell_type": "code", "execution_count": n, "metadata": {}, "outputs": outputs or [],
            "source": src.splitlines(keepends=True)}


def notebook(cells: list[dict]) -> dict:
    return {"cells": cells, "metadata": {}, "nbformat": 4, "nbformat_minor": 5}


def build_notebook(root: Path) -> Path:
    import pyogrio
    leg = root / "legacy_nb"
    (leg / "data").mkdir(parents=True, exist_ok=True)
    pyogrio.write_dataframe(zones_gdf(), leg / "data" / "in.shp", driver="ESRI Shapefile")
    cells = [
        nb_cell("%matplotlib inline\nimport geopandas as gpd\nBASE = 'data'\n", 1),
        nb_cell("!ogr2ogr data/out.gpkg data/in.shp\n", 2),
        nb_cell("g = gpd.read_file(BASE + '/out.gpkg')\n", 5,
                [{"output_type": "stream", "name": "stdout", "text": ["wrote data/res.gpkg\n"]}]),
        nb_cell("g.to_file(BASE + '/res.gpkg')\n", 3),
        nb_cell("print('not run yet')\n", None),
    ]
    (leg / "analysis.ipynb").write_text(json.dumps(notebook(cells), indent=1))
    ck = [dict(c) for c in cells]
    ck[3] = nb_cell("g.to_file(BASE + '/res_v2.gpkg')\n", 3)
    (leg / ".ipynb_checkpoints").mkdir()
    (leg / ".ipynb_checkpoints" / "analysis-checkpoint.ipynb").write_text(json.dumps(notebook(ck), indent=1))
    return leg


def build_aprx(root: Path) -> Path:
    import numpy as np
    leg = root / "legacy_aprx"
    (leg / "data").mkdir(parents=True, exist_ok=True)
    write_raster(leg / "data" / "dem.tif", np.ones((H, W)))
    gis = {"type": "CIMGISProject", "mapDefinitions": ["CIMPATH=map/map.json"]}
    mp = {"type": "CIMMapDocument", "mapDefinition": {"name": "Main", "spatialReference": {"wkid": 4326},
                                                       "layers": ["CIMPATH=map/dem.json"]}}
    lay = {"type": "CIMRasterLayer", "name": "dem",
           "dataConnection": {"type": "CIMStandardDataConnection", "workspaceConnectionString": "DATABASE=..\\data",
                              "workspaceFactory": "Raster", "dataset": "dem.tif", "datasetType": "esriDTRasterDataset"}}
    (leg / "maps").mkdir()
    with zipfile.ZipFile(leg / "maps" / "proj.aprx", "w") as z:
        z.writestr("GISProject.json", json.dumps(gis))
        z.writestr("map/map.json", json.dumps(mp))
        z.writestr("map/dem.json", json.dumps(lay))
    return leg


MODEL3 = '''<!DOCTYPE model>
<Option type="Map">
  <Option name="model_name" type="QString" value="demo"/>
  <Option name="children" type="Map">
    <Option name="native:buffer_1" type="Map">
      <Option name="alg_id" type="QString" value="native:buffer"/>
      <Option name="params" type="Map">
        <Option name="INPUT" type="List"><Option type="Map">
          <Option name="source" type="int" value="2"/>
          <Option name="static_value" type="QString" value="in/roads.gpkg"/></Option></Option>
        <Option name="DISTANCE" type="List"><Option type="Map">
          <Option name="source" type="int" value="2"/>
          <Option name="static_value" type="double" value="25"/></Option></Option>
      </Option>
      <Option name="outputs" type="Map">
        <Option name="OUTPUT" type="Map">
          <Option name="name" type="QString" value="buffered"/>
          <Option name="default_value" type="QString" value="out/buffered.gpkg"/>
        </Option>
      </Option>
    </Option>
  </Option>
</Option>
'''


def build_model3(root: Path) -> Path:
    import pyogrio
    leg = root / "legacy_model"
    (leg / "in").mkdir(parents=True, exist_ok=True)
    (leg / "out").mkdir(parents=True, exist_ok=True)
    pyogrio.write_dataframe(zones_gdf(), leg / "in" / "roads.gpkg", layer="roads", driver="GPKG")
    pyogrio.write_dataframe(zones_gdf(), leg / "out" / "buffered.gpkg", layer="buffered", driver="GPKG")
    (leg / "demo.model3").write_text(MODEL3)
    return leg


def run(script: Path, *args, env=None, check=True) -> subprocess.CompletedProcess:
    e = dict(os.environ)
    e.update(env or {})
    r = subprocess.run([*AXPY_CMD, str(script), *map(str, args)], capture_output=True, text=True, env=e)
    if check and r.returncode != 0:
        raise AssertionError(f"{script.name} {args} exit {r.returncode}\n{r.stdout}\n{r.stderr}")
    return r


def scan(legacy: Path, out: Path, **kw) -> dict:
    run(SCAN, legacy, "--out", out, **kw)
    return json.loads((out / "inventory.json").read_text())


def tree_state(root: Path) -> dict:
    import hashlib
    out = {}
    for dp, dn, fn in os.walk(root):
        for f in fn:
            p = Path(dp) / f
            out[str(p.relative_to(root))] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def op_by(inv: dict, step_prefix: str, tool: str) -> list[dict]:
    return [o for o in inv["ops"] if o["id"].startswith(step_prefix) and o["tool"] == tool]


def ds(inv: dict, did: str) -> dict:
    return next(d for d in inv["datasets"] if d["id"] == did)


def copy_tree(src: Path, dst: Path) -> Path:
    shutil.copytree(src, dst, symlinks=True)
    return dst

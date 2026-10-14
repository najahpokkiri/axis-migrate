"""A tiny synthetic two-step project for the target tests (no fixture facts; CONTRACTS §0).

    root/legacy_tgt/03_INPUT/{pop.tif (16x16 float32, EPSG:32631), zones.gpkg (layer zones), rates.xlsx (sheet rates)}
    root/legacy_tgt/05_TEMP/01_pop_clip.tif               legacy intermediate
    root/legacy_tgt/04_FINAL/{zone_stats.gpkg, zone_report.xlsx}   legacy finals
    root/migration/{graph.json, params.yaml (with a target: block), inventory.json, pipeline/, tests/}

`reference_verified(root)` builds it and drives the real tools to a sealed, reference-verified run:
run_local.py → compare.py run → plot.py run → plot.py review-add (Q-all receipts; a test stand-in for the model's
visual review) → compare.py aggregate → seal.py. Steps write a raster, a GPKG and an XLSX (the formats that fail on
seek-less mounts), so the dry run exercises the runtime's staging helper.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import numpy as np

PLUGIN = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(PLUGIN / "lib"))
import axis_common as ac  # noqa: E402

AXPY = str(PLUGIN / "bin" / "axpy")
LEGACY = "legacy_tgt"
CRS = "EPSG:32631"
NODATA_IN = -99999.0
CLIP_NODATA = -9999.0
SHAPE = (16, 16)

S01 = '''\
"""Clip population to the zones (cell-centre rule) on the population grid."""
import numpy as np
from rasterio.features import geometry_mask

STEP_ID = "s01_clip_population"
INPUTS = ["ds.pop", "ds.zones"]
OUTPUTS = ["ds.01_pop_clip"]
PARAMS = ["p.clip_nodata"]


def run(ctx):
    pop, _ = ctx.read_raster("ds.pop")
    grid = ctx.grid_of("ds.pop")
    zones = ctx.read_vector("ds.zones").to_crs(grid.crs)
    inside = geometry_mask(list(zones.geometry), out_shape=grid.shape, transform=grid.transform, invert=True)
    nodata = ctx.param("p.clip_nodata")
    out = np.where(inside & ~np.ma.getmaskarray(pop), np.ma.getdata(pop), np.float32(nodata)).astype("float32")
    ctx.write_raster("ds.01_pop_clip", out, grid=grid, dtype="float32", nodata=nodata)
    ctx.log.info("clipped %d cells", int(inside.sum()))
'''

S02 = '''\
"""Zone population sums and value (sum x rate) as a GPKG layer and an XLSX report."""
import numpy as np
from rasterio.features import rasterize

STEP_ID = "s02_zone_report"
INPUTS = ["ds.01_pop_clip", "ds.rates", "ds.zones"]
OUTPUTS = ["ds.zone_report", "ds.zone_stats"]
PARAMS = ["p.rates_sheet", "p.report_sheet"]


def run(ctx):
    pop, _ = ctx.read_raster("ds.01_pop_clip")
    grid = ctx.grid_of("ds.01_pop_clip")
    zones = ctx.read_vector("ds.zones")
    ids = rasterize(((g, i + 1) for i, g in enumerate(zones.to_crs(grid.crs).geometry)), out_shape=grid.shape,
                    transform=grid.transform, fill=0, dtype="int32")
    vals = np.ma.filled(pop.astype("float64"), 0.0)
    sums = np.bincount(ids.ravel(), weights=vals.ravel(), minlength=len(zones) + 1)[1:]
    wb = ctx.open_workbook("ds.rates")
    ws = wb[ctx.param("p.rates_sheet")]
    rate = {r[0]: float(r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    zones["pop_sum"] = sums
    zones["value"] = [s * rate[n] for n, s in zip(zones["name"], sums)]
    ctx.write_vector("ds.zone_stats", zones)
    import openpyxl
    out = openpyxl.Workbook()
    sh = out.active
    sh.title = ctx.param("p.report_sheet")
    sh.append(["name", "pop_sum", "value"])
    for n, s, v in zip(zones["name"], sums, zones["value"]):
        sh.append([n, float(s), float(v)])
    ctx.save_workbook("ds.zone_report", out)
'''

SMOKE = '''\
import axis_pipeline.s01_clip_population as s01
import axis_pipeline.s02_zone_report as s02


def test_contract_constants():
    assert s01.STEP_ID == "s01_clip_population" and callable(s01.run)
    assert set(s02.OUTPUTS) == {"ds.zone_report", "ds.zone_stats"}
'''

PARAMS = """\
version: 1
params:
  - id: p.clip_nodata
    node: s01_clip_population
    kind: nodata
    value: -9999.0
    resolution: resolved
    source: code
    status: confirmed
  - id: p.rates_sheet
    node: s02_zone_report
    kind: string
    value: rates
    resolution: resolved
    source: input_data
    status: confirmed
  - id: p.report_sheet
    node: s02_zone_report
    kind: string
    value: report
    resolution: resolved
    source: code
    status: confirmed
  - id: p.measure_crs
    node: s02_zone_report
    kind: crs
    value: EPSG:32631
    resolution: resolved
    source: code
    status: confirmed
target:
  platform: databricks
  runtime: dbr-17.3-lts
  compute: job_cluster_single_node
  node_memory_gb: 32
  volume: {catalog: null, schema: null, volume: null}
  io: {no_seek_roots: []}
"""


def _grid():
    from rasterio.transform import from_origin
    return dict(crs=CRS, transform=from_origin(500000.0, 5600000.0, 100.0, 100.0), width=SHAPE[1], height=SHAPE[0])


def _write_tif(path: Path, arr, nodata) -> None:
    import rasterio
    path.parent.mkdir(parents=True, exist_ok=True)
    with rasterio.open(path, "w", driver="GTiff", dtype=arr.dtype.name, count=1, nodata=nodata, **_grid()) as dst:
        dst.write(arr, 1)


def _zones():
    import geopandas as gpd
    from shapely.geometry import box
    a = box(500000, 5600000 - 1000, 500700, 5600000)
    b = box(500700, 5600000 - 1400, 501400, 5600000 - 300)
    return gpd.GeoDataFrame({"name": ["Alpha", "Beta"], "zone_id": [1, 2]}, geometry=[a, b], crs=CRS)


RATES = {"Alpha": 2.5, "Beta": 4.0}


def legacy_outputs(pop, zones):
    from rasterio.features import geometry_mask, rasterize
    g = _grid()
    inside = geometry_mask(list(zones.geometry), out_shape=SHAPE, transform=g["transform"], invert=True)
    clip = np.where(inside & (pop != NODATA_IN), pop, np.float32(CLIP_NODATA)).astype("float32")
    ids = rasterize(((geom, i + 1) for i, geom in enumerate(zones.geometry)), out_shape=SHAPE,
                    transform=g["transform"], fill=0, dtype="int32")
    vals = np.where(clip == CLIP_NODATA, 0.0, clip.astype("float64"))
    sums = np.bincount(ids.ravel(), weights=vals.ravel(), minlength=3)[1:]
    return clip, sums


def _ds(i, path, layer, kind, role, *, producer=None, new=None, sheet=None, compare=None):
    return {"id": i, "title": i[3:], "kind": kind, "role": role, "persistence": "file",
            "legacy": {"path": path, "layer": layer, "sheet": sheet, "exists": True},
            "new": new, "produced_by": {"step": producer, "op": f"{producer}.op1"} if producer else None,
            "consumed_by": [],
            "reference": {"source": "legacy" if role != "input" else "none", "match_method": "lineage"},
            "legacy_in_memory": False, "no_reference_reason": None,
            "compare": dict({"may_be_empty": False, "measure_crs_param": "p.measure_crs" if kind == "vector" else None},
                            **(compare or {}))}


def graph_doc() -> dict:
    steps = [
        {"id": "s01_clip_population", "title": "clip", "kind": "script", "tool": "python", "scope": "migrate",
         "scope_reason": None, "params": ["p.clip_nodata"], "inputs": ["ds.pop", "ds.zones"],
         "outputs": ["ds.01_pop_clip"], "timeout_s": 120,
         "new": {"module": "pipeline/s01_clip_population.py", "function": "run"}},
        {"id": "s02_zone_report", "title": "report", "kind": "script", "tool": "python", "scope": "migrate",
         "scope_reason": None, "params": ["p.rates_sheet", "p.report_sheet"],
         "inputs": ["ds.01_pop_clip", "ds.rates", "ds.zones"], "outputs": ["ds.zone_report", "ds.zone_stats"],
         "timeout_s": 120, "new": {"module": "pipeline/s02_zone_report.py", "function": "run"}},
    ]
    datasets = [
        _ds("ds.pop", "03_INPUT/pop.tif", None, "raster", "input"),
        _ds("ds.zones", "03_INPUT/zones.gpkg", "zones", "vector", "input"),
        _ds("ds.rates", "03_INPUT/rates.xlsx", None, "table", "input", sheet="rates"),
        _ds("ds.01_pop_clip", "05_TEMP/01_pop_clip.tif", None, "raster", "intermediate",
            producer="s01_clip_population", new={"path": "s01_clip_population/01_pop_clip.tif", "layer": None}),
        _ds("ds.zone_stats", "04_FINAL/zone_stats.gpkg", "zone_stats", "vector", "final",
            producer="s02_zone_report", new={"path": "s02_zone_report/zone_stats.gpkg", "layer": "zone_stats"},
            compare={"key": ["name"], "match": "key"}),
        _ds("ds.zone_report", "04_FINAL/zone_report.xlsx", None, "table", "final", sheet="report",
            producer="s02_zone_report", new={"path": "s02_zone_report/zone_report.xlsx", "layer": None},
            compare={"key": ["name"], "match": "key", "header_row": 1, "sheet": "report"}),
    ]
    for d in datasets:
        d["consumed_by"] = sorted(s["id"] for s in steps if d["id"] in s["inputs"])
    edges = []
    for s in steps:
        for i in s["inputs"]:
            edges.append({"id": f"e.{i}->{s['id']}", "from": i, "to": s["id"], "kind": "reads", "state": "verified",
                          "evidence": [{"state": "directly_observed", "ref": "scripts/legacy.py:5", "summary": "reads"}]})
        for o in s["outputs"]:
            edges.append({"id": f"e.{s['id']}->{o}", "from": s["id"], "to": o, "kind": "writes", "state": "verified",
                          "evidence": [{"state": "directly_observed", "ref": "scripts/legacy.py:9",
                                        "summary": "writes"}]})
    for s in steps:                     # the plan approval the runner requires (PLAN_NOT_APPROVED otherwise)
        s["plan"] = {"approved": {"by": "user:synthetic", "at": "2026-10-13T00:00:00Z", "channel": "test",
                                  "verbatim": "approve"}}
    return {"schema": "axis-migrate.graph.v1",
            "project": {"name": "synth target", "legacy_root": f"../{LEGACY}", "legacy_tree_sha256": None,
                        "plugin_version": ac.PLUGIN_VERSION, "created_at": "2026-10-13T00:00:00Z"},
            "revision": 1, "steps": steps, "datasets": datasets, "edges": sorted(edges, key=lambda e: e["id"]),
            "questions": [], "invariants": [], "legacy_output_dirs": ["04_FINAL", "05_TEMP"]}


def build_project(root: Path) -> Path:
    import openpyxl
    import pyogrio
    legacy = root / LEGACY
    rng = np.random.default_rng(11)
    pop = (rng.random(SHAPE) * 5).astype("float32")
    pop[0, 15] = NODATA_IN
    _write_tif(legacy / "03_INPUT" / "pop.tif", pop, NODATA_IN)
    zones = _zones()
    pyogrio.write_dataframe(zones, legacy / "03_INPUT" / "zones.gpkg", layer="zones", driver="GPKG")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "rates"
    ws.append(["name", "rate"])
    for n, r in RATES.items():
        ws.append([n, r])
    wb.save(legacy / "03_INPUT" / "rates.xlsx")
    clip, sums = legacy_outputs(pop, zones)
    _write_tif(legacy / "05_TEMP" / "01_pop_clip.tif", clip, CLIP_NODATA)
    (legacy / "04_FINAL").mkdir(parents=True)
    z2 = zones.copy()
    z2["pop_sum"] = sums
    z2["value"] = [s * RATES[n] for n, s in zip(z2["name"], sums)]
    pyogrio.write_dataframe(z2, legacy / "04_FINAL" / "zone_stats.gpkg", layer="zone_stats", driver="GPKG")
    rep = openpyxl.Workbook()
    sh = rep.active
    sh.title = "report"
    sh.append(["name", "pop_sum", "value"])
    for n, s, v in zip(z2["name"], sums, z2["value"]):
        sh.append([n, float(s), float(v)])
    rep.save(legacy / "04_FINAL" / "zone_report.xlsx")
    (legacy / "scripts").mkdir()
    (legacy / "scripts" / "legacy.py").write_text("# synthetic legacy script (never run)\n" * 10)

    mig = root / "migration"
    pipe = mig / "pipeline"
    pipe.mkdir(parents=True)
    for f in ("_axis_runtime.py", "__main__.py", "__init__.py"):
        (pipe / f).write_bytes((PLUGIN / "skills" / "migrate" / "runtime" / f).read_bytes())
    (pipe / "s01_clip_population.py").write_text(S01)
    (pipe / "s02_zone_report.py").write_text(S02)
    (mig / "tests").mkdir()
    (mig / "tests" / "test_smoke.py").write_text(SMOKE)
    graph_py().init(mig, graph_doc())        # graph.json is only ever written by graph.py (GRAPH_OPS.md §8)
    (mig / "params.yaml").write_text(PARAMS)
    (mig / "tolerances.yaml").write_text("version: 1\ndefaults: {}\nrecords: []\n")
    (mig / "decisions.md").write_text("# decisions\n")
    (mig / "known-differences.md").write_text("# known differences\n")
    files = []
    for p in sorted(legacy.rglob("*")):
        if p.is_file():
            r = p.relative_to(legacy).as_posix()
            files.append({"path": r, "size": p.stat().st_size, "sha256": ac.sha256_file(p),
                          "content_sha256": ac.content_sha256(p)})
    roles = [{"id": d["id"], "legacy_path": d["legacy"]["path"], "role": d["role"], "kind": d["kind"],
              "exists": True} for d in graph_doc()["datasets"]]
    (mig / "inventory.json").write_text(json.dumps({"schema": "axis-migrate.inventory.v1",
                                                    "legacy_root": f"../{LEGACY}", "files": files,
                                                    "datasets": roles}, indent=1))
    return mig


def graph_py():
    sys.path.insert(0, str(PLUGIN / "skills" / "discover"))
    import graph
    return graph


def sh(*args, ok=(0,), env=None) -> subprocess.CompletedProcess:
    cp = subprocess.run([AXPY, *map(str, args)], capture_output=True, text=True, timeout=900, env=env)
    if cp.returncode not in ok:
        raise AssertionError(f"{[str(a) for a in args[:3]]} exit {cp.returncode}\n{cp.stdout[-3000:]}\n"
                             f"{cp.stderr[-3000:]}")
    return cp


def receipts(run_dir: Path) -> None:
    for pj in sorted((run_dir / "verify" / "plots").glob("*/plots.json")):
        doc = json.loads(pj.read_text())
        ds = pj.parent.name
        seen = [{"plot": f"plots/{ds}/{e['name']}", "plot_sha256": e["sha256"]} for e in doc.get("png") or []]
        if not seen:
            continue
        f = {"dataset_id": ds, "question_id": "Q-all", "category": "none", "answer": "no", "effect": "none",
             "plot": seen[0]["plot"], "plot_sha256": seen[0]["plot_sha256"], "panel": "all",
             "note": "test stand-in receipt", "model": "test-stand-in", "plots_seen": seen,
             "finding_id": f"vf-{ds[3:]}-all", "run_id": run_dir.name}
        fp = run_dir.parent.parent / "analysis" / f"receipt-{ds}.json"
        fp.parent.mkdir(exist_ok=True)
        fp.write_text(json.dumps(f))
        sh(PLUGIN / "skills/verify/plot.py", "review-add", "--run", run_dir, "--json", fp)


def reference_verified(root: Path) -> tuple[Path, str]:
    mig = build_project(root)
    out = json.loads(sh(PLUGIN / "skills/run/run_local.py", mig, "--json").stdout.strip().splitlines()[-1])
    run_dir = Path(out["run_dir"])
    sh(PLUGIN / "skills/verify/compare.py", "run", "--run", run_dir)
    sh(PLUGIN / "skills/verify/plot.py", "run", "--run", run_dir)
    receipts(run_dir)
    sh(PLUGIN / "skills/verify/compare.py", "aggregate", "--run", run_dir)
    sh(PLUGIN / "skills/report/seal.py", mig, "--run", run_dir.name)
    return mig, run_dir.name

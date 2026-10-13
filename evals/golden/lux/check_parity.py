"""Quick, independent parity check of a lux run against the legacy files (dev/eval only; not compare.py).

    axpy check_parity.py <run_dir> [--legacy fixtures/lux-exposure-legacy] [--json]

Pairs come from the run's snapshot graph (dataset id → legacy path, new path). Rasters: CRS, transform, shape, dtype,
nodata and every cell exactly. Vectors: layer, CRS, schema, row order and every attribute exactly, geometry by
normalised WKB (exact) else symmetric-difference area in EPSG:2169. CSV: bytes. XLSX: cell values (data_only).
Exit 0 only when every pair is exact.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

os.environ.setdefault("GDAL_PAM_ENABLED", "NO")
import numpy as np  # noqa: E402

REPO = Path(__file__).resolve().parents[3]


def raster(leg: Path, new: Path) -> dict:
    import rasterio
    with rasterio.open(leg) as a, rasterio.open(new) as b:
        r = {"crs": a.crs == b.crs, "transform": tuple(a.transform) == tuple(b.transform), "shape": a.shape == b.shape,
             "dtype": a.dtypes == b.dtypes, "nodata": (a.nodata == b.nodata) or (a.nodata is None and b.nodata is None),
             "count": a.count == b.count}
        if r["shape"]:
            x, y = a.read(), b.read()
            same = np.array_equal(x, y, equal_nan=True)
            r["values"] = bool(same)
            if not same:
                d = np.abs(x.astype("float64") - y.astype("float64"))
                r["cells_differ"] = int((x != y).sum())
                r["max_abs_diff"] = float(np.nanmax(d))
        r["exact"] = all(v for k, v in r.items() if isinstance(v, bool))
        return r


def vector(leg: Path, new: Path, layer: str | None, new_layer: str | None) -> dict:
    import geopandas as gpd
    import pyogrio
    import shapely
    la = [lyr for lyr, _ in pyogrio.list_layers(leg)]
    lb = [lyr for lyr, _ in pyogrio.list_layers(new)]
    a = gpd.read_file(leg, layer=layer) if layer else gpd.read_file(leg)
    b = gpd.read_file(new, layer=new_layer or layer) if (new_layer or layer) else gpd.read_file(new)
    r = {"layers": la == lb, "crs": bool(a.crs == b.crs), "rows": len(a) == len(b),
         "columns": list(a.columns) == list(b.columns),
         "dtypes": [str(t) for t in a.dtypes] == [str(t) for t in b.dtypes]}
    if not r["layers"]:
        r["layer_names"] = {"legacy": la, "new": lb}
    if not r["columns"]:
        r["column_names"] = {"legacy": list(a.columns), "new": list(b.columns)}
    if r["rows"] and r["columns"]:
        attrs = [c for c in a.columns if c != a.geometry.name]
        diff_cols = []
        for c in attrs:
            x, y = a[c].to_numpy(), b[c].to_numpy()
            if x.dtype.kind == "f" and y.dtype.kind == "f":
                ok = np.array_equal(x, y, equal_nan=True)
                if not ok:
                    diff_cols.append({"column": c, "rows": int((x != y).sum()),
                                      "max_rel": float(np.nanmax(np.abs(x - y) / np.maximum(np.abs(x), 1e-300)))})
            elif not (a[c].fillna("<NA>") == b[c].fillna("<NA>")).all():
                diff_cols.append({"column": c, "rows": int((a[c].fillna("<NA>") != b[c].fillna("<NA>")).sum())})
        r["attributes"] = not diff_cols
        if diff_cols:
            r["attribute_diffs"] = diff_cols
        wa = shapely.to_wkb(shapely.normalize(a.geometry.values))
        wb = shapely.to_wkb(shapely.normalize(b.geometry.values))
        neq = int(sum(1 for p, q in zip(wa, wb) if p != q))
        r["geometry_exact"] = neq == 0
        if neq:
            r["geometries_differ"] = neq
            ga = a.to_crs(2169).geometry.values
            gb = b.to_crs(2169).geometry.values
            sd = shapely.area(shapely.symmetric_difference(ga, gb))
            r["max_symdiff_m2"] = float(np.max(sd))
            r["topologically_equal"] = bool(np.all(shapely.equals(ga, gb)))
            r["geom_types"] = sorted(set(a.geom_type) ^ set(b.geom_type))
    r["exact"] = all(v for k, v in r.items() if isinstance(v, bool))
    return r


def csv_pair(leg: Path, new: Path) -> dict:
    return {"bytes": leg.read_bytes() == new.read_bytes(), "exact": leg.read_bytes() == new.read_bytes()}


def xlsx(leg: Path, new: Path) -> dict:
    import openpyxl
    a = openpyxl.load_workbook(leg, data_only=True)
    b = openpyxl.load_workbook(new, data_only=True)
    va = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in a.worksheets}
    vb = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in b.worksheets}
    r = {"sheets": list(va) == list(vb), "values": va == vb}
    r["exact"] = r["sheets"] and r["values"]
    return r


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("run_dir")
    ap.add_argument("--legacy", default=str(REPO / "fixtures" / "lux-exposure-legacy"))
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    run = Path(a.run_dir)
    graph = json.loads((run / "snapshot" / "graph.json").read_text())
    man = json.loads((run / "manifest.json").read_text())
    outs = {o["dataset_id"]: o for o in man["outputs"]}
    legacy = Path(a.legacy)
    results = {}
    for d in graph["datasets"]:
        if d["id"] not in outs or not d["legacy"]["exists"]:
            continue
        leg, new = legacy / d["legacy"]["path"], run / outs[d["id"]]["path"]
        kind = d["kind"]
        if kind == "raster":
            r = raster(leg, new)
        elif kind == "vector":
            r = vector(leg, new, d["legacy"]["layer"], (d.get("new") or {}).get("layer"))
        elif leg.suffix == ".csv":
            r = csv_pair(leg, new)
        else:
            r = xlsx(leg, new)
        results[d["id"]] = r
    n_exact = sum(1 for r in results.values() if r["exact"])
    if a.json:
        print(json.dumps({"run": man["run_id"], "pairs": len(results), "exact": n_exact, "results": results},
                         indent=1, sort_keys=True))
    else:
        for k, r in sorted(results.items()):
            extra = {kk: v for kk, v in r.items() if not isinstance(v, bool)}
            print(f"{'EXACT' if r['exact'] else 'DIFF '}  {k:30s} {extra if extra else ''}")
        print(f"{n_exact}/{len(results)} pairs exact")
    return 0 if n_exact == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())

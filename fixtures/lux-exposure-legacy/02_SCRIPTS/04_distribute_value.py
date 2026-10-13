# 04 - canton value = pop x value per capita, spread over residential cells by pop
# outputs -> 04_FINAL
import os, sys, csv
import openpyxl

#BASE = r"C:\GIS\LUX_exposure"
BASE = "/home/migrate/axis-migrate-dev/fixtures/lux-exposure-legacy"
QGIS_PREFIX = "/home/migrate/micromamba/envs/qgis"

VPC_XLSX = BASE + "/03_INPUT/value_per_capita.xlsx"
TMP = BASE + "/04_FINAL/tmp"
OUT = BASE + "/04_FINAL"
CANTON_POP = TMP + "/canton_pop.gpkg|layername=canton_pop"
POP_RES = TMP + "/pop_res.tif"

OUT_TIF = OUT + "/lux_res_exposure_2020.tif"
OUT_GPKG = OUT + "/exposure_by_canton.gpkg"
OUT_CSV = OUT + "/exposure_by_canton.csv"
OUT_XLSX = OUT + "/exposure_by_canton.xlsx"

# --- qgis standalone boilerplate ---
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from qgis.core import QgsApplication, QgsVectorLayer, QgsRasterLayer, QgsField
from qgis.PyQt.QtCore import QMetaType
QgsApplication.setPrefixPath(QGIS_PREFIX, True)
qgs = QgsApplication([], False)
qgs.initQgis()
sys.path.append(QGIS_PREFIX + "/share/qgis/python/plugins")
import processing
from processing.core.Processing import Processing
Processing.initialize()
# -----------------------------------

# value per capita table (kEUR!)
wb = openpyxl.load_workbook(VPC_XLSX, data_only=True)
ws = wb["vpc"]
vpc = {}
for row in ws.iter_rows(min_row=5, max_row=16, values_only=True):
    if row[0] is None:
        continue
    vpc[row[0].strip()] = float(row[1]) * 1000.0   # -> EUR
print(vpc)

lyr = QgsVectorLayer(CANTON_POP, "canton_pop", "ogr")
pr = lyr.dataProvider()
for fname in ["vpc_eur", "tot_value", "factor"]:
    if lyr.fields().indexOf(fname) == -1:
        pr.addAttributes([QgsField(fname, QMetaType.Type.Double)])
lyr.updateFields()
i_vpc = lyr.fields().indexOf("vpc_eur")
i_tot = lyr.fields().indexOf("tot_value")
i_fac = lyr.fields().indexOf("factor")

# per canton
changes = {}
for f in lyr.getFeatures():
    name = f["canton"]
    if name not in vpc:
        print("!! no vpc for", name)
        continue
    tot = f["pop_sum"] * vpc[name]
    if f["res_sum"] and f["res_sum"] > 0:
        fac = tot / f["res_sum"]
    else:
        print("!! no residential pop in", name)
        fac = 0
    print(name, round(f["pop_sum"]), round(f["res_sum"]), tot, fac)
    changes[f.id()] = {i_vpc: vpc[name], i_tot: tot, i_fac: fac}
pr.changeAttributeValues(changes)
del lyr

# factor raster on the pop grid
popres = QgsRasterLayer(POP_RES, "pop_res")
ext = popres.extent()
extent = "%.12f,%.12f,%.12f,%.12f [EPSG:4326]" % (ext.xMinimum(), ext.xMaximum(), ext.yMinimum(), ext.yMaximum())
processing.run("gdal:rasterize", {
    'INPUT': CANTON_POP,
    'FIELD': 'factor',
    'BURN': 0,
    'USE_Z': False,
    'UNITS': 0,
    'WIDTH': popres.width(),
    'HEIGHT': popres.height(),
    'EXTENT': extent,
    'NODATA': None,
    'CREATION_OPTIONS': 'COMPRESS=LZW',
    'DATA_TYPE': 6,
    'INIT': 0,
    'INVERT': False,
    'EXTRA': '',
    'OUTPUT': TMP + "/factor.tif"})

# exposure = residential pop x factor (EUR per cell)
processing.run("gdal:rastercalculator", {
    'INPUT_A': POP_RES, 'BAND_A': 1,
    'INPUT_B': TMP + "/factor.tif", 'BAND_B': 1,
    'FORMULA': 'A*B',
    'NO_DATA': -9999,
    'EXTENT_OPT': 0,
    'RTYPE': 5,
    'CREATION_OPTIONS': 'COMPRESS=LZW',
    'EXTRA': '',
    'OUTPUT': OUT_TIF})

# check: sum back per canton
processing.run("native:zonalstatisticsfb", {
    'INPUT': CANTON_POP,
    'INPUT_RASTER': OUT_TIF,
    'RASTER_BAND': 1,
    'COLUMN_PREFIX': 'exp_',
    'STATISTICS': [1],
    'OUTPUT': OUT_GPKG})

out = QgsVectorLayer(OUT_GPKG, "exp", "ogr")
rows = []
for f in out.getFeatures():
    rows.append([f["canton"], f["pop_sum"], f["res_sum"], f["vpc_eur"], f["tot_value"], f["exp_sum"],
                 f["exp_sum"] - f["tot_value"]])
rows.sort(key=lambda r: r[0])
hdr = ["canton", "pop_total", "pop_residential", "value_per_capita_eur", "total_value_eur",
       "exposure_raster_sum_eur", "diff_eur"]

with open(OUT_CSV, "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(hdr)
    for r in rows:
        w.writerow(r)

wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "by_canton"
ws2.append(hdr)
for r in rows:
    ws2.append(r)
ws2.append([])
ws2.append(["TOTAL", sum(r[1] for r in rows), sum(r[2] for r in rows), None,
            sum(r[4] for r in rows), sum(r[5] for r in rows), None])
wb2.save(OUT_XLSX)

for r in rows:
    print(r)
print("done 04")
qgs.exitQgis()

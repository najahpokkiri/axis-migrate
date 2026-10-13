"""Content comparison of two 04_FINAL folders for files whose bytes differ
(GPKG: per-table rows incl. geometry blobs; XLSX: cell values + zip members)."""
import sys, sqlite3, zipfile, hashlib, glob, os
import openpyxl
A, B = sys.argv[1], sys.argv[2]

def gpkg_dump(p):
    con = sqlite3.connect(p)
    out = {}
    for (t,) in con.execute("select name from sqlite_master where type='table' order by name"):
        rows = con.execute(f'select * from "{t}"').fetchall()
        out[t] = hashlib.sha256(repr(sorted(rows, key=repr)).encode()).hexdigest()
    con.close()
    return out

for pa in sorted(glob.glob(A + "/**/*.gpkg", recursive=True)):
    rel = os.path.relpath(pa, A); da, db = gpkg_dump(pa), gpkg_dump(os.path.join(B, rel))
    diff = [t for t in sorted(set(da) | set(db)) if da.get(t) != db.get(t)]
    print(f"{rel}: differing tables = {diff}")
    if diff:
        for p in (pa, os.path.join(B, rel)):
            con = sqlite3.connect(p)
            for t in diff:
                if t in ("gpkg_contents", "gpkg_metadata", "gpkg_metadata_reference"):
                    print("   ", t, con.execute(f'select * from "{t}"').fetchall()[:2] if t != "gpkg_contents" else con.execute("select table_name,last_change from gpkg_contents").fetchall())
            con.close()

for pa in sorted(glob.glob(A + "/**/*.xlsx", recursive=True)):
    rel = os.path.relpath(pa, A); pb = os.path.join(B, rel)
    va = [r for ws in openpyxl.load_workbook(pa).worksheets for r in ws.iter_rows(values_only=True)]
    vb = [r for ws in openpyxl.load_workbook(pb).worksheets for r in ws.iter_rows(values_only=True)]
    za, zb = zipfile.ZipFile(pa), zipfile.ZipFile(pb)
    mdiff = [n for n in za.namelist() if za.read(n) != zb.read(n)]
    print(f"{rel}: cell values identical = {va == vb}; differing zip members = {mdiff}")

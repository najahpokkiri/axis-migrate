"""Per-step inventory of intermediates/outputs: sha256 + raster grid or vector schema,
plus exact grid-alignment and no-resampling checks against the raw WorldPop grid.
Prints Markdown for PROVENANCE.md."""
import sys, os, glob, hashlib, sqlite3
import numpy as np, rasterio, pyogrio
FX = sys.argv[1]
ref = rasterio.open(FX + "/03_INPUT/lux_ppp_2020.tif")
sha = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()
files = sorted(glob.glob(FX + "/05_TEMP/*")) + sorted(glob.glob(FX + "/04_FINAL/*"))
files = [f for f in files if not f.endswith(".aux.xml")]
print("| step | file | sha256 | CRS | grid / schema |\n|---|---|---|---|---|")
for p in files:
    rel = os.path.relpath(p, FX); name = os.path.basename(p)
    step = name[:2] if name[:2].isdigit() else "04 (final)"
    if p.endswith(".tif"):
        r = rasterio.open(p); t = r.transform
        exact = (t == ref.transform) and r.shape == ref.shape
        dmax = max(abs(a - b) for a, b in zip(t[:6], ref.transform[:6]))
        desc = (f"{r.width}×{r.height}, {r.dtypes[0]}, nodata={r.nodata}, transform=({t.a!r}, {t.b}, {t.c!r}, {t.d}, {t.e!r}, {t.f!r}); "
                f"{'identical to WorldPop transform' if exact else f'max |Δ| vs WorldPop transform {dmax:.1e}'}; compress={r.profile.get('compress')}")
        crs = r.crs.to_string()
    elif p.endswith(".gpkg"):
        parts = []
        for L, _ in pyogrio.list_layers(p):
            i = pyogrio.read_info(p, layer=L)
            props = ", ".join(f"{k}:{t}" for k, t in zip(i["fields"], i["dtypes"]))
            parts.append(f"layer `{L}` ({i['features']} × {i['geometry_type']}): {props}")
            crs = i["crs"]
        desc = "; ".join(parts)
    elif p.endswith(".csv"):
        crs = "–"; desc = "header: " + open(p).readline().strip() + f"; {sum(1 for _ in open(p)) - 1} rows"
    elif p.endswith(".xlsx"):
        import openpyxl
        ws = openpyxl.load_workbook(p).active
        crs = "–"; desc = f"sheet `{ws.title}`, {ws.max_row} rows × {ws.max_column} cols (header + 12 cantons + blank + TOTAL)"
    print(f"| {step} | `{rel}` | `{sha(p)}` | {crs} | {desc} |")

# no-resampling check: clipped pop equals raw WorldPop wherever it is not nodata
a = rasterio.open(FX + "/05_TEMP/01_pop_lux.tif").read(1); b = ref.read(1)
v = a != -99999
print(f"\n01_pop_lux vs raw WorldPop on {int(v.sum())} valid cells: values identical = {bool(np.array_equal(a[v], b[v]))}; "
      f"cells set to nodata by the clip = {int((~v & (b != -99999)).sum())}")

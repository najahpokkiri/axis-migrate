"""Write the SYNTHETIC value-per-capita table (made-up, plausible numbers; not from any
real source). Layout deliberately messy: title row, units cell, blank row, header on
row 4, one name with a trailing space, a national check row below the data."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

out = sys.argv[1]
rows = [  # canton (xlsx spelling), kEUR per inhabitant, comment
    ("Capellen", 182.0, ""),
    ("Clervaux", 141.5, ""),
    ("Diekirch", 148.0, ""),
    ("Echternach", 151.0, ""),
    ("Esch/Alzette", 158.5, "south, older stock"),
    ("Grevenmacher", 167.0, ""),
    ("Luxembourg", 214.0, "incl. city"),
    ("Mersch", 171.0, ""),
    ("Redange", 156.0, ""),
    ("Remich", 169.5, ""),
    ("Vianden", 139.0, "small, check"),
    ("Wiltz ", 143.0, ""),
]
wb = Workbook()
ws = wb.active
ws.title = "vpc"
ws["A1"] = "Residential value per capita by canton (working table)"
ws["A1"].font = Font(bold=True)
ws["A2"] = "units:"
ws["B2"] = "kEUR / inh."
ws["D2"] = "2020 prices, rough"
ws.append([])  # row 3 blank
ws.append(["Canton", "Value per capita", "comment"])  # row 4
for r in ws[4]:
    r.font = Font(bold=True)
for name, v, c in rows:
    ws.append([name, v, c or None])
ws.append([])
ws.append(["national (avg, check)", "=AVERAGE(B5:B16)", None])
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 22
wb.save(out)

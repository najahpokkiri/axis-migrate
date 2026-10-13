"""Independent rasterio/geopandas/numpy implementation of the lux legacy pipeline (dev code, devtests only).

Used for the L2 correct-spread measurement (QGIS reference vs GDAL-via-rasterio) and, with one-line variants, for the
nearest-wrong recipes of the trap catalogue (L2b, L6, L7, L8, L9, L10, L11, L12). It reads only 03_INPUT files plus
the outputs of its own earlier steps (or whatever `pop`/`mask`/... paths the caller passes, e.g. a mutated s01
output for chained-inheritance tests). It is never the reference: references are the legacy run's own files.

Knowledge used here comes from PROVENANCE.md and the legacy scripts (allowed in devtests, never in the plugin).
"""
from __future__ import annotations

import csv
import os
from pathlib import Path

os.environ["GDAL_PAM_ENABLED"] = "NO"
os.environ.setdefault("OGR_CURRENT_DATE", "2000-01-01T00:00:00Z")

import numpy as np  # noqa: E402

POP_NODATA = -99999.0
POPRES_NODATA = float(np.finfo("float32").min)        # -3.4028234663852886e+38, QGIS rastercalc default
EXPO_NODATA = -9999.0
MIN_AREA = 2500.0
POP_MIN = 0.5


def _write_raster(path: Path, arr: np.ndarray, profile: dict, *, dtype: str, nodata, compress: str | None):
    import rasterio
    p = {"driver": "GTiff", "width": profile["width"], "height": profile["height"], "count": 1, "dtype": dtype,
         "crs": profile["crs"], "transform": profile["transform"], "nodata": nodata}
    if compress:
        p["compress"] = compress
    path.parent.mkdir(parents=True, exist_ok=True)
    with rasterio.open(path, "w", **p) as dst:
        dst.write(arr.astype(dtype), 1)


def _write_gpkg(path: Path, gdf, layer: str):
    import pyogrio
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()
    pyogrio.write_dataframe(gdf, path, layer=layer, driver="GPKG")


def _to_multi(geoms):
    from shapely.geometry import MultiPolygon, Polygon
    out = []
    for g in geoms:
        if isinstance(g, Polygon):
            out.append(MultiPolygon([g]))
        else:
            out.append(g)
    return out


def _zone_ids(cantons, profile, all_touched=False):
    from rasterio.features import rasterize
    return rasterize(((g, i + 1) for i, g in enumerate(cantons.geometry)), out_shape=(profile["height"], profile["width"]),
                     transform=profile["transform"], fill=0, all_touched=all_touched, dtype="int32")


def read(path):
    import rasterio
    with rasterio.open(path) as ds:
        return ds.read(1), ds.profile.copy(), ds.nodata


# ------------------------------------------------------------------------------------------------------- s01
def s01(inp: Path, out: Path, *, clip_all_touched=False):
    import geopandas as gpd
    from rasterio.features import geometry_mask
    cantons = gpd.read_file(inp / "cantons_lux_edit.gpkg", layer="cantons")
    outline = cantons.dissolve()                                   # attributes of the first feature, like QGIS
    outline = outline[["shapeName", "shapeISO", "shapeID", "canton", "geometry"]]
    outline["geometry"] = _to_multi(outline.geometry)
    _write_gpkg(out / "s01_prepare_population" / "01_lux_outline.gpkg", outline, "01_lux_outline")
    pop, prof, nd = read(inp / "lux_ppp_2020.tif")
    inside = geometry_mask(list(outline.geometry), out_shape=pop.shape, transform=prof["transform"], invert=True,
                           all_touched=clip_all_touched)
    res = np.where(inside & (pop != nd), pop, POP_NODATA).astype("float32")
    _write_raster(out / "s01_prepare_population" / "01_pop_lux.tif", res, prof, dtype="float32", nodata=POP_NODATA,
                  compress="lzw")


# ------------------------------------------------------------------------------------------------------- s02
def s02(inp: Path, out: Path, *, pop: Path | None = None, all_touched=True, area_crs="EPSG:2169",
        make_valid_method="structure", min_area=MIN_AREA, simplify_m=0.0):
    import geopandas as gpd
    import shapely
    from rasterio.features import rasterize
    pop = pop or out / "s01_prepare_population" / "01_pop_lux.tif"
    lu = gpd.read_file(inp / "osm_landuse_lux.gpkg", layer="landuse").to_crs("EPSG:2169")
    if make_valid_method == "structure":
        lu["geometry"] = shapely.make_valid(lu.geometry.values, method="structure", keep_collapsed=False)
    else:
        lu["geometry"] = shapely.make_valid(lu.geometry.values)          # default "linework": the wrong recipe
    if area_crs == "EPSG:2169":
        lu["area_m2"] = lu.geometry.area
    elif area_crs == "ellipsoidal":                    # QGIS area() with a project ellipsoid set (WGS84)
        from pyproj import Geod
        geod = Geod(ellps="WGS84")
        lu["area_m2"] = [abs(geod.geometry_area_perimeter(g)[0]) for g in lu.geometry.to_crs("EPSG:4326")]
    else:
        lu["area_m2"] = lu.geometry.to_crs(area_crs).area
    res = lu[(lu["landuse"] == "residential") & (lu["area_m2"] >= min_area)].copy()
    res = res[["osm_id", "osm_way_id", "landuse", "name", "area_m2", "geometry"]]
    if simplify_m:
        res["geometry"] = res.geometry.simplify(simplify_m)
    if make_valid_method == "structure":
        res["geometry"] = _to_multi(res.geometry)
    o = out / "s02_landuse_mask"
    _write_gpkg(o / "02_res_2169.gpkg", res, "02_res_2169")
    r4326 = res.to_crs("EPSG:4326")
    _write_gpkg(o / "02_res_4326.gpkg", r4326, "02_res_4326")
    _, prof, _ = read(pop)
    geoms = [g for g in r4326.geometry if g is not None and not g.is_empty]
    mask = rasterize(((g, 1) for g in geoms), out_shape=(prof["height"], prof["width"]), transform=prof["transform"],
                     fill=0, all_touched=all_touched, dtype="uint8") if geoms else \
        np.zeros((prof["height"], prof["width"]), dtype="uint8")
    _write_raster(o / "02_res_mask.tif", mask, prof, dtype="uint8", nodata=0.0, compress="lzw")


# ------------------------------------------------------------------------------------------------------- s03
def _zone_sum(V, Z, valid, i, how):
    if how == "bincount":
        return float(np.bincount(Z[valid], weights=V[valid].astype("float64"), minlength=i + 1)[i])
    return float(V[(Z == i) & valid].astype("float64").sum())


def s03(inp: Path, out: Path, *, pop: Path | None = None, mask: Path | None = None, arcpy_semantics=False,
        zones_all_touched=False, zonal="masked"):
    import geopandas as gpd
    pop = pop or out / "s01_prepare_population" / "01_pop_lux.tif"
    mask = mask or out / "s02_landuse_mask" / "02_res_mask.tif"
    P, prof, pnd = read(pop)
    M, _, mnd = read(mask)
    pv = P != pnd
    mv = M != mnd
    if arcpy_semantics:
        # ArcPy 03b: Con(IsNull(mask), 0, mask); weights 0 (not NoData) outside residential
        m0 = np.where(mv, M, 0).astype("float64")
        w = np.where(P >= POP_MIN, P, 0).astype("float64") * m0
        out_arr = np.where(pv, w, POPRES_NODATA).astype("float32")
    else:
        valid = pv & mv
        w = (P >= POP_MIN).astype("float32") * P * M.astype("float32")
        out_arr = np.where(valid, w, POPRES_NODATA).astype("float32")
    o = out / "s03_zonal_population"
    _write_raster(o / "03_pop_res.tif", out_arr, prof, dtype="float32", nodata=POPRES_NODATA, compress=None)
    cantons = gpd.read_file(inp / "cantons_lux_edit.gpkg", layer="cantons")
    Z = _zone_ids(cantons, prof, all_touched=zones_all_touched)
    R = out_arr
    rv = R != np.float32(POPRES_NODATA)
    rows = {"pop_count": [], "pop_sum": [], "res_count": [], "res_sum": []}
    for i in range(len(cantons)):
        z = Z == i + 1
        rows["pop_count"].append(float((z & pv).sum()))
        rows["pop_sum"].append(_zone_sum(P, Z, pv, i + 1, zonal))
        rows["res_count"].append(float((z & rv).sum()))
        rows["res_sum"].append(_zone_sum(R, Z, rv, i + 1, zonal))
    for k, v in rows.items():
        cantons[k] = v
    _write_gpkg(o / "03_canton_pop.gpkg", cantons, "03_canton_pop")


# ------------------------------------------------------------------------------------------------------- s04
def s04(inp: Path, out: Path, *, pop_res: Path | None = None, canton_pop: Path | None = None, strip=True,
        scale=1000.0, factor_all_touched=False, exposure_dtype="float32", factor_shift_cols=0, zonal="masked"):
    import geopandas as gpd
    import openpyxl
    from rasterio.features import rasterize
    pop_res = pop_res or out / "s03_zonal_population" / "03_pop_res.tif"
    canton_pop = canton_pop or out / "s03_zonal_population" / "03_canton_pop.gpkg"
    wb = openpyxl.load_workbook(inp / "value_per_capita.xlsx", data_only=True)
    vpc = {}
    for row in wb["vpc"].iter_rows(min_row=5, max_row=16, values_only=True):
        if row[0] is None:
            continue
        vpc[row[0].strip() if strip else row[0]] = float(row[1]) * scale
    cf = gpd.read_file(canton_pop, layer="03_canton_pop")
    vpc_eur, tot, fac = [], [], []
    for _, f in cf.iterrows():
        name = f["canton"]
        if name not in vpc:
            vpc_eur.append(None)
            tot.append(None)
            fac.append(None)
            continue
        t = f["pop_sum"] * vpc[name]
        vpc_eur.append(vpc[name])
        tot.append(t)
        fac.append(t / f["res_sum"] if f["res_sum"] and f["res_sum"] > 0 else 0.0)
    cf["vpc_eur"] = np.array(vpc_eur, dtype="float64")
    cf["tot_value"] = np.array(tot, dtype="float64")
    cf["factor"] = np.array(fac, dtype="float64")
    o = out / "s04_distribute_value"
    _write_gpkg(o / "04_canton_factor.gpkg", cf, "canton_factor")
    R, prof, rnd = read(pop_res)
    shapes = [(g, v) for g, v in zip(cf.geometry, cf["factor"]) if v is not None and np.isfinite(v)]
    F = rasterize(shapes, out_shape=R.shape, transform=prof["transform"], fill=0, all_touched=factor_all_touched,
                  dtype="float64")
    if factor_shift_cols:
        F = np.roll(F, factor_shift_cols, axis=1)
    _write_raster(o / "04_factor.tif", F, prof, dtype="float64", nodata=0.0, compress="lzw")
    valid = (R != rnd) & (F != 0.0)
    E = np.where(valid, R.astype("float64") * F, EXPO_NODATA)
    E = E.astype(exposure_dtype)
    _write_raster(o / "lux_res_exposure_2020.tif", E, prof, dtype=exposure_dtype, nodata=EXPO_NODATA, compress="lzw")
    cantons = gpd.read_file(inp / "cantons_lux_edit.gpkg", layer="cantons")
    Z = _zone_ids(cantons, prof)
    ev = E != EXPO_NODATA
    exp_by = {}
    for i, name in enumerate(cantons["canton"]):
        exp_by[name] = _zone_sum(E, Z, ev, i + 1, zonal)
    eg = cf.copy()
    eg["exp_sum"] = [exp_by.get(n) if f is not None else exp_by.get(n) for n, f in zip(eg["canton"], fac)]
    _write_gpkg(o / "exposure_by_canton.gpkg", eg, "exposure_by_canton")
    rows = []
    for _, f in eg.iterrows():
        d = (f["exp_sum"] - f["tot_value"]) if (f["tot_value"] is not None and np.isfinite(f["tot_value"])) else None
        rows.append([f["canton"], f["pop_sum"], f["res_sum"], _none(f["vpc_eur"]), _none(f["tot_value"]),
                     f["exp_sum"], _none(d)])
    rows.sort(key=lambda r: r[0])
    hdr = ["canton", "pop_total", "pop_residential", "value_per_capita_eur", "total_value_eur",
           "exposure_raster_sum_eur", "diff_eur"]
    with open(o / "exposure_by_canton.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(hdr)
        for r in rows:
            w.writerow(["" if x is None else x for x in r])
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "by_canton"
    ws2.append(hdr)
    for r in rows:
        ws2.append(r)
    ws2.append([])
    ws2.append(["TOTAL", sum(r[1] for r in rows), sum(r[2] for r in rows), None,
                sum(r[4] for r in rows if r[4] is not None), sum(r[5] for r in rows), None])
    wb2.save(o / "exposure_by_canton.xlsx")


def _none(x):
    if x is None:
        return None
    try:
        return None if not np.isfinite(x) else float(x)
    except TypeError:
        return x


def run_all(inp: Path, out: Path, **variants) -> Path:
    s01(inp, out, **variants.get("s01", {}))
    s02(inp, out, **variants.get("s02", {}))
    s03(inp, out, **variants.get("s03", {}))
    s04(inp, out, **variants.get("s04", {}))
    return out
